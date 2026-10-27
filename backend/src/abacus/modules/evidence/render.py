"""Deterministic trial balance rendering (ADR-042). PROTECTED. TASK-009 design §6.

Identical data → byte-identical files:
- lines are sorted by account code, then name;
- workbook properties are fixed (creator `platform`; created and modified = the pull time);
- the ZIP container is rebuilt with a fixed timestamp and fixed compression.

Account names come from the client's system and are hostile (AGENTS.md #8). They're written as
explicit strings, never interpreted as formulas. Totals are computed here (code computes,
ADR-050).
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import cast
from uuid import UUID

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_AMOUNT = "#,##0.00;-#,##0.00"
_BOLD = Font(bold=True)


@dataclass(frozen=True)
class TrialBalanceLine:
    account_code: str
    account_name: str
    debit: Decimal
    credit: Decimal


@dataclass(frozen=True)
class TrialBalance:
    """A validated trial balance (TASK-010's ledger snapshot maps onto this)."""

    client_entity_id: UUID
    entity_name: str
    period_start: date
    period_end: date
    pulled_at: datetime
    snapshot_id: UUID
    source: str
    source_fingerprint: str
    lines: tuple[TrialBalanceLine, ...]
    # SPEC-028 (TASK-052): `imported` for an uploaded export, with its details in the footer.
    method: str = "retrieved"
    notes: tuple[tuple[str, str], ...] = ()


_MAX_CELL = 32_767  # Excel's limit; longer values would corrupt the file
_FORMULA_LEADS = ("=", "+", "-", "@")


def _text(sheet: Worksheet, row: int, column: int, value: str) -> None:
    if len(value) > _MAX_CELL:
        raise ValueError("text longer than a spreadsheet cell can hold")
    cell = cast(Cell, sheet.cell(row=row, column=column))
    cell.value = value
    cell.data_type = "s"  # always a string: a leading "=" never becomes a formula
    if value.startswith(_FORMULA_LEADS):
        cell.quotePrefix = True  # and stays text if copied or exported elsewhere


def _amount(sheet: Worksheet, row: int, column: int, value: Decimal) -> None:
    if not value.is_finite():
        raise ValueError("amounts must be finite")
    cell = cast(Cell, sheet.cell(row=row, column=column))
    cell.value = value
    cell.number_format = _AMOUNT


def _utc(moment: datetime) -> datetime:
    """Naive datetimes would be read in the host's zone, so output would vary by machine."""
    if moment.tzinfo is None or moment.utcoffset() is None:
        raise ValueError("pulled_at must be timezone-aware")
    return moment.astimezone(UTC)


# The trial-balance sheet layout, shared with readers of rendered evidence (screening): line rows
# start at FIRST_LINE_ROW; every line has an account code; one Total row (no code, TOTAL_LABEL in
# the name column) follows the last line; only footer text follows it.
CODE_COLUMN, NAME_COLUMN, DEBIT_COLUMN, CREDIT_COLUMN = 1, 2, 3, 4
FIRST_LINE_ROW = 2
TOTAL_LABEL = "Total"


def _workbook(tb: TrialBalance) -> Workbook:
    pulled_at = _utc(tb.pulled_at)
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "Trial balance"
    for column, header in enumerate(("Account", "Name", "Debit", "Credit"), start=CODE_COLUMN):
        _text(sheet, 1, column, header)
        cast(Cell, sheet.cell(row=1, column=column)).font = _BOLD
    lines = sorted(tb.lines, key=lambda line: (line.account_code, line.account_name))
    row = FIRST_LINE_ROW
    for line in lines:
        _text(sheet, row, CODE_COLUMN, line.account_code)
        _text(sheet, row, NAME_COLUMN, line.account_name)
        _amount(sheet, row, DEBIT_COLUMN, line.debit)
        _amount(sheet, row, CREDIT_COLUMN, line.credit)
        row += 1
    _text(sheet, row, NAME_COLUMN, TOTAL_LABEL)
    _amount(sheet, row, DEBIT_COLUMN, sum((line.debit for line in lines), Decimal(0)))
    _amount(sheet, row, CREDIT_COLUMN, sum((line.credit for line in lines), Decimal(0)))
    row += 2
    footer = (
        ("Source", tb.source),
        ("Method", tb.method),
        ("Pulled at", pulled_at.isoformat()),
        ("Period", f"{tb.period_start.isoformat()} to {tb.period_end.isoformat()}"),
        ("Entity", tb.entity_name),
        ("Entity ID", str(tb.client_entity_id)),
        ("Snapshot ID", str(tb.snapshot_id)),
        ("Source fingerprint", tb.source_fingerprint),
        *tb.notes,
    )
    for label, value in footer:
        _text(sheet, row, 1, label)
        _text(sheet, row, 2, value)
        row += 1
    for letter, width in (("A", 16), ("B", 48), ("C", 18), ("D", 18)):
        sheet.column_dimensions[letter].width = width
    properties = workbook.properties
    properties.creator = "platform"
    properties.lastModifiedBy = "platform"
    properties.created = pulled_at.replace(tzinfo=None)
    properties.modified = pulled_at.replace(tzinfo=None)
    return workbook


_MODIFIED = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _normalised(raw: bytes, pulled_at: datetime) -> bytes:
    """Rebuild the ZIP with fixed entry timestamps, attributes and compression. openpyxl stamps
    the save time into `dcterms:modified`; it is set back to the pull time here."""
    stamp = _utc(pulled_at).strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as source,
        zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target,
    ):
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            entry.create_system = 3  # the same on every platform
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data, count = _MODIFIED.subn(lambda m: m.group(1) + stamp + m.group(2), data)
                if count != 1:
                    raise RuntimeError("docProps/core.xml has no modified stamp to pin")
            target.writestr(entry, data)
    return out.getvalue()


def render_trial_balance(tb: TrialBalance) -> bytes:
    buffer = io.BytesIO()
    _workbook(tb).save(buffer)
    return _normalised(buffer.getvalue(), tb.pulled_at)


# --- Open in Excel: the trial balance of record at one layer (A-02-T22; AC-8.1; ADR-042) --------


@dataclass(frozen=True)
class LayerLine:
    """One line of the export: the amounts the ledger computed for the layer, and the line's
    review marks and citations already rendered as text by the caller (the screen's words)."""

    account_code: str
    account_name: str
    area: str
    wp_ref: str
    boy: Decimal | None
    eoy: Decimal
    aje: Decimal | None
    rje: Decimal | None
    final: Decimal | None
    ticks: str
    citations: str


@dataclass(frozen=True)
class LayerExport:
    """A trial balance of record at one layer and version, as `Open in Excel` writes it. Every
    total is the caller's (code computes, ADR-050) and is written as given."""

    title: str
    layer: str  # the layer's word ("Unadjusted", "Final")
    adjusted: bool  # AJE, RJE and Final columns
    lines: tuple[LayerLine, ...]
    total_boy: Decimal | None
    total_eoy: Decimal
    debit: Decimal
    credit: Decimal
    difference: Decimal
    pulled_at: datetime  # the file's fixed stamp: the version's snapshot pull time
    footer: tuple[tuple[str, str], ...]  # provenance rows, in order


def _layer_workbook(tb: LayerExport) -> Workbook:
    stamp = _utc(tb.pulled_at)
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "Trial balance"
    prior = tb.total_boy is not None
    headers = ["Account", "Name", "Area"]
    if prior:
        headers.append("Beginning of Year")
    headers.append("End of Year" if not tb.adjusted else "Unadjusted")
    if tb.adjusted:
        headers += ["AJE", "RJE", "Final"]
    headers += ["WP Ref", "Ticks", "Citations"]
    _text(sheet, 1, 1, f"{tb.title} · {tb.layer}")
    cast(Cell, sheet.cell(row=1, column=1)).font = _BOLD
    for column, header in enumerate(headers, start=1):
        _text(sheet, 3, column, header)
        cast(Cell, sheet.cell(row=3, column=column)).font = _BOLD
    row = 4
    for line in tb.lines:
        values: list[str | Decimal | None] = [line.account_code, line.account_name, line.area]
        if prior:
            values.append(line.boy)
        values.append(line.eoy)
        if tb.adjusted:
            values += [line.aje, line.rje, line.final]
        values += [line.wp_ref, line.ticks, line.citations]
        for column, value in enumerate(values, start=1):
            if isinstance(value, Decimal):
                _amount(sheet, row, column, value)
            elif value is not None:
                _text(sheet, row, column, value)
        row += 1
    _text(sheet, row, 2, TOTAL_LABEL)
    column = 4
    if prior and tb.total_boy is not None:
        _amount(sheet, row, column, tb.total_boy)
        column += 1
    if tb.adjusted:
        column += 3
    _amount(sheet, row, column, tb.total_eoy)
    row += 2
    balance = (
        ("Debits", tb.debit),
        ("Credits", tb.credit),
        ("Difference", tb.difference),
    )
    for label, value in balance:
        _text(sheet, row, 1, label)
        _amount(sheet, row, 2, value)
        row += 1
    row += 1
    for label, value in tb.footer:
        _text(sheet, row, 1, label)
        _text(sheet, row, 2, value)
        row += 1
    for index, width in enumerate((14, 44, 22, 18, 18, 18, 18, 18, 12, 28, 36), start=1):
        sheet.column_dimensions[chr(ord("A") + index - 1)].width = width
    properties = workbook.properties
    properties.creator = "platform"
    properties.lastModifiedBy = "platform"
    properties.created = stamp.replace(tzinfo=None)
    properties.modified = stamp.replace(tzinfo=None)
    return workbook


def render_trial_balance_layers(tb: LayerExport) -> bytes:
    """AC-8.1: the same export gives the same bytes; client strings stay text (ADR-052)."""
    buffer = io.BytesIO()
    _layer_workbook(tb).save(buffer)
    return _normalised(buffer.getvalue(), tb.pulled_at)
