"""Reading an upload's text, tables and fields (SPEC-029 goal 3; TASK-054; ADR-110). PROTECTED.

The file is hostile (ADR-052) and was malware-scanned before it was stored. `extract_isolated`
runs `extract` in a short-lived child process with limits (`kernel.isolation`); a crash, a hang
or a limit hit is `unreadable`, never a worker failure (AC-2).

- PDF (pdfplumber): page text and simple tables, first `MAX_PAGES` pages. When pdfplumber finds
  no text, or can't open the file, PDFium (pypdfium2) reads the text a second way: many card and
  bank e-statements use fonts pdfminer can't map (TASK-323). A file neither reads is
  `not_extracted/needs_ocr`: kept, never a problem for the client, waiting for OCR (P-32). A
  password-protected file is `unreadable/encrypted` (AC-3).
- XLSX: cached values only (`data_only`), read-only, VBA dropped: formulas are never evaluated
  and macros never run (AC-4). CSV: UTF-8 or Windows-1252.
- Images, Word documents and legacy Excel/Word files: `not_extracted` (HEIC with the reason
  `format_not_supported`). A PDF with active content: `not_extracted/active_content`, never
  opened (DEC-296).

The result is a derived record: it never alters the evidence and can be rebuilt (ADR-110).
"""

from __future__ import annotations

import csv
import io
import unicodedata
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from typing import Final, Literal, Protocol, cast

from abacus.kernel.isolation import IsolationFailed, run_isolated
from abacus.kernel.pdf_pages import problem_of
from abacus.modules.evidence.active_content import is_active
from abacus.modules.evidence.document_fields import (
    Aging,
    Statement,
    aging_fields,
    amounts_in,
    dates_in,
    statement_fields,
)

# v2 (TASK-323): PDFium text fallback; no text is needs_ocr. v3 (TASK-334): each PDF text line's
# box on its page, so a value can point back to where it was printed.
EXTRACTOR_VERSION: Final = "v3"
MAX_PAGE_LINES: Final = 2000
MAX_PAGES: Final = 200
MAX_ROWS: Final = 60_000
MAX_CELL: Final = 500
MAX_PAGE_TEXT: Final = 50_000
PDF: Final = "application/pdf"
XLSX: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV: Final = "text/csv"
HEIC: Final = "image/heic"
HEIF: Final = "image/heif"
ACTIVE_CONTENT: Final = "active_content"
_REFUSED: Final = frozenset({"Cc", "Cf", "Cs", "Co"})

Status = Literal["extracted", "unreadable", "not_extracted"]


@dataclass(frozen=True)
class Table:
    sheet: str | None  # a spreadsheet's sheet name
    page: int | None  # a PDF's page number, from 1
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class TextLine:
    """A printed line and its box on the page as fractions of the page (x0, top, x1, bottom),
    so it maps onto the page image at any size."""

    text: str
    box: tuple[float, float, float, float]


@dataclass(frozen=True)
class Extraction:
    status: Status
    reason: str | None = None  # scanned, encrypted, damaged, too_large, crashed, legacy_format…
    page_count: int | None = None
    truncated: bool = False
    pages: tuple[str, ...] = ()
    tables: tuple[Table, ...] = ()
    dates: tuple[date, ...] = ()
    amounts: tuple[Decimal, ...] = field(default=())
    statement: Statement | None = None
    aging: Aging | None = None
    page_lines: tuple[
        tuple[TextLine, ...], ...
    ] = ()  # TASK-334: PDF only, when pdfplumber read it


def _text(value: object) -> str:
    """Plain text: control characters dropped, length capped (ADR-052)."""
    text = "" if value is None else str(value)
    kept = "".join(c for c in text if c in "\n\t" or unicodedata.category(c) not in _REFUSED)
    return kept[:MAX_CELL].strip()


def _reason(exc: BaseException) -> str:
    """`encrypted` when a password is needed (pdfplumber wraps pdfminer's error), else
    `damaged` (K-01-T20: the kernel's reading, shared with the methodology importer)."""
    return problem_of(exc)


class _TextPage(Protocol):  # what pypdfium2's untyped text page offers
    def get_text_range(self) -> str: ...
    def close(self) -> None: ...


class _Page(Protocol):
    def get_textpage(self) -> _TextPage: ...
    def close(self) -> None: ...


def _pdfium_pages(content: bytes) -> tuple[int, list[str]] | None:
    """The page count and each page's text by PDFium, or None when PDFium can't open the file
    (damaged, or a password it doesn't have)."""
    import pypdfium2

    try:
        document = pypdfium2.PdfDocument(content)
    except Exception:  # a damaged, hostile or password-protected file
        return None
    try:
        count = len(document)
        pages: list[str] = []
        for number in range(min(count, MAX_PAGES)):
            page = cast(_Page, document[number])
            textpage = page.get_textpage()
            pages.append(textpage.get_text_range()[:MAX_PAGE_TEXT].replace("\r\n", "\n"))
            textpage.close()
            page.close()
        return count, pages
    except Exception:
        return None
    finally:
        document.close()


def _from_pages(
    count: int,
    pages: Sequence[str],
    tables: Sequence[Table],
    page_lines: Sequence[tuple[TextLine, ...]] = (),
) -> Extraction:
    text = "\n".join(pages)
    return Extraction(
        "extracted",
        page_count=count,
        truncated=count > MAX_PAGES,
        pages=tuple(pages),
        tables=tuple(tables),
        dates=dates_in(text),
        amounts=amounts_in(text),
        statement=statement_fields(text),
        page_lines=tuple(page_lines),
    )


def _lines_of(page: object) -> tuple[TextLine, ...]:
    """pdfplumber's text lines with their boxes as fractions of the page."""
    width = float(getattr(page, "width", 0) or 0)
    height = float(getattr(page, "height", 0) or 0)
    if width <= 0 or height <= 0:
        return ()
    found = cast(list[dict[str, object]], page.extract_text_lines())  # type: ignore[attr-defined]  # pdfplumber Page
    out: list[TextLine] = []
    for line in found[:MAX_PAGE_LINES]:
        x0, top = float(cast(float, line["x0"])), float(cast(float, line["top"]))
        x1, bottom = float(cast(float, line["x1"])), float(cast(float, line["bottom"]))
        box = (
            round(max(0.0, x0 / width), 4),
            round(max(0.0, top / height), 4),
            round(min(1.0, x1 / width), 4),
            round(min(1.0, bottom / height), 4),
        )
        out.append(TextLine(_text(line.get("text")), box))
    return tuple(out)


def _fallback(content: bytes, reason: str, count: int | None) -> Extraction:
    """PDFium's reading when pdfplumber got nothing; else the file waits for OCR (no text) or is
    unreadable (encrypted, damaged)."""
    found = _pdfium_pages(content)
    if found is not None and any(p.strip() for p in found[1]):
        return _from_pages(found[0], found[1], ())
    if reason == "no_text":
        return Extraction(
            "not_extracted",
            "needs_ocr",
            count if count is not None else (found[0] if found else None),
        )
    return Extraction("unreadable", reason, count)


def _pdf(content: bytes) -> Extraction:
    import pdfplumber

    try:
        pdf = pdfplumber.open(io.BytesIO(content))
    except Exception as exc:  # a damaged or hostile file: pdfminer raises many kinds
        return _fallback(content, _reason(exc), None)
    with pdf:
        try:
            encrypted = getattr(pdf.doc, "encryption", None) is not None
            count = len(pdf.pages)
            pages: list[str] = []
            tables: list[Table] = []
            lines: list[tuple[TextLine, ...]] = []
            for number, page in enumerate(pdf.pages[:MAX_PAGES], start=1):
                pages.append((page.extract_text() or "")[:MAX_PAGE_TEXT])
                lines.append(_lines_of(page))
                for found in page.extract_tables():
                    rows = tuple(tuple(_text(c) for c in row) for row in found)
                    tables.append(Table(None, number, rows))
        except Exception as exc:
            return _fallback(content, _reason(exc), None)
    if not any(p.strip() for p in pages):
        return _fallback(content, "encrypted" if encrypted else "no_text", count)
    return _from_pages(count, pages, tables, lines)


def _sheet_rows_xlsx(content: bytes) -> list[tuple[str, list[list[object]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_vba=False)
    sheets: list[tuple[str, list[list[object]]]] = []
    total = 0
    for sheet in workbook.worksheets:
        rows: list[list[object]] = []
        for values in sheet.iter_rows(values_only=True):
            rows.append([cast(object, v) for v in values])
            total += 1
            if total > MAX_ROWS:
                raise OverflowError
        sheets.append((sheet.title, rows))
    workbook.close()
    return sheets


def _sheet_rows_csv(content: bytes) -> list[tuple[str, list[list[object]]]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("encoding")
    csv.field_size_limit(MAX_CELL * 20)
    rows: list[list[object]] = []
    for row in csv.reader(io.StringIO(text)):
        rows.append(list(row))
        if len(rows) > MAX_ROWS:
            raise OverflowError
    return [("", rows)]


def _sheets(content: bytes, media_type: str) -> Extraction:
    try:
        sheets = _sheet_rows_xlsx(content) if media_type == XLSX else _sheet_rows_csv(content)
    except OverflowError:
        return Extraction("unreadable", "too_large")
    except Exception:  # a damaged or hostile workbook: openpyxl raises many kinds
        return Extraction("unreadable", "damaged")
    tables = tuple(
        Table(name or None, None, tuple(tuple(_text(c) for c in row) for row in rows))
        for name, rows in sheets
    )
    text = "\n".join(" ".join(row) for table in tables for row in table.rows)
    aging = next((a for _, rows in sheets if (a := aging_fields(rows)) is not None), None)
    return Extraction(
        "extracted",
        tables=tables,
        dates=dates_in(text),
        amounts=amounts_in(text),
        aging=aging,
    )


def extract(content: bytes, media_type: str) -> Extraction:
    """In-process extraction (tests, and inside the isolated child). A PDF with active content is
    never read (DEC-296, ED-021); HEIC waits for a reader no allowlisted dependency provides."""
    if media_type == PDF:
        if is_active(content):
            return Extraction("not_extracted", ACTIVE_CONTENT)
        return _pdf(content)
    if media_type in (XLSX, CSV):
        return _sheets(content, media_type)
    if media_type in (HEIC, HEIF):
        return Extraction("not_extracted", "format_not_supported")
    reason = "legacy_format" if media_type == "application/x-ole-storage" else None
    return Extraction("not_extracted", reason)


def extract_isolated(content: bytes, media_type: str) -> Extraction:
    """`extract` in a short-lived child process with limits (AC-2; TASK-054 D2)."""
    try:
        return run_isolated(extract, content, media_type)
    except IsolationFailed:
        return Extraction("unreadable", "crashed")
    except Exception:  # anything the child didn't turn into a result
        return Extraction("unreadable", "crashed")
