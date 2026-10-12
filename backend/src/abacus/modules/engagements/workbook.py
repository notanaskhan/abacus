"""Methodology workbooks (SPEC-008 Q1): parsed and validated into typed rows, deterministically.

A workbook is hostile input (ADR-052): opened read-only with cached values (formulas are never
evaluated, macros and external links ignored), bounded by settings, and every problem is reported
as sheet, row, column and a fixed code, never a cell's contents.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from dataclasses import dataclass
from itertools import pairwise
from typing import Final, Literal, cast

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from abacus.kernel.config import settings

Tier = Literal["A", "B", "C", "D", "E"]
TIERS: Final = frozenset({"A", "B", "C", "D", "E"})
AREAS: Final = "Areas"
REQUESTS: Final = "Requests"
RULES: Final = "Account rules"
HEADERS: Final[dict[str, tuple[str, ...]]] = {
    AREAS: ("code", "name"),
    REQUESTS: ("area_code", "description", "tier"),
    RULES: ("area_code", "account_from", "account_to"),
}
_MAX_LENGTH: Final = {
    "code": 20,
    "name": 100,
    "area_code": 20,
    "description": 2000,
    "tier": 1,
    "account_from": 50,
    "account_to": 50,
}
MAX_PROBLEMS: Final = 100


@dataclass(frozen=True)
class Problem:
    sheet: str | None
    row: int | None
    column: str | None
    code: str


class TemplateInvalid(Exception):
    """The workbook breaks the layout or the limits (AC-2): nothing is stored."""

    def __init__(self, problems: list[Problem]) -> None:
        super().__init__(f"{len(problems)} problem(s)")
        self.problems = problems[:MAX_PROBLEMS]


@dataclass(frozen=True)
class Area:
    code: str
    name: str


Effort = Literal["small", "medium", "large"]
EFFORTS: Final = frozenset({"small", "medium", "large"})
# SPEC-031 (TASK-060): optional request columns after the required ones.
OPTIONAL_REQUEST_COLUMNS: Final = ("why", "good_looks_like", "blocks_fieldwork", "effort")
_GUIDANCE_LENGTH: Final = 1000


@dataclass(frozen=True)
class TemplateItem:
    area_code: str
    description: str
    tier: Tier
    why: str | None = None
    good_looks_like: str | None = None
    blocks_fieldwork: bool = False
    effort: Effort | None = None


@dataclass(frozen=True)
class AccountRule:
    area_code: str
    account_from: str
    account_to: str


@dataclass(frozen=True)
class Methodology:
    areas: tuple[Area, ...]
    items: tuple[TemplateItem, ...]
    rules: tuple[AccountRule, ...]


def _text(value: object) -> str:
    """A cell as text: numeric account codes keep their digits (`1000.0` is `1000`)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet(workbook: Workbook, name: str) -> Iterator[tuple[object, ...]]:
    sheet = cast(ReadOnlyWorksheet, workbook[name])
    return iter(sheet.iter_rows(values_only=True))


def _rows(
    rows: Iterator[tuple[object, ...]], name: str, problems: list[Problem]
) -> Iterator[tuple[int, dict[str, str]]]:
    header = tuple(_text(v) for v in next(rows, ()))
    expected = HEADERS[name]
    if header[: len(expected)] != expected:
        problems.append(Problem(name, 1, None, "missing_header"))
        return
    for number, values in enumerate(rows, start=2):
        cells = [_text(v) for v in values[: len(expected)]]
        if not any(cells):
            continue  # blank rows are allowed anywhere
        cells += [""] * (len(expected) - len(cells))
        row = dict(zip(expected, cells, strict=True))
        bad = False
        for column, value in row.items():
            if not value:
                problems.append(Problem(name, number, column, "empty_value"))
                bad = True
            elif len(value) > _MAX_LENGTH[column]:
                problems.append(Problem(name, number, column, "too_long"))
                bad = True
        if not bad:
            yield number, row


def _optional(
    rows: Iterator[tuple[object, ...]], problems: list[Problem]
) -> dict[int, dict[str, str]]:
    """SPEC-031: the requests sheet's optional guidance columns, by row number, validated."""
    header = [_text(v).casefold() for v in next(rows, ())]
    where = {c: header.index(c) for c in OPTIONAL_REQUEST_COLUMNS if c in header}
    found: dict[int, dict[str, str]] = {}
    if not where:
        return found
    for number, values in enumerate(rows, start=2):
        cells = {c: _text(values[i]) if i < len(values) else "" for c, i in where.items()}
        for column in ("why", "good_looks_like"):
            if len(cells.get(column, "")) > _GUIDANCE_LENGTH:
                problems.append(Problem(REQUESTS, number, column, "too_long"))
        if cells.get("effort") and cells["effort"].casefold() not in EFFORTS:
            problems.append(Problem(REQUESTS, number, "effort", "invalid_effort"))
        blocks = cells.get("blocks_fieldwork", "").casefold()
        if blocks and blocks not in ("yes", "no", "y", "n", "true", "false"):
            problems.append(Problem(REQUESTS, number, "blocks_fieldwork", "invalid_yes_no"))
        found[number] = cells
    return found


def parse_workbook(data: bytes) -> Methodology:
    """The workbook's areas, request items and account rules, or `TemplateInvalid`."""
    s = settings()
    if len(data) > s.methodology_max_bytes:
        raise TemplateInvalid([Problem(None, None, None, "too_large")])
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # not a zip, not a workbook, or malformed parts: all the same to the caller
        raise TemplateInvalid([Problem(None, None, None, "not_a_workbook")]) from None
    problems: list[Problem] = []
    try:
        missing = [name for name in HEADERS if name not in workbook.sheetnames]
        if missing:
            raise TemplateInvalid([Problem(name, None, None, "missing_sheet") for name in missing])
        areas: list[Area] = []
        codes: set[str] = set()
        for number, row in _rows(_sheet(workbook, AREAS), AREAS, problems):
            if row["code"] in codes:
                problems.append(Problem(AREAS, number, "code", "duplicate_area"))
                continue
            codes.add(row["code"])
            areas.append(Area(row["code"], row["name"]))
        items: list[TemplateItem] = []
        extras = _optional(_sheet(workbook, REQUESTS), problems)
        for number, row in _rows(_sheet(workbook, REQUESTS), REQUESTS, problems):
            extra = extras.get(number, {})
            if row["area_code"] not in codes:
                problems.append(Problem(REQUESTS, number, "area_code", "unknown_area"))
            elif row["tier"].upper() not in TIERS:
                problems.append(Problem(REQUESTS, number, "tier", "invalid_tier"))
            else:
                tier = cast(Tier, row["tier"].upper())
                items.append(
                    TemplateItem(
                        row["area_code"],
                        row["description"],
                        tier,
                        extra.get("why") or None,
                        extra.get("good_looks_like") or None,
                        extra.get("blocks_fieldwork", "").casefold() in ("yes", "y", "true"),
                        cast(Effort, extra["effort"].casefold()) if extra.get("effort") else None,
                    )
                )
        rules: list[tuple[int, AccountRule]] = []
        for number, row in _rows(_sheet(workbook, RULES), RULES, problems):
            if row["area_code"] not in codes:
                problems.append(Problem(RULES, number, "area_code", "unknown_area"))
            elif row["account_from"] > row["account_to"]:
                problems.append(Problem(RULES, number, "account_to", "invalid_range"))
            else:
                rule = AccountRule(row["area_code"], row["account_from"], row["account_to"])
                rules.append((number, rule))
    finally:
        workbook.close()
    problems += _overlaps(rules)
    if not areas and not any(p.code == "missing_header" and p.sheet == AREAS for p in problems):
        problems.append(Problem(AREAS, None, None, "no_areas"))
    for sheet, count, limit in (
        (AREAS, len(areas), s.methodology_max_areas),
        (REQUESTS, len(items), s.methodology_max_items),
        (RULES, len(rules), s.methodology_max_rules),
    ):
        if count > limit:
            problems.append(Problem(sheet, None, None, "too_many_rows"))
    if problems:
        raise TemplateInvalid(problems)
    return Methodology(tuple(areas), tuple(items), tuple(rule for _, rule in rules))


def _overlaps(rules: list[tuple[int, AccountRule]]) -> list[Problem]:
    """Account ranges may not overlap (compared as text), so mapping never depends on ties."""
    ordered = sorted(rules, key=lambda r: (r[1].account_from, r[1].account_to))
    problems: list[Problem] = []
    for (_, before), (number, after) in pairwise(ordered):
        if after.account_from <= before.account_to:
            problems.append(Problem(RULES, number, "account_from", "overlapping_range"))
    return problems


def area_for(code: str, rules: tuple[AccountRule, ...]) -> str | None:
    """The area of an account code: the first rule in rule order whose range holds it (AC-7)."""
    for rule in rules:
        if rule.account_from <= code <= rule.account_to:
            return rule.area_code
    return None
