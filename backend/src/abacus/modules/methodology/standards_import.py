"""Firm-standard content as workbooks (K-02-T11; K-02 §3 F5, US-5.1; KNW-04; SPEC-500 AC-45 to
AC-48).

**The format** (export writes it, import reads it back, so a round trip is an empty diff):
- `Rows`: the payload's list field (`rows`, `items`, `tiers` or `applies_to`), one row per
  entry, the header naming the entry's fields; a list inside an entry (`required_levels`) is
  written comma-separated.
- `Settings`: `field` and `value` columns for the payload's other fields; a nested field is
  written `parent.child` (`performance_materiality.low_pct`, `applies_to.form`).
The kind comes from the standard; a workbook never names it. Values are text in the file and
parse as the kind's payload model (`standards.validate_payload`).

**Hostile input** (ADR-052; `workbook.py`'s rules): read-only, cached values only (no formula is
evaluated), bounded rows, columns and bytes; every problem is sheet, row, column and a fixed code,
never a cell's content. A CSV is the `Rows` sheet alone.

**Onboarding workbook** (`POST /v1/standards/imports`): a `Standards` sheet (`key`, `name`,
`kind`, `owner_email`) and, per standard, a sheet named by its key (its `Rows`) and optionally
`<key>.settings`. Each row creates the standard and its first draft; an owner who is not a
partner member is that row's problem (`owner_not_partner`), and the other rows still import.

The export is deterministic (ADR-042): fixed core properties and the ZIP rebuilt with sorted
entries and a fixed timestamp.
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from typing import Final, cast

from openpyxl import Workbook, load_workbook

from abacus.kernel.errors import DomainInvalid

ROWS: Final = "Rows"
SETTINGS: Final = "Settings"
STANDARDS: Final = "Standards"
MAX_BYTES: Final = 5 * 1024 * 1024
MAX_ROWS: Final = 2_000
MAX_COLUMNS: Final = 30
# Which payload field the Rows sheet carries, per kind.
LIST_FIELD: Final[dict[str, str]] = {
    "aje_thresholds": "rows",
    "variance_thresholds": "rows",
    "materiality_ranges": "rows",
    "review_checklist": "items",
    "review_levels": "tiers",
    "return_approval_policy": "applies_to",
}
# Entry fields that are lists, written comma-separated.
LIST_CELLS: Final = frozenset({"required_levels"})
_ZIP_TIME: Final = (1980, 1, 1, 0, 0, 0)
_EPOCH: Final = datetime(2026, 1, 1, tzinfo=UTC)


@dataclass(frozen=True)
class Problem:
    sheet: str
    row: int | None
    column: str | None
    code: str


class ImportInvalid(DomainInvalid):
    """The workbook has problems; nothing was drafted."""

    code = "import_invalid"

    def __init__(self, problems: Sequence[Problem]) -> None:
        super().__init__(facts={"problems": len(problems)})
        self.problems = list(problems)


def _letter(n: int) -> str:
    out = ""
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    return text or None


def _scalar(value: str) -> object:
    """A cell's text as the payload expects it: booleans and numbers by their spelling; every
    other value stays text for the payload model to parse (decimals stay exact)."""
    lowered = value.lower()
    if lowered in ("true", "false"):
        return lowered == "true"
    if value.lstrip("-").isdigit():
        return int(value)
    return value


def _set(target: dict[str, object], dotted: str, value: object) -> None:
    head, _, rest = dotted.partition(".")
    if rest:
        nested = target.setdefault(head, {})
        if isinstance(nested, dict):
            _set(cast(dict[str, object], nested), rest, value)
        return
    target[head] = value


def _table(
    name: str, raw: Sequence[Sequence[object]], problems: list[Problem]
) -> list[dict[str, object]]:
    if not raw:
        return []
    header = [_text(h) or "" for h in raw[0]][:MAX_COLUMNS]
    for i, h in enumerate(header, start=1):
        if not h:
            problems.append(Problem(name, 1, _letter(i), "header_missing"))
    if len(set(header)) != len(header):
        problems.append(Problem(name, 1, None, "header_duplicate"))
    out: list[dict[str, object]] = []
    for r, line in enumerate(raw[1:], start=2):
        if r > MAX_ROWS + 1:
            problems.append(Problem(name, r, None, "too_many_rows"))
            break
        cells = [_text(v) for v in line]
        if not any(cells):
            continue
        entry: dict[str, object] = {}
        for c, field in enumerate(header):
            value = cells[c] if c < len(cells) else None
            if value is None or not field:
                continue
            parsed: object = (
                [v.strip() for v in value.split(",") if v.strip()]
                if field in LIST_CELLS
                else _scalar(value)
            )
            _set(entry, field, parsed)
        out.append(entry)
    return out


def _sheets(data: bytes, *, csv_file: bool) -> dict[str, list[list[object]]]:
    if len(data) > MAX_BYTES:
        raise ImportInvalid([Problem("", None, None, "too_large")])
    if csv_file:
        return {ROWS: [list(r) for r in csv.reader(io.StringIO(data.decode("utf-8-sig")))]}
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception:  # any unreadable workbook is one problem, never its content
        raise ImportInvalid([Problem("", None, None, "not_a_workbook")]) from None
    try:
        return {
            ws.title: [
                list(r)
                for r in ws.iter_rows(
                    max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1, values_only=True
                )
            ]
            for ws in book.worksheets
        }
    finally:
        book.close()


def payload_from_sheets(
    kind: str,
    rows_sheet: Sequence[Sequence[object]],
    settings_sheet: Sequence[Sequence[object]],
    *,
    rows_name: str = ROWS,
    settings_name: str = SETTINGS,
) -> tuple[dict[str, object], list[Problem]]:
    problems: list[Problem] = []
    payload: dict[str, object] = {"kind": kind}
    for i, entry in enumerate(_table(settings_name, settings_sheet, problems), start=2):
        field, value = entry.get("field"), entry.get("value")
        if not isinstance(field, str):
            problems.append(Problem(settings_name, i, "A", "field_missing"))
            continue
        _set(payload, field, value)
    list_field = LIST_FIELD.get(kind)
    if list_field is not None:
        payload[list_field] = _table(rows_name, rows_sheet, problems)
    return payload, problems


def read_payload(kind: str, data: bytes, *, csv_file: bool) -> dict[str, object]:
    """One standard's payload from its workbook (or CSV); `ImportInvalid` with every problem."""
    sheets = _sheets(data, csv_file=csv_file)
    payload, problems = payload_from_sheets(kind, sheets.get(ROWS, []), sheets.get(SETTINGS, []))
    if problems:
        raise ImportInvalid(problems)
    return payload


def _cell(value: object) -> object:
    if isinstance(value, list):
        return ", ".join(str(v) for v in cast(list[object], value))
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, Decimal)):
        return str(value)
    return value


def _flatten(prefix: str, value: object, out: list[tuple[str, object]]) -> None:
    if isinstance(value, dict):
        for key in sorted(cast(dict[str, object], value)):
            _flatten(
                f"{prefix}.{key}" if prefix else key, cast(dict[str, object], value)[key], out
            )
    else:
        out.append((prefix, value))


def _rebuild(raw: bytes) -> bytes:
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            entry = zipfile.ZipInfo(name, date_time=_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            entry.create_system = 3
            target.writestr(entry, source.read(name))
    return out.getvalue()


def export_workbook(kind: str, title: str, payload: Mapping[str, object]) -> bytes:
    """The payload in the import format, byte-stable for the same payload (ADR-042)."""
    book = Workbook()
    rows_ws = book.active
    assert rows_ws is not None
    rows_ws.title = ROWS
    list_field = LIST_FIELD.get(kind)
    entries = cast(list[dict[str, object]], payload.get(list_field, []) if list_field else [])
    columns: list[str] = []
    flat_entries: list[dict[str, object]] = []
    for entry in entries:
        flat: list[tuple[str, object]] = []
        _flatten("", entry, flat)
        flat_entries.append(dict(flat))
        for key, _ in flat:
            if key not in columns:
                columns.append(key)
    rows_ws.append(columns)
    for flat_entry in flat_entries:
        rows_ws.append([_cell(flat_entry.get(c)) for c in columns])
    settings_ws = book.create_sheet(SETTINGS)
    settings_ws.append(["field", "value"])
    rest: list[tuple[str, object]] = []
    for key in sorted(payload):
        if key in ("kind", list_field):
            continue
        _flatten(key, payload[key], rest)
    for key, value in rest:
        settings_ws.append([key, _cell(value)])
    props = book.properties
    props.creator = "abacus"
    props.lastModifiedBy = "abacus"
    props.title = title
    props.created = _EPOCH.replace(tzinfo=None)
    props.modified = _EPOCH.replace(tzinfo=None)
    raw = io.BytesIO()
    book.save(raw)
    return _rebuild(raw.getvalue())


@dataclass(frozen=True)
class OnboardingRow:
    row: int
    key: str
    name: str
    kind: str
    owner_email: str
    payload: dict[str, object]


def read_onboarding(data: bytes) -> tuple[list[OnboardingRow], list[Problem]]:
    """The onboarding workbook: every standard's row and payload, and every problem."""
    sheets = _sheets(data, csv_file=False)
    problems: list[Problem] = []
    if STANDARDS not in sheets:
        return [], [Problem(STANDARDS, None, None, "sheet_missing")]
    out: list[OnboardingRow] = []
    for i, entry in enumerate(_table(STANDARDS, sheets[STANDARDS], problems), start=2):
        key, name, kind, owner = (entry.get(f) for f in ("key", "name", "kind", "owner_email"))
        if not all(isinstance(v, str) for v in (key, name, kind, owner)):
            problems.append(Problem(STANDARDS, i, None, "row_incomplete"))
            continue
        k = cast(str, key)
        payload, found = payload_from_sheets(
            cast(str, kind),
            sheets.get(k, []),
            sheets.get(f"{k}.settings", []),
            rows_name=k,
            settings_name=f"{k}.settings",
        )
        problems.extend(found)
        out.append(
            OnboardingRow(i, k, cast(str, name), cast(str, kind), cast(str, owner), payload)
        )
    return out, problems
