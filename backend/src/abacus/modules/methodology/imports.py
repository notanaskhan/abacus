"""Methodology import review (K-01-T18; K-01 §3.6, §3.7 AC7.9, §4.2 migration C, §4.3; KNW-02 c,
d; DEC-1755).

- **An import** (`methodology.manage`, a person): created once per Idempotency-Key for a named
  template; it waits for its files (`awaiting_files`).
- **Files** arrive one role at a time as the raw request body (no multipart, TASK-023 D3):
  `document_docx`, `document_pdf`, `workbook`. Each is scanned and typed before anything keeps
  it, then written once to the evidence store under S3 Object Lock (through the platform's file
  store, which evidence registers: DEC-1759); a refused file stores nothing. A file's extension
  names its role (`.docx`, `.pdf`, `.xlsx`; the screen checks the bytes agree, DEC-296). One
  file per role; the storage key and its object version are kept and never returned.
- **Start**: a workbook-only import parses deterministically into the template's next version
  (K-01-T05's import) and lands Needs review; an import with a document needs the Methodology
  importer (`methodology.importer`, K-01-T21) and answers 409 `importer_not_enabled` until it is
  on and its workflow exists. **Cancel** stops an import that has not finished.
- **Diff** (`version_diff`): code against another version of the same template (default: the
  version drafted from, else the previous number), by lineage (`diff.lineage_diff`).
- **Standards mapping** (`methodology.manage`): the standards set (a platform row; an unknown id
  is 422 `standards_set_invalid`), the QM reference, the engagement types and whether it is the
  default for new engagements, while the version is a draft or in review.
- **Flag resolution** (`methodology.manage`, a person): confirm or edit a step's parse flag; an
  edit changes the text and keeps the step's source (its citation stays); the event names them.
- **Activity**: the version's events, oldest first, with each actor's name.
- **Download**: the original file of an import whose version this is; audited.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from typing import Final, Literal, cast
from uuid import UUID

from sqlalchemy import func, insert, select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from abacus.kernel._flags import METHODOLOGY_IMPORTER
from abacus.kernel.config import settings
from abacus.kernel.db import tenant_session
from abacus.kernel.errors import DomainConflict, DomainInvalid, NotFound
from abacus.kernel.flags import flag_enabled
from abacus.kernel.uow import Ref, Target, uow
from abacus.modules.identity.api import AuthContext, Forbidden, Resource, authorise, names_of
from abacus.modules.methodology import lifecycle
from abacus.modules.methodology.diff import StepRow, VersionDiff, lineage_diff
from abacus.modules.methodology.document_import import document_importer
from abacus.modules.methodology.models import (
    MethodologyImport,
    MethodologySourceFile,
    MethodologyTemplate,
    MethodologyVersion,
    MethodologyVersionEvent,
    MethodologyVersionState,
    Program,
    ProgramStep,
    ProgramStepSource,
)
from abacus.modules.methodology.repository import lock_or_insert_template
from abacus.modules.methodology.service import EngagementType, import_template
from abacus.modules.platform.api import read_stored, screen_and_store

MANAGE: Final = "methodology.manage"
READ: Final = "methodology.read"
SUBJECT: Final = "methodology_import"
FileRole = Literal["document_docx", "document_pdf", "workbook"]
DOCUMENTS: Final = frozenset({"document_docx", "document_pdf"})
EXTENSION_OF: Final[dict[str, str]] = {
    "document_docx": ".docx",
    "document_pdf": ".pdf",
    "workbook": ".xlsx",
}
OPEN: Final = frozenset({"awaiting_files", "scanning", "parsing"})
DIFFED: Final = (
    "text",
    "assertions",
    "standard_refs",
    "level",
    "guidance",
    "mandatory",
    "lock_note",
)


class ImportFrozen(DomainConflict):
    """The import is not waiting for files (or has finished): this step does not apply."""

    code = "import_frozen"


class FileExists(DomainConflict):
    """A file of this role is already on the import; files are written once."""

    code = "file_exists"


class FilesMissing(DomainConflict):
    """The import has no workbook and no document to read yet."""

    code = "files_missing"


class ImporterNotEnabled(DomainConflict):
    """Reading a Word or PDF methodology needs the Methodology importer, which is not on for this
    firm (K-01 AC7.7): import the workbook alone."""

    code = "importer_not_enabled"


class FileTooLarge(DomainInvalid):
    """The file is over the methodology size limit."""

    code = "file_too_large"


class StandardsSetInvalid(DomainInvalid):
    """The standards set is not a published platform set."""

    code = "standards_set_invalid"


class AgainstOtherTemplate(DomainInvalid):
    """A version is diffed only against another version of its own template."""

    code = "against_other_template"


class TextRequired(DomainInvalid):
    """An edit that resolves a flag carries the step's new text."""

    code = "text_required"


class FileRoleMismatch(DomainInvalid):
    """The file's extension is not its role's (`.docx`, `.pdf`, `.xlsx`)."""

    code = "file_role_mismatch"


class FlagNotOpen(DomainConflict):
    """The step has no open parse flag."""

    code = "flag_not_open"


async def _read(tenant_id: UUID, key: str, version: str, sha256: str, size: int) -> bytes:
    """A stored source file's verified bytes, through the platform's file store (DEC-1759)."""
    return await read_stored(tenant_id, key, version, sha256, size)


def _person(ctx: AuthContext, action: str) -> None:
    if ctx.tenant.actor_kind != "human" or ctx.client or ctx.is_support:
        raise Forbidden(action, "role")


@dataclass(frozen=True)
class SourceFileView:
    id: UUID
    role: str
    filename: str
    media_type: str
    bytes: int
    sha256: str
    scan_status: str
    created_at: datetime


@dataclass(frozen=True)
class ImportView:
    id: UUID
    template_id: UUID
    version_id: UUID | None
    status: str
    pages_done: int
    pages_total: int | None
    failure_code: str | None
    started_at: datetime
    files: tuple[SourceFileView, ...]


@dataclass(frozen=True)
class ActivityEntry:
    seq: int
    kind: str
    actor_kind: str
    actor_user_id: UUID | None
    actor_name: str | None
    at: datetime
    facts: dict[str, object]


async def _import(
    session: AsyncSession, import_id: UUID, *, lock: bool = False
) -> MethodologyImport:
    stmt = select(MethodologyImport).where(MethodologyImport.id == import_id)
    if lock:
        stmt = stmt.with_for_update()
    found = (await session.execute(stmt)).scalar_one_or_none()
    if found is None:
        raise NotFound(SUBJECT)
    return found


async def _view(session: AsyncSession, import_id: UUID) -> ImportView:
    row = await _import(session, import_id)
    files = (
        await session.execute(
            select(MethodologySourceFile)
            .where(MethodologySourceFile.import_id == import_id)
            .order_by(MethodologySourceFile.created_at, MethodologySourceFile.role)
        )
    ).scalars()
    return ImportView(
        id=row.id,
        template_id=row.template_id,
        version_id=row.version_id,
        status=row.status,
        pages_done=row.pages_done,
        pages_total=row.pages_total,
        failure_code=row.failure_code,
        started_at=row.started_at,
        files=tuple(
            SourceFileView(
                f.id,
                f.role,
                f.filename,
                f.media_type,
                f.bytes,
                f.sha256,
                f.scan_status,
                f.created_at,
            )
            for f in files
        ),
    )


async def import_view(ctx: AuthContext, import_id: UUID) -> ImportView:
    await authorise(ctx, MANAGE, Resource.firm(ctx.tenant_id))
    async with tenant_session(ctx.tenant) as session:
        return await _view(session, import_id)


async def version_import(ctx: AuthContext, version_id: UUID) -> ImportView:
    """K-01-T23 (KNW-02 a, a2, b): the newest import that made this version, for a reader of the
    version (the RunCard polls it; the viewer reads its files)."""
    await authorise(ctx, READ, Resource.firm(ctx.tenant_id))
    async with tenant_session(ctx.tenant) as session:
        found = await session.scalar(
            select(MethodologyImport.id)
            .where(MethodologyImport.version_id == version_id)
            .order_by(MethodologyImport.started_at.desc(), MethodologyImport.id)
            .limit(1)
        )
        if found is None:
            raise NotFound(SUBJECT)
        return await _view(session, found)


@dataclass(frozen=True)
class StepSourceView:
    step_id: UUID
    file_id: UUID
    page_from: int
    page_to: int
    bbox: list[float] | None
    quote: str | None
    verified: bool


async def step_sources(
    ctx: AuthContext, step_ids: Sequence[UUID]
) -> dict[UUID, list[StepSourceView]]:
    """K-01-T23: where each parsed step came from (0580's `program_step_sources`), for a reader
    of the version: KNW-02's viewer lights the step's region and the region lights the row."""
    await authorise(ctx, READ, Resource.firm(ctx.tenant_id))
    out: dict[UUID, list[StepSourceView]] = {}
    if not step_ids:
        return out
    async with tenant_session(ctx.tenant) as session:
        rows = (
            await session.execute(
                select(ProgramStepSource)
                .where(ProgramStepSource.step_id.in_(list(step_ids)))
                .order_by(ProgramStepSource.step_id, ProgramStepSource.page_from)
            )
        ).scalars()
        for r in rows:
            out.setdefault(r.step_id, []).append(
                StepSourceView(
                    r.step_id,
                    r.source_file_id,
                    r.page_from,
                    r.page_to,
                    [float(x) for x in cast(list[float], r.bbox)] if r.bbox else None,
                    r.quote,
                    r.verified,
                )
            )
    return out


async def create_import(
    ctx: AuthContext,
    *,
    template_name: str,
    engagement_types: list[str],
    idempotency_key: str,
) -> ImportView:
    """A new import for the named template (created if new), once per Idempotency-Key: the same
    key answers the same import."""
    _person(ctx, MANAGE)
    await authorise(ctx, MANAGE, Resource.firm(ctx.tenant_id))
    name = template_name.strip()
    # A replay reads outside the unit of work: it changes nothing, so it records nothing.
    async with tenant_session(ctx.tenant) as session:
        existing = (
            await session.execute(
                select(MethodologyImport.id).where(
                    MethodologyImport.idempotency_key == idempotency_key
                )
            )
        ).scalar_one_or_none()
        if existing is not None:
            return await _view(session, existing)
    async with uow(ctx.tenant) as tx:
        template_id, _ = await lock_or_insert_template(
            tx.session,
            tenant_id=ctx.tenant_id,
            name=name,
            created_by=ctx.user_id,
            engagement_types=sorted(set(engagement_types)) or ["fs_audit"],
        )
        import_id = (
            await tx.session.execute(
                pg_insert(MethodologyImport)
                .values(
                    tenant_id=ctx.tenant_id,
                    template_id=template_id,
                    started_by=ctx.user_id,
                    idempotency_key=idempotency_key,
                )
                .returning(MethodologyImport.id)
            )
        ).scalar_one()
        tx.record(
            "methodology_import.created",
            target=Target(SUBJECT, import_id),
            after=Ref(template_id=template_id),
        )
        return await _view(tx.session, import_id)


def _sheets(data: bytes) -> int | None:
    """The workbook's sheet count, read without evaluating anything; None when it does not open."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # an unreadable workbook is the parse's problem, at start
        return None
    try:
        return len(book.sheetnames)
    finally:
        book.close()


async def upload_file(
    ctx: AuthContext, import_id: UUID, role: FileRole, filename: str, data: bytes
) -> ImportView:
    """Scan, type and store one file for a role, once. Nothing is kept when the scan or the type
    check refuses it (`file_refused`, `type_refused`, …, the evidence upload's codes)."""
    _person(ctx, MANAGE)
    await authorise(ctx, MANAGE, Resource.firm(ctx.tenant_id))
    if len(data) > settings().methodology_max_bytes:
        raise FileTooLarge(facts={"limit": settings().methodology_max_bytes})
    if not filename.lower().endswith(EXTENSION_OF[role]):
        raise FileRoleMismatch(facts={"role": role})
    async with tenant_session(ctx.tenant) as session:
        row = await _import(session, import_id)
        if row.status != "awaiting_files":
            raise ImportFrozen(facts={"status": row.status})
    screened = await screen_and_store(ctx.tenant_id, data, filename)
    sheets = _sheets(data) if role == "workbook" else None
    async with uow(ctx.tenant) as tx:
        row = await _import(tx.session, import_id, lock=True)
        if row.status != "awaiting_files":
            raise ImportFrozen(facts={"status": row.status})
        taken = (
            await tx.session.execute(
                select(MethodologySourceFile.id).where(
                    MethodologySourceFile.import_id == import_id,
                    MethodologySourceFile.role == role,
                )
            )
        ).scalar_one_or_none()
        if taken is not None:
            raise FileExists(facts={"role": role})
        file_id = (
            await tx.session.execute(
                insert(MethodologySourceFile)
                .values(
                    tenant_id=ctx.tenant_id,
                    import_id=import_id,
                    role=role,
                    filename=screened.file_name,
                    media_type=screened.media_type,
                    storage_key=screened.key,
                    storage_version_id=screened.version_id,
                    sha256=screened.sha256,
                    bytes=len(data),
                    sheet_count=sheets,
                    scan_status="clean",
                    created_by=ctx.user_id,
                )
                .returning(MethodologySourceFile.id)
            )
        ).scalar_one()
        tx.record(
            "methodology_source_file.stored",
            target=Target("methodology_source_file", file_id),
            after=Ref(import_id=import_id, sha256=screened.sha256),
        )
        return await _view(tx.session, import_id)


async def _finish(ctx: AuthContext, import_id: UUID, *, failed: bool) -> None:
    """The import moves to parsing, or to failed with the workbook's problem code."""
    async with uow(ctx.tenant) as tx:
        if failed:
            await tx.session.execute(
                update(MethodologyImport)
                .where(MethodologyImport.id == import_id)
                .values(status="failed", failure_code="workbook_invalid")
            )
            tx.record("methodology_import.failed", target=Target(SUBJECT, import_id))
        else:
            await tx.session.execute(
                update(MethodologyImport)
                .where(MethodologyImport.id == import_id)
                .values(status="parsing")
            )
            tx.record("methodology_import.parsing", target=Target(SUBJECT, import_id))


async def start_import(ctx: AuthContext, import_id: UUID) -> ImportView:
    """Read the files: a workbook alone parses now into the template's next version, in review;
    any document needs the Methodology importer (409 `importer_not_enabled`)."""
    _person(ctx, MANAGE)
    await authorise(ctx, MANAGE, Resource.firm(ctx.tenant_id))
    async with tenant_session(ctx.tenant) as session:
        row = await _import(session, import_id)
        if row.status != "awaiting_files":
            raise ImportFrozen(facts={"status": row.status})
        files = list(
            (
                await session.execute(
                    select(MethodologySourceFile).where(
                        MethodologySourceFile.import_id == import_id
                    )
                )
            ).scalars()
        )
        template = (
            await session.execute(
                select(MethodologyTemplate).where(MethodologyTemplate.id == row.template_id)
            )
        ).scalar_one()
    if not files:
        raise FilesMissing(facts={"import_id": str(import_id)})
    if any(f.role in DOCUMENTS for f in files):
        # K-01-T21: the Methodology importer reads a PDF; a Word file is not parsed yet. Off, not
        # registered, or a firm it may not read (not synthetic, agent not cleared): 409, and the
        # UI offers the workbook path (AC7.7).
        importer = document_importer()
        if (
            importer is None
            or not any(f.role == "document_pdf" for f in files)
            or not await flag_enabled(ctx.tenant, METHODOLOGY_IMPORTER)
            or not await importer.admitted(ctx)
        ):
            raise ImporterNotEnabled(facts={"import_id": str(import_id)})
        await _finish(ctx, import_id, failed=False)
        await importer.start(ctx, import_id)
        async with tenant_session(ctx.tenant) as session:
            return await _view(session, import_id)
    workbook = next(f for f in files if f.role == "workbook")
    data = await _read(
        ctx.tenant_id,
        workbook.storage_key,
        workbook.storage_version_id,
        workbook.sha256,
        workbook.bytes,
    )
    await _finish(ctx, import_id, failed=False)
    try:
        summary = await import_template(
            ctx,
            template.name,
            data,
            tuple(cast("list[EngagementType]", template.engagement_types)),
            None,
        )
    except Exception:
        await _finish(ctx, import_id, failed=True)
        raise
    async with uow(ctx.tenant) as tx:
        await tx.session.execute(
            update(MethodologyImport)
            .where(MethodologyImport.id == import_id)
            .values(status="needs_review", version_id=summary.version_id)
        )
        await tx.session.execute(
            insert(MethodologyVersionEvent).values(
                tenant_id=ctx.tenant_id,
                version_id=summary.version_id,
                kind="import_started",
                actor_user_id=ctx.user_id,
                facts={"import_id": str(import_id), "files": len(files)},
            )
        )
        tx.record(
            "methodology_import.needs_review",
            target=Target(SUBJECT, import_id),
            after=Ref(version_id=summary.version_id),
        )
        return await _view(tx.session, import_id)


async def cancel_import(ctx: AuthContext, import_id: UUID) -> ImportView:
    _person(ctx, MANAGE)
    await authorise(ctx, MANAGE, Resource.firm(ctx.tenant_id))
    async with uow(ctx.tenant) as tx:
        row = await _import(tx.session, import_id, lock=True)
        if row.status not in OPEN:
            raise ImportFrozen(facts={"status": row.status})
        await tx.session.execute(
            update(MethodologyImport)
            .where(MethodologyImport.id == import_id)
            .values(status="cancelled")
        )
        tx.record("methodology_import.cancelled", target=Target(SUBJECT, import_id))
        return await _view(tx.session, import_id)


@dataclass(frozen=True)
class SourceDownload:
    filename: str
    media_type: str
    content: bytes


async def download_file(ctx: AuthContext, version_id: UUID, file_id: UUID) -> SourceDownload:
    """The original file of an import whose version this is (audited)."""
    await authorise(ctx, READ, Resource.firm(ctx.tenant_id))
    async with tenant_session(ctx.tenant) as session:
        found = (
            await session.execute(
                select(MethodologySourceFile)
                .join(MethodologyImport, MethodologyImport.id == MethodologySourceFile.import_id)
                .where(
                    MethodologySourceFile.id == file_id,
                    func.coalesce(MethodologySourceFile.version_id, MethodologyImport.version_id)
                    == version_id,
                )
            )
        ).scalar_one_or_none()
    if found is None:
        raise NotFound("methodology_source_file")
    content = await _read(
        ctx.tenant_id, found.storage_key, found.storage_version_id, found.sha256, found.bytes
    )
    async with uow(ctx.tenant) as tx:
        tx.record(
            "methodology_source_file.downloaded",
            target=Target("methodology_source_file", file_id),
            after=Ref(version_id=version_id),
        )
    return SourceDownload(found.filename, found.media_type, content)


# --- the diff -----------------------------------------------------------------------------------


async def _rows(session: AsyncSession, version_id: UUID) -> list[StepRow]:
    found = await session.execute(
        select(Program.code, ProgramStep)
        .join(Program, Program.id == ProgramStep.program_id)
        .where(ProgramStep.version_id == version_id)
    )
    return [
        StepRow(step.step_key, code, step.number, {f: getattr(step, f) for f in DIFFED})
        for code, step in found.tuples()
    ]


@dataclass(frozen=True)
class VersionDiffView:
    version_id: UUID
    against_version_id: UUID | None
    against_label: str | None
    diff: VersionDiff


async def version_diff(
    ctx: AuthContext, version_id: UUID, against: UUID | None = None
) -> VersionDiffView:
    """The version against another of its template: `against`, else the version it was drafted
    from, else the version numbered before it. 422 when `against` is another template's."""
    await authorise(ctx, READ, Resource.firm(ctx.tenant_id))
    async with tenant_session(ctx.tenant) as session:
        version = (
            await session.execute(
                select(MethodologyVersion).where(MethodologyVersion.id == version_id)
            )
        ).scalar_one_or_none()
        if version is None:
            raise NotFound(lifecycle.SUBJECT)
        if against is None:
            against = version.drafted_from_version_id
        if against is None:
            against = (
                await session.execute(
                    select(MethodologyVersion.id).where(
                        MethodologyVersion.template_id == version.template_id,
                        MethodologyVersion.version == version.version - 1,
                    )
                )
            ).scalar_one_or_none()
        other: MethodologyVersion | None = None
        if against is not None:
            other = (
                await session.execute(
                    select(MethodologyVersion).where(MethodologyVersion.id == against)
                )
            ).scalar_one_or_none()
            if other is None:
                raise NotFound(lifecycle.SUBJECT)
            if other.template_id != version.template_id:
                raise AgainstOtherTemplate(facts={"against": str(against)})
        before = await _rows(session, other.id) if other is not None else []
        after = await _rows(session, version_id)
    return VersionDiffView(
        version_id,
        other.id if other else None,
        other.label if other else None,
        lineage_diff(before, after),
    )


# --- standards mapping, flags, activity ---------------------------------------------------------


async def map_standards(
    ctx: AuthContext,
    version_id: UUID,
    *,
    standards_set_id: UUID | None,
    quality_management_ref: str | None,
    engagement_types: list[str],
    default_for_new_engagements: bool,
    row_version: int,
) -> lifecycle.VersionState:
    _person(ctx, MANAGE)
    await authorise(ctx, MANAGE, Resource.firm(ctx.tenant_id))
    try:
        async with uow(ctx.tenant) as tx:
            state = (
                await tx.session.execute(
                    select(MethodologyVersionState)
                    .where(MethodologyVersionState.version_id == version_id)
                    .with_for_update()
                )
            ).scalar_one_or_none()
            if state is None:
                raise NotFound(lifecycle.SUBJECT)
            if state.row_version != row_version:
                raise lifecycle.VersionStale(facts={"version_id": str(version_id)})
            if state.status not in ("draft", "needs_review"):
                raise lifecycle.VersionFrozen(facts={"status": state.status})
            await tx.session.execute(
                update(MethodologyVersionState)
                .where(MethodologyVersionState.version_id == version_id)
                .values(
                    standards_set_id=standards_set_id,
                    quality_management_ref=quality_management_ref,
                    engagement_types=sorted(set(engagement_types)),
                    default_for_new_engagements=default_for_new_engagements,
                    row_version=state.row_version + 1,
                )
            )
            await tx.session.execute(
                insert(MethodologyVersionEvent).values(
                    tenant_id=ctx.tenant_id,
                    version_id=version_id,
                    kind="standards_mapped",
                    actor_user_id=ctx.user_id,
                    facts={
                        "standards_set_id": str(standards_set_id) if standards_set_id else None,
                        "engagement_types": len(set(engagement_types)),
                        "default": default_for_new_engagements,
                    },
                )
            )
            tx.record(
                "methodology_version.standards_mapped",
                target=Target(lifecycle.SUBJECT, version_id),
            )
            return await lifecycle.view_in(tx.session, version_id)
    except IntegrityError as error:
        # The only foreign key this write can miss is the standards set's (a platform row).
        if getattr(getattr(error, "orig", None), "sqlstate", None) == "23503":
            raise StandardsSetInvalid(facts={"standards_set_id": str(standards_set_id)}) from None
        raise


@dataclass(frozen=True)
class FlagResolved:
    step_id: UUID
    step_key: str
    parse_flag: str
    resolution: str
    text: str
    row_version: int


async def resolve_flag(
    ctx: AuthContext,
    version_id: UUID,
    step_id: UUID,
    *,
    action: Literal["confirm", "edit"],
    text: str | None,
    row_version: int,
) -> FlagResolved:
    """A person confirms the parse or edits the text; the step keeps its key, origin and source
    (its citation), and `edited_by` names an editor."""
    _person(ctx, MANAGE)
    await authorise(ctx, MANAGE, Resource.firm(ctx.tenant_id))
    async with uow(ctx.tenant) as tx:
        step = (
            await tx.session.execute(
                select(ProgramStep)
                .where(ProgramStep.id == step_id, ProgramStep.version_id == version_id)
                .with_for_update()
            )
        ).scalar_one_or_none()
        if step is None:
            raise NotFound("program_step")
        if step.row_version != row_version:
            raise lifecycle.VersionStale(facts={"step_id": str(step_id)})
        if step.parse_flag is None or step.flag_resolution is not None:
            raise FlagNotOpen(facts={"step_id": str(step_id)})
        state = (
            await tx.session.execute(
                select(MethodologyVersionState.status).where(
                    MethodologyVersionState.version_id == version_id
                )
            )
        ).scalar_one()
        if state not in ("draft", "needs_review", "parsing"):
            raise lifecycle.VersionFrozen(facts={"status": state})
        values: dict[str, object] = {
            "flag_resolution": "confirmed" if action == "confirm" else "edited",
            "flag_resolved_by": ctx.user_id,
            "flag_resolved_at": func.abacus_now(),
            "row_version": step.row_version + 1,
        }
        if action == "edit":
            if text is None or not text.strip():
                raise TextRequired(facts={"step_id": str(step_id)})
            values |= {
                "text": text.strip(),
                "edited_by": ctx.user_id,
                "edited_at": func.abacus_now(),
            }
        await tx.session.execute(
            update(ProgramStep).where(ProgramStep.id == step_id).values(**values)
        )
        await tx.session.execute(
            insert(MethodologyVersionEvent).values(
                tenant_id=ctx.tenant_id,
                version_id=version_id,
                kind="flag_resolved",
                actor_user_id=ctx.user_id,
                facts={"step_key": step.step_key, "flag": step.parse_flag, "action": action},
            )
        )
        tx.record(
            "methodology_step.flag_resolved",
            target=Target("methodology_step", step_id),
            after=Ref(version_id=version_id),
        )
        fresh = (
            await tx.session.execute(select(ProgramStep).where(ProgramStep.id == step_id))
        ).scalar_one()
        return FlagResolved(
            step_id,
            fresh.step_key,
            step.parse_flag,
            str(fresh.flag_resolution),
            fresh.text,
            fresh.row_version,
        )


async def version_activity(ctx: AuthContext, version_id: UUID) -> list[ActivityEntry]:
    await authorise(ctx, READ, Resource.firm(ctx.tenant_id))
    async with tenant_session(ctx.tenant) as session:
        found = (
            await session.execute(
                select(MethodologyVersionEvent)
                .where(MethodologyVersionEvent.version_id == version_id)
                .order_by(MethodologyVersionEvent.seq)
            )
        ).scalars()
        events = list(found)
    people = {e.actor_user_id for e in events if e.actor_user_id is not None}
    names = await names_of(list(people)) if people else {}
    return [
        ActivityEntry(
            e.seq,
            e.kind,
            e.actor_kind,
            e.actor_user_id,
            names.get(e.actor_user_id) if e.actor_user_id else None,
            e.at,
            dict(e.facts),
        )
        for e in events
    ]


__all__ = [
    "ActivityEntry",
    "AgainstOtherTemplate",
    "FileExists",
    "FileRole",
    "FileTooLarge",
    "FilesMissing",
    "FlagNotOpen",
    "FlagResolved",
    "ImportFrozen",
    "ImportView",
    "ImporterNotEnabled",
    "SourceDownload",
    "SourceFileView",
    "StandardsSetInvalid",
    "TextRequired",
    "VersionDiffView",
    "cancel_import",
    "create_import",
    "download_file",
    "import_view",
    "map_standards",
    "resolve_flag",
    "start_import",
    "upload_file",
    "version_activity",
    "version_diff",
]
