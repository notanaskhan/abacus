"""Methodology workbooks (SPEC-008 Q1; workbook v2, K-01-T05): parsed and validated into typed
rows, deterministically.

A workbook is hostile input (ADR-052): opened read-only with cached values (formulas are never
evaluated, macros and external links ignored), bounded by settings and row limits, and every
problem is reported as sheet, row, column and a fixed code, never a cell's contents.

Workbook v2 (K-01 §3.2, AC2.1 to AC2.4): `Areas` and `Requests` are required; `Account rules`
is optional (a legacy three-sheet workbook imports as before); `Programs` and `Steps` are optional
and come together. A step's `mandatory` and `lock_note` come only from its columns, its `level` is
Foundation, Planning, Testing or Completion, an optional `step_id` is its stable lineage key, and
`requests` names the Requests rows (1 = the first request) the step links to.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterator
from dataclasses import dataclass
from itertools import pairwise
from typing import Final, Literal, cast
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from abacus.kernel.config import settings

Tier = Literal["A", "B", "C", "D", "E"]
TIERS: Final = frozenset({"A", "B", "C", "D", "E"})
AREAS: Final = "Areas"
REQUESTS: Final = "Requests"
RULES: Final = "Account rules"
PROGRAMS: Final = "Programs"
STEPS: Final = "Steps"
REQUIRED_SHEETS: Final = (AREAS, REQUESTS)
HEADERS: Final[dict[str, tuple[str, ...]]] = {
    AREAS: ("code", "name"),
    REQUESTS: ("area_code", "description", "tier"),
    RULES: ("area_code", "account_from", "account_to"),
}
_MAX_LENGTH: Final = {
    "code": 20,
    "name": 100,
    "area_code": 20,
    "description": 2000,
    "tier": 1,
    "account_from": 50,
    "account_to": 50,
}
MAX_PROBLEMS: Final = 100


@dataclass(frozen=True)
class Problem:
    sheet: str | None
    row: int | None
    column: str | None
    code: str


class TemplateInvalid(Exception):
    """The workbook breaks the layout or the limits (AC-2): nothing is stored."""

    def __init__(self, problems: list[Problem]) -> None:
        super().__init__(f"{len(problems)} problem(s)")
        self.problems = problems[:MAX_PROBLEMS]


@dataclass(frozen=True)
class Area:
    code: str
    name: str


Effort = Literal["small", "medium", "large"]
EFFORTS: Final = frozenset({"small", "medium", "large"})
# SPEC-031 (TASK-060): optional request columns after the required ones.
OPTIONAL_REQUEST_COLUMNS: Final = ("why", "good_looks_like", "blocks_fieldwork", "effort")
_GUIDANCE_LENGTH: Final = 1000


@dataclass(frozen=True)
class TemplateItem:
    area_code: str
    description: str
    tier: Tier
    why: str | None = None
    good_looks_like: str | None = None
    blocks_fieldwork: bool = False
    effort: Effort | None = None
    # SPEC-035 Q3 (TASK-074): `drafted` when the text is approved drafter text, with who approved.
    guidance_source: str | None = None
    guidance_approved_by: UUID | None = None


@dataclass(frozen=True)
class AccountRule:
    area_code: str
    account_from: str
    account_to: str


Level = Literal["foundation", "planning", "testing", "completion"]
LEVELS: Final = frozenset({"foundation", "planning", "testing", "completion"})
PROGRAM_COLUMNS: Final = ("code", "name")
STEP_COLUMNS: Final = ("program_code", "number", "text", "level")
STEP_OPTIONAL: Final = (
    "step_id",
    "assertions",
    "standard_refs",
    "mandatory",
    "lock_note",
    "templates",  # A-04-T07: `request:<key>` or `workpaper:<kind>`, separated by ";"
)
STEP_OPTIONAL_LINKS: Final = "requests"
MAX_PROGRAMS: Final = 100
MAX_STEPS: Final = 3000
MAX_LIST_ITEMS: Final = 20
_STEP_LENGTH: Final = {
    "program_code": 40,
    "text": 4000,
    "level": 20,
    "step_id": 80,
    "lock_note": 200,
    "code": 40,
    "name": 200,
}
_PROGRAM_CODE: Final = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,39}")
STEP_KEY: Final = re.compile(r"[a-z0-9][a-z0-9_.-]{0,79}")
_LIST_ITEM: Final = re.compile(r"[A-Za-z0-9][A-Za-z0-9 _.,()§-]{0,79}")
TEMPLATE_REF: Final = re.compile(r"(request|workpaper):[a-z][a-z0-9_]{0,63}")


def template_refs(cell: str) -> tuple[str, ...] | None:
    """A `templates` cell's references, or None when one is malformed (A-04-T07)."""
    found = tuple(p.strip() for p in cell.split(";") if p.strip())
    if len(found) > MAX_LIST_ITEMS or not all(TEMPLATE_REF.fullmatch(f) for f in found):
        return None
    return tuple(dict.fromkeys(found))


# Text that a spreadsheet would read as a formula, or that hides or reorders what is shown (control
# and bidi characters): refused at import (ADR-052; security review L3, DEC-577).
_FORMULA_LEADS: Final = ("=", "+", "-", "@", "\t", "\r")
_HIDDEN: Final = re.compile(
    "[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\u200e\u200f\u202a-\u202e\u2066-\u2069]"
)
_YES: Final = ("yes", "y", "true")
_YES_NO: Final = (*_YES, "no", "n", "false")


@dataclass(frozen=True)
class ProgramRow:
    code: str
    name: str


@dataclass(frozen=True)
class StepRow:
    program_code: str
    number: int
    text: str
    level: Level
    mandatory: bool = False
    lock_note: str | None = None
    step_id: str | None = None  # the stable lineage key, when the workbook names one
    assertions: tuple[str, ...] = ()
    standard_refs: tuple[str, ...] = ()
    request_positions: tuple[int, ...] = ()  # the linked Requests rows, 0-based
    template_refs: tuple[str, ...] = ()  # A-04-T07: the templates the step points at


@dataclass(frozen=True)
class Methodology:
    areas: tuple[Area, ...]
    items: tuple[TemplateItem, ...]
    rules: tuple[AccountRule, ...]
    programs: tuple[ProgramRow, ...] = ()
    steps: tuple[StepRow, ...] = ()


def _text(value: object) -> str:
    """A cell as text: numeric account codes keep their digits (`1000.0` is `1000`)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet(workbook: Workbook, name: str) -> Iterator[tuple[object, ...]]:
    sheet = cast(ReadOnlyWorksheet, workbook[name])
    return iter(sheet.iter_rows(values_only=True))


# Raw rows read per sheet, blank ones included (ED-163): a sheet of a million empty rows is
# refused, never walked. Far above every row limit, so a real workbook never meets it.
MAX_RAW_ROWS: Final = 20_000


def _bounded(
    rows: Iterator[tuple[object, ...]], sheet: str, problems: list[Problem] | None
) -> Iterator[tuple[int, tuple[object, ...]]]:
    """The data rows after the header, numbered from 2, at most MAX_RAW_ROWS of them; past that
    the sheet is `too_many_rows` (reported once, by the caller that passes `problems`)."""
    for number, values in enumerate(rows, start=2):
        if number - 1 > MAX_RAW_ROWS:
            if problems is not None:
                problems.append(Problem(sheet, None, None, "too_many_rows"))
            return
        yield number, values


def _rows(
    rows: Iterator[tuple[object, ...]], name: str, problems: list[Problem]
) -> Iterator[tuple[int, dict[str, str]]]:
    header = tuple(_text(v) for v in next(rows, ()))
    expected = HEADERS[name]
    if header[: len(expected)] != expected:
        problems.append(Problem(name, 1, None, "missing_header"))
        return
    for number, values in _bounded(rows, name, problems):
        cells = [_text(v) for v in values[: len(expected)]]
        if not any(cells):
            continue  # blank rows are allowed anywhere
        cells += [""] * (len(expected) - len(cells))
        row = dict(zip(expected, cells, strict=True))
        bad = False
        for column, value in row.items():
            if not value:
                problems.append(Problem(name, number, column, "empty_value"))
                bad = True
            elif len(value) > _MAX_LENGTH[column]:
                problems.append(Problem(name, number, column, "too_long"))
                bad = True
        if not bad:
            yield number, row


def _optional(
    rows: Iterator[tuple[object, ...]], problems: list[Problem]
) -> dict[int, dict[str, str]]:
    """SPEC-031: the requests sheet's optional guidance columns, by row number, validated."""
    header = [_text(v).casefold() for v in next(rows, ())]
    where = {c: header.index(c) for c in OPTIONAL_REQUEST_COLUMNS if c in header}
    found: dict[int, dict[str, str]] = {}
    if not where:
        return found
    for number, values in _bounded(rows, REQUESTS, None):
        cells = {c: _text(values[i]) if i < len(values) else "" for c, i in where.items()}
        for column in ("why", "good_looks_like"):
            if len(cells.get(column, "")) > _GUIDANCE_LENGTH:
                problems.append(Problem(REQUESTS, number, column, "too_long"))
        if cells.get("effort") and cells["effort"].casefold() not in EFFORTS:
            problems.append(Problem(REQUESTS, number, "effort", "invalid_effort"))
        blocks = cells.get("blocks_fieldwork", "").casefold()
        if blocks and blocks not in ("yes", "no", "y", "n", "true", "false"):
            problems.append(Problem(REQUESTS, number, "blocks_fieldwork", "invalid_yes_no"))
        found[number] = cells
    return found


def parse_workbook(data: bytes, *, programs: bool = False) -> Methodology:
    """The workbook's areas, request items and account rules, and with `programs` (the
    `methodology.programs` flag) its programs and steps; or `TemplateInvalid`. Without it the
    Programs and Steps sheets are not read."""
    s = settings()
    if len(data) > s.methodology_max_bytes:
        raise TemplateInvalid([Problem(None, None, None, "too_large")])
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # not a zip, not a workbook, or malformed parts: all the same to the caller
        raise TemplateInvalid([Problem(None, None, None, "not_a_workbook")]) from None
    problems: list[Problem] = []
    try:
        missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
        if programs and (PROGRAMS in workbook.sheetnames) != (STEPS in workbook.sheetnames):
            missing += [n for n in (PROGRAMS, STEPS) if n not in workbook.sheetnames]
        if missing:
            raise TemplateInvalid([Problem(name, None, None, "missing_sheet") for name in missing])
        areas: list[Area] = []
        codes: set[str] = set()
        for number, row in _rows(_sheet(workbook, AREAS), AREAS, problems):
            if row["code"] in codes:
                problems.append(Problem(AREAS, number, "code", "duplicate_area"))
                continue
            codes.add(row["code"])
            areas.append(Area(row["code"], row["name"]))
        items: list[TemplateItem] = []
        extras = _optional(_sheet(workbook, REQUESTS), problems)
        for number, row in _rows(_sheet(workbook, REQUESTS), REQUESTS, problems):
            extra = extras.get(number, {})
            if row["area_code"] not in codes:
                problems.append(Problem(REQUESTS, number, "area_code", "unknown_area"))
            elif row["tier"].upper() not in TIERS:
                problems.append(Problem(REQUESTS, number, "tier", "invalid_tier"))
            else:
                tier = cast(Tier, row["tier"].upper())
                items.append(
                    TemplateItem(
                        row["area_code"],
                        row["description"],
                        tier,
                        extra.get("why") or None,
                        extra.get("good_looks_like") or None,
                        extra.get("blocks_fieldwork", "").casefold() in ("yes", "y", "true"),
                        cast(Effort, extra["effort"].casefold()) if extra.get("effort") else None,
                    )
                )
        rules: list[tuple[int, AccountRule]] = []
        # Optional since workbook v2: without the sheet there are no rules.
        sheet_rules = (
            _sheet(workbook, RULES) if RULES in workbook.sheetnames else iter([HEADERS[RULES]])
        )
        for number, row in _rows(sheet_rules, RULES, problems):
            if row["area_code"] not in codes:
                problems.append(Problem(RULES, number, "area_code", "unknown_area"))
            elif row["account_from"] > row["account_to"]:
                problems.append(Problem(RULES, number, "account_to", "invalid_range"))
            else:
                rule = AccountRule(row["area_code"], row["account_from"], row["account_to"])
                rules.append((number, rule))
        program_rows: list[ProgramRow] = []
        step_rows: list[StepRow] = []
        if programs and PROGRAMS in workbook.sheetnames and STEPS in workbook.sheetnames:
            program_rows = _programs(_sheet(workbook, PROGRAMS), problems)
            step_rows = _steps(
                _sheet(workbook, STEPS), {p.code for p in program_rows}, len(items), problems
            )
    finally:
        workbook.close()
    problems += _overlaps(rules)
    if not areas and not any(p.code == "missing_header" and p.sheet == AREAS for p in problems):
        problems.append(Problem(AREAS, None, None, "no_areas"))
    for sheet, count, limit in (
        (AREAS, len(areas), s.methodology_max_areas),
        (REQUESTS, len(items), s.methodology_max_items),
        (RULES, len(rules), s.methodology_max_rules),
        (PROGRAMS, len(program_rows), MAX_PROGRAMS),
        (STEPS, len(step_rows), MAX_STEPS),
    ):
        if count > limit:
            problems.append(Problem(sheet, None, None, "too_many_rows"))
    if problems:
        raise TemplateInvalid(problems)
    return Methodology(
        tuple(areas),
        tuple(items),
        tuple(rule for _, rule in rules),
        tuple(program_rows),
        tuple(step_rows),
    )


def _overlaps(rules: list[tuple[int, AccountRule]]) -> list[Problem]:
    """Account ranges may not overlap (compared as text), so mapping never depends on ties."""
    ordered = sorted(rules, key=lambda r: (r[1].account_from, r[1].account_to))
    problems: list[Problem] = []
    for (_, before), (number, after) in pairwise(ordered):
        if after.account_from <= before.account_to:
            problems.append(Problem(RULES, number, "account_from", "overlapping_range"))
    return problems


def area_for(code: str, rules: tuple[AccountRule, ...]) -> str | None:
    """The area of an account code: the first rule in rule order whose range holds it (AC-7)."""
    for rule in rules:
        if rule.account_from <= code <= rule.account_to:
            return rule.area_code
    return None


# --- Workbook v2: Programs and Steps (K-01-T05) ----------------------------------------------


def _named(
    rows: Iterator[tuple[object, ...]],
    sheet: str,
    required: tuple[str, ...],
    optional: tuple[str, ...],
    problems: list[Problem],
) -> Iterator[tuple[int, dict[str, str]]]:
    """A sheet's rows by header name (case-folded): the required columns anywhere, the optional
    ones when present, anything else ignored. Bounded: rows past the limit are not read."""
    header = [_text(v).casefold() for v in next(rows, ())]
    where = {c: header.index(c) for c in (*required, *optional) if c in header}
    absent = [c for c in required if c not in where]
    if absent:
        problems.append(Problem(sheet, 1, absent[0], "missing_header"))
        return
    limit = MAX_STEPS if sheet == STEPS else MAX_PROGRAMS
    seen = 0
    for number, values in _bounded(rows, sheet, problems):
        cells = {c: _text(values[i]) if i < len(values) else "" for c, i in where.items()}
        if not any(cells.values()):
            continue
        seen += 1
        if seen > limit + 1:  # one past the limit is enough to report it; read no further
            return
        yield number, cells


def _safe(sheet: str, number: int, row: dict[str, str], problems: list[Problem]) -> bool:
    """No cell reads as a formula or carries a control or bidi character (`unsafe_text`)."""
    ok = True
    for column, value in row.items():
        if value.startswith(_FORMULA_LEADS) or _HIDDEN.search(value) is not None:
            problems.append(Problem(sheet, number, column, "unsafe_text"))
            ok = False
    return ok


def _length_ok(sheet: str, number: int, row: dict[str, str], problems: list[Problem]) -> bool:
    ok = True
    for column, value in row.items():
        limit = _STEP_LENGTH.get(column)
        if limit is not None and len(value) > limit:
            problems.append(Problem(sheet, number, column, "too_long"))
            ok = False
    return ok


def _programs(rows: Iterator[tuple[object, ...]], problems: list[Problem]) -> list[ProgramRow]:
    found: list[ProgramRow] = []
    codes: set[str] = set()
    for number, row in _named(rows, PROGRAMS, PROGRAM_COLUMNS, (), problems):
        empty = [c for c in PROGRAM_COLUMNS if not row.get(c)]
        if empty:
            problems += [Problem(PROGRAMS, number, c, "empty_value") for c in empty]
            continue
        if not _length_ok(PROGRAMS, number, row, problems) or not _safe(
            PROGRAMS, number, row, problems
        ):
            continue
        if not _PROGRAM_CODE.fullmatch(row["code"]):
            problems.append(Problem(PROGRAMS, number, "code", "invalid_code"))
        elif row["code"] in codes:
            problems.append(Problem(PROGRAMS, number, "code", "duplicate_program"))
        else:
            codes.add(row["code"])
            found.append(ProgramRow(row["code"], row["name"]))
    return found


def _list(value: str) -> tuple[str, ...] | None:
    """A `;`-separated list of short codes (assertions, standards references); None if bad."""
    items = tuple(part.strip() for part in value.split(";") if part.strip())
    if len(items) > MAX_LIST_ITEMS or not all(_LIST_ITEM.fullmatch(i) for i in items):
        return None
    return items


def _positions(value: str, requests: int) -> tuple[int, ...] | None:
    """`requests`: Requests rows by number (1 = the first request), `;` or `,` separated."""
    parts = [p.strip() for p in re.split(r"[;,]", value) if p.strip()]
    if len(parts) > MAX_LIST_ITEMS:
        return None
    found: list[int] = []
    for part in parts:
        if not part.isdigit() or not 1 <= int(part) <= requests:
            return None
        if int(part) - 1 not in found:
            found.append(int(part) - 1)
    return tuple(found)


def _number(value: str) -> int | None:
    text = value[:-2] if value.endswith(".0") else value
    return int(text) if text.isdigit() and 1 <= int(text) <= 9999 else None


def _steps(
    rows: Iterator[tuple[object, ...]],
    programs: set[str],
    requests: int,
    problems: list[Problem],
) -> list[StepRow]:
    found: list[StepRow] = []
    numbers: dict[str, dict[int, int]] = {}  # program -> number -> row
    keys: set[str] = set()
    optional = (*STEP_OPTIONAL, STEP_OPTIONAL_LINKS)
    for number, row in _named(rows, STEPS, STEP_COLUMNS, optional, problems):
        empty = [c for c in STEP_COLUMNS if not row.get(c)]
        if empty:
            problems += [Problem(STEPS, number, c, "empty_value") for c in empty]
            continue
        if not _length_ok(STEPS, number, row, problems) or not _safe(STEPS, number, row, problems):
            continue
        bad = False
        if row["program_code"] not in programs:
            problems.append(Problem(STEPS, number, "program_code", "unknown_program"))
            bad = True
        step_no = _number(row["number"])
        if step_no is None:
            problems.append(Problem(STEPS, number, "number", "invalid_number"))
            bad = True
        level = row["level"].casefold()
        if level not in LEVELS:
            problems.append(Problem(STEPS, number, "level", "invalid_level"))
            bad = True
        mandatory_cell = row.get("mandatory", "").casefold()
        if mandatory_cell and mandatory_cell not in _YES_NO:
            problems.append(Problem(STEPS, number, "mandatory", "invalid_yes_no"))
            bad = True
        mandatory = mandatory_cell in _YES
        lock_note = row.get("lock_note") or None
        if lock_note is not None and not mandatory:
            problems.append(Problem(STEPS, number, "lock_note", "lock_note_without_mandatory"))
            bad = True
        step_id = (row.get("step_id") or "").casefold() or None
        if step_id is not None:
            if not STEP_KEY.fullmatch(step_id):
                problems.append(Problem(STEPS, number, "step_id", "invalid_step_id"))
                bad = True
            elif step_id in keys:
                problems.append(Problem(STEPS, number, "step_id", "duplicate_step_id"))
                bad = True
            else:
                keys.add(step_id)
        assertions = _list(row.get("assertions", ""))
        if assertions is None:
            problems.append(Problem(STEPS, number, "assertions", "invalid_list"))
            bad = True
        refs = _list(row.get("standard_refs", ""))
        if refs is None:
            problems.append(Problem(STEPS, number, "standard_refs", "invalid_list"))
            bad = True
        templates = template_refs(row.get("templates", ""))
        if templates is None:
            problems.append(Problem(STEPS, number, "templates", "invalid_template_ref"))
            bad = True
        links = _positions(row.get(STEP_OPTIONAL_LINKS, ""), requests)
        if links is None:
            problems.append(Problem(STEPS, number, STEP_OPTIONAL_LINKS, "unknown_request"))
            bad = True
        if step_no is not None and row["program_code"] in programs:
            taken = numbers.setdefault(row["program_code"], {})
            if step_no in taken:
                problems.append(Problem(STEPS, number, "number", "duplicate_number"))
                bad = True
            else:
                taken[step_no] = number
        if bad or step_no is None or assertions is None or refs is None or links is None:
            continue
        if templates is None:  # pragma: no cover -- `bad` above
            continue
        found.append(
            StepRow(
                row["program_code"],
                step_no,
                row["text"],
                cast(Level, level),
                mandatory,
                lock_note,
                step_id,
                assertions,
                refs,
                links,
                templates,
            )
        )
    # Steps are numbered 1 to n in each program, with no gap.
    for code in sorted(numbers):
        taken = numbers[code]
        for expected, step_no in enumerate(sorted(taken), start=1):
            if step_no != expected:
                problems.append(Problem(STEPS, taken[step_no], "number", "numbering_gap"))
                break
    return found
