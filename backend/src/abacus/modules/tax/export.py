"""The engine export, file mode (T-01-T25; T-01 §3.9; `docs/contracts/tax-engine-bridge.md` §2, §6;
ADR-150 W2 `tax_engine_file_bridge`; DEC-1462). RED: what leaves the firm for the engine.

- **Export** (`tax_return.export`, a person): only after the signer's final approval, of the
  checkpoint they approved. The final approval issued one grant (`approval.py`,
  `tax_return.approve`, bound to the checkpoint's content hash, held for the 10 s Undo window); the
  export consumes it
  through `identity.api.consume_grant`, so a grant is spent once and only on that content. The
  consumer reference is derived from the return and checkpoint, so a retry of the same export is
  the same consumption, and a second export from the same grant for anything else is refused. The
  same checkpoint answers the same export (one per checkpoint, the bytes identical).
- **The workbook** is code's: sheet `Return` (`return_id`, `tax_year`, `form`) and sheet `Fields`
  (`field_key`, `form`, `line`, `value`, sorted), openpyxl rebuilt with fixed ZIP timestamps and a
  pinned `modified` stamp as `evidence/render.py` does, so the same checkpoint is byte-identical.
  The engine return id comes from ADM-08's template and rules (`integrations.api`); once the
  engine has a link, its id and version are used. Stored write-once and encrypted (evidence's
  staging) before the unit of work that records it.
- **Kill switches**: the W2 capability must be enabled in the write-capability registry and the
  firm's file-bridge install connected; otherwise 409 `export_disabled` (`facts.switch`).
- **Download** is audited; **Mark exported** is a person's, once, at the version seen: the sync
  chip then reads Synced (`engine_import.sync_view`).
- **Human only** (INV-1) at five layers: the matrix (agent and system deny), the live-person check
  here, the database (`write_operations` and the `engine_exports` guard need a human actor), the
  human-only audit register (`engine_export.generated`, `.marked_exported`, 0561), and the grant
  itself (issued only by a person's decision).

Values are TRI: never logged, never in an error's facts.
"""

from __future__ import annotations

import io
import re
import uuid
import zipfile
from dataclasses import dataclass, replace
from datetime import UTC, datetime
from typing import Annotated, ClassVar, Final, cast
from uuid import UUID

from openpyxl import Workbook
from sqlalchemy import text

from abacus.kernel.classification import classified
from abacus.kernel.db import TenantContext, tenant_session
from abacus.kernel.errors import DomainConflict, NotFound
from abacus.kernel.uow import DomainEvent, Ref, Target, uow
from abacus.modules.engagements.api import work_item_facts
from abacus.modules.evidence.api import read_staged, screen_and_stage
from abacus.modules.identity.api import (
    AuthContext,
    Forbidden,
    GrantNotYetDue,
    Resource,
    authorise,
    consume_grant,
    engagement_resource,
    latest_grant,
)
from abacus.modules.integrations.api import engine_return_id, export_settings, write_capability
from abacus.modules.review.api import current_checkpoint, current_values
from abacus.modules.tax import engine_repository as repo
from abacus.modules.tax.engine_sync import ReadyToTransmit

TARGET: Final = "tax_return"
EXPORT: Final = "tax_return.export"
APPROVE: Final = "tax_return.approve"
CAPABILITY: Final = "tax_engine_file_bridge"
EXPORTABLE: Final = frozenset({"approved", "awaiting_signature", "filed", "accepted"})
_ZIP_TIME: Final = (1980, 1, 1, 0, 0, 0)
_MODIFIED: Final = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")
_CREATED: Final = re.compile(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)")
_NS: Final = uuid.UUID("5d0f1f4e-2b51-4c6a-9b63-0e6f2a4c7d25")
_MEDIA: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportDisabled(DomainConflict):
    """A kill switch is off: the W2 capability, or the firm's file-bridge install."""

    code = "export_disabled"


class NotFinalApproved(DomainConflict):
    """The return has no final approval to export, or the checkpoint is not the approved one."""

    code = "not_final_approved"


class GrantRefused(DomainConflict):
    """The approval's grant cannot be consumed: spent on another export, voided, expired, for other
    content, or its approver no longer holds the action."""

    code = "grant_refused"


class UndoWindowOpen(DomainConflict):
    """The final approval's Undo window has not closed yet: try again in a few seconds."""

    code = "undo_window_open"


class AlreadyMarked(DomainConflict):
    """The export was already marked exported."""

    code = "already_marked"


class ExportStale(DomainConflict):
    """`seen_version` is not the export's version."""

    code = "stale_version"


class EngineSyncExported(DomainEvent):
    """An export's workbook was generated (identifiers only)."""

    event_type: ClassVar[str] = "engine_sync.exported"

    work_item_id: Annotated[UUID, classified("internal")]
    export_id: Annotated[UUID, classified("internal")]


@dataclass(frozen=True)
class ExportView:
    id: UUID
    return_id: UUID
    checkpoint_id: UUID
    file_name: str
    sha256: str
    generated_at: datetime
    marked_exported_at: datetime | None
    version: int


# --- the workbook, code's and byte-stable ---------------------------------------------------


def _pinned(raw: bytes, at: datetime) -> bytes:
    stamp = at.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as source,
        zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target,
    ):
        for name in sorted(source.namelist()):
            entry = zipfile.ZipInfo(name, date_time=_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            entry.create_system = 3
            data = source.read(name)
            if name == "docProps/core.xml":
                data = _MODIFIED.sub(lambda m: m.group(1) + stamp + m.group(2), data)
                data = _CREATED.sub(lambda m: m.group(1) + stamp + m.group(2), data)
            target.writestr(entry, data)
    return out.getvalue()


def _text(value: object) -> str:
    """A field value as the engine reads it; a formula-like string is prefixed (ADR-052)."""
    out = "" if value is None else str(value)
    return "'" + out if out[:1] in ("=", "+", "-", "@") and not _number(out) else out


def _number(value: str) -> bool:
    return re.fullmatch(r"-?[0-9]+(\.[0-9]+)?", value) is not None


def workbook(
    *,
    engine_id: str,
    tax_year: int,
    form: str,
    values: dict[str, dict[str, object]],
    at: datetime,
) -> bytes:
    """Sheet `Return` (a key, value header, then `return_id` without its version, `tax_year`,
    `form`) and sheet `Fields` (contract §6), the same bytes for the same input; the fake engine's
    `read_export` reads it as the engine would."""
    book = Workbook()
    head = book.active
    if head is None:  # openpyxl always makes one sheet
        head = book.create_sheet()
    head.title = "Return"
    head.append(["key", "value"])
    for row in (("return_id", engine_id), ("tax_year", str(tax_year)), ("form", form)):
        head.append(list(row))
    fields = book.create_sheet("Fields")
    fields.append(["field_key", "form", "line", "value"])
    for key in sorted(values):
        line = key.rsplit(".", 1)[-1]
        fields.append([key, form, line, _text(values[key].get("v"))])
    book.properties.creator = ""
    book.properties.lastModifiedBy = ""
    raw = io.BytesIO()
    book.save(raw)
    return _pinned(raw.getvalue(), at)


def client_code(rule: str, client_name: str) -> str:
    """ADM-08's client-code rule over the client's name: `last_name_first_initial` ("Jonah &
    Claire Feld" → "FeldJ"); `entity_name` and, until client numbers exist, `client_number` are
    the name's letters and digits (20 at most)."""
    words = [w for w in re.split(r"[^A-Za-z0-9]+", client_name) if w]
    if rule == "last_name_first_initial" and len(words) >= 2:
        code = words[-1] + words[0][0].upper()
    else:
        code = "".join(words)
    return code[:20] or "Client"


# --- helpers --------------------------------------------------------------------------------


def _person(ctx: AuthContext) -> None:
    """INV-1: a live person; a client principal never sees a return (INV-4)."""
    if ctx.client:
        raise NotFound(TARGET)
    if ctx.tenant.actor_kind != "human" or ctx.is_support:
        raise Forbidden(EXPORT, "role")


async def _resource(ctx: AuthContext, return_id: UUID) -> tuple[Resource, UUID, str]:
    for item in await work_item_facts(ctx, TARGET):
        if item.work_item_id == return_id:
            base = await engagement_resource(ctx.tenant, item.engagement_id)
            return replace(base, work_item_id=return_id), item.engagement_id, item.client_name
    raise NotFound(TARGET)


def export_id_for(return_id: UUID, checkpoint_id: UUID) -> UUID:
    """One export per approved checkpoint: the id is the grant's consumer reference too, so a
    retry is the same consumption."""
    return uuid.uuid5(_NS, f"{return_id}:{checkpoint_id}")


_EXPORT_ROW = text(
    "SELECT id, return_id, checkpoint_id, file_name, sha256, generated_at, marked_exported_at, "
    "row_version, file_key, file_version, file_size, engine_return_id, version, generated_by "
    "FROM engine_exports WHERE id = :id"
)


def _view(row: repo.TextRow) -> ExportView:
    return ExportView(
        row.id,
        row.return_id,
        row.checkpoint_id,
        row.file_name,
        row.sha256,
        row.generated_at,
        row.marked_exported_at,
        int(row.row_version),
    )


async def _export_row(tenant: TenantContext, export_id: UUID) -> repo.TextRow | None:
    async with tenant_session(tenant) as session:
        return (await session.execute(_EXPORT_ROW, {"id": export_id})).first()


async def _switches(tenant: TenantContext) -> None:
    entry = write_capability(CAPABILITY)
    if entry is None or not entry.enabled:
        raise ExportDisabled(facts={"switch": "capability"})
    if not (await export_settings(tenant)).enabled:
        raise ExportDisabled(facts={"switch": "firm"})


# --- export, download, mark exported --------------------------------------------------------


async def export_return(ctx: AuthContext, return_id: UUID, checkpoint_id: UUID) -> ExportView:
    """Generate the workbook of the approved checkpoint, consuming the final approval's grant."""
    _person(ctx)
    resource, _engagement_id, client_name = await _resource(ctx, return_id)
    await authorise(ctx, EXPORT, resource)
    await _switches(ctx.tenant)
    export_id = export_id_for(return_id, checkpoint_id)
    existing = await _export_row(ctx.tenant, export_id)
    if existing is not None:
        return _view(existing)  # the same checkpoint answers the same export
    async with tenant_session(ctx.tenant) as session:
        facts = (
            await session.execute(
                text("SELECT status, tax_year, form FROM tax_returns WHERE work_item_id = :r"),
                {"r": return_id},
            )
        ).first()
        link = (
            await session.execute(
                text(
                    "SELECT engine_return_id, version FROM engine_return_links "
                    "WHERE return_id = :r AND engine = 'cch_axcess'"
                ),
                {"r": return_id},
            )
        ).first()
    if facts is None:
        raise NotFound(TARGET)
    checkpoint = await current_checkpoint(ctx.tenant, TARGET, return_id)
    if str(facts.status) not in EXPORTABLE or checkpoint is None or checkpoint[0] != checkpoint_id:
        raise NotFinalApproved(facts={"work_item_id": str(return_id)})
    content = checkpoint[1]
    grant = await latest_grant(ctx.tenant, Target(TARGET, return_id), APPROVE)
    if grant is None:
        raise NotFinalApproved(facts={"work_item_id": str(return_id)})
    try:
        consumed = await consume_grant(
            ctx.tenant_id, grant, export_id, content.removeprefix("sha256:")
        )
    except GrantNotYetDue:
        raise UndoWindowOpen(facts={"work_item_id": str(return_id)}) from None
    if not consumed:
        raise GrantRefused(facts={"work_item_id": str(return_id)})
    settings = await export_settings(ctx.tenant)
    form = str(facts.form)
    if link is not None:
        engine_id, version = str(link.engine_return_id), max(int(link.version), 1)
    else:
        engine_id = engine_return_id(
            settings.return_id_template,
            tax_year=int(facts.tax_year),
            type_letter=settings.type_letters.get(form, "I"),
            client_code=client_code(settings.client_code_rule, client_name),
            version=None,
        )
        version = 1
    at = datetime(int(facts.tax_year) + 1, 1, 1, tzinfo=UTC)
    values = await current_values(ctx.tenant, TARGET, return_id)
    content_bytes = workbook(
        engine_id=engine_id,
        tax_year=int(facts.tax_year),
        form=form,
        values=values,
        at=at,
    )
    file_name = f"{engine_id.replace(':', '_')}_V{version}.xlsx"
    screened = await screen_and_stage(ctx.tenant_id, content_bytes, file_name, None)
    async with uow(ctx.tenant) as tx:
        operation_id = uuid.uuid4()
        await tx.session.execute(
            text(
                "INSERT INTO write_operations (id, tenant_id, capability, subject_type, "
                "subject_id, grant_id, payload_sha256, performed_by) VALUES (:id, :t, :c, "
                ":st, :s, :g, :h, :u)"
            ),
            {
                "id": operation_id,
                "t": ctx.tenant_id,
                "c": CAPABILITY,
                "st": TARGET,
                "s": return_id,
                "g": grant,
                "h": screened.sha256,
                "u": ctx.user_id,
            },
        )
        await tx.session.execute(
            text(
                "INSERT INTO engine_exports (id, tenant_id, return_id, checkpoint_id, "
                "content_hash, engine_return_id, version, file_name, file_key, file_version, "
                "file_size, sha256, write_operation_id, generated_by) VALUES (:id, :t, :r, :cp, "
                ":ch, :e, :v, :n, :k, :fv, :sz, :h, :op, :u)"
            ),
            {
                "id": export_id,
                "t": ctx.tenant_id,
                "r": return_id,
                "cp": checkpoint_id,
                "ch": content,
                "e": engine_id,
                "v": version,
                "n": screened.file_name,
                "k": screened.stored.key,
                "fv": screened.stored.version_id,
                "sz": screened.stored.size,
                "h": screened.sha256,
                "op": operation_id,
                "u": ctx.user_id,
            },
        )
        await repo.insert_event(
            tx.session,
            {
                "tenant_id": ctx.tenant_id,
                "return_id": return_id,
                "kind": "export_generated",
                "engine_return_id": engine_id,
                "version": version,
                "file_name": screened.file_name,
                "file_key": screened.stored.key,
                "file_version": screened.stored.version_id,
                "file_size": screened.stored.size,
                "payload_sha256": screened.sha256,
                "actor_user_id": ctx.user_id,
                "write_operation_id": operation_id,
            },
        )
        tx.record(
            "engine_export.generated",
            target=Target(TARGET, return_id),
            after=Ref(export=export_id, checkpoint=checkpoint_id, sha256=screened.sha256),
        )
        tx.emit(EngineSyncExported(work_item_id=return_id, export_id=export_id))
    found = await _export_row(ctx.tenant, export_id)
    if found is None:  # written in the unit of work just committed
        raise NotFound("engine_export")
    return _view(found)


async def _owned(ctx: AuthContext, export_id: UUID) -> tuple[repo.TextRow, UUID]:
    _person(ctx)
    row = await _export_row(ctx.tenant, export_id)
    if row is None:
        raise NotFound("engine_export")
    resource, engagement_id, _c = await _resource(ctx, cast(UUID, row.return_id))
    await authorise(ctx, EXPORT, resource)
    return row, engagement_id


async def download(ctx: AuthContext, export_id: UUID) -> tuple[str, bytes]:
    """The workbook, fingerprint-checked on read, and an audit event naming who took it."""
    row, _engagement_id = await _owned(ctx, export_id)
    content = await read_staged(
        ctx.tenant_id, row.file_key, row.file_version, row.sha256, int(row.file_size)
    )
    async with uow(ctx.tenant) as tx:
        tx.record(
            "engine_export.downloaded",
            target=Target(TARGET, row.return_id),
            after=Ref(export=export_id, sha256=row.sha256),
        )
    return str(row.file_name), content


async def mark_exported(ctx: AuthContext, export_id: UUID, seen_version: int) -> ExportView:
    """A person confirms the file reached the engine: the chip reads Synced."""
    row, engagement_id = await _owned(ctx, export_id)
    if row.marked_exported_at is not None:
        raise AlreadyMarked(facts={"export_id": str(export_id)})
    async with uow(ctx.tenant) as tx:
        moved = (
            await tx.session.execute(
                text(
                    "UPDATE engine_exports SET marked_exported_at = abacus_now(), "
                    "marked_exported_by = :u, row_version = row_version + 1 WHERE id = :id "
                    "AND row_version = :seen AND marked_exported_at IS NULL RETURNING id"
                ),
                {"u": ctx.user_id, "id": export_id, "seen": seen_version},
            )
        ).first()
        if moved is None:
            raise ExportStale(facts={"export_id": str(export_id)})
        link = await repo.link_of(tx.session, row.return_id)
        if link is None:
            await repo.insert_link(tx.session, ctx.tenant_id, row.return_id, row.engine_return_id)
        await repo.set_link_version(tx.session, row.return_id, int(row.version))
        await repo.insert_event(
            tx.session,
            {
                "tenant_id": ctx.tenant_id,
                "return_id": row.return_id,
                "kind": "marked_exported",
                "engine_return_id": row.engine_return_id,
                "version": int(row.version),
                "actor_user_id": ctx.user_id,
            },
        )
        tx.record(
            "engine_export.marked_exported",
            target=Target(TARGET, row.return_id),
            after=Ref(export=export_id),
        )
        # T-01-T27: `ready_to_transmit` to whoever generated the file, if someone else marked it.
        tx.emit(
            ReadyToTransmit(
                work_item_id=row.return_id,
                engagement_id=engagement_id,
                by_user=ctx.user_id,
                recipient_user_ids=[row.generated_by],
            )
        )
    found = await _export_row(ctx.tenant, export_id)
    if found is None:
        raise NotFound("engine_export")
    return _view(found)


__all__ = [
    "AlreadyMarked",
    "ExportDisabled",
    "ExportStale",
    "ExportView",
    "GrantRefused",
    "NotFinalApproved",
    "UndoWindowOpen",
    "client_code",
    "download",
    "export_id_for",
    "export_return",
    "mark_exported",
    "workbook",
]
