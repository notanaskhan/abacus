"""Organizer templates: the import format and its dry-run diff, pure (T-02-T16; T-02 §3.5;
DEC-1369; ADR-050, ADR-052).

A template is the firm's database content (AGENTS 16), imported as CSV, XLSX (first sheet; a
formula cell is read as its cached value) or JSON, and parsed in the isolated child: the file is
client-of-the-firm content and hostile. The rows (CSV and XLSX) are:

    section, icon_key, key, type, client_wording, help_text, conditional_on, conditional_value

JSON is `{"sections": [{"name", "icon_key", "questions": [{"key", "type", "client_wording",
"help_text", "conditional_on", "conditional_value"}]}]}`. Every question needs a key (stable across
versions: the diff and the prior-year copy match on it), a type and the client's wording; a
condition names an earlier key. Problems are codes with where they are, never the cell's text.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from typing import Final, Literal, cast

from openpyxl import load_workbook

Format = Literal["csv", "xlsx", "json"]
TYPES: Final = frozenset({"yes_no", "amount_group", "upload_if_yes", "long_text", "choice"})
COLUMNS: Final = (
    "section",
    "icon_key",
    "key",
    "type",
    "client_wording",
    "help_text",
    "conditional_on",
    "conditional_value",
)
MAX_BYTES: Final = 2 * 1024 * 1024
MAX_QUESTIONS: Final = 500
_KEY: Final = re.compile(r"[a-z0-9][a-z0-9_.-]{0,79}")
_ICON: Final = re.compile(r"[a-z][a-z0-9-]{0,39}")


class TemplateInvalid(ValueError):
    def __init__(self, problems: list[str]) -> None:
        super().__init__(", ".join(problems))
        self.problems = problems


@dataclass(frozen=True)
class Question:
    key: str
    type: str
    client_wording: str
    help_text: str | None
    conditional_on: str | None
    conditional_value: str | None


@dataclass(frozen=True)
class Group:
    name: str
    icon_key: str | None
    questions: tuple[Question, ...]


@dataclass(frozen=True)
class Template:
    sections: tuple[Group, ...]

    def content(self) -> dict[str, object]:
        """The stored form (`questionnaire_template_versions.content`)."""
        return {
            "sections": [
                {
                    "name": s.name,
                    "icon_key": s.icon_key,
                    "questions": [
                        {
                            "key": q.key,
                            "type": q.type,
                            "client_wording": q.client_wording,
                            "help_text": q.help_text,
                            "conditional_on": q.conditional_on,
                            "conditional_value": q.conditional_value,
                        }
                        for q in s.questions
                    ],
                }
                for s in self.sections
            ]
        }

    def questions(self) -> dict[str, Question]:
        return {q.key: q for s in self.sections for q in s.questions}


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _rows_to_template(rows: list[dict[str, str]]) -> Template:
    problems: list[str] = []
    order: list[str] = []
    sections: dict[str, tuple[str | None, list[Question]]] = {}
    seen: set[str] = set()
    for n, row in enumerate(rows, start=2):
        section, key, kind = row.get("section", ""), row.get("key", ""), row.get("type", "")
        wording = row.get("client_wording", "")
        icon = row.get("icon_key") or None
        cond = row.get("conditional_on") or None
        if not section or len(section) > 80:
            problems.append(f"row {n}: section")
        if not _KEY.fullmatch(key):
            problems.append(f"row {n}: key")
        elif key in seen:
            problems.append(f"row {n}: duplicate_key")
        if kind not in TYPES:
            problems.append(f"row {n}: type")
        if not wording or len(wording) > 2000:
            problems.append(f"row {n}: client_wording")
        if icon is not None and not _ICON.fullmatch(icon):
            problems.append(f"row {n}: icon_key")
        if cond is not None and cond not in seen:
            problems.append(f"row {n}: conditional_on")
        if len(row.get("help_text", "")) > 2000:
            problems.append(f"row {n}: help_text")
        seen.add(key)
        if section not in sections:
            order.append(section)
            sections[section] = (icon, [])
        sections[section][1].append(
            Question(
                key,
                kind,
                wording,
                row.get("help_text") or None,
                cond,
                row.get("conditional_value") or None,
            )
        )
    if not rows:
        problems.append("empty")
    if len(seen) > MAX_QUESTIONS:
        problems.append("too_many_questions")
    if problems:
        raise TemplateInvalid(problems)
    return Template(
        tuple(Group(name, sections[name][0], tuple(sections[name][1])) for name in order)
    )


def _table(header: list[str], body: list[list[object]]) -> list[dict[str, str]]:
    names = [_text(h).lower() for h in header]
    missing = [c for c in ("section", "key", "type", "client_wording") if c not in names]
    if missing:
        raise TemplateInvalid([f"missing_column: {c}" for c in missing])
    return [
        {name: _text(cell) for name, cell in zip(names, row, strict=False) if name in COLUMNS}
        for row in body
        if any(_text(c) for c in row)
    ]


def parse(content: bytes, fmt: Format) -> Template:
    """The template, or `TemplateInvalid` with every problem found (run in the isolated child)."""
    if len(content) > MAX_BYTES:
        raise TemplateInvalid(["too_large"])
    if fmt == "json":
        try:
            parsed: object = json.loads(content.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise TemplateInvalid(["not_json"]) from None
        if not isinstance(parsed, dict):
            raise TemplateInvalid(["sections"])
        doc = cast(dict[str, object], parsed)
        if not isinstance(doc.get("sections"), list):
            raise TemplateInvalid(["sections"])
        rows: list[dict[str, str]] = []
        for section in cast(list[object], doc["sections"]):
            if not isinstance(section, dict):
                raise TemplateInvalid(["sections"])
            s = cast(dict[str, object], section)
            for q in cast(list[object], s.get("questions") or []):
                if not isinstance(q, dict):
                    raise TemplateInvalid(["questions"])
                question = cast(dict[str, object], q)
                rows.append(
                    {
                        "section": _text(s.get("name")),
                        "icon_key": _text(s.get("icon_key")),
                        **{c: _text(question.get(c)) for c in COLUMNS[2:]},
                    }
                )
        return _rows_to_template(rows)
    if fmt == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise TemplateInvalid(["not_utf8"]) from None
        table = list(csv.reader(io.StringIO(text)))
        if not table:
            raise TemplateInvalid(["empty"])
        return _rows_to_template(_table(table[0], cast(list[list[object]], table[1:])))
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # openpyxl raises several types for a file that is not a workbook
        raise TemplateInvalid(["not_xlsx"]) from None
    sheet = book.worksheets[0]
    values = [list(r) for r in sheet.iter_rows(values_only=True)]
    if not values:
        raise TemplateInvalid(["empty"])
    return _rows_to_template(
        _table([_text(h) for h in values[0]], cast(list[list[object]], values[1:]))
    )


@dataclass(frozen=True)
class Changes:
    added: tuple[str, ...]
    changed: tuple[str, ...]
    removed: tuple[str, ...]


def changes_between(before: Template | None, after: Template) -> Changes:
    """Questions by key: added, changed (any field) and removed, in key order."""
    old = before.questions() if before else {}
    new = after.questions()
    return Changes(
        tuple(sorted(new.keys() - old.keys())),
        tuple(sorted(k for k in new.keys() & old.keys() if new[k] != old[k])),
        tuple(sorted(old.keys() - new.keys())),
    )


def from_content(content: object) -> Template:
    """A stored version back as a template (the stored form is the JSON format)."""
    return parse(json.dumps(content).encode(), "json")


__all__ = [
    "COLUMNS",
    "MAX_BYTES",
    "TYPES",
    "Changes",
    "Format",
    "Group",
    "Question",
    "Template",
    "TemplateInvalid",
    "changes_between",
    "from_content",
    "parse",
]
