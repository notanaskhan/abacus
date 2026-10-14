"""Parse a general ledger export, at scale, in an isolated child (SPEC-034; TASK-069 D2, D3;
ADR-050, ADR-052, ADR-110). PROTECTED.

The file is hostile and was malware-scanned first. The child streams its rows: account headings,
opening balance rows, postings, totals. It writes each posting, normalised, to a private file the
parent gave it, and returns only a summary (per account opening, debits, credits; months; totals),
so a million lines never cross the process boundary. Text is cleaned (control characters dropped,
lengths capped); amounts accept signs, separators, currency symbols and parentheses. Anything
malformed is a `NormaliseError` with a short code; a crash or hang is `unreadable`.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Final, cast

from abacus.kernel.isolation import IsolationFailed, Limits, run_isolated_with
from abacus.modules.connections.gl_formats import GL_FORMATS, GlFormat
from abacus.modules.connections.import_parse import (
    _amount,  # pyright: ignore[reportPrivateUsage] -- the same amount rules as trial balances
)
from abacus.modules.ledger.api import GlAccount, GlEntry, GlSummary, NormaliseError

MAX_GL_BYTES: Final = 100 * 1024 * 1024  # SPEC-034 Q2
MAX_GL_LINES: Final = 1_000_000
GL_LIMITS: Final = Limits(timeout_seconds=300, cpu_seconds=240, memory_bytes=2 * 1024**3)  # D3
_REFUSED: Final = frozenset({"Cc", "Cf", "Cs", "Co", "Zl", "Zp"})
_LEADING: Final = re.compile(r"^([0-9][0-9.\-]{0,19})\s+(.+)$")
_TRAILING: Final = re.compile(r"^(.+?)\s*\(([0-9A-Za-z.\-]{1,20})\)$")
_SEPARATOR: Final = re.compile("^[\u00b7\u2013:-]\\s*")
_DATES: Final = ("%m/%d/%Y", "%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%b %d, %Y", "%B %d, %Y")
_FIELDS: Final = (
    "line_no",
    "account_code",
    "account_name",
    "posted_on",
    "reference",
    "kind",
    "counterparty",
    "description",
    "debit",
    "credit",
    "source_ref",
)

Cell = str | int | float | Decimal | datetime | date | None


def _clean(value: Cell, limit: int) -> str:
    text = "" if value is None else str(value)
    kept = "".join(c for c in text if unicodedata.category(c) not in _REFUSED)
    return " ".join(kept.split())[:limit]


def _date(value: Cell) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value, 40)
    for pattern in _DATES:
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _account(heading: str, fmt: GlFormat) -> tuple[str, str]:
    if fmt.code_style == "parenthesised":
        found = _TRAILING.match(heading)
        if found is not None:
            return found.group(2), found.group(1).strip()
    else:
        found = _LEADING.match(heading)
        if found is not None:
            return found.group(1), _SEPARATOR.sub("", found.group(2).strip()) or found.group(2)
    return heading[:50], heading


def _rows(content: bytes, fmt: GlFormat) -> Iterator[list[Cell]]:
    if fmt.file_type == "xlsx":
        from openpyxl import load_workbook  # only in the child process

        try:
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True, keep_vba=False
            )
        except Exception:
            raise NormaliseError("unreadable") from None
        for values in workbook.worksheets[0].iter_rows(values_only=True):
            yield [cast(Cell, v) for v in values]
        workbook.close()
        return
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise NormaliseError("unreadable")
    csv.field_size_limit(10_000)
    try:
        for row in csv.reader(io.StringIO(text), strict=True):
            yield list(row)
    except csv.Error:
        raise NormaliseError("unreadable") from None


@dataclass
class _Account:
    code: str
    name: str
    opening: Decimal | None = None
    debit: Decimal = Decimal(0)
    credit: Decimal = Decimal(0)


@dataclass
class _State:
    accounts: dict[str, _Account] = field(default_factory=dict[str, _Account])
    months: set[tuple[int, int]] = field(default_factory=set[tuple[int, int]])
    lines: int = 0
    outside: int = 0
    debit: Decimal = Decimal(0)
    credit: Decimal = Decimal(0)


def _cell(row: list[Cell], columns: dict[str, int], role: str) -> Cell:
    index = columns.get(role)
    return row[index] if index is not None else None


def _header(row: list[Cell], fmt: GlFormat) -> dict[str, int] | None:
    labels = {_clean(c, 60).casefold(): i for i, c in enumerate(row) if _clean(c, 60)}
    wanted = {
        "date": fmt.date_header,
        "kind": fmt.kind_header,
        "reference": fmt.reference_header,
        "counterparty": fmt.counterparty_header,
        "description": fmt.description_header,
        "amount": fmt.amount_header,
        "debit": fmt.debit_header,
        "credit": fmt.credit_header,
        "balance": fmt.balance_header,
    }
    found: dict[str, int] = {}
    for role, label in wanted.items():
        if label is None:
            continue
        index = labels.get(label.casefold())
        if index is None and role != "balance":
            return None
        if index is not None:
            found[role] = index
    return found


def parse_gl_to_file(
    content: bytes, format_id: str, period_start: date, period_end: date, out: str
) -> GlSummary:
    """In the child (or a test): write the postings to `out` as CSV and return the summary."""
    fmt = GL_FORMATS.get(format_id)
    if fmt is None:
        raise NormaliseError("unknown_format")
    if len(content) > MAX_GL_BYTES:
        raise NormaliseError("payload_too_large")
    if (fmt.file_type == "xlsx") != (content[:4] == b"PK\x03\x04"):
        raise NormaliseError("unrecognised_format")
    state = _State()
    columns: dict[str, int] | None = None
    current: _Account | None = None
    with Path(out).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        for n, raw in enumerate(_rows(content, fmt), start=1):
            if columns is None:
                if n > 50:
                    raise NormaliseError("unrecognised_format")
                columns = _header(raw, fmt)
                continue
            width = max(columns.values()) + 1
            row = [*raw, *([None] * (width - len(raw)))]

            first = _clean(row[0], 200)
            kind = _clean(_cell(row, columns, "kind"), 60)
            posted = _date(_cell(row, columns, "date"))
            if first and any(first.startswith(p) for p in fmt.total_prefixes):
                current = None
                continue
            if first in fmt.skip_labels:
                continue
            if first in fmt.opening_labels or kind in fmt.opening_labels:
                if current is None:
                    raise NormaliseError("malformed_payload")
                balance = _cell(row, columns, "balance")
                if fmt.amount_mode == "split" and balance in (None, ""):
                    balance_value = _amount(_cell(row, columns, "debit")) - _amount(
                        _cell(row, columns, "credit")
                    )
                else:
                    balance_value = _amount(
                        balance if balance not in (None, "") else _cell(row, columns, "amount")
                    )
                current.opening = balance_value
                continue
            if posted is None:
                others = [c for i, c in enumerate(row) if i != 0 and _clean(c, 10)]
                if first and not others:  # an account heading
                    code, name = _account(first, fmt)
                    current = state.accounts.setdefault(code, _Account(code, name[:200]))
                continue  # a blank or note row
            if current is None:
                raise NormaliseError("malformed_payload")  # a posting outside any account
            if fmt.amount_mode == "signed":
                amount = _amount(_cell(row, columns, "amount"))
                debit, credit = (amount, Decimal(0)) if amount >= 0 else (Decimal(0), -amount)
            else:
                debit, credit = (
                    _amount(_cell(row, columns, "debit")),
                    _amount(_cell(row, columns, "credit")),
                )
                net = debit - credit
                debit, credit = (net, Decimal(0)) if net >= 0 else (Decimal(0), -net)
            state.lines += 1
            if state.lines > MAX_GL_LINES:
                raise NormaliseError("payload_too_large")
            if not (period_start <= posted <= period_end):
                state.outside += 1
            state.months.add((posted.year, posted.month))
            current.debit += debit
            current.credit += credit
            state.debit += debit
            state.credit += credit
            writer.writerow(
                [
                    state.lines,
                    current.code,
                    current.name,
                    posted.isoformat(),
                    _clean(_cell(row, columns, "reference"), 100),
                    kind,
                    _clean(_cell(row, columns, "counterparty"), 200),
                    _clean(_cell(row, columns, "description"), 500),
                    str(debit),
                    str(credit),
                    f"row-{n}",
                ]
            )
    if columns is None:
        raise NormaliseError("unrecognised_format")
    if state.lines == 0:
        raise NormaliseError("empty")
    return GlSummary(
        line_count=state.lines,
        total_debit=state.debit,
        total_credit=state.credit,
        accounts=tuple(
            GlAccount(a.code, a.name, a.opening, a.debit, a.credit)
            for a in state.accounts.values()
        ),
        months=frozenset(state.months),
        outside_period=state.outside,
        has_openings=any(a.opening is not None for a in state.accounts.values()),
    )


def read_entries(path: str) -> Iterator[GlEntry]:
    """The postings the child wrote, back as `GlEntry`, in line order."""
    with Path(path).open(encoding="utf-8", newline="") as handle:
        for row in csv.reader(handle):
            values = dict(zip(_FIELDS, row, strict=True))
            yield GlEntry(
                int(values["line_no"]),
                values["account_code"],
                values["account_name"],
                date.fromisoformat(values["posted_on"]),
                values["reference"] or None,
                values["kind"] or None,
                values["counterparty"] or None,
                values["description"] or None,
                Decimal(values["debit"]),
                Decimal(values["credit"]),
                values["source_ref"],
            )


def parse_gl_isolated(
    content: bytes, format_id: str, period_start: date, period_end: date, out: str
) -> GlSummary:
    """`parse_gl_to_file` in a short-lived child with the GL limits (D3)."""
    try:
        return run_isolated_with(
            GL_LIMITS, parse_gl_to_file, content, format_id, period_start, period_end, out
        )
    except IsolationFailed:
        raise NormaliseError("unreadable") from None
