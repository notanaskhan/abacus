"""Parse an uploaded ledger export into the common ledger model (SPEC-028; TASK-052; ADR-038,
ADR-050, ADR-110). PROTECTED.

The file is hostile (ADR-052) and was malware-scanned before reaching here. Parsing runs in a
short-lived child process with CPU, memory and time limits (`parse_isolated`, TASK-052 D2): a
crash or a hang is `unreadable`, never a worker crash. Spreadsheets are read for cached values
only; formulas are never evaluated and macros never run. Amounts accept signs, thousands
separators, currency symbols and parentheses for negatives; each line's net (debit minus credit)
decides which side it's on. Rows without an account (sub-totals, notes) are skipped. Control
totals come from the file's total row, or else from the lines (SPEC-028 Q2). Anything unexpected
is a `NormaliseError` with a short code; `looks_like_<system>` when the file matches another
format.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from typing import Final, cast

from abacus.kernel.isolation import IsolationFailed, run_isolated
from abacus.modules.connections.import_formats import FORMATS, Format
from abacus.modules.ledger.api import LedgerLine, NormalisedTrialBalance, NormaliseError

MAX_BYTES: Final = 20 * 1024 * 1024
MAX_ROWS: Final = 60_000
MAX_LINES: Final = 50_000
MAX_TEXT: Final = 200
CURRENCY: Final = "USD"  # SPEC-028 Q3: US exports rarely say; anything else is refused
_CENT: Final = Decimal("0.01")
_AMOUNT: Final = re.compile(r"-?[0-9]{1,15}(?:\.[0-9]{1,2})?")
_REFUSED: Final = frozenset({"Cc", "Cf", "Cs", "Co", "Zl", "Zp"})
_LEADING_CODE: Final = re.compile(r"^([0-9][0-9.\-]{0,19})\s+(.+)$")
_TRAILING_CODE: Final = re.compile(r"^(.+?)\s*\(([0-9A-Za-z.\-]{1,20})\)$")
_DATE_FORMATS: Final = ("%B %d, %Y", "%d %B %Y", "%m/%d/%Y", "%Y-%m-%d", "%b %d, %Y", "%d %b %Y")
# TASK-068 D2: a period by month ("Dec 2026") ends on that month's last day.
_MONTH_FORMATS: Final = ("%b %Y", "%B %Y")
_SEPARATOR: Final = re.compile(r"^[\u00b7\-\u2013:]\s*")

Cell = str | int | float | Decimal | datetime | date | None


@dataclass(frozen=True)
class ParsedImport:
    trial_balance: NormalisedTrialBalance
    totals_from_lines: bool


def _rows_csv(content: bytes) -> list[list[Cell]]:
    for encoding in ("utf-8-sig", "cp1252"):  # SPEC-028 Q6
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise NormaliseError("unreadable")
    csv.field_size_limit(MAX_TEXT * 50)
    try:
        rows: list[list[Cell]] = []
        for row in csv.reader(io.StringIO(text), strict=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS:
                raise NormaliseError("payload_too_large")
        return rows
    except csv.Error:
        raise NormaliseError("unreadable") from None


def _rows_xlsx(content: bytes) -> list[list[Cell]]:
    from openpyxl import load_workbook  # only in the child process

    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_vba=False
        )
        sheet = workbook.worksheets[0]
        rows: list[list[Cell]] = []
        for values in sheet.iter_rows(values_only=True):
            rows.append([cast(Cell, v) for v in values])
            if len(rows) > MAX_ROWS:
                raise NormaliseError("payload_too_large")
        workbook.close()
        return rows
    except NormaliseError:
        raise
    except Exception:
        raise NormaliseError("unreadable") from None


def _text(value: Cell) -> str:
    text = "" if value is None else str(value).strip()
    if len(text) > MAX_TEXT or any(unicodedata.category(c) in _REFUSED for c in text):
        raise NormaliseError("malformed_payload")
    return text


def _amount(value: Cell) -> Decimal:
    if value is None or value == "":
        return Decimal(0)
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            return Decimal(str(value)).quantize(_CENT)
        except InvalidOperation:
            raise NormaliseError("invalid_amount") from None
    text = str(value).strip().replace("$", "").replace(",", "").replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    if text in ("", "-"):
        return Decimal(0)
    if not _AMOUNT.fullmatch(text):
        raise NormaliseError("invalid_amount")
    found = Decimal(text).quantize(_CENT)
    return -found if negative else found


def _header_index(rows: list[list[Cell]], fmt: Format) -> tuple[int, dict[str, int]] | None:
    """The header row and the columns it names, if this file has the format's header."""
    for n, row in enumerate(rows[:50]):
        labels = {_text(c).casefold(): i for i, c in enumerate(row) if _text(c)}
        debit = next(
            (labels[h.casefold()] for h in fmt.debit_headers if h.casefold() in labels), None
        )
        credit = next(
            (labels[h.casefold()] for h in fmt.credit_headers if h.casefold() in labels), None
        )
        if debit is None or credit is None:
            continue
        if fmt.account_header:
            account = labels.get(fmt.account_header.casefold())
            if account is None:
                continue
        else:
            account = 0
            if row and _text(row[0]):
                continue  # the account column's header is blank in this format
        found = {"account": account, "debit": debit, "credit": credit}
        if fmt.code_style == "column":
            code = labels.get(fmt.code_header.casefold())
            if code is None:
                continue
            found["code"] = code
        return n, found
    return None


def _period_end(rows: list[list[Cell]], before: int, fmt: Format) -> date | None:
    for row in rows[:before]:
        for cell in row:
            text = _text(cell)
            for prefix in fmt.period_prefixes:
                if text.casefold().startswith(prefix.casefold()):
                    rest = text[len(prefix) :].strip()
                    for pattern in _DATE_FORMATS:
                        try:
                            return datetime.strptime(rest, pattern).date()
                        except ValueError:
                            continue
                    for pattern in _MONTH_FORMATS:
                        try:
                            first = datetime.strptime(rest, pattern).date()
                        except ValueError:
                            continue
                        following = first.replace(year=first.year + first.month // 12)
                        following = following.replace(month=first.month % 12 + 1)
                        return following - timedelta(days=1)
    return None


def _account(name: str, fmt: Format, code: str = "") -> tuple[str, str]:
    if fmt.code_style == "column":
        return (code or name), name  # TASK-068: the code has its own column
    pattern = _LEADING_CODE if fmt.code_style == "leading" else _TRAILING_CODE
    found = pattern.match(name)
    if found is None:
        return name, name  # no account numbers: the name is the code
    if fmt.code_style == "leading":
        # "1000 · Checking", "1000 - Checking" (TASK-068): the separator isn't the name.
        return found.group(1), _SEPARATOR.sub("", found.group(2).strip()) or found.group(2)
    return found.group(2), found.group(1).strip()


def parse_export(
    content: bytes, format_id: str, period_start: date, period_end: date
) -> ParsedImport:
    """In-process parse (tests, and inside the isolated child)."""
    fmt = FORMATS.get(format_id)
    if fmt is None:
        raise NormaliseError("unknown_format")
    if len(content) > MAX_BYTES:
        raise NormaliseError("payload_too_large")
    is_zip = content[:4] == b"PK\x03\x04"
    if (fmt.file_type == "xlsx") != is_zip:
        raise NormaliseError(_mismatch(content, fmt))
    rows = _rows_xlsx(content) if is_zip else _rows_csv(content)
    found = _header_index(rows, fmt)
    if found is None:
        raise NormaliseError(_mismatch(content, fmt, rows))
    header, columns = found
    declared_end = _period_end(rows, header, fmt)
    lines: list[LedgerLine] = []
    total: tuple[Decimal, Decimal] | None = None
    for n, row in enumerate(rows[header + 1 :], start=header + 2):
        cells = [*row, *([None] * (max(columns.values()) + 1 - len(row)))]
        name = _text(cells[columns["account"]])
        debit, credit = _amount(cells[columns["debit"]]), _amount(cells[columns["credit"]])
        if not name:
            continue  # spacer or sub-total without a label
        if name in fmt.total_labels:
            total = (debit, credit)
            break
        if debit == 0 and credit == 0 and all(_text(c) == "" for c in cells[1:]):
            continue  # a section heading ("Assets")
        net = debit - credit
        code, account_name = _account(
            name, fmt, _text(cells[columns["code"]]) if "code" in columns else ""
        )
        lines.append(
            LedgerLine(
                account_code=code,
                account_name=account_name,
                debit=net if net > 0 else Decimal(0),
                credit=-net if net < 0 else Decimal(0),
                source_ref=f"row-{n}",
            )
        )
        if len(lines) > MAX_LINES:
            raise NormaliseError("payload_too_large")
    sum_debit = sum((line.debit for line in lines), Decimal(0))
    sum_credit = sum((line.credit for line in lines), Decimal(0))
    return ParsedImport(
        NormalisedTrialBalance(
            period_start=period_start,
            period_end=declared_end or period_end,
            currency=CURRENCY,
            lines=tuple(lines),
            declared_debit=total[0] if total is not None else sum_debit,
            declared_credit=total[1] if total is not None else sum_credit,
        ),
        totals_from_lines=total is None,
    )


def _mismatch(content: bytes, chosen: Format, rows: list[list[Cell]] | None = None) -> str:
    """`looks_like_<system>` when another format fits, else `unrecognised_format`."""
    is_zip = content[:4] == b"PK\x03\x04"
    for other in FORMATS.values():
        if other.id == chosen.id or (other.file_type == "xlsx") != is_zip:
            continue
        try:
            candidate = (
                rows
                if rows is not None
                else (_rows_xlsx(content) if is_zip else _rows_csv(content))
            )
        except NormaliseError:
            return "unrecognised_format"
        if _header_index(candidate, other) is not None:
            return f"looks_like_{other.system_key}"
    return "unrecognised_format"


def parse_isolated(
    content: bytes, format_id: str, period_start: date, period_end: date
) -> ParsedImport:
    """`parse_export` in a short-lived child process with limits (ADR-110; TASK-052 D2)."""
    try:
        return run_isolated(parse_export, content, format_id, period_start, period_end)
    except IsolationFailed:
        raise NormaliseError("unreadable") from None
