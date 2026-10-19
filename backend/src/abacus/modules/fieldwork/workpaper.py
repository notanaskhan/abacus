"""The journal entry testing workpaper, rendered by code (SPEC-037 AC-5; TASK-084; ADR-131).

Deterministic (ADR-042): the same records always give the same bytes. openpyxl's write-only
workbook, every text cell an explicit string (never a formula), fixed document properties, and a
ZIP rebuilt with fixed timestamps. A signed version's fingerprint covers its content without the
Sign-off sheet (`with_signoff=False`), so later signatures don't change what was signed; the
download adds the Sign-off sheet (D1). Pure: no database, no clock.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Final

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

_ZIP_TIME: Final = (1980, 1, 1, 0, 0, 0)
_MODIFIED = re.compile(
    rb"(<dcterms:(?:modified|created)[^>]*>)[^<]*(</dcterms:(?:modified|created)>)"
)
_FORMULA_LEADS: Final = ("=", "+", "-", "@", "\t", "\r")
_MAX_CELL: Final = 32_767
_AMOUNT: Final = "#,##0.00;(#,##0.00)"
_BOLD: Final = Font(bold=True)
RESULT_LABELS: Final = {
    "supported": "Supported",
    "exception": "Exception",
    "not_tested": "Not tested",
    None: "No result",
}
RULE_LABELS: Final = {
    "large": "At or above the selection threshold",
    "round": "Round amount",
    "near_year_end": "Posted near year end",
    "manual": "Manual entry type",
    "rare_account": "Rarely used account",
    "weekend": "Posted on a weekend",
    "keywords": "Description contains a flagged word",
    "random": "Random selection",
    "picked": "Selected by the team",
    "band": "At or above its timing band's threshold",
}
NOT_PERFORMED: Final = {
    "after_year_end": "Entries posted after year end: needs the subsequent-period ledger.",
    "entered_by": "Who entered each entry: the ledger export doesn't include it.",
}


@dataclass(frozen=True)
class Row:
    posted_on: date
    kind: str
    reference: str
    counterparty: str
    accounts: tuple[str, ...]
    amount: Decimal
    rules: tuple[str, ...]
    reason: str | None
    result: str | None
    note: str | None
    evidence: tuple[str, ...]
    support_requested: bool
    band: int | None = None  # SPEC-038
    recorded_bill: str | None = None
    relates_to_year: str | None = None


@dataclass(frozen=True)
class Signature:
    kind: str  # ready | reviewed
    name: str
    role: str
    signed_at: datetime
    version: int


@dataclass(frozen=True)
class Workpaper:
    client_name: str
    engagement_name: str
    year_end: date
    overall: Decimal
    performance: Decimal
    trivial: Decimal
    benchmark: str
    definition: str
    definition_version: int
    gl_snapshot_id: str
    entry_count: int
    line_count: int
    total_debit: Decimal
    total_credit: Decimal
    criteria_version: int
    random_method: str
    seed: int
    threshold: Decimal
    parameters: Sequence[tuple[str, str]]
    by_rule: Sequence[tuple[str, int]]
    rows: Sequence[Row]
    conclusion: str
    not_performed: Sequence[str]
    stamp: datetime  # the newest record included: the document's fixed date
    title: str = "Journal entry testing"
    population_rows: Sequence[tuple[str, str]] = ()  # SPEC-038: replaces the journal entries block
    notes: Sequence[str] = ()


Sheet = Worksheet


def _text(sheet: Sheet, value: str, *, bold: bool = False) -> Cell:
    if len(value) > _MAX_CELL:
        value = value[:_MAX_CELL]
    cell = WriteOnlyCell(sheet, value=value)
    cell.data_type = "s"  # always a string: a leading "=" never becomes a formula
    if value.startswith(_FORMULA_LEADS):
        cell.quotePrefix = True
    if bold:
        cell.font = _BOLD
    return cell


def _amount(sheet: Sheet, value: Decimal) -> Cell:
    cell = WriteOnlyCell(sheet)
    cell.value = value
    cell.number_format = _AMOUNT
    return cell


def _number(sheet: Sheet, value: int) -> Cell:
    return WriteOnlyCell(sheet, value=value)


def _pair(sheet: Sheet, label: str, value: Cell | str) -> list[Cell]:
    return [
        _text(sheet, label, bold=True),
        value if isinstance(value, Cell) else _text(sheet, value),
    ]


def _summary(book: Workbook, w: Workpaper) -> None:
    s = book.create_sheet("Summary")
    rows: list[list[Cell]] = [
        [_text(s, w.title, bold=True)],
        _pair(s, "Client", w.client_name),
        _pair(s, "Engagement", w.engagement_name),
        _pair(s, "Year end", w.year_end.isoformat()),
        [],
        [_text(s, "Materiality (entered by the engagement team)", bold=True)],
        _pair(s, "Overall", _amount(s, w.overall)),
        _pair(s, "Performance", _amount(s, w.performance)),
        _pair(s, "Clearly trivial", _amount(s, w.trivial)),
        _pair(s, "Benchmark", w.benchmark),
        [],
        [_text(s, "Population", bold=True)],
        *(
            [_pair(s, k, v) for k, v in w.population_rows]
            if w.population_rows
            else [
                _pair(s, "Source", f"General ledger snapshot {w.gl_snapshot_id}"),
                _pair(
                    s,
                    "Full listing",
                    "The journal listing workbook held as evidence for this ledger",
                ),
                _pair(s, "Journal entries", _number(s, w.entry_count)),
                _pair(s, "Ledger lines", _number(s, w.line_count)),
                _pair(s, "Total debits", _amount(s, w.total_debit)),
                _pair(s, "Total credits", _amount(s, w.total_credit)),
                _pair(s, "Tie-out", "Agrees to the ledger snapshot (lines and debits)"),
            ]
        ),
        [],
        [_text(s, "Selection", bold=True)],
        _pair(s, "Procedure", f"{w.definition} v{w.definition_version}"),
        _pair(s, "Criteria version", _number(s, w.criteria_version)),
        _pair(s, "Selection threshold", _amount(s, w.threshold)),
        *[_pair(s, k, v) for k, v in w.parameters],
        _pair(s, "Random method", w.random_method),
        _pair(s, "Seed", _number(s, w.seed)),
        [],
        [_text(s, "Entries selected, by rule", bold=True)],
        *[_pair(s, RULE_LABELS.get(r, r), _number(s, n)) for r, n in w.by_rule],
        [],
        [_text(s, "Results", bold=True)],
    ]
    for key in ("supported", "exception", "not_tested", None):
        count = sum(1 for r in w.rows if r.result == key)
        rows.append(_pair(s, RESULT_LABELS[key], _number(s, count)))
    if w.not_performed:
        rows += [[], [_text(s, "Not performed", bold=True)]]
        rows += [[_text(s, NOT_PERFORMED.get(n, n))] for n in w.not_performed]
    if w.notes:
        rows += [[], [_text(s, "Notes", bold=True)]]
        rows += [[_text(s, n)] for n in w.notes]
    rows += [[], [_text(s, "Conclusion", bold=True)], [_text(s, w.conclusion)]]
    for row in rows:
        s.append(row)


_HEAD: Final = (
    "Date",
    "Type",
    "Reference",
    "Name",
    "Accounts",
    "Amount",
    "Why selected",
    "Team's reason",
    "Result",
    "Note",
    "Evidence",
    "Support requested",
)


_BANDED: Final = ("Band", "Bill recorded at year end", "Relates to the year")
RELATES: Final = {"yes": "Yes", "no": "No", "unclear": "Unclear"}


def _rows(book: Workbook, title: str, rows: Sequence[Row], banded: bool = False) -> None:
    s = book.create_sheet(title)
    s.append([_text(s, h, bold=True) for h in (*_HEAD, *(_BANDED if banded else ()))])
    for r in rows:
        extra = (
            [
                _number(s, r.band or 0),
                _text(s, r.recorded_bill or "None found"),
                _text(s, RELATES.get(r.relates_to_year or "", "")),
            ]
            if banded
            else []
        )
        s.append(
            [
                _text(s, r.posted_on.isoformat()),
                _text(s, r.kind),
                _text(s, r.reference),
                _text(s, r.counterparty),
                _text(s, ", ".join(r.accounts)),
                _amount(s, r.amount),
                _text(s, "; ".join(RULE_LABELS.get(x, x) for x in r.rules)),
                _text(s, r.reason or ""),
                _text(s, RESULT_LABELS.get(r.result, r.result or "")),
                _text(s, r.note or ""),
                _text(s, "; ".join(r.evidence)),
                _text(s, "Yes" if r.support_requested else ""),
                *extra,
            ]
        )


def _signoff(book: Workbook, signatures: Sequence[Signature]) -> None:
    s = book.create_sheet("Sign-off")
    s.append([_text(s, h, bold=True) for h in ("Signed as", "By", "Role", "When", "Version")])
    for g in signatures:
        s.append(
            [
                _text(s, "Prepared (ready)" if g.kind == "ready" else "Reviewed"),
                _text(s, g.name),
                _text(s, g.role),
                _text(s, g.signed_at.astimezone(UTC).strftime("%Y-%m-%d %H:%M UTC")),
                _number(s, g.version),
            ]
        )


def _normalised(raw: bytes, stamp: datetime) -> bytes:
    pinned = stamp.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as source,
        zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target,
    ):
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=_ZIP_TIME)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            entry.create_system = 3
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = _MODIFIED.sub(lambda m: m.group(1) + pinned + m.group(2), data)
            target.writestr(entry, data)
    return out.getvalue()


def render_workpaper(
    w: Workpaper, signatures: Sequence[Signature] = (), *, with_signoff: bool = True
) -> bytes:
    book = Workbook(write_only=True)
    book.properties.creator = "Abacus"
    book.properties.created = w.stamp.astimezone(UTC).replace(tzinfo=None)
    _summary(book, w)
    banded = bool(w.population_rows)
    _rows(book, "Selections", w.rows, banded)
    _rows(book, "Exceptions", [r for r in w.rows if r.result == "exception"], banded)
    if with_signoff:
        _signoff(book, signatures)
    buffer = io.BytesIO()
    book.save(buffer)
    return _normalised(buffer.getvalue(), w.stamp)
