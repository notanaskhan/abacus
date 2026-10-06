"""Citation verification by code (ADR-066). PROTECTED. TASK-011 design §5.

A cited cell must exist in the evidence spreadsheet; a quoted text must equal the cell's text; a
cited value must equal the cell's number. Anything else is unverified (AC-15), with a reason.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import cast

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet

from abacus.modules.agents.handoff import Citation, VerifiedCitation
from abacus.modules.evidence.api import (
    CODE_COLUMN,
    CREDIT_COLUMN,
    DEBIT_COLUMN,
    FIRST_LINE_ROW,
    NAME_COLUMN,
    TOTAL_LABEL,
)


@dataclass(frozen=True)
class SheetFacts:
    """What code computes from the rendered trial balance for screening (ADR-050): cell positions,
    totals summed here from the line rows, whether the sheet's own Total row agrees, and the
    account names (client content: untrusted). Never rows of amounts."""

    total_row: int
    last_line_row: int
    max_row: int
    line_count: int
    total_debit: Decimal
    total_credit: Decimal
    total_row_matches: bool
    account_names: tuple[str, ...]


def _sheet(content: bytes) -> Worksheet:
    workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    return cast(Worksheet, workbook.worksheets[0])


def _as_decimal(value: object) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation:
        return None


def _value(sheet: Worksheet, row: int, column: int) -> object:
    return cast(Cell, sheet.cell(row=row, column=column)).value


def _amount(sheet: Worksheet, row: int, column: int) -> Decimal:
    found = _as_decimal(_value(sheet, row, column))
    if found is None:
        raise SheetLayoutError(f"row {row} column {column} is not an amount")
    return found


class SheetLayoutError(ValueError):
    """The evidence isn't a trial balance in the platform's rendered layout: it isn't screened."""


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def facts(content: bytes) -> SheetFacts:
    """Strict: every row from FIRST_LINE_ROW to the Total row is a line with an account code and
    two amounts; exactly one row is the Total row (no code, TOTAL_LABEL); nothing after it holds an
    amount. Anything else is refused, so a client line can't pose as the Total row."""
    sheet = _sheet(content)
    candidates = [
        r
        for r in range(FIRST_LINE_ROW, sheet.max_row + 1)
        if _blank(_value(sheet, r, CODE_COLUMN)) and _value(sheet, r, NAME_COLUMN) == TOTAL_LABEL
    ]
    if len(candidates) != 1:
        raise SheetLayoutError("the sheet must have exactly one Total row")
    total_row = candidates[0]
    rows = range(FIRST_LINE_ROW, total_row)
    if any(_blank(_value(sheet, r, CODE_COLUMN)) for r in rows):
        raise SheetLayoutError("every line above the Total row has an account code")
    after = range(total_row + 1, sheet.max_row + 1)
    if any(
        not _blank(_value(sheet, r, column))
        for r in after
        for column in (DEBIT_COLUMN, CREDIT_COLUMN)
    ):
        raise SheetLayoutError("no amounts follow the Total row")
    debit = sum((_amount(sheet, r, DEBIT_COLUMN) for r in rows), Decimal(0))
    credit = sum((_amount(sheet, r, CREDIT_COLUMN) for r in rows), Decimal(0))
    return SheetFacts(
        total_row=total_row,
        last_line_row=total_row - 1,
        max_row=sheet.max_row,
        line_count=len(rows),
        total_debit=debit,
        total_credit=credit,
        total_row_matches=(
            _amount(sheet, total_row, DEBIT_COLUMN) == debit
            and _amount(sheet, total_row, CREDIT_COLUMN) == credit
        ),
        account_names=tuple(str(_value(sheet, r, NAME_COLUMN) or "") for r in rows),
    )


def verify(content: bytes, citations: list[Citation]) -> list[VerifiedCitation]:
    sheet = _sheet(content)
    verified: list[VerifiedCitation] = []
    for citation in citations:
        column, row = coordinate_from_string(citation.cell)
        reason: str | None = None
        if row > sheet.max_row or row < 1:
            reason = "cell_not_found"
        else:
            raw = cast(Cell, sheet[f"{column}{row}"]).value
            if raw is None:
                reason = "cell_not_found"
            elif citation.quote is not None and str(raw) != citation.quote:
                reason = "quote_mismatch"
            elif citation.value is not None and _as_decimal(raw) != citation.value:
                reason = "value_mismatch"
        verified.append(
            VerifiedCitation(
                cell=citation.cell,
                quote=citation.quote,
                value=None if citation.value is None else str(citation.value),
                verified=reason is None,
                reason=reason,
            )
        )
    return verified
