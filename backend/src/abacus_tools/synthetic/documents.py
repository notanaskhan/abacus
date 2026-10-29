"""Synthetic documents for upload processing (SPEC-029; TASK-054 D4). Synthetic data only.

- `bank_statement_pdf`: a text PDF in one of three layouts (labels as US banks print them),
  written by a tiny PDF writer (no dependency).
- `scanned_pdf`: a page with a drawing and no text layer, as a scanner makes.
- `encrypted_pdf`: a text PDF whose trailer asks for a password nobody has.
- `password_pdf`: a text PDF that really opens with a password (RC4, standard security handler
  revision 2, written here with no dependency): the Halvard consolidated 1099 the unlock opens
  (P-03-T19).
- `aging_schedule_xlsx`: an accounts receivable aging, optionally with formulas (no cached values)
  and a total row that doesn't add up.
- `pii_corpus`: a W-2, a K-1, a 1099-INT and a direct-deposit form with synthetic TINs, accounts
  and a date of birth where real forms print them (P-03-T13).
"""

from __future__ import annotations

import hashlib
import io
import struct
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Literal

from openpyxl import Workbook

Layout = Literal["a", "b", "c"]


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


# ISO 32000-1 §7.6.3.3: the padding a password is completed with.
_PAD = bytes.fromhex("28bf4e5e4e758a4164004e56fffa01082e2e00b6d0683e802f0ca9fe6453697a")
_PERMISSIONS = -4  # every permission (the reserved high bits set)
_FILE_ID = b"\x33" * 16


def _rc4(key: bytes, data: bytes) -> bytes:
    box = list(range(256))
    j = 0
    for i in range(256):
        j = (j + box[i] + key[i % len(key)]) % 256
        box[i], box[j] = box[j], box[i]
    out = bytearray()
    i = j = 0
    for byte in data:
        i = (i + 1) % 256
        j = (j + box[i]) % 256
        box[i], box[j] = box[j], box[i]
        out.append(byte ^ box[(box[i] + box[j]) % 256])
    return bytes(out)


def _padded(password: str) -> bytes:
    return (password.encode("latin-1") + _PAD)[:32]


def _standard_security(password: str) -> tuple[bytes, bytes, bytes]:
    """Revision 2 (40-bit RC4), owner password = user password: the O and U entries and the file
    key (ISO 32000-1 algorithms 2, 3 and 4)."""
    owner = _rc4(
        hashlib.md5(_padded(password), usedforsecurity=False).digest()[:5], _padded(password)
    )
    key = hashlib.md5(
        _padded(password) + owner + struct.pack("<i", _PERMISSIONS) + _FILE_ID,
        usedforsecurity=False,
    ).digest()[:5]
    return owner, _rc4(key, _PAD), key


def _object_key(key: bytes, number: int) -> bytes:
    return hashlib.md5(
        key + struct.pack("<i", number)[:3] + b"\x00\x00", usedforsecurity=False
    ).digest()[:10]


def _pdf(pages: Sequence[bytes], *, encrypt: bool = False, password: str | None = None) -> bytes:
    """A minimal PDF: one Helvetica font, one content stream per page. `password` encrypts the
    content streams so the file opens with that password only (`password_pdf`)."""
    objects: list[bytes] = []
    count = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(count))
    font = 3 + 2 * count
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode())
    security = _standard_security(password) if password is not None else None
    for i, stream in enumerate(pages):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
        )
        if security is not None:
            stream = _rc4(_object_key(security[2], 4 + 2 * i), stream)
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    trailer_extra = b""
    if security is not None:
        owner, user, _key = security
        objects.append(
            b"<< /Filter /Standard /V 1 /R 2 /P %d /O <%s> /U <%s> >>"
            % (_PERMISSIONS, owner.hex().encode(), user.hex().encode())
        )
        trailer_extra = b" /Encrypt %d 0 R /ID [<%s> <%s>]" % (
            len(objects),
            _FILE_ID.hex().encode(),
            _FILE_ID.hex().encode(),
        )
    elif encrypt:
        objects.append(
            b"<< /Filter /Standard /V 1 /R 2 /P -4 /O <"
            + b"11" * 32
            + b"> /U <"
            + b"22" * 32
            + b"> >>"
        )
        trailer_extra = b" /Encrypt %d 0 R /ID [<%s> <%s>]" % (
            len(objects),
            b"33" * 16,
            b"33" * 16,
        )
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R%s >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, trailer_extra, xref)
    )
    return out.getvalue()


def _text_page(lines: Sequence[str]) -> bytes:
    body = "".join(f"({_escape(line)}) Tj T* " for line in lines)
    return f"BT /F1 10 Tf 14 TL 50 750 Td {body}ET".encode("latin-1")


def positioned_pdf(
    pages: Sequence[Sequence[tuple[float, float, str]]], *, password: str | None = None
) -> bytes:
    """A text PDF whose every run sits at a fixed place: `(x, baseline, text)` in page points,
    bottom-left origin, Helvetica 10 on a US Letter page. Scenario documents use it so that each
    cited value has a fixed page region (F-04-T25). Text is ASCII."""
    streams = [
        "".join(
            f"BT /F1 10 Tf 1 0 0 1 {x:.2f} {y:.2f} Tm ({_escape(text)}) Tj ET "
            for x, y, text in runs
        ).encode("ascii")
        for runs in pages
    ]
    return _pdf(streams, password=password)


def _money(value: Decimal) -> str:
    return f"${value:,.2f}" if value >= 0 else f"(${-value:,.2f})"


def bank_statement_pdf(
    *,
    layout: Layout = "a",
    holder: str = "Example Company Inc",
    last4: str = "1234",
    start: date = date(2026, 12, 1),
    end: date = date(2026, 12, 31),
    opening: Decimal = Decimal("10000.00"),
    credits: Decimal = Decimal("5250.00"),
    debits: Decimal = Decimal("3125.50"),
    closing: Decimal | None = None,
    encrypt: bool = False,
) -> bytes:
    """`closing` defaults to opening + credits - debits; pass another to make it not add up."""
    ending = closing if closing is not None else opening + credits - debits
    if layout == "a":
        lines = [
            "First Example Bank",
            f"Account holder: {holder}",
            f"Account number: ****{last4}",
            f"Statement period: {start:%B} {start.day}, {start.year} to "
            f"{end:%B} {end.day}, {end.year}",
            f"Beginning balance {_money(opening)}",
            f"Total deposits and other credits {_money(credits)}",
            f"Total withdrawals and other debits {_money(debits)}",
            f"Ending balance {_money(ending)}",
            "Date Description Amount",
            f"{end:%m/%d} Deposit {_money(credits)}",
        ]
    elif layout == "b":
        lines = [
            "Second Example Savings",
            f"Statement for: {holder}",
            f"Account no. XXXX-XXXX-{last4}",
            f"For the period {start:%m/%d/%Y} through {end:%m/%d/%Y}",
            f"Opening balance {opening:,.2f}",
            f"Total credits {credits:,.2f}",
            f"Total debits {debits:,.2f}",
            f"Closing balance {ending:,.2f}",
        ]
    else:
        lines = [
            "Third Example Credit Union",
            f"Prepared for {holder}",
            f"Account # {last4}",
            f"Statement date: {end:%b} {end.day}, {end.year}",
            f"Previous balance {_money(opening)}",
            f"Deposits and credits {_money(credits)}",
            f"Checks and other debits {_money(debits)}",
            f"New balance {_money(ending)}",
        ]
    return _pdf([_text_page(lines)], encrypt=encrypt)


def scanned_pdf(pages: int = 1) -> bytes:
    """Pages with a drawn box and no text, as a scanner (without OCR) makes them."""
    return _pdf([b"0.5 g 50 50 500 700 re f" for _ in range(pages)])


def encrypted_pdf() -> bytes:
    return bank_statement_pdf(encrypt=True)


AGING_HEADER = ("Customer", "Current", "1-30", "31-60", "61-90", "Over 90", "Total")


def aging_schedule_xlsx(
    rows: Sequence[tuple[str, Sequence[Decimal]]],
    *,
    company: str = "Example Company Inc",
    as_of: date = date(2026, 12, 31),
    formulas: bool = False,
    total_off_by: Decimal = Decimal(0),
    payables: bool = False,
) -> bytes:
    """Each row: a customer (or, with `payables`, a vendor) and its five bucket amounts. With
    `formulas`, totals are written as formulas with no cached value (as a script, not Excel,
    would save them)."""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("openpyxl made a workbook without a sheet")
    sheet.title = "AP Aging" if payables else "AR Aging"
    side = "Accounts payable" if payables else "Accounts receivable"
    sheet.append([f"{company} - {side} aging"])
    sheet.append([f"As of {as_of:%B} {as_of.day}, {as_of.year}"])
    sheet.append([])
    sheet.append(["Vendor", *AGING_HEADER[1:]] if payables else list(AGING_HEADER))
    first = sheet.max_row + 1
    for name, buckets in rows:
        n = sheet.max_row + 1
        total: object = f"=SUM(B{n}:F{n})" if formulas else float(sum(buckets, Decimal(0)))
        sheet.append([name, *(float(b) for b in buckets), total])
    last = sheet.max_row
    sums = [sum((r[1][i] for r in rows), Decimal(0)) for i in range(5)]
    grand = sum(sums, Decimal(0)) + total_off_by
    if formulas:
        cols = "BCDEFG"
        sheet.append(["Total", *(f"=SUM({c}{first}:{c}{last})" for c in cols)])
    else:
        sheet.append(["Total", *(float(s) for s in sums), float(grand)])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


# --- Upload reviewer cases (SPEC-030; TASK-058) -------------------------------------------------


def detailed_statement_pdf(
    *,
    holder: str = "Example Company Inc",
    last4: str = "1234",
    start: date = date(2026, 12, 1),
    end: date = date(2026, 12, 31),
    opening: Decimal = Decimal("10000.00"),
    header_lines: tuple[str, ...] = (),
    footer_lines: tuple[str, ...] = (),
    tamper_line: int | None = None,
    second_page_closing: Decimal | None = None,
    outside_period_line: bool = False,
    duplicate_check: bool = False,
) -> bytes:
    """A statement with transaction lines and running balances. `header_lines` and
    `footer_lines` add text (a watermark, a signature line, an instruction); `outside_period_line`
    dates the third line in the month before the period and `duplicate_check` reuses check 1041
    for the last payment (both editing signs that need no arithmetic); `tamper_line` makes
    that line's running balance not follow from the one before (opening, credits, debits and
    closing still agree, so code's totals check passes); `second_page_closing` adds a summary
    page stating a different closing balance."""
    moves = (
        ("Customer deposit", Decimal("2500.00")),
        ("Check 1041", Decimal("-820.00")),
        ("ACH payroll", Decimal("-1405.50")),
        ("Customer deposit", Decimal("2750.00")),
        ("Check 1041" if duplicate_check else "Wire to supplier", Decimal("-900.00")),
    )
    credits = sum((m for _, m in moves if m > 0), Decimal(0))
    debits = -sum((m for _, m in moves if m < 0), Decimal(0))
    closing = opening + credits - debits
    lines = [
        *header_lines,
        "First Example Bank",
        f"Account holder: {holder}",
        f"Account number: ****{last4}",
        f"Statement period: {start:%B} {start.day}, {start.year} to "
        f"{end:%B} {end.day}, {end.year}",
        f"Beginning balance {_money(opening)}",
        f"Total deposits and other credits {_money(credits)}",
        f"Total withdrawals and other debits {_money(debits)}",
        f"Ending balance {_money(closing)}",
        "Date Description Amount Balance",
    ]
    balance = opening
    for n, (what, amount) in enumerate(moves, start=1):
        balance += amount
        shown = balance + Decimal("500.00") if tamper_line == n else balance
        day = min(start.day + 4 * n, 28)
        month = (end.month - 2) % 12 + 1 if outside_period_line and n == 3 else end.month
        lines.append(f"{month:02d}/{day:02d} {what} {_money(amount)} {_money(shown)}")
    lines.extend(footer_lines)
    pages = [_text_page(lines)]
    if second_page_closing is not None:
        pages.append(
            _text_page(
                [
                    "Account summary",
                    f"Account number: ****{last4}",
                    f"Closing balance as of {end:%B} {end.day}, {end.year} "
                    f"{_money(second_page_closing)}",
                ]
            )
        )
    return _pdf(pages)


def invoice_pdf(
    *,
    vendor: str = "Northwind Supplies LLC",
    total: Decimal = Decimal("4210.00"),
    dated: date = date(2026, 12, 15),
) -> bytes:
    """An invoice: the wrong kind of document for a bank statement request."""
    return _pdf(
        [
            _text_page(
                [
                    "INVOICE",
                    vendor,
                    "Invoice number: INV-20931",
                    f"Invoice date: {dated:%B} {dated.day}, {dated.year}",
                    "Bill to: Example Company Inc",
                    "Description Qty Unit price Amount",
                    f"Office chairs 10 $421.00 {_money(total)}",
                    f"Amount due {_money(total)}",
                    "Payment terms: Net 30",
                ]
            )
        ]
    )


def vendor_statement_pdf(
    *,
    vendor: str = "Granite Rock Logistics",
    customer: str = "Example Company Inc",
    dated: date = date(2026, 8, 31),
    footer_lines: tuple[str, ...] = (),
) -> bytes:
    """A vendor's monthly statement of account (P-06-T21; J22): it lists invoices and a balance
    due, so it is the wrong document for a request for one invoice. `footer_lines` adds text (an
    instruction carried in the file)."""
    month = dated.replace(day=1)
    return _pdf(
        [
            _text_page(
                [
                    "STATEMENT OF ACCOUNT",
                    vendor,
                    f"Customer: {customer}",
                    f"Statement date: {dated:%B} {dated.day}, {dated.year}",
                    f"Activity from {month:%B} {month.day} to "
                    f"{dated:%B} {dated.day}, {dated.year}",
                    "Date Reference Charges Payments Balance",
                    f"{month:%m}/04 Invoice GRL-5512 $1,840.00 $1,840.00",
                    f"{month:%m}/11 Payment received $1,840.00 $0.00",
                    f"{month:%m}/18 Invoice GRL-5547 $2,215.00 $2,215.00",
                    f"{month:%m}/25 Invoice GRL-5590 $960.00 $3,175.00",
                    "Balance due $3,175.00",
                    "This statement lists invoices issued; it is not an invoice.",
                    *footer_lines,
                ]
            )
        ]
    )


def tax_form_pdf(
    *,
    form: Literal["W-2", "1099-NEC"] = "W-2",
    tax_year: int = 2025,
    payer: str = "Example Company Inc",
    recipient: str = "Sam Example",
    amount: Decimal = Decimal("48250.00"),
) -> bytes:
    """A wage or nonemployee-compensation statement (P-06-T21; J06). Identification numbers are
    masked letters, never digits shaped like a taxpayer ID."""
    if form == "W-2":
        lines = [
            f"Form W-2 Wage and Tax Statement {tax_year}",
            f"Employer: {payer}",
            "Employer identification number: XX-XXXXXXX",
            f"Employee: {recipient}",
            "Employee's social security number: XXX-XX-XXXX",
            f"Box 1 Wages, tips, other compensation {_money(amount)}",
            f"Box 2 Federal income tax withheld {_money(amount * Decimal('0.12'))}",
        ]
    else:
        lines = [
            f"Form 1099-NEC Nonemployee Compensation {tax_year}",
            f"Payer: {payer}",
            "Payer's TIN: XX-XXXXXXX",
            f"Recipient: {recipient}",
            "Recipient's TIN: XXX-XX-XXXX",
            f"Box 1 Nonemployee compensation {_money(amount)}",
            "Box 4 Federal income tax withheld $0.00",
        ]
    return _pdf([_text_page(lines)])


def form_1095c_pdf(
    *,
    tax_year: int = 2025,
    employer: str = "Example Company Inc",
    employee: str = "Sam Example",
) -> bytes:
    """An employer-provided health insurance offer and coverage statement (R-01-T11; J06's
    1095-C). Identification numbers are masked letters; no coverage figures are printed."""
    return _pdf(
        [
            _text_page(
                [
                    f"Form 1095-C {tax_year}",
                    "Employer-Provided Health Insurance Offer and Coverage",
                    f"Employer: {employer}",
                    "Employer identification number: XX-XXXXXXX",
                    f"Employee: {employee}",
                    "Employee's social security number: XXX-XX-XXXX",
                ]
            )
        ]
    )


def estimated_payment_pdf(
    *,
    taxpayer: str = "Sam Example",
    tax_year: int = 2026,
    quarter: int = 3,
    paid_on: date = date(2026, 9, 15),
    amount: Decimal = Decimal("3200.00"),
) -> bytes:
    """A confirmation of an estimated tax payment (P-06-T21): the right document for its
    request, with nothing to flag."""
    return _pdf(
        [
            _text_page(
                [
                    "Payment confirmation",
                    f"Taxpayer: {taxpayer}",
                    f"Payment type: Form 1040-ES estimated tax, {tax_year} quarter {quarter}",
                    f"Payment date: {paid_on:%B} {paid_on.day}, {paid_on.year}",
                    f"Amount paid {_money(amount)}",
                    "Confirmation number: EXAMPLE-CONF-0001",
                    "Status: Accepted for processing",
                ]
            )
        ]
    )


def letter_pdf(dated: date = date(2027, 1, 10)) -> bytes:
    """A cover letter: the wrong kind of document for a statement or an aging."""
    return _pdf(
        [
            _text_page(
                [
                    "Example Company Inc",
                    f"{dated:%B} {dated.day}, {dated.year}",
                    "Dear audit team,",
                    "Please find our documents for the year-end audit attached separately.",
                    "We will send the bank statements once we receive them from the bank.",
                    "Kind regards,",
                    "Rachel, Controller",
                ]
            )
        ]
    )


def receipt_pdf(dated: date = date(2026, 12, 18)) -> bytes:
    """A store receipt: the wrong kind of document for a bank statement request."""
    return _pdf(
        [
            _text_page(
                [
                    "OFFICE DEPOT STORE 2231",
                    f"{dated:%m/%d/%Y} 14:32",
                    "Printer paper 5 x $8.99 $44.95",
                    "Sales tax $3.71",
                    "TOTAL $48.66",
                    "VISA ****9921 APPROVED",
                    "Thank you for shopping with us",
                ]
            )
        ]
    )


# --- Content matching cases (SPEC-033; TASK-065 D2) ---------------------------------------------


def _long_date(day: date) -> str:
    return f"{day:%B} {day.day}, {day.year}"


def loan_agreement_pdf(
    *,
    borrower: str = "Example Company Inc",
    dated: date = date(2026, 3, 1),
    title: str = "TERM LOAN AGREEMENT",
    notes: tuple[str, ...] = (),
) -> bytes:
    return _pdf(
        [
            _text_page(
                [
                    *notes,
                    title,
                    f"This Term Loan Agreement is dated {_long_date(dated)}",
                    f"between {borrower} (the Borrower) and First Harbor Bank, N.A. (the Lender).",
                    "1. The Loan. The Lender agrees to lend the Borrower $2,500,000.00.",
                    "2. Interest. Interest accrues at SOFR plus 2.25% per annum.",
                    "3. Repayment. Quarterly installments of principal beginning June 30.",
                    "4. Covenants. The Borrower shall maintain a debt service coverage ratio",
                    "of not less than 1.25 to 1.00, tested at each fiscal year end.",
                    "IN WITNESS WHEREOF the parties have executed this Agreement.",
                ]
            )
        ]
    )


def lease_pdf(
    *,
    tenant: str = "Example Company Inc",
    dated: date = date(2025, 7, 1),
    notes: tuple[str, ...] = (),
) -> bytes:
    return _pdf(
        [
            _text_page(
                [
                    *notes,
                    "COMMERCIAL LEASE AGREEMENT",
                    f"Commencement date: {_long_date(dated)}",
                    "Landlord: Riverside Properties LLC",
                    f"Tenant: {tenant}",
                    "Premises: Suite 400, 1200 Market Street",
                    "Term: Sixty (60) months from the commencement date",
                    "Base rent: $18,500.00 per month, increasing 3% annually",
                    "Security deposit: $37,000.00",
                ]
            )
        ]
    )


def insurance_certificate_pdf(
    *, insured: str = "Example Company Inc", dated: date = date(2026, 1, 1)
) -> bytes:
    return _pdf(
        [
            _text_page(
                [
                    "CERTIFICATE OF LIABILITY INSURANCE",
                    f"Date issued: {_long_date(dated)}",
                    f"Insured: {insured}",
                    "Insurer A: Granite Mutual Insurance Company",
                    "Commercial general liability - each occurrence $1,000,000",
                    "Policy number: CGL-8830-2026",
                    f"Policy period: {_long_date(dated)} to "
                    f"{_long_date(dated.replace(year=dated.year + 1))}",
                ]
            )
        ]
    )


def board_minutes_pdf(
    *,
    company: str = "Example Company Inc",
    dated: date = date(2026, 11, 18),
    notes: tuple[str, ...] = (),
) -> bytes:
    return _pdf(
        [
            _text_page(
                [
                    f"MINUTES OF A MEETING OF THE BOARD OF DIRECTORS OF {company.upper()}",
                    f"Held on {_long_date(dated)} at the company's offices",
                    "Present: J. Alvarez (Chair), M. Chen, R. Okafor",
                    "1. The minutes of the previous meeting were approved.",
                    "2. The board approved the fiscal year budget.",
                    "3. The board authorised the CFO to renew the revolving credit facility.",
                    "There being no further business, the meeting was adjourned.",
                    *notes,
                ]
            )
        ]
    )


def flyer_pdf() -> bytes:
    """A marketing flyer: fits no audit request."""
    return _pdf(
        [
            _text_page(
                [
                    "SPRING OPEN HOUSE",
                    "Join us for coffee and pastries!",
                    "Saturday 10am - 2pm at our new showroom",
                    "Free parking - bring a friend",
                    "RSVP to events@example.test",
                ]
            )
        ]
    )


def _sheet(title: str, rows: Sequence[Sequence[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("openpyxl made a workbook without a sheet")
    sheet.title = title
    for row in rows:
        sheet.append(list(row))
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def fixed_asset_schedule_xlsx(*, as_of: date = date(2026, 12, 31)) -> bytes:
    return _sheet(
        "Fixed assets",
        [
            [f"Fixed asset rollforward for the year ended {_long_date(as_of)}"],
            [],
            ["Asset class", "Opening cost", "Additions", "Disposals", "Closing cost"],
            ["Machinery", 410000.0, 52000.0, 0.0, 462000.0],
            ["Vehicles", 96000.0, 0.0, 18000.0, 78000.0],
            ["Furniture", 34000.0, 4500.0, 0.0, 38500.0],
        ],
    )


def payroll_register_xlsx(*, period_end: date = date(2026, 12, 31)) -> bytes:
    return _sheet(
        "Payroll register",
        [
            [f"Payroll register - pay period ending {_long_date(period_end)}"],
            [],
            ["Employee", "Gross pay", "Federal tax", "FICA", "Net pay"],
            ["A. Rivera", 5200.0, 780.0, 397.8, 4022.2],
            ["B. Nguyen", 4800.0, 690.0, 367.2, 3742.8],
            ["C. Patel", 6100.0, 950.0, 466.65, 4683.35],
        ],
    )


CardLine = tuple[date, str, Decimal]  # posted date, description, amount (+ charge, - credit)


def card_statement_pdf(
    *,
    issuer: str = "Example Card Services",
    last4: str = "4471",
    start: date = date(2026, 11, 22),
    end: date = date(2026, 12, 21),
    previous: Decimal = Decimal("1250.00"),
    lines: Sequence[CardLine] = (
        (date(2026, 11, 24), "PAYMENT THANK YOU", Decimal("-1250.00")),
        (date(2026, 11, 28), "OFFICE SUPPLY CO #221", Decimal("184.37")),
        (date(2026, 12, 2), "CLOUD HOSTING INC", Decimal("612.00")),
        (date(2026, 12, 9), "RETURN OFFICE SUPPLY CO", Decimal("-24.99")),
        (date(2026, 12, 15), "AIRLINE 0167 SEATTLE", Decimal("438.20")),
    ),
    new_balance: Decimal | None = None,
    pages: int = 1,
) -> bytes:
    """A credit card e-statement as US issuers print it (summary, then payments and credits,
    then purchases). `new_balance` defaults to what the lines add up to; pass another to make
    it not reconcile. With `pages=2` the purchases continue on a second page."""
    credits = [line for line in lines if line[2] < 0]
    charges = [line for line in lines if line[2] >= 0]
    total = previous + sum((line[2] for line in lines), Decimal(0))
    ending = new_balance if new_balance is not None else total
    head = [
        issuer,
        f"Card ending in {last4}",
        f"Opening/Closing Date {start:%m/%d/%y} - {end:%m/%d/%y}",
        f"Previous Balance {_money(previous)}",
        f"Payments, Credits -{_money(-sum((c[2] for c in credits), Decimal(0)))}",
        f"Purchases +{_money(sum((c[2] for c in charges), Decimal(0)))}",
        f"New Balance {_money(ending)}",
        "Minimum Payment Due $35.00",
        "PAYMENTS AND OTHER CREDITS",
        *(f"{d:%m/%d} {text} -{abs(amount):,.2f}" for d, text, amount in credits),
        "PURCHASES",
    ]
    rows = [f"{d:%m/%d} {text} {amount:,.2f}" for d, text, amount in charges]
    if pages == 1:
        return _pdf([_text_page([*head, *rows])])
    half = len(rows) // 2
    return _pdf([_text_page([*head, *rows[:half]]), _text_page(["PURCHASES", *rows[half:]])])


# --- PII corpus (P-03-T13; P-03 §4.4) ------------------------------------------------------------
# Synthetic only: SSNs in the 000 and 9xx areas, EINs with the never-assigned 00 prefix, routing
# numbers that fail the ABA checksum, `SYN-` accounts. Each value sits under or beside the label
# a real form prints, as the detector must find it (recall 1.0 on this set).


@dataclass(frozen=True)
class PiiDocument:
    """A synthetic form and what the detector must find on it: (page, kind, last four)."""

    name: str
    content: bytes
    expected: tuple[tuple[int, str, str | None], ...]


def w2_pdf() -> bytes:
    """A W-2: box a holds the employee's SSN 000-00-6081, box b the employer's EIN."""
    return positioned_pdf(
        [
            [
                (36.0, 740.0, "a Employee's social security number"),
                (36.0, 726.0, "000-00-6081"),
                (300.0, 740.0, "b Employer identification number (EIN)"),
                (300.0, 726.0, "00-1230287"),
                (36.0, 690.0, "c Employer's name, address, and ZIP code"),
                (36.0, 676.0, "Brightwater Hale (Synthetic)"),
                (300.0, 690.0, "1 Wages, tips, other compensation"),
                (300.0, 676.0, "84,250.00"),
            ]
        ]
    )


def k1_pdf() -> bytes:
    """A Schedule K-1 (Form 1065): the partnership's EIN in item A, the partner's SSN in E."""
    return positioned_pdf(
        [
            [
                (36.0, 740.0, "A Partnership's employer identification number"),
                (36.0, 726.0, "00-4567890"),
                (36.0, 690.0, "E Partner's identifying number"),
                (36.0, 676.0, "000-00-4412"),
                (300.0, 690.0, "1 Ordinary business income (loss)"),
                (300.0, 676.0, "12,480"),
            ]
        ]
    )


def form_1099_pdf() -> bytes:
    """A 1099-INT: the payer's TIN, the recipient's TIN (900-12-3456) and the account number."""
    return positioned_pdf(
        [
            [
                (36.0, 740.0, "PAYER'S TIN"),
                (200.0, 740.0, "RECIPIENT'S TIN"),
                (36.0, 726.0, "00-7654321"),
                (200.0, 726.0, "900-12-3456"),
                (36.0, 690.0, "Account number (see instructions)"),
                (36.0, 676.0, "SYN-4410-2291"),
                (300.0, 690.0, "1 Interest income"),
                (300.0, 676.0, "2,310.00"),
            ]
        ]
    )


def direct_deposit_pdf() -> bytes:
    """A direct-deposit form: routing and account numbers and a date of birth, with an
    unlabelled number that fails every rule (never masked) and an invoice number."""
    return positioned_pdf(
        [
            [
                (36.0, 740.0, "Routing number"),
                (36.0, 726.0, "000000001"),
                (220.0, 740.0, "Account number"),
                (220.0, 726.0, "SYN-0001-7731"),
                (36.0, 690.0, "Date of birth"),
                (36.0, 676.0, "01/01/1980"),
                (36.0, 640.0, "Reference 000-00-1111 Invoice 12345 Total 4,210.55"),
            ]
        ]
    )


def password_pdf(password: str) -> bytes:
    """The Halvard 2025 consolidated 1099 (SHL-08d), opening only with `password`: the
    recipient's TIN 900-12-3456 (a 9xx fake) under its label, and the interest it reports."""
    return positioned_pdf(
        [
            [
                (36.0, 760.0, "Halvard Securities (Synthetic) 2025 Consolidated Form 1099"),
                (36.0, 740.0, "PAYER'S TIN"),
                (200.0, 740.0, "RECIPIENT'S TIN"),
                (36.0, 726.0, "00-7654321"),
                (200.0, 726.0, "900-12-3456"),
                (300.0, 690.0, "1 Interest income"),
                (300.0, 676.0, "2,310.00"),
            ]
        ],
        password=password,
    )


def pii_corpus() -> tuple[PiiDocument, ...]:
    """The seeded PII corpus (P-03-T13 acceptance: recall 1.0)."""
    return (
        PiiDocument("W-2", w2_pdf(), ((1, "ssn", "6081"), (1, "ein", "0287"))),
        PiiDocument("K-1", k1_pdf(), ((1, "ein", "7890"), (1, "ssn", "4412"))),
        PiiDocument(
            "1099-INT",
            form_1099_pdf(),
            ((1, "ein", "4321"), (1, "ssn", "3456"), (1, "bank_account", "2291")),
        ),
        PiiDocument(
            "Direct deposit",
            direct_deposit_pdf(),
            ((1, "routing_number", "0001"), (1, "bank_account", "7731"), (1, "dob", None)),
        ),
    )


# --- Intake matching cases (N-01-T21; matcher v3) ----------------------------------------------


def w2c_pdf(
    *,
    tax_year: int = 2025,
    employer: str = "Example Company Inc",
    employee: str = "Sam Example",
    previous: Decimal = Decimal("48250.00"),
    corrected: Decimal = Decimal("49100.00"),
) -> bytes:
    """A corrected wage statement (Form W-2c): it answers the W-2 request, as a correction.
    Identification numbers are masked letters."""
    return _pdf(
        [
            _text_page(
                [
                    f"Form W-2c Corrected Wage and Tax Statement {tax_year}",
                    f"Employer: {employer}",
                    "Employer identification number: XX-XXXXXXX",
                    f"Employee: {employee}",
                    "Employee's social security number: XXX-XX-XXXX",
                    "Box 1 Wages, tips, other compensation",
                    f"Previously reported {_money(previous)}",
                    f"Correct information {_money(corrected)}",
                ]
            )
        ]
    )


def form_1099_div_pdf(
    *,
    tax_year: int = 2025,
    payer: str = "Example Securities LLC",
    recipient: str = "Sam Example",
    corrected: bool = False,
    ordinary: Decimal = Decimal("1840.22"),
    qualified: Decimal = Decimal("1512.90"),
    title: str | None = None,
    notes: Sequence[str] = (),
) -> bytes:
    """A dividends statement (Form 1099-DIV), optionally marked CORRECTED; `title` replaces the
    form's heading (a unicode deception case) and `notes` add lines (an instruction in the
    file)."""
    head = title or f"Form 1099-DIV Dividends and Distributions {tax_year}"
    return _pdf(
        [
            _text_page(
                [
                    *(["CORRECTED"] if corrected else []),
                    head,
                    f"Payer: {payer}",
                    "Payer's TIN: XX-XXXXXXX",
                    f"Recipient: {recipient}",
                    "Recipient's TIN: XXX-XX-XXXX",
                    f"Box 1a Total ordinary dividends {_money(ordinary)}",
                    f"Box 1b Qualified dividends {_money(qualified)}",
                    *notes,
                ]
            )
        ]
    )


def census_xlsx(*, plan: str = "Example Company 401(k) Plan", plan_year: int = 2025) -> bytes:
    """A plan's employee census: one row per participant, synthetic names, masked identifiers."""
    return _sheet(
        "Census",
        [
            [f"{plan} employee census, plan year {plan_year}"],
            ["Employee", "Identifier", "Hire date", "Status", "Compensation", "Deferrals"],
            ["Avery Example", "XXX-XX-XXXX", "2019-04-01", "Active", 72500, 4350],
            ["Jordan Example", "XXX-XX-XXXX", "2021-08-16", "Active", 58200, 2910],
            ["Riley Example", "XXX-XX-XXXX", "2024-02-05", "Terminated", 21400, 0],
        ],
    )
