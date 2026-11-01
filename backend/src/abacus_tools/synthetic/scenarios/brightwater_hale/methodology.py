"""The `methodology` slice (K-01-T10; K-01 §5): Brightwater Hale's methodology content, loaded the
way a firm loads its own: each version's workbook imported through `methodology.api`, landing in
Needs review, then approved by the QM owner with its effective date, which supersedes the version
before it (K-01-T06). Synthetic, Brightwater Hale only; nothing here is platform content.

- **BH Audit Methodology · EBP**: v1 (effective Jan 1, 2024) and v2 (Jan 1, 2025), 64 steps each,
  then v3 (Jan 1, 2026): seven programs of 16 + 8 + 6 + 8 + 10 + 7 + 12 = 67 steps, 21 of them
  mandatory (KNW-02). v3 against v2: four steps added (BH-EBP-110, 140, 410, 510), nine changed
  (BH-EBP-110 steps 7, 9 and 10, KNW-03 e, and one step in each of the six other programs,
  `EBP_V3_CHANGED`) and one removed (BH-EBP-310): KNW-02 c's 4 · 9 · 1 · 54 (K-01-T18).
  Each EBP version carries the 24-request PBC template.
- **BH Audit Methodology · FS** v1 to v5 (effective each Jan 1, 2022 to 2026) and **· SSARS** v1
  and v2 (2025, 2026): their shapes are the seed's (K-01 names only the versions).
- **The flags** the import path sits behind, `methodology.programs` and `methodology.lifecycle`,
  on first.

The firm exists from Sep 1, 2026 on the timeline, so every import and approval happens at its
setup; the effective dates are data, and `approved_at` reads the load instant. KNW-02's canon
("Approved · Jan 12, 2026 2:31 PM") stays in the fixtures (DEC-652, as DEC-363 did for FIL-02).
Opt-in (`--slice methodology`, which brings `people` for Grace's and Margaret's roles).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, timedelta
from typing import Final

from openpyxl import Workbook

from abacus_tools.synthetic.documents import positioned_pdf
from abacus_tools.synthetic.scenarios.brightwater_hale.scenario import SETUP_AT
from abacus_tools.synthetic.scenarios.brightwater_hale.slices import register_slice
from abacus_tools.synthetic.scenarios.brightwater_hale.timeline import Event

SLICE: Final = "methodology"
NEEDS: Final = ("people",)
IMPORTER: Final = "GH"  # Grace Holloway, firm admin (`methodology.manage`)
APPROVER: Final = "MO"  # Margaret Okonkwo, Managing Partner and QM owner (`methodology.approve`)
FLAGS: Final = ("methodology.programs", "methodology.lifecycle")
TEMPLATE: Final = "BH Audit Methodology"
EBP_PROGRAMS: Final = (110, 140, 210, 310, 410, 510, 610)
EBP_V3: Final = (16, 8, 6, 8, 10, 7, 12)  # 67
EBP_V2: Final = (15, 7, 6, 9, 9, 6, 12)  # 64: v3 less 110, 140, 410, 510's new steps, plus 310's
EBP_MANDATORY: Final = (5, 3, 2, 3, 3, 2, 3)  # 21: each program's first steps
# K-01-T18 (KNW-02 c, AC6.1): the steps v3 rewords, by program and number: the seven KNW-02c draws
# (110 steps 7, 9, 10; 140·3; 210·2; 310·1; 410·6) and two below its fold.
EBP_V3_CHANGED: Final = frozenset(
    {(110, 7), (110, 9), (110, 10), (140, 3), (210, 2), (310, 1), (410, 6), (510, 2), (610, 8)}
)
PBC_REQUESTS: Final = 24
LEVELS: Final = ("Foundation", "Planning", "Testing", "Completion")


@dataclass(frozen=True)
class VersionSpec:
    family: str  # ebp | fs | ssars
    number: int
    shape: tuple[int, ...]
    mandatory: tuple[int, ...]
    effective_from: date

    @property
    def name(self) -> str:
        return f"{TEMPLATE} · {self.family.upper()}"

    @property
    def engagement_type(self) -> str:
        return {"ebp": "ebp_audit", "fs": "fs_audit", "ssars": "review"}[self.family]

    @property
    def codes(self) -> tuple[str, ...]:
        if self.family == "ebp":
            return tuple(f"BH-EBP-{p}" for p in EBP_PROGRAMS[: len(self.shape)])
        prefix = {"fs": "BH-FS", "ssars": "BH-SR"}[self.family]
        return tuple(f"{prefix}-{100 + 10 * n}" for n in range(1, len(self.shape) + 1))


VERSIONS: Final[tuple[VersionSpec, ...]] = (
    VersionSpec("ebp", 1, EBP_V2, EBP_MANDATORY, date(2024, 1, 1)),
    VersionSpec("ebp", 2, EBP_V2, EBP_MANDATORY, date(2025, 1, 1)),
    VersionSpec("ebp", 3, EBP_V3, EBP_MANDATORY, date(2026, 1, 1)),
    *(
        VersionSpec("fs", n, (12, 10, 8, 9), (3, 2, 2, 2), date(2021 + n, 1, 1))
        for n in range(1, 6)
    ),
    *(VersionSpec("ssars", n, (6, 5, 4), (2, 1, 1), date(2024 + n, 1, 1)) for n in (1, 2)),
)


def workbook(spec: VersionSpec) -> bytes:
    """The version's workbook in K-01-T05's layout (Areas, Requests, Account rules, Programs,
    Steps); the same rows for a spec (the loader runs each import once, by its trace)."""
    book = Workbook()
    areas = book.active
    if areas is None:  # pragma: no cover -- openpyxl always opens a sheet
        raise RuntimeError("no sheet")
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["PLAN", "Planning"])
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    for n in range(PBC_REQUESTS):
        requests.append(["PLAN", f"{spec.name} request {n + 1}", "B"])
    rules = book.create_sheet("Account rules")
    rules.append(["area_code", "account_from", "account_to"])
    programs = book.create_sheet("Programs")
    programs.append(["code", "name"])
    for code in spec.codes:
        programs.append([code, f"{code} procedures"])
    steps = book.create_sheet("Steps")
    steps.append(["program_code", "number", "text", "level", "mandatory", "requests"])
    for code, count, locked in zip(spec.codes, spec.shape, spec.mandatory, strict=True):
        for n in range(1, count + 1):
            program = int(code.rsplit("-", 1)[1])
            revised = spec.family == "ebp" and spec.number >= 3 and (program, n) in EBP_V3_CHANGED
            steps.append(
                [
                    code,
                    n,
                    f"{code} step {n}" + (" (revised in v3)" if revised else ""),
                    LEVELS[n % 4],
                    "yes" if n <= locked else "no",
                    str(n) if spec.family == "ebp" and n <= PBC_REQUESTS else "",
                ]
            )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _key(spec: VersionSpec) -> str:
    return f"{spec.family}:v{spec.number}"


@register_slice(SLICE)
def methodology_events() -> list[Event]:
    events = [
        Event(f"flag:{name}", SETUP_AT, "platform.flag_set", "system", SLICE, {"flag": name})
        for name in FLAGS
    ]
    for n, spec in enumerate(VERSIONS):
        data: dict[str, object] = {
            "family": spec.family,
            "number": spec.number,
            "name": spec.name,
            "engagement_type": spec.engagement_type,
            "effective_from": spec.effective_from.isoformat(),
        }
        events.append(
            Event(
                f"methodology-import:{_key(spec)}",
                SETUP_AT + timedelta(minutes=10, seconds=n),
                "methodology.version_imported",
                IMPORTER,
                SLICE,
                data,
            )
        )
        events.append(
            Event(
                f"methodology-approve:{_key(spec)}",
                SETUP_AT + timedelta(minutes=20, seconds=n),
                "methodology.version_approved",
                APPROVER,
                SLICE,
                data,
            )
        )
    return events


def spec_of(family: str, number: int) -> VersionSpec:
    return next(v for v in VERSIONS if (v.family, v.number) == (family, number))


__all__ = ["EBP_V3_CHANGED", "NEEDS", "SLICE", "VERSIONS", "VersionSpec", "spec_of", "workbook"]


# --- K-01-T21: the EBP v3 methodology as a PDF, for the importer on the fake model ---------------

LINES_PER_PAGE: Final = 9
# The three steps a reviewer is asked to check (KNW-02 a's 3 flags): one runs onto the next page,
# one is worded "(as applicable)" (the fake importer is unsure of it), one cites a step the
# document does not have.
SPLIT_STEP: Final = (210, 3)
UNSURE_STEP: Final = (410, 4)
DANGLING_STEP: Final = (610, 5)


def ebp_v3_step_text(program: int, n: int) -> str:
    text = f"Perform procedure {n} of BH-EBP-{program} and document the result"
    if (program, n) == UNSURE_STEP:
        return f"{text} (as applicable)"
    if (program, n) == DANGLING_STEP:
        return f"{text} and agree it to BH-EBP-999 step 1"
    if (program, n) == (310, 2):
        return f"{text} and agree it to BH-EBP-110 step 7"
    return text


def ebp_v3_document_pdf() -> bytes:
    """BH Audit Methodology EBP v3 (67 steps in seven programs) as a text PDF: a heading line
    `BH-EBP-110 Procedures` per program, a line `n. text` per step, nine lines a page; the split
    step's last words open the next page. Synthetic, Brightwater Hale only."""
    lines: list[tuple[str, bool]] = []  # (text, break the page after it)
    for program, count in zip(EBP_PROGRAMS, EBP_V3, strict=True):
        lines.append((f"BH-EBP-{program} Procedures", False))
        for n in range(1, count + 1):
            text = ebp_v3_step_text(program, n)
            if (program, n) == SPLIT_STEP:
                head, _, tail = text.rpartition(" and ")
                lines.append((f"{n}. {head}", True))
                lines.append((f"and {tail}", False))
            else:
                lines.append((f"{n}. {text}", False))
    pages: list[list[tuple[float, float, str]]] = [[]]
    for text, breaks in lines:
        if len(pages[-1]) >= LINES_PER_PAGE:
            pages.append([])
        pages[-1].append((72.0, 720.0 - 20.0 * len(pages[-1]), text))
        if breaks:
            pages.append([])
    return positioned_pdf([p for p in pages if p])
