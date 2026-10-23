"""The E0 scenario documents, with fixed citation regions (F-04-T25; F-04.26; ADR-066).

Three documents for the Northwind Robotics 401(k) Plan audit (A1), each answering a request on
its PBC list or the Trial Balance workpaper. Every figure is blueprint 01 §9.4's (the statement of
net assets, the figure book and the loan rollforward), on the trust basis the plan's records have
at E0: planning excludes the 48,600 contributions receivable that fieldwork finds later, so net
assets read 16,812,400.00.

- `northwind_net_trust_assets`: Harborline Trust Company's summary of net trust assets (PDF).
- `northwind_loan_rollforward`: Keystone Retirement Services' annual loan balance roll-forward
  (PDF).
- `northwind_trial_balance`: the plan's trial balance on the trust basis, in the platform's
  rendered layout, which the screener reads (XLSX).

Each citeable value has a region: on a PDF page a rectangle in page points with a top-left origin
(the C4 `Rect`, `[x0, top, x1, bottom]`), on a sheet a cell. The regions are computed from where
the writer puts the text, never measured afterwards, and a test re-reads every region from the
bytes. The recorded model responses (`responses/*.json`) cite only these regions, except the one
planted unsupported value (`responses.PLANTED_UNSUPPORTED`). Synthetic; deterministic: the
same bytes, and so the same SHA-256, on every build.
"""

from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from functools import cache
from typing import Final
from uuid import UUID

import pdfplumber
from openpyxl import load_workbook

from abacus.modules.evidence.api import (
    CREDIT_COLUMN,
    DEBIT_COLUMN,
    FIRST_LINE_ROW,
    NAME_COLUMN,
    TrialBalance,
    TrialBalanceLine,
    render_trial_balance,
)
from abacus_tools.synthetic.documents import positioned_pdf
from abacus_tools.synthetic.scenarios.brightwater_hale.epochs import CT
from abacus_tools.synthetic.scenarios.brightwater_hale.prng import stable_id

PAGE_HEIGHT: Final = 792.0
_FIRST_BASELINE: Final = 720.0
_LINE_STEP: Final = 18.0
# Columns: a label, then two amounts (the statement's years, or a count and an amount).
_LABEL_X, _FIRST_X, _SECOND_X = 50.0, 380.0, 480.0
_BOUNDS: Final = ((48.0, 372.0), (376.0, 472.0), (476.0, 566.0))
# A region's band around a baseline: Helvetica 10 sits within 8 points above it and 2.1 below.
_ABOVE, _BELOW = 11.0, 4.0

PDF: Final = "application/pdf"
XLSX: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class PageRegion:
    """Text at a fixed place on a PDF page."""

    id: str
    page: int
    bbox: tuple[float, float, float, float]  # [x0, top, x1, bottom], points, top-left origin
    text: str


@dataclass(frozen=True)
class CellRegion:
    """A sheet cell; `text` is the cell's value as the screener is shown it."""

    id: str
    sheet: str
    cell: str
    text: str


Region = PageRegion | CellRegion


@dataclass(frozen=True)
class ScenarioDocument:
    key: str
    engagement: str  # the audit label ("A1")
    request: str | None  # the PBC request it answers; None for the Trial Balance workpaper
    filename: str
    media_type: str  # PDF or XLSX
    content: bytes
    regions: tuple[Region, ...]

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.content).hexdigest()

    def region(self, region_id: str) -> Region:
        for found in self.regions:
            if found.id == region_id:
                return found
        raise KeyError(f"{self.key} has no region {region_id!r}")


def _money(value: int | Decimal) -> str:
    return f"{Decimal(value):,.2f}"


Row = tuple[str, str, str | None, str | None]  # region id prefix, label, first, second


def _page(
    heading: tuple[str, ...], rows: tuple[Row, ...], page: int = 1
) -> tuple[list[tuple[float, float, str]], list[PageRegion]]:
    """The runs for one page, and a region for every non-empty run of a row."""
    runs: list[tuple[float, float, str]] = []
    regions: list[PageRegion] = []
    baseline = _FIRST_BASELINE
    for n, text in enumerate(heading):
        runs.append((_LABEL_X, baseline, text))
        regions.append(_region(f"heading-{n + 1}", page, baseline, _BOUNDS[0], text))
        baseline -= _LINE_STEP
    baseline -= _LINE_STEP  # a blank line under the heading
    for key, label, first, second in rows:
        for column, (x, text, suffix) in enumerate(
            ((_LABEL_X, label, "label"), (_FIRST_X, first, "first"), (_SECOND_X, second, "second"))
        ):
            if text:
                runs.append((x, baseline, text))
                regions.append(_region(f"{key}.{suffix}", page, baseline, _BOUNDS[column], text))
        baseline -= _LINE_STEP
    return runs, regions


def _region(
    region_id: str, page: int, baseline: float, bounds: tuple[float, float], text: str
) -> PageRegion:
    top = PAGE_HEIGHT - baseline - _ABOVE
    return PageRegion(region_id, page, (bounds[0], top, bounds[1], top + _ABOVE + _BELOW), text)


# --- Harborline Trust Company: summary of net trust assets ---

_NET_ASSETS: Final[tuple[Row, ...]] = (
    ("columns", "", "12/31/2025", "1/1/2025"),
    ("krs-500", "KRS 500 Index Fund", _money(6_120_450), _money(5_102_300)),
    ("krs-2035", "KRS Target 2035 Fund", _money(2_384_910), _money(2_011_640)),
    ("krs-2045", "KRS Target 2045 Fund", _money(2_915_330), _money(2_388_970)),
    ("krs-bond", "KRS Bond Index Fund", _money(1_408_220), _money(1_276_090)),
    ("krs-intl", "KRS International Index Fund", _money(902_610), _money(868_000)),
    ("stable-value", "Harborline Stable Value Fund", _money(1_163_480), _money(1_036_000)),
    ("mutual-funds", "Total mutual funds & CITs", _money(14_895_000), _money(12_683_000)),
    (
        "brokerage",
        "Self-directed brokerage (BrokerWindow)",
        _money(1_605_000),
        _money(1_257_000),
    ),
    ("investments", "Total investments at fair value", _money(16_500_000), _money(13_940_000)),
    ("loans", "Notes receivable from participants", _money(312_400), _money(288_500)),
    ("net-assets", "Net assets held in trust", _money(16_812_400), _money(14_228_500)),
)


def _net_trust_assets() -> ScenarioDocument:
    runs, regions = _page(
        (
            "Harborline Trust Company",
            "Northwind Robotics 401(k) Plan",
            "Summary of net trust assets (trust basis)",
            "Plan year January 1, 2025 to December 31, 2025",
        ),
        _NET_ASSETS,
    )
    return ScenarioDocument(
        "northwind_net_trust_assets",
        "A1",
        "Summary of net trust assets",
        "Northwind 401k - Summary of net trust assets 2025.pdf",
        PDF,
        positioned_pdf([runs]),
        tuple(regions),
    )


# --- Keystone Retirement Services: annual loan balance roll-forward ---

_LOANS: Final[tuple[Row, ...]] = (
    ("columns", "", "Loans", "Amount"),
    ("beginning", "Balance at January 1, 2025", "41", _money(288_500)),
    ("new", "New loans", "5", _money(96_200)),
    ("repayments", "Principal repayments", "7", f"({_money(57_450)})"),
    ("offsets", "Loan offsets (terminated participants)", "3", f"({_money(14_850)})"),
    ("ending", "Balance at December 31, 2025", "36", _money(312_400)),
)


def _loan_rollforward() -> ScenarioDocument:
    runs, regions = _page(
        (
            "Keystone Retirement Services",
            "Northwind Robotics 401(k) Plan",
            "Annual loan balance roll-forward",
            "Plan year January 1, 2025 to December 31, 2025",
        ),
        _LOANS,
    )
    return ScenarioDocument(
        "northwind_loan_rollforward",
        "A1",
        "Annual loan balance roll-forward",
        "Northwind 401k - Loan roll-forward 2025.pdf",
        PDF,
        positioned_pdf([runs]),
        tuple(regions),
    )


# --- The plan's trial balance (trust basis) ---

# (account code, name, debit, credit). The codes are the scenario's chart, not a figure; every
# amount is blueprint §9.4's (statement of net assets, 12/31/2025; figure book, trust basis).
_TB_LINES: Final[tuple[tuple[str, str, int, int], ...]] = (
    ("1110", "KRS 500 Index Fund", 6_120_450, 0),
    ("1120", "KRS Target 2035 Fund", 2_384_910, 0),
    ("1130", "KRS Target 2045 Fund", 2_915_330, 0),
    ("1140", "KRS Bond Index Fund", 1_408_220, 0),
    ("1150", "KRS International Index Fund", 902_610, 0),
    ("1160", "Harborline Stable Value Fund", 1_163_480, 0),
    ("1200", "Self-directed brokerage (BrokerWindow)", 1_605_000, 0),
    ("1300", "Notes receivable from participants", 312_400, 0),
    ("3000", "Net assets available for benefits, January 1, 2025", 0, 14_228_500),
    ("4100", "Contributions received", 0, 1_318_300),
    ("4200", "Net appreciation in fair value of investments", 0, 1_755_990),
    ("4300", "Interest and dividends", 0, 342_860),
    ("4400", "Interest income on notes receivable from participants", 0, 17_800),
    ("5100", "Benefits paid to participants", 812_650, 0),
    ("5200", "Administrative expenses", 38_400, 0),
)
_TB_SHEET: Final = "Trial balance"
# The workpaper's status on AUD-02 at E0: "Complete · 9:26 AM".
_TB_PULLED_AT: Final = datetime(2026, 9, 14, 9, 26, tzinfo=CT)


def _column(index: int) -> str:
    return "ABCDEFGHIJ"[index - 1]


def _trial_balance() -> ScenarioDocument:
    entity = "Northwind Robotics 401(k) Plan"
    tb = TrialBalance(
        client_entity_id=UUID(stable_id(f"entity:{entity}")),
        entity_name=entity,
        period_start=date(2025, 1, 1),
        period_end=date(2025, 12, 31),
        pulled_at=_TB_PULLED_AT,
        snapshot_id=UUID(stable_id("ledger-snapshot:northwind-tb-2025")),
        source="Harborline Trust Company trust report",
        source_fingerprint=stable_id("source:northwind-tb-2025"),
        lines=tuple(
            TrialBalanceLine(code, name, Decimal(debit), Decimal(credit))
            for code, name, debit, credit in _TB_LINES
        ),
        method="imported",
    )
    content = render_trial_balance(tb)
    ordered = sorted(_TB_LINES)  # the renderer orders lines by account code
    regions: list[Region] = []
    for row, (code, name, debit, credit) in enumerate(ordered, start=FIRST_LINE_ROW):
        regions.append(CellRegion(f"{code}.name", _TB_SHEET, f"{_column(NAME_COLUMN)}{row}", name))
        amount, column = (debit, DEBIT_COLUMN) if debit else (credit, CREDIT_COLUMN)
        regions.append(
            CellRegion(f"{code}.amount", _TB_SHEET, f"{_column(column)}{row}", _money(amount))
        )
    total_row = FIRST_LINE_ROW + len(ordered)
    debit_total = sum(d for _, _, d, _ in _TB_LINES)
    credit_total = sum(c for _, _, _, c in _TB_LINES)
    regions += [
        CellRegion("total.label", _TB_SHEET, f"{_column(NAME_COLUMN)}{total_row}", "Total"),
        CellRegion(
            "total.debit", _TB_SHEET, f"{_column(DEBIT_COLUMN)}{total_row}", _money(debit_total)
        ),
        CellRegion(
            "total.credit",
            _TB_SHEET,
            f"{_column(CREDIT_COLUMN)}{total_row}",
            _money(credit_total),
        ),
    ]
    return ScenarioDocument(
        "northwind_trial_balance",
        "A1",
        None,
        "Northwind 401k - Trial balance 2025 (trust basis).xlsx",
        XLSX,
        content,
        tuple(regions),
    )


@cache
def documents() -> tuple[ScenarioDocument, ...]:
    """The E0 scenario documents, in a fixed order."""
    return (_net_trust_assets(), _loan_rollforward(), _trial_balance())


def by_sha256(sha256: str) -> ScenarioDocument | None:
    return next((d for d in documents() if d.sha256 == sha256), None)


# --- reading a region back from the bytes (verification by code) ---


def page_text(document: ScenarioDocument, page: int) -> str:
    """A PDF page's extracted text, as the evidence pipeline's extraction reads it."""
    with pdfplumber.open(io.BytesIO(document.content)) as pdf:
        return pdf.pages[page - 1].extract_text() or ""


def region_text(document: ScenarioDocument, region: Region) -> str:
    """What the bytes hold at the region: the characters wholly inside a page rectangle, or a
    cell's value (an amount formatted as the screener is shown it)."""
    if isinstance(region, PageRegion):
        with pdfplumber.open(io.BytesIO(document.content)) as pdf:
            cropped = pdf.pages[region.page - 1].within_bbox(region.bbox)
            return " ".join((cropped.extract_text() or "").split())
    workbook = load_workbook(io.BytesIO(document.content), read_only=True, data_only=True)
    value = workbook[region.sheet][region.cell].value
    if isinstance(value, int | float | Decimal) and not isinstance(value, bool):
        return _money(Decimal(str(value)))
    return "" if value is None else str(value)
