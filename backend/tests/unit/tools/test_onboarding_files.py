"""Q-04-T018: the onboarding rehearsal's input files are synthetic, importable and reproducible.

Acceptance: every content step completes through its screen with timings recorded. These files
are what steps 1 to 4 upload: the client list must map on CLI-06's own headings and entity words,
hold no TIN and no deliverable address, send each `--near` row to review, and the PBC template
must carry Workbook v2's required sheets and columns with valid areas and tiers.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from openpyxl import load_workbook

from abacus.kernel.tin_patterns import find_tins
from abacus.modules.engagements.client_import_parse import ENTITY_TYPE_WORDS, HEADINGS
from abacus.modules.methodology.workbook import TIERS
from abacus_tools.synthetic import onboarding_files as of
from abacus_tools.synthetic.names import MARKER


def _rows(count: int = 200, near: tuple[str, ...] = ()) -> list[dict[str, str]]:
    body = of.client_csv(of.client_rows(count, 2027, near)).decode("utf-8")
    return list(csv.DictReader(io.StringIO(body)))


def test_every_heading_maps_and_every_entity_word_is_known() -> None:
    norm = {h: " ".join(h.casefold().split()) for h in of.CLIENT_HEADINGS}
    assert all(n in HEADINGS for n in norm.values())
    assert {r["Entity Type"].casefold() for r in _rows()} <= set(ENTITY_TYPE_WORDS)


def test_two_hundred_distinct_synthetic_clients_with_no_tin_and_no_real_mail() -> None:
    rows = _rows()
    assert len(rows) == 200 and len({r["Client Name"] for r in rows}) == 200
    assert all(r["Client Name"].endswith(MARKER) for r in rows)
    assert all(r["Email Address"].endswith("@example.test") for r in rows)
    body = of.client_csv(of.client_rows(200, 2027))
    assert find_tins(body.decode("utf-8")) == []


def test_each_near_name_is_added_once_more_for_review() -> None:
    rows = _rows(200, ("Okafor Household (Synthetic)",))
    assert len(rows) == 201 and rows[-1]["Client Name"] == "Okafor Household (Synthetic)"


def test_the_same_seed_gives_the_same_bytes(tmp_path: Path) -> None:
    first = of.write(tmp_path / "a", clients=200, seed=2027)
    second = of.write(tmp_path / "b", clients=200, seed=2027)
    assert first == second
    other = of.write(tmp_path / "c", clients=200, seed=2028)
    assert other != first


def test_the_pbc_template_is_workbook_v2(tmp_path: Path) -> None:
    body, count = of.pbc_workbook()
    path = tmp_path / "pbc.xlsx"
    path.write_bytes(body)
    book = load_workbook(path, read_only=True)
    assert set(of.XLSX_SHEETS) <= set(book.sheetnames)
    areas = [r for r in book["Areas"].iter_rows(values_only=True)]
    requests = [r for r in book["Requests"].iter_rows(values_only=True)]
    assert areas[0] == ("code", "name") and requests[0] == ("area_code", "description", "tier")
    codes = {r[0] for r in areas[1:]}
    assert len(requests) - 1 == count
    assert all(r[0] in codes and str(r[2]) in TIERS for r in requests[1:])


def test_the_command_writes_the_manifest_and_refuses_an_oversized_list(tmp_path: Path) -> None:
    assert of.main(["--out", str(tmp_path), "--clients", "5"]) == 0
    manifest = json.loads((tmp_path / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["files"]["clients-5.csv"]["rows"] == 5
    assert of.main(["--out", str(tmp_path), "--clients", "2001"]) == 2
