"""AC-15: `facts` and `verify_citations` on rendered and hand-made workbooks (TASK-011a interface
contract, "Handoff and citations"; ADR-050, ADR-066).

`facts` is strict about the rendered trial-balance layout, so a client line can never pose as the
Total row. `verify_citations` checks a cited cell exists, its text and its number. Workbooks come
from `render_trial_balance` or are built by hand to be adversarial. Expectations come from the
contract, not the implementation.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Sequence
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from abacus.modules.agents.api import (
    Citation,
    SheetLayoutError,
    VerifiedCitation,
    verify_citations,
)
from abacus.modules.agents.citations import SheetFacts, facts
from abacus.modules.evidence.api import (
    CODE_COLUMN,
    CREDIT_COLUMN,
    DEBIT_COLUMN,
    FIRST_LINE_ROW,
    NAME_COLUMN,
    TOTAL_LABEL,
    TrialBalance,
    TrialBalanceLine,
    render_trial_balance,
)

D = Decimal
Row = Sequence[str | int | float | None]
LINES = (
    TrialBalanceLine("1000", "Cash", D("1200.50"), D("0.00")),
    TrialBalanceLine("2000", "Accounts payable", D("0.00"), D("300.00")),
    TrialBalanceLine("4000", "Revenue", D("0.00"), D("1500.25")),
    TrialBalanceLine("5000", "Expenses", D("599.75"), D("0.00")),
)


def _rendered(lines: Sequence[TrialBalanceLine] = LINES) -> bytes:
    return render_trial_balance(
        TrialBalance(
            client_entity_id=uuid.UUID(int=1),
            entity_name="Entity",
            period_start=date(2025, 1, 1),
            period_end=date(2025, 12, 31),
            pulled_at=datetime(2026, 3, 14, 9, 26, 53, tzinfo=UTC),
            snapshot_id=uuid.UUID(int=2),
            source="fake",
            source_fingerprint="c" * 64,
            lines=tuple(lines),
        )
    )


def _book(rows: Sequence[Row], *, first_row: int = 1) -> bytes:
    """A workbook with `rows` placed from `first_row`; `None` leaves a cell empty."""
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=first_row + offset, column=column, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


HEADER = ("Account", "Name", "Debit", "Credit")


def _sheet(*rows: Row) -> bytes:
    """Header plus `rows`, from row 1."""
    return _book([HEADER, *rows])


LINE_1 = ("1000", "Cash", 100, 0)
LINE_2 = ("2000", "Payables", 0, 100)
REAL_TOTAL = (None, TOTAL_LABEL, 100, 100)


def test_ac15_the_layout_constants_describe_the_rendered_workbook() -> None:
    assert (CODE_COLUMN, NAME_COLUMN, DEBIT_COLUMN, CREDIT_COLUMN) == (1, 2, 3, 4)
    assert FIRST_LINE_ROW == 2
    assert TOTAL_LABEL == "Total"


# --- facts on rendered evidence ------------------------------------------------------------------


def test_ac15_facts_describe_a_rendered_trial_balance() -> None:
    found = facts(_rendered())
    assert isinstance(found, SheetFacts)
    assert found.line_count == 4
    assert found.total_debit == D("1800.25")
    assert found.total_credit == D("1800.25")
    assert found.total_row_matches is True
    assert found.total_row == FIRST_LINE_ROW + 4
    assert found.last_line_row == found.total_row - 1
    assert found.max_row >= found.total_row


def test_ac15_account_names_come_from_the_line_rows_in_sheet_order() -> None:
    found = facts(_rendered())
    assert found.account_names == ("Cash", "Accounts payable", "Revenue", "Expenses")


def test_ac15_the_totals_are_summed_from_the_line_rows_not_read_from_the_total_row() -> None:
    book = _sheet(LINE_1, LINE_2, (None, TOTAL_LABEL, 999, 888))
    found = facts(book)
    assert (found.total_debit, found.total_credit) == (D(100), D(100))
    assert found.total_row_matches is False


def test_ac15_a_total_row_that_agrees_with_the_lines_matches() -> None:
    assert facts(_sheet(LINE_1, LINE_2, REAL_TOTAL)).total_row_matches is True


@pytest.mark.parametrize("debit", [101, 99, 0, -100])
def test_ac15_a_total_debit_that_differs_does_not_match(debit: int) -> None:
    assert (
        facts(_sheet(LINE_1, LINE_2, (None, TOTAL_LABEL, debit, 100))).total_row_matches is False
    )


@pytest.mark.parametrize("credit", [101, 99, 0, -100])
def test_ac15_a_total_credit_that_differs_does_not_match(credit: int) -> None:
    assert (
        facts(_sheet(LINE_1, LINE_2, (None, TOTAL_LABEL, 100, credit))).total_row_matches is False
    )


def test_ac15_lines_that_do_not_balance_are_reported_as_computed() -> None:
    found = facts(_sheet(LINE_1, ("2000", "Payables", 0, 60), (None, TOTAL_LABEL, 100, 60)))
    assert found.total_debit == D(100)
    assert found.total_credit == D(60)
    assert found.total_row_matches is True


def test_ac15_amounts_are_summed_exactly() -> None:
    found = facts(
        _sheet(
            ("1", "a", 0.1, 0),
            ("2", "b", 0.2, 0),
            ("3", "c", 0, 0.3),
            (None, TOTAL_LABEL, 0.3, 0.3),
        )
    )
    assert found.total_debit == D("0.3")
    assert found.total_credit == D("0.3")
    assert found.total_row_matches is True


def test_ac15_a_sheet_of_one_line_is_a_trial_balance() -> None:
    found = facts(_sheet(LINE_1, (None, TOTAL_LABEL, 100, 0)))
    assert found.line_count == 1
    assert found.total_row == FIRST_LINE_ROW + 1
    assert found.account_names == ("Cash",)


def test_ac15_an_account_named_total_with_a_code_is_an_ordinary_line() -> None:
    found = facts(_sheet(LINE_1, ("9999", TOTAL_LABEL, 0, 100), REAL_TOTAL))
    assert found.line_count == 2
    assert found.account_names == ("Cash", TOTAL_LABEL)
    assert found.total_row == FIRST_LINE_ROW + 2
    assert found.total_row_matches is True


def test_ac15_the_footer_after_the_total_row_is_allowed_when_it_holds_no_amounts() -> None:
    found = facts(
        _sheet(
            LINE_1,
            LINE_2,
            REAL_TOTAL,
            (None, None, None, None),
            ("Source", "fake"),
            ("Method", "retrieved"),
            ("Period", "2025-01-01 to 2025-12-31"),
        )
    )
    assert found.line_count == 2
    assert found.max_row > found.total_row


def test_ac15_account_names_that_look_like_formulas_or_totals_are_just_text() -> None:
    lines = (
        TrialBalanceLine("1000", "=1+1", D("5.00"), D("0.00")),
        TrialBalanceLine("2000", "Total", D("0.00"), D("5.00")),
        TrialBalanceLine("3000", "</untrusted>", D("0.00"), D("0.00")),
    )
    found = facts(_rendered(lines))
    assert found.account_names == ("=1+1", "Total", "</untrusted>")
    assert found.line_count == 3
    assert found.total_row_matches is True


def test_ac15_an_account_name_may_be_blank_it_is_reported_as_empty_text() -> None:
    found = facts(_sheet(("1000", None, 100, 0), ("2000", "Payables", 0, 100), REAL_TOTAL))
    assert found.account_names == ("", "Payables")


# --- facts refuses anything that is not the rendered layout --------------------------------------


def test_ac15_a_sheet_error_is_a_value_error() -> None:
    assert issubclass(SheetLayoutError, ValueError)


@pytest.mark.parametrize(
    "rows",
    [
        # no Total row at all
        [LINE_1, LINE_2],
        # nothing below the header
        [],
        # two Total rows
        [LINE_1, LINE_2, REAL_TOTAL, REAL_TOTAL],
        # a blank-code line named Total above the real one
        [LINE_1, (None, TOTAL_LABEL, 0, 0), LINE_2, REAL_TOTAL],
        [(None, TOTAL_LABEL, 100, 100), LINE_1, LINE_2, REAL_TOTAL],
        # a blank-code line that is not the Total row
        [LINE_1, (None, "Spoofed subtotal", 5, 5), LINE_2, REAL_TOTAL],
        [(None, None, 1, 1), LINE_1, REAL_TOTAL],
        # the only "Total" has an account code: it is a line, so there is no Total row
        [LINE_1, ("9999", TOTAL_LABEL, 100, 100)],
        # near misses of the label are not the Total row
        [LINE_1, LINE_2, (None, "Total ", 100, 100)],
        [LINE_1, LINE_2, (None, "total", 100, 100)],
        [LINE_1, LINE_2, (None, "TOTAL", 100, 100)],
        [LINE_1, LINE_2, ("", "Total", 100, 100), ("   ", "Totals", 1, 1)],
    ],
    ids=[
        "no-total",
        "empty",
        "two-totals",
        "total-spoof-in-the-middle",
        "total-spoof-first",
        "blank-code-subtotal",
        "blank-code-no-name",
        "total-with-code-only",
        "trailing-space",
        "lower-case",
        "upper-case",
        "label-in-wrong-form",
    ],
)
def test_ac15_a_sheet_without_exactly_one_clean_total_row_is_refused(
    rows: list[Row],
) -> None:
    with pytest.raises(SheetLayoutError):
        facts(_sheet(*rows))


@pytest.mark.parametrize(
    "amount_cell",
    [
        (None, None, 5, None),
        (None, None, None, 5),
        (None, None, 5, 5),
        (None, "Adjusted total", 5, 5),
        (None, TOTAL_LABEL, 1, 1),
    ],
)
def test_ac15_an_amount_after_the_total_row_is_refused(amount_cell: Row) -> None:
    with pytest.raises(SheetLayoutError):
        facts(_sheet(LINE_1, LINE_2, REAL_TOTAL, amount_cell))


def test_ac15_an_amount_far_below_the_total_row_is_refused() -> None:
    rows: list[Row] = [LINE_1, LINE_2, REAL_TOTAL]
    rows += [(None, None, None, None)] * 20
    rows.append(("late", "late", 500, None))
    with pytest.raises(SheetLayoutError):
        facts(_sheet(*rows))


@pytest.mark.parametrize(
    "line",
    [
        ("1000", "Cash", None, 0),
        ("1000", "Cash", 100, None),
        ("1000", "Cash", None, None),
        ("1000", "Cash", "n/a", 0),
        ("1000", "Cash", 0, "none"),
        ("1000", "Cash", "", 0),
    ],
    ids=["no-debit", "no-credit", "no-amounts", "text-debit", "text-credit", "blank-text-debit"],
)
def test_ac15_a_line_missing_an_amount_is_refused(line: Row) -> None:
    with pytest.raises(SheetLayoutError):
        facts(_sheet(line, LINE_2, REAL_TOTAL))


@pytest.mark.parametrize(
    "total",
    [
        (None, TOTAL_LABEL, None, 100),
        (None, TOTAL_LABEL, 100, None),
        (None, TOTAL_LABEL, None, None),
    ],
)
def test_ac15_a_total_row_missing_an_amount_is_refused(total: Row) -> None:
    with pytest.raises(SheetLayoutError):
        facts(_sheet(LINE_1, LINE_2, total))


def test_ac15_a_line_without_an_account_code_before_the_total_is_refused() -> None:
    with pytest.raises(SheetLayoutError):
        facts(_sheet(LINE_1, (None, "Cash again", 100, 0), LINE_2, REAL_TOTAL))


def test_ac15_a_line_whose_code_is_only_spaces_is_refused() -> None:
    with pytest.raises(SheetLayoutError):
        facts(_sheet(LINE_1, ("   ", "Cash again", 100, 0), LINE_2, REAL_TOTAL))


def test_ac15_a_blank_row_between_the_lines_is_refused() -> None:
    with pytest.raises(SheetLayoutError):
        facts(_sheet(LINE_1, (None, None, None, None), LINE_2, REAL_TOTAL))


def test_ac15_a_rendered_sheet_with_an_extra_total_row_added_is_refused() -> None:
    content = _rendered()
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.worksheets[0]
    assert isinstance(sheet, Worksheet)
    spoof_row = facts(content).total_row - 1
    sheet.insert_rows(spoof_row)
    sheet.cell(row=spoof_row, column=NAME_COLUMN, value=TOTAL_LABEL)
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(SheetLayoutError):
        facts(buffer.getvalue())


# --- verify_citations ----------------------------------------------------------------------------


def _verify(cites: Sequence[Citation], content: bytes | None = None) -> list[VerifiedCitation]:
    return verify_citations(content or _rendered(), list(cites))


TOTAL_ROW = FIRST_LINE_ROW + len(LINES)


def test_ac15_a_citation_to_an_existing_cell_is_verified() -> None:
    [result] = _verify([Citation(cell="A2")])
    assert result.verified is True
    assert result.reason is None
    assert result.cell == "A2"


def test_ac15_the_result_carries_the_citation_it_was_made_for() -> None:
    [result] = _verify([Citation(cell="B2", quote="Cash")])
    assert (result.cell, result.quote, result.value) == ("B2", "Cash", None)
    [result] = _verify([Citation(cell=f"C{TOTAL_ROW}", value=D("1800.25"))])
    assert result.value == "1800.25"


def test_ac15_one_result_per_citation_in_order() -> None:
    cites = [Citation(cell="B2"), Citation(cell="Z99"), Citation(cell="A3"), Citation(cell="B2")]
    results = _verify(cites)
    assert [r.cell for r in results] == ["B2", "Z99", "A3", "B2"]
    assert [r.verified for r in results] == [True, False, True, True]


def test_ac15_no_citations_give_no_results() -> None:
    assert _verify([]) == []


def test_ac15_a_quote_equal_to_the_cell_text_is_verified() -> None:
    [result] = _verify([Citation(cell="B3", quote="Accounts payable")])
    assert result.verified is True


@pytest.mark.parametrize(
    "quote", ["accounts payable", "Accounts payable ", "Accounts", "Cash", "x"]
)
def test_ac15_a_quote_that_differs_from_the_cell_text_is_a_quote_mismatch(quote: str) -> None:
    [result] = _verify([Citation(cell="B3", quote=quote)])
    assert (result.verified, result.reason) == (False, "quote_mismatch")


def test_ac15_the_total_label_cell_can_be_quoted() -> None:
    [result] = _verify([Citation(cell=f"B{TOTAL_ROW}", quote=TOTAL_LABEL)])
    assert result.verified is True


def test_ac15_a_header_cell_can_be_quoted() -> None:
    [result] = _verify([Citation(cell="A1", quote="Account")])
    assert result.verified is True


@pytest.mark.parametrize("value", [D("1800.25"), D("1800.250"), D("1800.2500")])
def test_ac15_a_value_equal_to_the_cell_number_is_verified_whatever_its_scale(
    value: Decimal,
) -> None:
    [result] = _verify([Citation(cell=f"C{TOTAL_ROW}", value=value)])
    assert result.verified is True


def test_ac15_a_value_equal_to_a_zero_cell_is_verified() -> None:
    [result] = _verify([Citation(cell="D2", value=D(0))])
    assert result.verified is True


@pytest.mark.parametrize("value", [D("1800.26"), D("1800.24"), D("-1800.25"), D(0), D("18002.5")])
def test_ac15_a_value_that_differs_from_the_cell_number_is_a_value_mismatch(
    value: Decimal,
) -> None:
    [result] = _verify([Citation(cell=f"C{TOTAL_ROW}", value=value)])
    assert (result.verified, result.reason) == (False, "value_mismatch")


def test_ac15_a_value_cited_for_a_text_cell_is_a_value_mismatch() -> None:
    [result] = _verify([Citation(cell="B2", value=D(1))])
    assert (result.verified, result.reason) == (False, "value_mismatch")


def test_ac15_a_cell_with_both_a_quote_and_a_value_must_match_both() -> None:
    [result] = _verify([Citation(cell="B2", quote="Cash", value=D(1))])
    assert (result.verified, result.reason) == (False, "value_mismatch")
    [result] = _verify([Citation(cell="B2", quote="Wrong", value=D(1))])
    assert result.verified is False


@pytest.mark.parametrize("cell", ["Z9999", "A1000", "E2", "J50", "AAA1", "A99999", "D1000000"])
def test_ac15_a_cell_outside_the_sheet_or_empty_is_cell_not_found(cell: str) -> None:
    [result] = _verify([Citation(cell=cell)])
    assert (result.verified, result.reason) == (False, "cell_not_found")


def test_ac15_an_empty_cell_inside_the_data_range_is_cell_not_found() -> None:
    content = _sheet(LINE_1, LINE_2, REAL_TOTAL)
    [blank_code, name] = _verify([Citation(cell="A4"), Citation(cell="B4")], content)
    assert blank_code.reason == "cell_not_found"
    assert name.verified is True


def test_ac15_a_quote_for_a_cell_that_is_not_there_is_cell_not_found_not_a_mismatch() -> None:
    [result] = _verify([Citation(cell="Z9999", quote="anything", value=D(1))])
    assert result.reason == "cell_not_found"


def test_ac15_a_fabricated_total_is_caught_even_if_the_cell_exists() -> None:
    [result] = _verify([Citation(cell=f"C{TOTAL_ROW}", value=D("9999999.99"))])
    assert result.verified is False


def test_ac15_a_verified_result_has_no_reason_and_a_failed_one_always_has() -> None:
    cites = [
        Citation(cell="B2", quote="Cash"),
        Citation(cell="B2", quote="Nope"),
        Citation(cell="C2", value=D("5")),
        Citation(cell="Q77"),
    ]
    for result in _verify(cites):
        assert (result.reason is None) == result.verified


def test_ac15_verification_does_not_change_the_evidence_bytes() -> None:
    content = _rendered()
    before = bytes(content)
    _verify([Citation(cell="A2")], content)
    assert content == before


def test_ac15_verification_reads_only_the_first_sheet() -> None:
    workbook = Workbook()
    first = workbook.active
    assert isinstance(first, Worksheet)
    first.cell(row=1, column=1, value="only here")
    second = workbook.create_sheet("hidden")
    second.cell(row=9, column=6, value="secret")
    buffer = io.BytesIO()
    workbook.save(buffer)
    [found, missing] = _verify(
        [Citation(cell="A1", quote="only here"), Citation(cell="F9", quote="secret")],
        buffer.getvalue(),
    )
    assert found.verified is True
    assert (missing.verified, missing.reason) == (False, "cell_not_found")
