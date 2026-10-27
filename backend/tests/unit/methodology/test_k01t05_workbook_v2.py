"""K-01-T05 (TASK-1427): workbook v2, parsed deterministically (K-01 §3.2, AC2.1 to AC2.4).

- optional Programs and Steps sheets (read with the programs flag), Account rules optional, a
  legacy three-sheet workbook unchanged;
- problems are sheet, row, column and a fixed code (ADR-052), never a cell's contents;
- `mandatory` and `lock_note` come only from their columns; levels are the four words;
- formulas are never evaluated;
- step keys: the workbook's `step_id`, else the previous version's key for the same program and
  number, else derived; deterministic and unique;
- the blank template's sheets and headers are what the parser reads, and its bytes are stable.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from abacus.modules.methodology.repository import assign_step_keys
from abacus.modules.methodology.workbook import StepRow, TemplateInvalid, parse_workbook
from abacus.modules.methodology.workbook_template import SHEETS, workbook_template

Row = list[object]


def _book(
    *,
    rules: bool = True,
    programs: list[Row] | None = None,
    steps: list[Row] | None = None,
    step_header: Row | None = None,
    requests: int = 3,
) -> bytes:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["PLAN", "Planning"])
    sheet = book.create_sheet("Requests")
    sheet.append(["area_code", "description", "tier"])
    for n in range(requests):
        sheet.append(["PLAN", f"Request {n + 1}", "B"])
    if rules:
        sheet = book.create_sheet("Account rules")
        sheet.append(["area_code", "account_from", "account_to"])
    if programs is not None:
        sheet = book.create_sheet("Programs")
        sheet.append(["code", "name"])
        for row in programs:
            sheet.append(row)
    if steps is not None:
        sheet = book.create_sheet("Steps")
        sheet.append(
            step_header
            or [
                "program_code",
                "number",
                "text",
                "level",
                "step_id",
                "mandatory",
                "lock_note",
                "assertions",
                "standard_refs",
                "requests",
            ]
        )
        for row in steps:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


PROGRAMS: list[Row] = [["P-110", "General planning"], ["P-120", "Risk assessment"]]
STEPS: list[Row] = [
    ["P-110", 1, "Read the plan document", "Foundation", "p110.read", "yes", "Reserved to EP"],
    ["P-110", 2, "Set materiality", "Planning", None, "no", None, "E;C", "AU-C 320", "1;3"],
    ["P-120", 1, "Assess risks", "testing"],
]


def _codes(data: bytes, *, programs: bool = True) -> set[tuple[str | None, int | None, str]]:
    with pytest.raises(TemplateInvalid) as caught:
        parse_workbook(data, programs=programs)
    return {(p.sheet, p.row, p.code) for p in caught.value.problems}


def test_ac21_programs_and_steps_come_with_the_flag() -> None:
    found = parse_workbook(_book(programs=PROGRAMS, steps=STEPS), programs=True)
    assert [p.code for p in found.programs] == ["P-110", "P-120"]
    assert [(s.program_code, s.number, s.level) for s in found.steps] == [
        ("P-110", 1, "foundation"),
        ("P-110", 2, "planning"),
        ("P-120", 1, "testing"),
    ]
    first, second, _third = found.steps
    assert (first.mandatory, first.lock_note, first.step_id) == (
        True,
        "Reserved to EP",
        "p110.read",
    )
    assert (second.mandatory, second.assertions, second.standard_refs) == (
        False,
        ("E", "C"),
        ("AU-C 320",),
    )
    assert second.request_positions == (0, 2)
    # Without the flag the sheets are not read: the PBC part imports as before.
    legacy = parse_workbook(_book(programs=PROGRAMS, steps=STEPS), programs=False)
    assert (legacy.programs, legacy.steps, len(legacy.items)) == ((), (), 3)


def test_ac21_ac22_a_legacy_or_pbc_only_workbook_still_imports() -> None:
    assert len(parse_workbook(_book()).items) == 3
    assert parse_workbook(_book(rules=False), programs=True).rules == ()


def test_ac23_problems_are_codes_never_contents() -> None:
    bad: list[Row] = [
        ["P-999", 1, "Unknown program", "Planning"],
        ["P-110", 1, "Bad level", "Fieldwork"],
        ["P-110", 3, "A gap", "Planning", "dup"],
        ["P-110", 4, "Same id", "Planning", "dup"],
        ["P-110", 5, "Note without lock", "Planning", None, "no", "Reserved"],
        ["P-110", 6, "Unknown request", "Planning", None, None, None, None, None, "9"],
    ]
    found = _codes(_book(programs=PROGRAMS, steps=bad))
    assert {
        ("Steps", 2, "unknown_program"),
        ("Steps", 3, "invalid_level"),
        ("Steps", 5, "duplicate_step_id"),
        ("Steps", 6, "lock_note_without_mandatory"),
        ("Steps", 7, "unknown_request"),
    } <= found
    assert any(code == "numbering_gap" for _s, _r, code in found)
    # One of the pair without the other is a missing sheet.
    assert ("Steps", None, "missing_sheet") in _codes(_book(programs=PROGRAMS))


def test_adr052_a_formula_is_never_evaluated() -> None:
    steps: list[Row] = [["P-110", 1, "=1+1", "Planning"], ["P-120", 1, "Assess", "Planning"]]
    # The formula has no cached value: read as empty, never computed.
    assert ("Steps", 2, "empty_value") in _codes(_book(programs=PROGRAMS, steps=steps))


def test_ac24_step_keys_are_stable_and_deterministic() -> None:
    steps = [
        StepRow("P-110", 1, "a", "planning", step_id="named"),
        StepRow("P-110", 2, "b", "planning"),
        StepRow("P-110", 3, "c", "planning"),
        StepRow("P-120", 1, "d", "testing"),
    ]
    prior = {("P-110", 2): "kept-from-v2", ("P-110", 3): "named"}
    keys = assign_step_keys(steps, prior)
    assert keys == ["named", "kept-from-v2", "p-110-3", "p-120-1"]
    assert assign_step_keys(steps, prior) == keys
    assert len(set(keys)) == len(keys)


def test_the_blank_template_is_what_the_parser_reads_and_stable() -> None:
    data = workbook_template()
    assert data == workbook_template()
    book = load_workbook(io.BytesIO(data), read_only=True)
    assert book.sheetnames == [name for name, _ in SHEETS]
    for name, columns in SHEETS:
        header = next(book[name].iter_rows(values_only=True))
        assert tuple(header) == columns
