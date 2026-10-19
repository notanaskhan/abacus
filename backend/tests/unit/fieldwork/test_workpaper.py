"""SPEC-037 AC-5, AC-6 (TASK-078; ADR-116): the workpaper renders deterministically; signing
fingerprints its content without the Sign-off sheet; the client's support request carries only
the entry's reference, date and amount."""

from __future__ import annotations

import io
from dataclasses import replace
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from abacus.modules.fieldwork.workpaper import Row, Signature, Workpaper, render_workpaper
from abacus.modules.fieldwork.workpapers import support_wording

STAMP = datetime(2026, 10, 19, 12, 0, tzinfo=UTC)


def _row(ref: str, result: str | None, note: str | None = None) -> Row:
    return Row(
        date(2026, 12, 30),
        "Journal Entry",
        ref,
        "",
        ("1000", "5000"),
        Decimal("43210.09"),
        ("near_year_end", "keywords"),
        None,
        result,
        note,
        ("upload v1",) if result == "supported" else (),
        result is None,
    )


DOC = Workpaper(
    client_name="Acme Holdings",
    engagement_name="FY2026 audit",
    year_end=date(2026, 12, 31),
    overall=Decimal(500_000),
    performance=Decimal(375_000),
    trivial=Decimal(25_000),
    benchmark="1% of revenue",
    definition="journal_entry_testing",
    definition_version=1,
    gl_snapshot_id="db09e01d-72a7-4174-a70d-a3a2e92b15bb",
    entry_count=380_036,
    line_count=760_072,
    total_debit=Decimal("123456789.10"),
    total_credit=Decimal("123456789.10"),
    criteria_version=2,
    random_method="seeded_v1",
    seed=989975274,
    threshold=Decimal(375_000),
    parameters=[("Random entries", "10")],
    by_rule=[("near_year_end", 2), ("random", 10)],
    rows=[
        _row("TE-YE", "supported"),
        _row("=cmd|' /C calc'!A0", "exception", "No approval"),
    ],
    conclusion="No material misstatement from management override was identified.",
    not_performed=("after_year_end", "entered_by"),
    stamp=STAMP,
)
SIGNED = [Signature("ready", "Sam Staff", "staff", STAMP, 1)]


def test_ac5_the_same_records_give_the_same_bytes() -> None:
    assert render_workpaper(DOC, SIGNED) == render_workpaper(DOC, SIGNED)
    assert render_workpaper(DOC, with_signoff=False) == render_workpaper(DOC, with_signoff=False)


def test_ac5_any_change_changes_the_bytes() -> None:
    changed = replace(DOC, conclusion="Exceptions found.")
    assert render_workpaper(changed, with_signoff=False) != render_workpaper(
        DOC, with_signoff=False
    )


def test_d1_signatures_never_change_the_signed_content() -> None:
    later = [*SIGNED, Signature("reviewed", "Dana Leader", "manager", STAMP, 1)]
    content = render_workpaper(DOC, with_signoff=False)
    assert content == render_workpaper(DOC, later, with_signoff=False)
    assert render_workpaper(DOC, SIGNED) != render_workpaper(DOC, later)


def test_ac5_its_sheets_hold_the_summary_selections_exceptions_and_sign_off() -> None:
    book = load_workbook(io.BytesIO(render_workpaper(DOC, SIGNED)))
    assert book.sheetnames == ["Summary", "Selections", "Exceptions", "Sign-off"]
    summary = [c for row in book["Summary"].iter_rows(values_only=True) for c in row if c]
    assert "Agrees to the ledger snapshot (lines and debits)" in summary
    assert DOC.conclusion in summary
    assert book["Selections"].max_row == 3
    assert book["Exceptions"].max_row == 2
    assert book["Sign-off"]["B2"].value == "Sam Staff"


def test_ac5_text_never_becomes_a_formula() -> None:
    book = load_workbook(io.BytesIO(render_workpaper(DOC, SIGNED)))
    cell = book["Exceptions"]["C2"]
    assert cell.data_type == "s" and cell.value == "=cmd|' /C calc'!A0"


def test_ac6_the_support_request_carries_reference_date_and_amount_only() -> None:
    wording = support_wording("TE-YE", "2026-12-30", Decimal("43210.09"))
    assert wording == (
        "Support for journal entry TE-YE of 2026-12-30, 43,210.09: the approved entry and "
        "the documents behind it"
    )
    assert support_wording("", "2026-12-30", Decimal(1)).startswith(
        "Support for the journal entry of 2026-12-30, 1.00"
    )
