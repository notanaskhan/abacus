"""A-03-T10 (TASK-2614; A-03 §3 F3; SPEC-580; ADR-340, ADR-042; DEC-1662): the sample-set
service's pure parts and the size table's deterministic function."""

from __future__ import annotations

import io
import uuid
import zipfile
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from abacus.kernel._statuses import SampleSetVersionStatus
from abacus.modules.methodology.api import (
    AttributeSampleSizesPayload,
    PolicyRef,
)
from abacus.modules.methodology.standards_contract import (
    StandardRowMissing,
)
from abacus.modules.methodology.standards_contract import (
    attribute_sample_size as size_from_table,
)
from abacus.modules.sampling.compute import Record, StratumRule, canonical, sampler
from abacus.modules.sampling.contract import (
    CitationOut,
    ParametersOut,
    PopulationOut,
    PopulationSourceOut,
    SampleItemOut,
    SampleSetOut,
    StratumIn,
)
from abacus.modules.sampling.sample_sets import (
    inputs_hash,
    seed_hex,
    stratified_size,
    tie_summary,
    workbook,
)

POLICY = PolicyRef(
    standard_key="bh-ebp-210-table-2", version_id=uuid.uuid4(), version=1, name="BH-EBP-210"
)
TABLE = AttributeSampleSizesPayload.model_validate(
    {
        "kind": "attribute_sample_sizes",
        "confidence_pct": "90",
        "rows": [
            {"rmm": "low", "tolerable_pct": "10", "expected_pct": "0", "size": 15},
            {"rmm": "moderate", "tolerable_pct": "10", "expected_pct": "0", "size": 25},
            {"rmm": "high", "tolerable_pct": "10", "expected_pct": "0", "size": 40},
        ],
    }
)


def test_ac_the_firm_table_gives_25_for_moderate_at_90_10_0() -> None:
    found = size_from_table(
        TABLE,
        POLICY,
        rmm="moderate",
        tolerable_pct=Decimal("10.0"),
        expected_pct=Decimal("0"),
        confidence_pct=Decimal("90.00"),
    )
    assert (found.size, found.policy) == (25, POLICY)


def test_a_question_the_table_does_not_answer_is_refused_never_guessed() -> None:
    with pytest.raises(StandardRowMissing):
        size_from_table(
            TABLE,
            POLICY,
            rmm="moderate",
            tolerable_pct=Decimal("5"),
            expected_pct=Decimal("0"),
            confidence_pct=Decimal("90"),
        )
    with pytest.raises(StandardRowMissing):
        size_from_table(
            TABLE,
            POLICY,
            rmm="moderate",
            tolerable_pct=Decimal("10"),
            expected_pct=Decimal("0"),
            confidence_pct=Decimal("95"),
        )


def test_the_seed_shows_as_six_hex_digits() -> None:
    assert seed_hex(0x5E2A91) == "0x5E2A91"
    assert seed_hex(7) == "0x000007"


def test_the_recorded_inputs_hash_is_the_samplers_own() -> None:
    population = canonical(Record(f"NW-{n}") for n in range(412))
    drawn = sampler("2.3").select(population, "attribute_nonstatistical", seed=0x5E2A91, size=25)
    assert drawn.inputs_hash == inputs_hash(
        "attribute_nonstatistical", 0x5E2A91, population.population_hash, 25, []
    )
    amounts = [Decimal(n * 100) for n in range(1, 41)]
    gl = canonical(Record(f"L{n:07d}", a) for n, a in enumerate(amounts))
    strata = [
        StratumIn(min_amount=Decimal(3000), test_rule="all"),
        StratumIn(
            min_amount=Decimal(1000), max_amount=Decimal(3000), test_rule="random", random_count=5
        ),
        StratumIn(max_amount=Decimal(1000), test_rule="none"),
    ]
    rules = [StratumRule(s.min_amount, s.max_amount, s.test_rule, s.random_count) for s in strata]
    stratified = sampler("2.3").select(
        gl, "threshold_random_stratified", seed=0x7C31B4, strata=rules
    )
    assert stratified.inputs_hash == inputs_hash(
        "threshold_random_stratified", 0x7C31B4, gl.population_hash, None, strata
    )
    assert stratified_size(amounts, strata) == len(stratified.items) == 11 + 5


def test_strata_that_do_not_partition_the_population_are_refused() -> None:
    from abacus.modules.sampling.sample_sets import ParametersInvalid

    overlapping = [
        StratumIn(min_amount=Decimal(100), test_rule="all"),
        StratumIn(max_amount=Decimal(200), test_rule="none"),
    ]
    with pytest.raises(ParametersInvalid):
        stratified_size([Decimal(150)], overlapping)
    with pytest.raises(ParametersInvalid):
        stratified_size([None], [StratumIn(test_rule="all")])


def test_ties_read_as_one_word_by_code() -> None:
    assert tie_summary([]) == ("check", None)
    assert tie_summary([{"result": "tie", "difference": None}] * 2) == ("tie", None)
    off = [
        {"result": "tie", "difference": None},
        {"result": "off", "difference": Decimal("-182300.00")},
    ]
    assert tie_summary(off) == ("off", Decimal("-182300.00"))
    assert tie_summary([{"result": "check", "difference": None}]) == ("check", None)


def _view(title: str = "Participants") -> SampleSetOut:
    at = datetime(2026, 9, 16, 21, 5, tzinfo=UTC)
    population = PopulationOut(
        id=uuid.uuid4(),
        engagement_id=uuid.uuid4(),
        work_item_id=uuid.uuid4(),
        source=PopulationSourceOut(kind="upload"),
        record_count=412,
        record_unit="participants",
        amount_total=Decimal("14560000.00"),
        sha256="a" * 64,
        sampling_unit="participant",
        deviation_definition=None,
        ties=[],
        created_by_name="Dana Kowalski",
        created_at=at,
    )
    return SampleSetOut(
        id=uuid.uuid4(),
        engagement_id=population.engagement_id,
        work_item_id=population.work_item_id,
        title=title,
        purpose=None,
        program_step_key=None,
        version_no=2,
        status=SampleSetVersionStatus.DRAFT,
        population=population,
        parameters=ParametersOut(
            method="attribute_nonstatistical",
            rmm="moderate",
            tolerable_rate=Decimal("0.10"),
            expected_rate=Decimal("0"),
            confidence=Decimal("0.90"),
            computed_size=25,
            statistical_crosscheck_size=22,
            override_size=None,
            override_rationale=None,
            final_size=25,
            seed=0x5E2A91,
            seed_hex="0x5E2A91",
            seed_source="server",
            sampler_version="2.3",
            citations=[CitationOut(label="BH-EBP-210 v1", citation_id=None, verified=True)],
        ),
        strata=[],
        items_total=1,
        items_amount=None,
        versions=[],
        inputs_hash="b" * 64,
        items_hash="c" * 64,
        content_hash="d" * 64,
    )


def test_ac_the_export_is_byte_identical_and_keeps_formula_text_as_text() -> None:
    items = [
        SampleItemOut(
            position=1,
            record_key="NW-1158",
            name='=HYPERLINK("http://x")',
            hire_date=None,
            amount=Decimal("82000.00"),
            selection_reason="random",
            stratum_position=None,
        )
    ]
    first = workbook(_view(), items)
    assert first == workbook(_view(), items)
    with zipfile.ZipFile(io.BytesIO(first)) as z:
        assert all(i.date_time == (1980, 1, 1, 0, 0, 0) for i in z.infolist())
        sheet = z.read("xl/worksheets/sheet2.xml").decode()
    assert "<f>" not in sheet
    book = load_workbook(io.BytesIO(first))
    cover = {r[0].value: r[1].value for r in book["Sample"].iter_rows()}
    assert cover["Seed"] == "0x5E2A91"
    assert cover["Population SHA-256"] == "a" * 64
    assert cover["Sampler version"] == "2.3"
    assert book["Items"]["C2"].value == '=HYPERLINK("http://x")'
    assert book["Items"]["C2"].data_type == "s"


async def test_ac_only_a_synthetic_firm_takes_a_scenario_seed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from abacus.kernel.db import TenantContext
    from abacus.modules.sampling import sample_sets

    tenant = TenantContext(uuid.uuid4(), "human", "u")
    synthetic = {"value": False}

    async def is_synthetic_firm(given: TenantContext) -> bool:
        return synthetic["value"]

    monkeypatch.setattr(sample_sets, "is_synthetic_firm", is_synthetic_firm)
    with pytest.raises(sample_sets.SeedNotAllowed):
        await sample_sets._seed(tenant, 0x5E2A91)  # pyright: ignore[reportPrivateUsage] -- the rule
    synthetic["value"] = True
    assert await sample_sets._seed(tenant, 0x5E2A91) == (0x5E2A91, "scenario")  # pyright: ignore[reportPrivateUsage] -- the rule
    seed, source = await sample_sets._seed(tenant, None)  # pyright: ignore[reportPrivateUsage] -- the rule
    assert source == "server" and 0 <= seed < 2**24
    with pytest.raises(sample_sets.SeedNotAllowed):
        await sample_sets._seed(tenant, 2**24)  # pyright: ignore[reportPrivateUsage] -- the rule


@pytest.mark.parametrize(
    ("confidence", "tolerable", "expected", "field"),
    [
        ("0.999", "0.05", "0", "confidence"),
        ("0.95", "0.005", "0", "tolerable_rate"),
        ("0.95", "0.05", "0.03", "expected_rate"),
    ],
)
async def test_parameters_outside_the_bounded_search_are_refused_before_any_work(
    confidence: str, tolerable: str, expected: str, field: str
) -> None:
    """Security review: 0.999 / 0.5% / 0.4% runs the size search for minutes; such a request is
    refused (422 `parameters_invalid`) before the search starts."""
    from abacus.kernel.db import TenantContext
    from abacus.modules.sampling import sample_sets
    from abacus.modules.sampling.contract import VersionIn

    body = VersionIn.model_validate(
        {
            "population_id": str(uuid.uuid4()),
            "method": "attribute_statistical",
            "parameters": {
                "confidence": confidence,
                "tolerable_rate": tolerable,
                "expected_rate": expected,
            },
        }
    )
    with pytest.raises(sample_sets.ParametersInvalid) as refused:
        await sample_sets._size(TenantContext(uuid.uuid4(), "human", "u"), body, [])  # pyright: ignore[reportPrivateUsage] -- the guard
    assert refused.value.facts == {"field": field}


async def test_more_than_twenty_strata_are_refused() -> None:
    from abacus.kernel.db import TenantContext
    from abacus.modules.sampling import sample_sets
    from abacus.modules.sampling.contract import VersionIn

    body = VersionIn.model_validate(
        {
            "population_id": str(uuid.uuid4()),
            "method": "threshold_random_stratified",
            "parameters": {},
            "strata": [{"min_amount": str(n), "test_rule": "none"} for n in range(21)],
        }
    )
    with pytest.raises(sample_sets.ParametersInvalid):
        await sample_sets._size(TenantContext(uuid.uuid4(), "human", "u"), body, [])  # pyright: ignore[reportPrivateUsage] -- the guard
