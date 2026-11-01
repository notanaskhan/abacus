"""A-05-T09 (TASK-2906; SPEC-641; DEC-1954): the EBP designed workpapers' cell addresses.

Row acceptance (export half): "every seeded anchor resolves to exactly one cell in the XLSX
export." Each anchor of each registered layout parses, resolves on its own sheets, and the
export writes the cell it names at that address on that sheet.
"""

from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import load_workbook

from abacus.modules.assurance.ebp_designed import DESIGNED
from abacus.modules.workpapers.contract import WorkpaperCellOut, WorkpaperRecordOut
from abacus.modules.workpapers.core import workbook
from abacus.modules.workpapers.layouts import AddressInvalid, parse, resolve


def test_addresses_parse_quoted_titles_and_ranges() -> None:
    assert parse("EBP-04!C14") == ("EBP-04", ("C14",))
    assert parse("'EBP-04 Participant data'!F12:F15") == (
        "EBP-04 Participant data",
        ("F12", "F13", "F14", "F15"),
    )
    assert parse("EBP-04!A1:B2") == ("EBP-04", ("A1", "B1", "A2", "B2"))
    for bad in ("EBP 04!C14", "EBP-04!C0", "EBP-04!C14:B2", "EBP-04", "'x'!AAA1"):
        with pytest.raises(AddressInvalid):
            parse(bad)


def test_nine_designed_papers_with_unique_codes_kinds_and_section_keys() -> None:
    assert [p.code for p in DESIGNED] == [f"EBP-0{n}" for n in range(1, 10)]
    assert len({p.kind for p in DESIGNED}) == 9
    assert len({p.section_key for p in DESIGNED}) == 9


def _cell(sheet: str, key: str, value: str) -> WorkpaperCellOut:
    return WorkpaperCellOut.model_validate(
        {
            "sheet_key": sheet,
            "cell_key": key,
            "section_key": None,
            "row_key": None,
            "col_key": None,
            "value": {"t": "text", "v": value},
            "value_source": "human",
            "value_state": "human",
            "suggestion_id": None,
            "tickmarks": [],
            "citation_ids": [],
            "updated_version_no": 1,
        }
    )


@pytest.mark.parametrize("paper", [p for p in DESIGNED if p.anchors], ids=lambda p: p.code)
def test_every_seeded_anchor_resolves_to_exactly_one_cell_in_the_export(paper: object) -> None:
    from abacus.modules.assurance.ebp_designed import DesignedPaper

    assert isinstance(paper, DesignedPaper)
    layout = paper.layout()
    cells: list[WorkpaperCellOut] = []
    expected: dict[tuple[str, str], str] = {}
    for name, address in paper.anchors.items():
        for n, (sheet_key, key) in enumerate(resolve(layout, address)):
            value = f"{name}-{n}"
            cells.append(_cell(sheet_key, key, value))
            expected[(sheet_key, key)] = value
    view = WorkpaperRecordOut.model_validate(
        {
            "id": str(uuid.uuid4()),
            "code": paper.code,
            "kind": paper.kind,
            "title": paper.title,
            "audit_area": None,
            "program_step_keys": list(paper.program_refs),
            "version_no": 1,
            "current_version_no": 1,
            "content_hash": "0" * 64,
            "tab": None,
            "sections": [],
            "cells": [c.model_dump() for c in cells],
            "flags": [],
            "sign_offs": [],
            "allowed_actions": [],
        }
    )
    book = load_workbook(io.BytesIO(workbook(view, {})))
    titles = {s.sheet_key: s.title for s in layout.sheets}
    for (sheet_key, key), value in expected.items():
        sheet = book[titles[sheet_key]]
        assert sheet[key].value == value, (paper.code, sheet_key, key)
    # Exactly one: no other cell on the paper's sheets holds an anchor's value.
    placed = [
        str(c.value)
        for title in titles.values()
        for row in book[title].iter_rows()
        for c in row
        if c.value is not None
    ]
    assert sorted(placed) == sorted(expected.values())


def test_the_dec_26_business_day_cell_is_f31_on_ebp_04() -> None:
    paper = next(p for p in DESIGNED if p.code == "EBP-04")
    assert resolve(paper.layout(), paper.anchors["dec_26_business_days"]) == (("main", "F31"),)
