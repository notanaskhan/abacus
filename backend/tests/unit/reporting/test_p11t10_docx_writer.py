"""P-11-T10 (TASK-2411; SPEC-520 AC-20; P-11 §4.4; ADR-320 §3): the DOCX writer.

The row's acceptance: the golden letter and report fixtures produce byte-identical DOCX across
runs, and their text equals the PDF's block text (the strings `build_html` places, which the
PDF prints). Also: fixed core and app properties, the DRAFT mark on a draft and none on a final,
the legend in the footer, clause states honoured, figures right-aligned, and the PDF's refusals.
No Pango is needed: python-docx and the pure HTML build only.
"""

from __future__ import annotations

import io
import uuid
import zipfile
from collections.abc import Callable
from datetime import UTC, datetime

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pydantic import TypeAdapter

from abacus.modules.reporting.contract import DocBlock
from abacus.modules.reporting.render.core import RenderInput, mark_for, render_docx
from abacus.modules.reporting.render.docx import docx_texts
from abacus.modules.reporting.render.head import RenderRefused
from abacus.modules.reporting.render.html import Letterhead, MissingVariable, document_html
from abacus.modules.reporting.render.layouts import LAYOUTS

AS_OF = datetime(2026, 9, 14, 13, 32, tzinfo=UTC)
Json = dict[str, object]
BLOCKS = TypeAdapter(list[DocBlock])
CITE = uuid.UUID("00000000-0000-4000-8000-000000000754")


def _text(text: str) -> Json:
    return {"type": "text", "text": text}


def _var(key: str) -> Json:
    return {"type": "var", "key": key}


def _para(*runs: Json) -> Json:
    return {"type": "paragraph", "runs": list(runs)}


LETTER: list[object] = [
    _para(_var("letter.date")),
    _para(_var("client.signer")),
    {"type": "heading", "level": 1, "text": "Re: Cost segregation and section 754 analysis"},
    {
        "type": "section",
        "id": "scope",
        "title": "1. Scope of services",
        "blocks": [
            {
                "type": "clause",
                "clause_id": "scope",
                "enabled": True,
                "non_standard": False,
                "blocks": [
                    _para(
                        _text("We will analyze whether the partnership should elect under "),
                        _text("section 754 for the June 30, 2025 transfer."),
                    )
                ],
            }
        ],
    },
    {
        "type": "section",
        "id": "timing",
        "title": "3. Timing and the section 754 deadline",
        "blocks": [
            {
                "type": "clause",
                "clause_id": "timing",
                "enabled": True,
                "non_standard": True,
                "blocks": [
                    _para(
                        _text("A section 754 election is made with the 2025 Form 1065, due "),
                        _var("deadline.date"),
                        _text("."),
                        {"type": "cite", "citation_id": str(CITE)},
                    )
                ],
            }
        ],
    },
    {
        "type": "section",
        "id": "reliance",
        "title": "5. Reliance on the engineer's study",
        "blocks": [
            {
                "type": "clause",
                "clause_id": "reliance",
                "enabled": True,
                "non_standard": True,
                "blocks": [_para(_text("The cost allocations come from the engineer's study."))],
            }
        ],
    },
    {
        "type": "section",
        "id": "fees",
        "title": "6. Fees and billing",
        "blocks": [_para(_text("A fixed fee of "), _var("fee.total"), _text("."))],
    },
    {
        "type": "signature_block",
        "signers": [
            {"label": "Partner", "name": [_var("firm.signer")]},
            {"label": "Managing Member", "name": [_text("Owen Calder")]},
        ],
    },
    {"type": "enclosures", "items": ["Consent to disclose tax return information"]},
]
VALUES = {
    "letter.date": "September 14, 2026",
    "client.signer": "Owen Calder, Managing Member, Ridgeline GP LLC",
    "deadline.date": "September 15, 2026",
    "fee.total": "$18,500",
    "firm.signer": "Daniel Reyes",
}
REPORT: list[object] = [
    {"type": "legend", "text": "No assurance is provided on these financial statements."},
    {"type": "heading", "level": 1, "text": "Copperleaf Coffee Roasters"},
    {
        "type": "section",
        "id": "balance",
        "title": "Balance sheet",
        "blocks": [
            {
                "type": "table",
                "columns": ["Account", "2026"],
                "rows": [["Cash", "41,200.00"], ["Receivables", "(1,050.00)"]],
            },
            {"type": "kv", "rows": [{"key": "Prepared by", "value": [_var("preparer")]}]},
        ],
    },
    {"type": "page_break"},
    {
        "type": "section",
        "id": "notes",
        "title": "Notes",
        "blocks": [{"type": "list", "ordered": True, "items": [[_text("Basis of accounting")]]}],
    },
]
LETTERHEAD = Letterhead("Brightwater Hale LLP", ("311 West Brannock Street, Chicago, IL",))


def _letter(**over: object) -> RenderInput:
    raw: dict[str, object] = {
        "layout_key": "letter",
        "blocks": LETTER,
        "values": VALUES,
        "title": "Engagement letter",
        "as_of": AS_OF,
        "mode": "final",
        "content_hash": "a" * 64,
        "letterhead": LETTERHEAD,
    }
    raw.update(over)
    return RenderInput(**raw)  # type: ignore[arg-type]  # test fixture assembled from a dict


def _report(**over: object) -> RenderInput:
    raw: dict[str, object] = {
        "layout_key": "package",
        "blocks": REPORT,
        "values": {"preparer": "Priya Raman"},
        "title": "Financial statements",
        "as_of": AS_OF,
        "mode": "final",
        "content_hash": "b" * 64,
    }
    raw.update(over)
    return RenderInput(**raw)  # type: ignore[arg-type]  # test fixture assembled from a dict


@pytest.mark.parametrize("fixture", [_letter, _report], ids=["letter", "report"])
def test_ac1_the_golden_fixtures_are_byte_identical_across_runs(
    fixture: Callable[[], RenderInput],
) -> None:
    first = render_docx(fixture())
    second = render_docx(fixture())
    assert first.docx == second.docx
    assert first.render_hash == second.render_hash


@pytest.mark.parametrize("fixture", [_letter, _report], ids=["letter", "report"])
def test_ac1_the_docx_text_equals_the_pdfs_block_text(
    fixture: Callable[[], RenderInput],
) -> None:
    request = fixture()
    html = document_html(
        LAYOUTS[request.layout_key],
        BLOCKS.validate_python(request.blocks),
        request.values,
        title=request.title,
        as_of=request.as_of,
        mark=mark_for(request.mode),
        letterhead=request.letterhead,
    )
    out = render_docx(request)
    assert out.texts == html.texts
    placed = "\n".join([*docx_texts(out.docx), _footer(out.docx)])
    assert [t for t in html.texts if t not in placed] == []


def _footer(docx: bytes) -> str:
    document = Document(io.BytesIO(docx))
    return "\n".join(p.text for p in document.sections[0].footer.paragraphs)


def _header(docx: bytes) -> str:
    document = Document(io.BytesIO(docx))
    return "\n".join(p.text for p in document.sections[0].header.paragraphs)


def test_fixed_core_and_app_properties_and_zip_entries() -> None:
    out = render_docx(_letter())
    document = Document(io.BytesIO(out.docx))
    props = document.core_properties
    assert (props.author, props.last_modified_by, props.revision) == ("", "", 1)
    assert props.identifier == "a" * 64
    assert props.created == AS_OF
    with zipfile.ZipFile(io.BytesIO(out.docx)) as package:
        infos = package.infolist()
        assert [i.filename for i in infos] == sorted(i.filename for i in infos)
        assert {i.date_time for i in infos} == {(1980, 1, 1, 0, 0, 0)}
        app = package.read("docProps/app.xml")
    assert b"Application" not in app
    assert b"Accrued" not in out.docx
    assert b"abacus" not in out.docx


def test_a_draft_carries_the_mark_and_a_final_none() -> None:
    assert _header(render_docx(_letter(mode="draft")).docx) == "DRAFT"
    assert _header(render_docx(_letter(mode="sample")).docx) == "SAMPLE"
    assert _header(render_docx(_letter()).docx) == ""


def test_the_legend_is_the_footer_and_figures_align_right() -> None:
    out = render_docx(_report())
    assert _footer(out.docx) == "No assurance is provided on these financial statements."
    table = Document(io.BytesIO(out.docx)).tables[0]
    assert table.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert table.rows[1].cells[0].paragraphs[0].alignment is None


def test_a_clause_switched_off_is_left_out_and_the_letterhead_leads() -> None:
    texts = docx_texts(render_docx(_letter(clause_states={"reliance": False})).docx)
    assert "The cost allocations come from the engineer's study." not in texts
    assert texts[:2] == ["Brightwater Hale LLP", "311 West Brannock Street, Chicago, IL"]
    assert "A fixed fee of $18,500." in texts


def test_a_docx_refuses_as_the_pdf_does() -> None:
    with pytest.raises(MissingVariable):
        render_docx(_letter(values={k: v for k, v in VALUES.items() if k != "fee.total"}))
    with pytest.raises(RenderRefused):
        render_docx(_letter(layout_key="nope"))
    with pytest.raises(RenderRefused):
        render_docx(_letter(blocks=[{"type": "script", "text": "x"}]))
