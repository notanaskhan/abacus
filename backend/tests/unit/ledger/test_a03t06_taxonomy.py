"""A-03-T06 (TASK-2506; SPEC-580 AC-17; A-03 F5 US-5.1): the taxonomy from a workbook.

A three-sheet workbook in BH FS groupings v1's shape imports with zero problems; an overlapping
range returns its row problem by sheet, row and code; unknown references and bad classes are
named; a formula-looking cell is text; ledger keeps an empty dependency set and imports no module.
Synthetic probe content only (AGENTS 16): BH's own set is loaded through the endpoint by T16.
"""

from __future__ import annotations

import ast
import io
import sys
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell

from abacus.modules.ledger.api import parse_grouping_set, parse_grouping_workbook
from abacus_tools.quality.banned_patterns import MODULE_DEPENDENCIES

LEDGER = Path(__file__).resolve().parents[3] / "src/abacus/modules/ledger"
GROUPINGS = [
    ("grouping_key", "grouping_label", "class"),
    ("cash", "Cash and equivalents", "assets"),
    ("ppe", "Property, plant & equipment", "assets"),
    ("payables", "Accounts payable", "liabilities"),
    ("revenue", "Revenue", "revenue"),
]
LINES = [
    ("fs_line_key", "fs_line_label", "grouping_key", "normal_balance"),
    ("cash_operating", "Operating cash", "cash", "debit"),
    ("ppe_deposits", "Deposits on equipment", "ppe", "debit"),
    ("ap_trade", "Trade payables", "payables", "credit"),
    ("sales", "Net sales", "revenue", "credit"),
]
RANGES = [
    ("account_from", "account_to", "fs_line_key"),
    ("1000", "1099", "cash_operating"),
    ("1400", "1499", "ppe_deposits"),
    ("2000", "2099", "ap_trade"),
    (4000, 4999, "sales"),  # numbers, as Excel stores them
]


def _book(
    groupings: Sequence[tuple[object, ...]] = GROUPINGS,
    lines: Sequence[tuple[object, ...]] = LINES,
    ranges: Sequence[tuple[object, ...]] = RANGES,
    *,
    drop: str | None = None,
) -> bytes:
    book = Workbook()
    first = book.active
    assert first is not None
    book.remove(first)
    for name, rows in (("Groupings", groupings), ("FS lines", lines), ("Number ranges", ranges)):
        if name == drop:
            continue
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_ac17_a_clean_workbook_imports_with_zero_problems() -> None:
    parsed = parse_grouping_workbook(_book())
    assert parsed.problems == []
    assert [g.key for g in parsed.groupings] == ["cash", "ppe", "payables", "revenue"]
    assert {line.key: line.normal_balance for line in parsed.lines}["ap_trade"] == "credit"
    assert ("4000", "4999", "sales") in parsed.ranges
    assert len(parsed.payload_hash) == 64
    assert parse_grouping_workbook(_book()).payload_hash == parsed.payload_hash


def test_ac17_an_overlapping_range_returns_its_row_problem() -> None:
    parsed = parse_grouping_workbook(_book(ranges=[*RANGES, ("1050", "1060", "cash_operating")]))
    assert [(p.sheet, p.row, p.code) for p in parsed.problems] == [
        ("Number ranges", 6, "overlapping_range")
    ]


def test_ac17_unknown_references_bad_classes_and_duplicates_are_named() -> None:
    parsed = parse_grouping_workbook(
        _book(
            groupings=[*GROUPINGS, ("misc", "Misc", "asset"), ("cash", "Cash again", "assets")],
            lines=[*LINES, ("orphan", "Orphan", "nowhere", "debit"), ("x", "X", "cash", "up")],
            ranges=[*RANGES, ("5000", "5999", "missing_line"), ("7000", "6000", "sales")],
        )
    )
    found = {(p.sheet, p.row, p.code) for p in parsed.problems}
    assert ("Groupings", 6, "class_invalid") in found
    assert ("Groupings", 7, "duplicate_key") in found
    assert ("FS lines", 6, "unknown_grouping") in found
    assert ("FS lines", 7, "normal_balance_invalid") in found
    assert ("Number ranges", 6, "unknown_fs_line") in found
    assert ("Number ranges", 7, "range_invalid") in found


def test_ac17_a_missing_sheet_a_bad_header_and_an_unreadable_file() -> None:
    missing = parse_grouping_workbook(_book(drop="FS lines"))
    assert [(p.sheet, p.code) for p in missing.problems] == [("FS lines", "sheet_missing")]
    header = parse_grouping_workbook(_book(groupings=[("key", "label", "class")]))
    assert ("Groupings", 1, "header") in {(p.sheet, p.row, p.code) for p in header.problems}
    assert [p.code for p in parse_grouping_workbook(b"not a workbook").problems] == [
        "workbook_unreadable"
    ]


def test_a_formula_looking_cell_is_kept_as_text() -> None:
    # A text cell that begins with "=" (Excel stores it as a string, not a formula).
    book = Workbook()
    first = book.active
    assert first is not None
    book.remove(first)
    for name, rows in (
        ("Groupings", [*GROUPINGS, ("x", "placeholder", "expenses")]),
        ("FS lines", LINES),
        ("Number ranges", RANGES),
    ):
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    label = book["Groupings"].cell(row=len(GROUPINGS) + 1, column=2)
    assert isinstance(label, Cell)
    label.value = "=HYPERLINK(1)"
    label.data_type = "s"
    out = io.BytesIO()
    book.save(out)
    parsed = parse_grouping_workbook(out.getvalue())
    assert parsed.problems == []
    assert {g.key: g.label for g in parsed.groupings}["x"] == "=HYPERLINK(1)"


def test_ac17_the_csv_path_reports_the_overlap_on_its_row_too() -> None:
    text = (
        "grouping_key,grouping_label,class,fs_line_key,fs_line_label,normal_balance,"
        "account_from,account_to\n"
        "cash,Cash,assets,cash,Cash,debit,1000,1099\n"
        "cash,Cash,assets,cash2,Cash 2,debit,1050,1060\n"
    )
    assert [(p.row, p.code) for p in parse_grouping_set(text).problems] == [
        (3, "overlapping_range")
    ]


def test_ac17_ledger_keeps_an_empty_dependency_set_and_imports_no_module() -> None:
    assert MODULE_DEPENDENCIES["ledger"] == frozenset()
    for path in sorted(LEDGER.rglob("*.py")):
        for node in ast.walk(ast.parse(path.read_text(encoding="utf-8"))):
            if isinstance(node, ast.ImportFrom) and node.module and node.level == 0:
                name = node.module
            elif isinstance(node, ast.Import):
                name = node.names[0].name
            else:
                continue
            if name.startswith("abacus.modules."):
                assert name.startswith("abacus.modules.ledger"), (path, name)
            assert not name.startswith("abacus_tools"), (path, name)
            assert name.split(".")[0] in {
                *sys.stdlib_module_names,
                "abacus",
                "sqlalchemy",
                "openpyxl",
                "pydantic",
                "__future__",
                "temporalio",
            }, (
                path,
                name,
            )
