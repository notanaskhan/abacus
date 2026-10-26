"""The engagement export package builder, offline (A-01-T19, TASK-1335; SPEC-344 AC-9 to AC-11,
AC-13; ADR-042, ADR-052): the same inputs give the same bytes, hostile names become safe paths and
literal cells, the checksums verify as `shasum -c` would, and a core section with no provider
stops the build before anything is read."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import uuid
import zipfile
from collections.abc import Sequence
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook

from abacus.kernel.db import TenantContext
from abacus.modules.evidence import export_package
from abacus.modules.evidence.export_activities import ExportBuildFailed, section_files
from abacus.modules.evidence.export_builder import (
    AnyFile,
    BuiltPackage,
    PackageFacts,
    PackageFile,
    PackageIntegrityError,
    PackageInvalid,
    StreamedFile,
    build_package,
)
from abacus.modules.evidence.export_manifest import (
    INDEX_COLUMNS,
    PROVENANCE_FILES,
    EngagementExportManifest,
    ManifestClient,
    ManifestEngagement,
    ManifestEvidence,
    ManifestPerson,
    ManifestScope,
    archive_path_problem,
)
from abacus.modules.evidence.export_sections import literal_cell, section_csv
from abacus.modules.evidence.export_selection import deduplicated, file_segment, segment
from abacus.modules.evidence.workflow_types import SECTION_MISSING

GENERATED = datetime(2026, 9, 14, 14, 40, tzinfo=UTC)
BINDER = uuid.UUID(int=7)
HOSTILE_FORMULA = '=HYPERLINK("http://x","y").pdf'
HOSTILE_TRAVERSAL = "../../etc/passwd"


def _facts(**overrides: object) -> PackageFacts:
    values: dict[str, object] = {
        "export_id": uuid.UUID(int=1),
        "kind": "binder",
        "generated_at": GENERATED,
        "as_of_seq": 48211,
        "tenant_id": uuid.UUID(int=2),
        "firm_name": "Brightwater Hale LLP",
        "client": ManifestClient(client_id=uuid.UUID(int=3), display_name="Northwind Labs, Inc."),
        "engagement": ManifestEngagement(
            engagement_id=uuid.UUID(int=4),
            work_item_id=uuid.UUID(int=5),
            display_name="Northwind FY2026 audit",
            period_end=date(2026, 6, 30),
        ),
        "requested_by": ManifestPerson(user_id=uuid.UUID(int=6), display_name="Hannah Leclerc"),
        "requested_at": datetime(2026, 9, 14, 14, 39, 50, tzinfo=UTC),
        "scope": ManifestScope(
            binder_ids=[BINDER], include_unaccepted=False, include_audit_trail=True
        ),
        "request_count": 2,
        "requests_without_file": 1,
    }
    return PackageFacts(**{**values, **overrides})  # type: ignore[arg-type]  # test table


def _evidence(path: str, content: bytes) -> PackageFile:
    return PackageFile(
        path,
        content,
        "evidence",
        evidence=ManifestEvidence(
            evidence_version_id=uuid.uuid5(BINDER, path),
            request_item_id=uuid.UUID(int=9),
            binder_id=BINDER,
            version_no=1,
            review_state="accepted",
            media_type="application/pdf",
        ),
    )


def _files() -> list[PackageFile]:
    index = section_csv(
        INDEX_COLUMNS,
        [
            ("01", "Cash", "01", HOSTILE_FORMULA, "Received ✓", "", None, None, 1, "p", "", None),
            (
                "01",
                "Cash",
                "02",
                "Bank statements",
                "Requested",
                "Dana",
                None,
                None,
                0,
                "",
                "",
                None,
            ),
        ],
    )
    return [
        PackageFile("index.csv", index, "index"),
        *(PackageFile(p, b"h\r\n", "provenance") for p in PROVENANCE_FILES),
        PackageFile("audit_trail/events.jsonl", b'{"seq":1}\n', "audit_trail"),
        _evidence("evidence/01 Cash/01 Statement/v1__statement.pdf", b"%PDF-1.7 one"),
    ]


def _build(files: Sequence[AnyFile] | None = None) -> tuple[bytes, BuiltPackage]:
    out = io.BytesIO()
    built = build_package(_facts(), files if files is not None else _files(), out)
    return out.getvalue(), built


def _streamed(f: PackageFile, read: bytes | None = None) -> StreamedFile:
    """`f` as the activity hands it over: named by its digest, read when written."""

    def read_it() -> bytes:
        return f.content if read is None else read

    return StreamedFile(
        f.path,
        hashlib.sha256(f.content).hexdigest(),
        len(f.content),
        f.role,
        read_it,
        evidence=f.evidence,
    )


# --- streaming (security review M4) --------------------------------------------------------------


def test_a_streamed_original_builds_the_same_bytes_as_one_in_memory() -> None:
    files = _files()
    streamed = [_streamed(f) if f.role == "evidence" else f for f in files]
    assert _build(streamed)[0] == _build(files)[0]


def test_a_streamed_original_that_differs_from_its_fingerprint_stops_the_build() -> None:
    files = _files()
    for wrong in (b"%PDF-1.7 two", b"%PDF-1.7 one!"):  # same size, then a different size
        streamed = [_streamed(f, wrong) if f.role == "evidence" else f for f in files]
        with pytest.raises(PackageIntegrityError) as failed:
            _build(streamed)
        assert failed.value.path.startswith("evidence/")


# --- AC-10: byte-identical -----------------------------------------------------------------------


def test_ac10_two_builds_of_the_same_data_are_byte_identical() -> None:
    first, built_first = _build()
    second, built_second = _build(list(reversed(_files())))  # order given never matters
    assert first == second
    assert built_first == built_second


def test_ac10_entries_are_sorted_with_fixed_timestamps_and_no_extra_fields() -> None:
    data, _ = _build()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
    names = [i.filename for i in infos]
    assert names == sorted(names)
    assert {i.date_time for i in infos} == {(2026, 9, 14, 14, 40, 0)}
    assert all(i.extra == b"" and i.create_system == 3 for i in infos)
    assert all(i.external_attr == 0o644 << 16 for i in infos)


# --- AC-9: verifiable with shasum ----------------------------------------------------------------


def test_ac9_the_checksums_and_the_manifest_verify() -> None:
    data, built = _build()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        contents = {i.filename: archive.read(i) for i in archive.infolist()}
    for line in contents["checksums.sha256"].decode().splitlines():
        digest, path = line.split("  ", 1)
        assert hashlib.sha256(contents[path]).hexdigest() == digest
    listed = {
        line.split("  ", 1)[1] for line in contents["checksums.sha256"].decode().splitlines()
    }
    assert listed == set(contents) - {"checksums.sha256", "manifest.sha256"}
    manifest_digest = contents["manifest.sha256"].decode().split("  ")[0]
    assert manifest_digest == hashlib.sha256(contents["manifest.json"]).hexdigest()
    assert built.manifest_sha256 == manifest_digest
    manifest = EngagementExportManifest.model_validate_json(contents["manifest.json"])
    assert manifest.counts.file_count == 1
    assert {f.path for f in manifest.files} == set(contents) - {
        "manifest.json",
        "manifest.sha256",
        "checksums.sha256",
    }
    raw = json.loads(contents["manifest.json"])
    assert (
        contents["manifest.json"]
        == json.dumps(raw, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()
    )


def test_ac9_the_xlsx_holds_the_csv_rows() -> None:
    data, _ = _build()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        rows = list(csv.reader(io.StringIO(archive.read("index.csv").decode())))
        sheet = load_workbook(io.BytesIO(archive.read("index.xlsx"))).active
    assert sheet is not None
    cells = [[c.value or "" for c in row] for row in sheet.iter_rows()]
    assert cells == rows


# --- AC-11: hostile names ------------------------------------------------------------------------


def test_ac11_formula_names_are_literal_cells() -> None:
    assert literal_cell(HOSTILE_FORMULA) == "'" + HOSTILE_FORMULA
    for lead in ("=1+1", "+1", "-1", "@SUM(A1)", "\tx", "\rx"):
        assert literal_cell(lead).startswith("'")
    assert literal_cell("Bank statements") == "Bank statements"
    data, _ = _build()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sheet = load_workbook(io.BytesIO(archive.read("index.xlsx"))).active
        index = archive.read("index.csv").decode()
    assert sheet is not None
    request = sheet.cell(row=2, column=4)
    assert request.data_type == "s" and request.value == "'" + HOSTILE_FORMULA
    assert "'=HYPERLINK" in index


def test_ac11_a_tin_in_a_cell_or_a_name_is_masked() -> None:
    assert "912-70-1234" not in literal_cell("W-9 for 912-70-1234")
    assert "912-70-1234" not in segment("W-9 912-70-1234", 56)


@pytest.mark.parametrize(
    "name",
    [
        HOSTILE_TRAVERSAL,
        HOSTILE_FORMULA,
        "..",
        ".",
        "C:\\Windows\\a.pdf",
        "a\u202eexe.pdf",
        "a\u0000b.pdf",
        "cafe\u0301.pdf",
        "x" * 400 + ".pdf",
        "",
    ],
)
def test_ac11_hostile_file_names_become_safe_paths(name: str) -> None:
    path = f"evidence/01 Cash/01 {segment(name, 56)}/{file_segment(name, 3)}"
    [safe] = deduplicated([path])
    assert archive_path_problem(safe) is None
    assert ".." not in safe.split("/")


def test_ac11_duplicates_are_suffixed() -> None:
    paths = [
        "evidence/01 Cash/01 Bank/v1__a.pdf",
        "evidence/01 Cash/01 Bank/v1__A.pdf",
        "evidence/01 Cash/01 Bank/v1__a.pdf",
    ]
    assert deduplicated(paths) == [
        "evidence/01 Cash/01 Bank/v1__a.pdf",
        "evidence/01 Cash/01 Bank/v1__A~2.pdf",
        "evidence/01 Cash/01 Bank/v1__a~3.pdf",
    ]


def test_ac11_an_unsafe_path_never_reaches_a_package() -> None:
    files = [*_files(), _evidence("evidence/../x.pdf", b"x")]
    with pytest.raises(PackageInvalid):
        _build(files)


# --- AC-13: a missing core section ---------------------------------------------------------------


async def test_ac13_a_core_section_without_a_provider_fails_the_build(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    registered = dict(export_package._SECTIONS)  # pyright: ignore[reportPrivateUsage] -- the slot
    registered.pop("screening", None)
    monkeypatch.setattr(export_package, "_SECTIONS", registered)
    request = export_package.SectionRequest(
        export_id=uuid.uuid4(),
        engagement_id=uuid.uuid4(),
        work_item_id=uuid.uuid4(),
        binder_ids=(BINDER,),
        as_of_seq=1,
        include_unaccepted=False,
        requested_by_user_id=uuid.uuid4(),
        membership_id=uuid.uuid4(),
    )
    with pytest.raises(ExportBuildFailed) as failed:
        await section_files(TenantContext(uuid.uuid4(), "system", "test"), request)
    assert failed.value.code == SECTION_MISSING
