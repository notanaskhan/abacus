"""SPEC-034 AC-6 (TASK-070): the general ledger and journal listing workbooks are deterministic,
reconcile on their first sheet, split before the row limit, group transactions, and never turn
client text into formulas."""

from __future__ import annotations

import io
import uuid
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from abacus.modules.evidence.render_gl import (
    GlBalanceRow,
    GlEntryRow,
    GlHeader,
    journals,
    render_general_ledger,
    render_journal_listing,
)

HEADER = GlHeader(
    uuid.UUID(int=1),
    "Acme Ltd",
    date(2026, 1, 1),
    date(2026, 12, 31),
    datetime(2026, 10, 1, 9, 30, tzinfo=UTC),
    uuid.UUID(int=2),
    "import_qbo_gl_csv",
    "a" * 64,
)
D = Decimal
BALANCES = [
    GlBalanceRow("4000", "Sales", D(0), D(0), D("20.00"), D("-20.00"), D("-20.00"), False),
    GlBalanceRow("1000", "Checking", D("5.00"), D("20.00"), D(0), D("25.00"), D("25.00"), True),
]
NAMES = {"1000": "Checking", "4000": "Sales"}


def _entries(n: int = 4) -> list[GlEntryRow]:
    rows: list[GlEntryRow] = []
    for i in range(n):
        debit = i % 2 == 0
        rows.append(
            GlEntryRow(
                i + 1,
                "1000" if debit else "4000",
                date(2026, 3, 15),
                "Invoice",
                f"INV-{i // 2}",
                '=HYPERLINK("http://x")',
                "Sale",
                Decimal("10.00") if debit else Decimal(0),
                Decimal(0) if debit else Decimal("10.00"),
            )
        )
    return rows


def test_ac6_the_same_snapshot_gives_the_same_bytes() -> None:
    first = render_general_ledger(HEADER, BALANCES, _entries(), NAMES)
    assert first == render_general_ledger(HEADER, BALANCES, _entries(), NAMES)
    assert render_journal_listing(HEADER, _entries(), NAMES) == render_journal_listing(
        HEADER, _entries(), NAMES
    )


def test_ac6_the_first_sheet_reconciles_every_account() -> None:
    book = load_workbook(io.BytesIO(render_general_ledger(HEADER, BALANCES, _entries(), NAMES)))
    sheet = book["Reconciliation"]
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2, max_row=4)]
    assert rows[0][:2] == ["1000", "Checking"] and rows[0][8] == "Opening derived"
    assert rows[1][:2] == ["4000", "Sales"]
    assert rows[2][1] == "Total" and Decimal(str(rows[2][7])) == 0  # no difference overall


def test_ac6_client_text_stays_text() -> None:
    book = load_workbook(io.BytesIO(render_general_ledger(HEADER, BALANCES, _entries(), NAMES)))
    cell = book["Detail"]["F2"]
    assert cell.data_type == "s" and str(cell.value).startswith("=HYPERLINK")


def test_ac6_detail_continues_on_a_new_sheet_before_the_limit() -> None:
    content = render_general_ledger(HEADER, BALANCES, _entries(25), NAMES, max_rows=10)
    names = load_workbook(io.BytesIO(content), read_only=True).sheetnames
    assert names == ["Reconciliation", "Detail", "Detail (2)", "Detail (3)"]


def test_d2_transactions_group_by_date_type_and_reference() -> None:
    groups = list(journals(_entries(6)))
    assert [key[2] for key, _ in groups] == ["INV-0", "INV-1", "INV-2"]
    assert all(len(lines) == 2 for _, lines in groups)
    book = load_workbook(io.BytesIO(render_journal_listing(HEADER, _entries(6), NAMES)))
    totals = [row for row in book["Journals"].iter_rows(values_only=True) if row[5] == "Total"]
    assert len(totals) == 3 and all(row[8] in (None, "") for row in totals)


def test_d2_a_transaction_that_does_not_balance_on_its_own_is_marked() -> None:
    lone = [_entries(1)[0]]
    book = load_workbook(io.BytesIO(render_journal_listing(HEADER, lone, NAMES)))
    totals = [row for row in book["Journals"].iter_rows(values_only=True) if row[5] == "Total"]
    assert totals[0][8] == "Doesn't balance on its own"
