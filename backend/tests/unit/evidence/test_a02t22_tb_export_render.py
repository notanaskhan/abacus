"""A-02-T22 (TASK-1470; A-02 AC-8.1; ADR-042, ADR-052): Open in Excel's renderer.

Acceptance: "a re-render is byte-identical". Around it: the layer's columns, ticks and
citations as text, a provenance footer, and a hostile account name stays text."""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from abacus.modules.evidence.api import LayerExport, LayerLine, render_trial_balance_layers

PULLED = datetime(2026, 9, 11, 15, 2, tzinfo=UTC)


def _export(name: str = "Cash", *, adjusted: bool = False) -> LayerExport:
    lines = (
        LayerLine(
            "1000",
            name,
            "Cash",
            "C-1",
            Decimal("1900000.00"),
            Decimal("2140000.00"),
            None,
            None,
            None,
            "L1 agreed Marcus Delaney 2026-09-14T15:02:00+00:00",
            "4a3c6a52-6a43-4a5e-9d55-0d6a7b0f1a21",
        ),
        LayerLine(
            "2000",
            "Accounts payable",
            "Liabilities",
            "",
            None,
            Decimal("-2140000.00"),
            None,
            None,
            None,
            "",
            "",
        ),
    )
    return LayerExport(
        title="Trial balance 2026-06-30",
        layer="Final" if adjusted else "Unadjusted",
        adjusted=adjusted,
        lines=lines,
        total_boy=Decimal("1900000.00"),
        total_eoy=Decimal("0.00"),
        debit=Decimal("2140000.00"),
        credit=Decimal("2140000.00"),
        difference=Decimal("0.00"),
        pulled_at=PULLED,
        footer=(("Source", "import_netsuite_tb_csv"), ("Snapshot ID", "s-1"), ("Version", "1")),
    )


def _cells(content: bytes) -> list[list[object]]:
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_ac8_1_a_re_render_is_byte_identical() -> None:
    assert render_trial_balance_layers(_export()) == render_trial_balance_layers(_export())
    assert render_trial_balance_layers(_export()) != render_trial_balance_layers(
        _export(adjusted=True)
    )


def test_the_layer_ticks_citations_and_provenance_are_written_as_text() -> None:
    rows = _cells(render_trial_balance_layers(_export()))
    assert rows[0][0] == "Trial balance 2026-06-30 · Unadjusted"
    header = [c for c in rows[2] if c is not None]
    assert header == [
        "Account",
        "Name",
        "Area",
        "Beginning of Year",
        "End of Year",
        "WP Ref",
        "Ticks",
        "Citations",
    ]
    cash = rows[3]
    assert cash[:3] == ["1000", "Cash", "Cash"]
    assert cash[6] == "L1 agreed Marcus Delaney 2026-09-14T15:02:00+00:00"
    assert cash[7] == "4a3c6a52-6a43-4a5e-9d55-0d6a7b0f1a21"
    flat = [c for row in rows for c in row if c is not None]
    for label in ("Debits", "Credits", "Difference", "Source", "Snapshot ID", "Version"):
        assert label in flat


def test_adjusted_layers_carry_aje_rje_and_final_columns() -> None:
    rows = _cells(render_trial_balance_layers(_export(adjusted=True)))
    assert {"AJE", "RJE", "Final"} <= {c for c in rows[2] if c is not None}


def test_adr052_a_formula_like_account_name_stays_text() -> None:
    hostile = '=HYPERLINK("http://x.example","pay")'
    content = render_trial_balance_layers(_export(hostile))
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    cell = sheet.cell(row=4, column=2)
    assert cell.value == hostile and cell.data_type == "s"
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")


FULL_WIDTH = [chr(c) for c in (0xFF1D, 0xFF0B, 0xFF0D, 0xFF20)]  # = + - @, full width


@pytest.mark.parametrize("lead", ["=", "+", "-", "@", "\t", "\r", *FULL_WIDTH])
def test_review_l2_every_formula_lead_is_quoted_full_width_tab_and_cr_too(lead: str) -> None:
    content = render_trial_balance_layers(_export(f"{lead}SUM(A1:A9)"))
    sheet = load_workbook(io.BytesIO(content)).active
    assert sheet is not None
    cell = sheet.cell(row=4, column=2)
    assert cell.data_type == "s" and cell.quotePrefix is True


def test_a_plain_account_name_is_not_quoted() -> None:
    sheet = load_workbook(io.BytesIO(render_trial_balance_layers(_export()))).active
    assert sheet is not None
    assert not sheet.cell(row=4, column=2).quotePrefix
