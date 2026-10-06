"""AC-12: deterministic trial-balance rendering (TASK-009 interface contract, "Renderer"; ADR-042).

Expectations come from the contract: identical data gives byte-identical files, across calls, line
orders and time; hostile account names are never formulas; the footer carries provenance.
"""

from __future__ import annotations

import hashlib
import io
import itertools
import re
import time
import uuid
import zipfile
from datetime import UTC, date, datetime, timedelta, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from abacus.modules.evidence.api import (
    XLSX_MEDIA_TYPE,
    TrialBalance,
    TrialBalanceLine,
    render_trial_balance,
)

PULLED_AT = datetime(2026, 3, 14, 9, 26, 53, tzinfo=UTC)
ENTITY_ID = uuid.UUID("11111111-2222-3333-4444-555555555555")
SNAPSHOT_ID = uuid.UUID("aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee")
SOURCE_FINGERPRINT = "c" * 64
LINES = (
    TrialBalanceLine("4000", "Revenue", Decimal("0.00"), Decimal("1500.25")),
    TrialBalanceLine("1000", "Cash", Decimal("1200.50"), Decimal("0.00")),
    TrialBalanceLine("2000", "Accounts payable", Decimal("0.00"), Decimal("300.00")),
    TrialBalanceLine("5000", "Expenses", Decimal("599.75"), Decimal("0.00")),
    TrialBalanceLine("1000", "Bank (secondary)", Decimal("0.00"), Decimal("0.00")),
)


def _tb(**overrides: object) -> TrialBalance:
    values: dict[str, object] = {
        "client_entity_id": ENTITY_ID,
        "entity_name": "Acme Holdings LLC",
        "period_start": date(2025, 1, 1),
        "period_end": date(2025, 12, 31),
        "pulled_at": PULLED_AT,
        "snapshot_id": SNAPSHOT_ID,
        "source": "quickbooks",
        "source_fingerprint": SOURCE_FINGERPRINT,
        "lines": LINES,
    }
    values.update(overrides)
    return TrialBalance(**values)  # pyright: ignore[reportArgumentType] -- test helper takes kwargs


def _sheet(content: bytes) -> Worksheet:
    sheet = load_workbook(io.BytesIO(content)).active
    assert isinstance(sheet, Worksheet)
    return sheet


def _cells(content: bytes) -> list[list[object]]:
    return [list(row) for row in _sheet(content).iter_rows(values_only=True)]


def _footer(content: bytes) -> dict[str, str]:
    found: dict[str, str] = {}
    for row in _cells(content):
        label, value = row[0], row[1]
        if isinstance(label, str) and isinstance(value, str) and label in LABELS:
            found[label] = value
    return found


LABELS = {
    "Source",
    "Method",
    "Pulled at",
    "Period",
    "Entity",
    "Entity ID",
    "Snapshot ID",
    "Source fingerprint",
}


def _core_xml(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return archive.read("docProps/core.xml").decode()


# --- AC-12: determinism ----------------------------------------------------------------------


def test_ac12_media_type_is_xlsx() -> None:
    assert XLSX_MEDIA_TYPE == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_ac12_rendering_twice_gives_identical_bytes_and_fingerprints() -> None:
    first, second = render_trial_balance(_tb()), render_trial_balance(_tb())
    assert first == second
    assert hashlib.sha256(first).hexdigest() == hashlib.sha256(second).hexdigest()


def test_ac12_rendering_is_identical_across_time() -> None:
    first = render_trial_balance(_tb())
    time.sleep(2.1)  # the save-time clock must never reach the file
    second = render_trial_balance(_tb())
    assert first == second


def test_ac12_rendering_is_identical_across_equal_but_separately_built_inputs() -> None:
    rebuilt = tuple(
        TrialBalanceLine(line.account_code, line.account_name, line.debit, line.credit)
        for line in LINES
    )
    assert render_trial_balance(_tb()) == render_trial_balance(_tb(lines=rebuilt))


def test_ac12_line_order_in_the_input_does_not_change_the_output() -> None:
    baseline = render_trial_balance(_tb())
    for order in itertools.islice(itertools.permutations(LINES), 0, None, 7):
        assert render_trial_balance(_tb(lines=order)) == baseline
    assert render_trial_balance(_tb(lines=tuple(reversed(LINES)))) == baseline


def test_ac12_changed_data_changes_the_output() -> None:
    changed = (*LINES[:-1], TrialBalanceLine("5000", "Expenses", Decimal("1.00"), Decimal("0.00")))
    assert render_trial_balance(_tb(lines=changed)) != render_trial_balance(_tb())


def test_ac12_a_later_pull_time_changes_the_output() -> None:
    later = PULLED_AT + timedelta(days=1)
    assert render_trial_balance(_tb(pulled_at=later)) != render_trial_balance(_tb())


def test_ac12_zip_entries_carry_no_wall_clock_time() -> None:
    with zipfile.ZipFile(io.BytesIO(render_trial_balance(_tb()))) as archive:
        stamps = {info.date_time for info in archive.infolist()}
    assert len(stamps) == 1
    assert next(iter(stamps))[0] <= 1980 + 1  # a fixed epoch, never the render time


def test_ac12_core_properties_have_creator_platform_and_pull_time() -> None:
    xml = _core_xml(render_trial_balance(_tb()))
    assert re.search(r"<dc:creator>platform</dc:creator>", xml)
    created = re.search(r"<dcterms:created[^>]*>([^<]*)</dcterms:created>", xml)
    modified = re.search(r"<dcterms:modified[^>]*>([^<]*)</dcterms:modified>", xml)
    assert created is not None
    assert modified is not None
    pulled = PULLED_AT.strftime("%Y-%m-%dT%H:%M:%S")
    assert created.group(1).startswith(pulled)
    assert modified.group(1).startswith(pulled)
    assert created.group(1) == modified.group(1)


def test_ac12_properties_follow_a_different_pull_time() -> None:
    other = datetime(2024, 7, 1, 23, 59, 59, tzinfo=UTC)
    xml = _core_xml(render_trial_balance(_tb(pulled_at=other)))
    assert xml.count("2024-07-01T23:59:59") == 2


# --- content ---------------------------------------------------------------------------------


def test_ac12_output_opens_and_rows_are_sorted_by_account_code() -> None:
    rows = _cells(render_trial_balance(_tb()))
    codes = [row[0] for row in rows if row[0] in {"1000", "2000", "4000", "5000"}]
    assert codes == ["1000", "1000", "2000", "4000", "5000"]


def test_ac12_equal_codes_are_ordered_by_name() -> None:
    rows = _cells(render_trial_balance(_tb()))
    names = [row[1] for row in rows if row[0] == "1000"]
    assert names == ["Bank (secondary)", "Cash"]


def test_ac12_total_row_has_computed_sums() -> None:
    rows = _cells(render_trial_balance(_tb()))
    total = next(row for row in rows if row[1] == "Total")
    assert Decimal(str(total[2])) == Decimal("1800.25")
    assert Decimal(str(total[3])) == Decimal("1800.25")


def test_ac12_total_is_a_value_not_a_formula() -> None:
    sheet = _sheet(render_trial_balance(_tb()))
    for row in sheet.iter_rows():
        for cell in row:
            assert not (isinstance(cell.value, str) and cell.value.startswith("="))
            assert cell.data_type != "f"


def test_ac12_an_empty_ledger_renders_with_zero_totals() -> None:
    rows = _cells(render_trial_balance(_tb(lines=())))
    total = next(row for row in rows if row[1] == "Total")
    assert Decimal(str(total[2] or 0)) == 0
    assert Decimal(str(total[3] or 0)) == 0


def test_ac12_footer_carries_the_provenance_labels() -> None:
    footer = _footer(render_trial_balance(_tb()))
    assert set(footer) == LABELS
    assert footer["Source"] == "quickbooks"
    assert footer["Method"] == "retrieved"
    assert footer["Entity"] == "Acme Holdings LLC"
    assert footer["Entity ID"] == str(ENTITY_ID)
    assert footer["Snapshot ID"] == str(SNAPSHOT_ID)
    assert footer["Source fingerprint"] == SOURCE_FINGERPRINT


def test_ac12_footer_pull_time_is_iso() -> None:
    parsed = datetime.fromisoformat(_footer(render_trial_balance(_tb()))["Pulled at"])
    assert parsed == PULLED_AT


def test_ac12_footer_period_names_both_dates() -> None:
    period = _footer(render_trial_balance(_tb()))["Period"]
    assert "2025-01-01" in period
    assert "2025-12-31" in period


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_ac12_hostile_account_names_are_string_cells_with_the_value_unchanged(prefix: str) -> None:
    name = f'{prefix}HYPERLINK("http://evil.test","x")'
    lines = (TrialBalanceLine("9000", name, Decimal("1.00"), Decimal("0.00")),)
    sheet = _sheet(render_trial_balance(_tb(lines=lines)))
    cells = [cell for row in sheet.iter_rows() for cell in row if cell.value == name]
    assert len(cells) == 1
    assert cells[0].data_type == "s"


def test_ac12_hostile_names_do_not_break_determinism() -> None:
    lines = (TrialBalanceLine("9000", "=1+1", Decimal("1.00"), Decimal("1.00")),)
    assert render_trial_balance(_tb(lines=lines)) == render_trial_balance(_tb(lines=lines))


# --- revision 1: hardening -------------------------------------------------------------------


def test_ac12_a_naive_pull_time_is_refused() -> None:
    with pytest.raises(ValueError):
        render_trial_balance(_tb(pulled_at=datetime(2026, 3, 14, 9, 26, 53)))


def test_ac12_a_non_utc_pull_time_is_converted_to_utc_once() -> None:
    plus_four = datetime(2026, 3, 14, 13, 26, 53, tzinfo=timezone(timedelta(hours=4)))
    content = render_trial_balance(_tb(pulled_at=plus_four))
    assert content == render_trial_balance(_tb(pulled_at=PULLED_AT))
    assert datetime.fromisoformat(_footer(content)["Pulled at"]) == PULLED_AT
    assert _footer(content)["Pulled at"].endswith(("+00:00", "Z"))
    xml = _core_xml(content)
    assert xml.count("2026-03-14T09:26:53") == 2


@pytest.mark.parametrize("bad", [Decimal("NaN"), Decimal("Infinity"), Decimal("-Infinity")])
def test_ac12_non_finite_amounts_are_refused(bad: Decimal) -> None:
    for line in (
        TrialBalanceLine("1000", "Cash", bad, Decimal(0)),
        TrialBalanceLine("1000", "Cash", Decimal(0), bad),
    ):
        with pytest.raises(ValueError):
            render_trial_balance(_tb(lines=(line,)))


def test_ac12_text_over_32767_characters_is_refused_and_the_limit_is_accepted() -> None:
    long_line = TrialBalanceLine("1000", "n" * 32_768, Decimal(0), Decimal(0))
    with pytest.raises(ValueError):
        render_trial_balance(_tb(lines=(long_line,)))
    ok_line = TrialBalanceLine("1000", "n" * 32_767, Decimal(0), Decimal(0))
    render_trial_balance(_tb(lines=(ok_line,)))


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_ac12_hostile_text_also_carries_the_quote_prefix(prefix: str) -> None:
    name = f"{prefix}cmd|' /C calc'!A0"
    lines = (TrialBalanceLine("9000", name, Decimal("1.00"), Decimal("0.00")),)
    sheet = _sheet(render_trial_balance(_tb(lines=lines)))
    [cell] = [c for row in sheet.iter_rows() for c in row if c.value == name]
    assert cell.quotePrefix is True


def test_ac12_zip_entries_are_marked_as_created_on_unix() -> None:
    with zipfile.ZipFile(io.BytesIO(render_trial_balance(_tb()))) as archive:
        assert {info.create_system for info in archive.infolist()} == {3}
