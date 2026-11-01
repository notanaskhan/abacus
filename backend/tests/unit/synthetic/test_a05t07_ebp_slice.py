"""A-05-T07 (TASK-2813; DEC-1862): Northwind PY2025's EBP figure book foots.

Row acceptance: "make seed-demo loads the slice, fixture mode shows identical figures, and every
figure-book total foots in the seed tests." This file is the footing half: every total in the
golden YAML is recomputed from the slice's rows, and the documents are deterministic. The load
through `evidence.api` runs in the loader's integration suite.
"""

from __future__ import annotations

import hashlib
import io
from decimal import Decimal
from itertools import pairwise
from pathlib import Path

import yaml
from openpyxl import load_workbook

from abacus_tools.synthetic.scenarios.brightwater_hale import ebp
from abacus_tools.synthetic.scenarios.brightwater_hale.slices import slice_events

GOLDEN = Path(__file__).resolve().parents[2] / "golden" / "ebp" / "northwind_py2025.yaml"
BOOK = yaml.safe_load(GOLDEN.read_text(encoding="utf-8"))


def test_census_foots_to_the_w3_and_the_participant_counts() -> None:
    rows = ebp.census()
    assert len(rows) == ebp.PARTICIPANTS_EOY == BOOK["census"]["rows"] == 301
    total = sum((r.compensation for r in rows), Decimal(0))
    assert total == ebp.W3_BOX1 + ebp.W3_BOX12D == Decimal("14560000.00")
    assert sum((r.deferrals for r in rows), Decimal(0)) == ebp.W3_BOX12D
    assert sum(r.status == "terminated" for r in rows) == 6  # DF-34
    assert len({r.key for r in rows}) == 301


def test_the_payroll_register_foots_and_the_late_deposit_has_both_counts() -> None:
    payrolls = ebp.payrolls()
    assert len(payrolls) == 26
    assert payrolls[0].pay_on.isoformat() == "2025-01-10"
    assert all((b.pay_on - a.pay_on).days == 14 for a, b in pairwise(payrolls))
    assert sum((p.wages for p in payrolls), Decimal(0)) == Decimal("14560000.00")
    assert sum((p.deferrals for p in payrolls), Decimal(0)) == ebp.W3_BOX12D
    late = payrolls[-1]
    assert (late.pay_on.isoformat(), late.deposited_on.isoformat()) == ("2025-12-26", "2026-01-09")
    assert late.deferrals == Decimal("33540.00")
    assert (late.business_days, late.calendar_days_late) == (9, 10)
    assert sum(p.business_days <= 2 for p in payrolls) == 25


def test_the_receivable_and_the_providers() -> None:
    assert ebp.RECEIVABLE == Decimal("15060.00") + Decimal("33540.00") == Decimal("48600.00")
    keystone = ebp.PROVIDERS[0]
    assert keystone.ein_last4 == "4420" and "4402" in (keystone.note or "")
    assert [p.role for p in ebp.PROVIDERS] == [
        "recordkeeper",
        "trustee_custodian",
        "payroll",
        "tpa",
    ]
    assert ebp.PLAN_NUMBER == "001" and ebp.RECORDKEEPER_PLAN_ID == "58217"


def test_the_golden_book_matches_the_slice() -> None:
    assert BOOK["census"]["compensation"] == str(ebp.COMPENSATION)
    assert BOOK["payroll"]["on_time"] == 25
    assert BOOK["late_remittance"]["business_days"] == 9
    assert BOOK["files"] == [f.file_name for f in ebp.FILES]


def test_the_documents_are_deterministic_and_readable() -> None:
    for file in ebp.FILES:
        assert hashlib.sha256(file.content()).digest() == hashlib.sha256(file.content()).digest()
    sheet = load_workbook(io.BytesIO(ebp.census_xlsx()), read_only=True)["Census"]
    assert sum(1 for _ in sheet.iter_rows(min_row=2)) == 301
    hr = load_workbook(io.BytesIO(ebp.hr_records_xlsx()), read_only=True)["HR records"]
    assert sum(1 for _ in hr.iter_rows(min_row=2)) == 25


def test_the_slice_uploads_once_on_northwind() -> None:
    [event] = slice_events("ebp")
    assert event.kind == "evidence.ebp_documents_uploaded"
    assert event.data["files"] == len(ebp.FILES) == 8
