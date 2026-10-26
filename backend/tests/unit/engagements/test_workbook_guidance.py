"""SPEC-031 AC-8 (TASK-060): the template workbook's optional guidance columns."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from abacus.modules.methodology.workbook import TemplateInvalid, parse_workbook


def _book(requests: list[list[object]]) -> bytes:
    workbook = Workbook()
    areas = workbook.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["CASH", "Cash"])
    sheet = workbook.create_sheet("Requests")
    for row in requests:
        sheet.append(row)
    rules = workbook.create_sheet("Account rules")
    rules.append(["area_code", "account_from", "account_to"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_ac8_new_columns_are_imported() -> None:
    data = _book(
        [
            [
                "area_code",
                "description",
                "tier",
                "why",
                "good_looks_like",
                "blocks_fieldwork",
                "effort",
            ],
            ["CASH", "Bank statements", "C", "To confirm cash", "The bank's PDF", "yes", "Large"],
        ]
    )
    (item,) = parse_workbook(data).items
    assert (item.why, item.good_looks_like, item.blocks_fieldwork, item.effort) == (
        "To confirm cash",
        "The bank's PDF",
        True,
        "large",
    )


def test_ac8_without_the_columns_the_template_still_loads() -> None:
    (item,) = parse_workbook(
        _book([["area_code", "description", "tier"], ["CASH", "Bank statements", "C"]])
    ).items
    assert (item.why, item.blocks_fieldwork, item.effort) == (None, False, None)


def test_ac8_invalid_values_are_reported() -> None:
    data = _book(
        [
            ["area_code", "description", "tier", "effort", "blocks_fieldwork"],
            ["CASH", "Bank statements", "C", "huge", "maybe"],
        ]
    )
    with pytest.raises(TemplateInvalid) as raised:
        parse_workbook(data)
    assert {p.code for p in raised.value.problems} == {"invalid_effort", "invalid_yes_no"}
