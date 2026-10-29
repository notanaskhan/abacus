"""K-01-T09: independent acceptance tests (ADR-078) for K-01-T02, T03, T05 and T06, written
from the contracts (`methodology-programs.md`, `knowledge-catalog.md`) and K-01 §3.1 to §3.3,
not from the implementing suites.

- AC2.3: a workbook's problems are `{sheet, row, column, code}`, never a cell's contents, and a
  formula is never evaluated.
- AC1.6 / AC1.9 / AC1.10: a staff reader never receives a draft (at the API); a usage count never
  counts a walled engagement (INV-3); a portal principal is refused every knowledge and
  methodology route (INV-4).
- AC3.3 (INV-1, INV-5): the agent principal and a system worker are refused approval and audited;
  a manager without `methodology.approve` is refused.
- AC3.5: a pin takes the version in force on the acceptance date, and pins are set once.
- AC3.6: once a version leaves draft, the app role cannot change its steps.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Sequence
from datetime import UTC, date, datetime, timedelta
from typing import cast

import pytest
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

from abacus.kernel.db import TenantContext
from abacus.kernel.db.session import ActorKind
from abacus.kernel.flags import forget_flags
from abacus.kernel.uow import Ref, Target, uow
from abacus.modules.engagements.api import MethodologyAlreadyApplied, pin_methodology_on
from abacus.modules.identity.api import AuthContext, Forbidden
from abacus.modules.methodology.lifecycle import approve_version

from .support import Json, Net, Person

pytestmark = pytest.mark.asyncio
SENTINEL = "SECRET-CELL-7f3a"  # a cell's contents must never come back


def _book(
    *, steps: Sequence[Sequence[object]] | None = None, request: str = "Plan document"
) -> bytes:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["PLAN", "Planning"])
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    requests.append(["PLAN", request, "B"])
    if steps is not None:
        programs = book.create_sheet("Programs")
        programs.append(["code", "name"])
        programs.append(["BH-EBP-110", "General planning"])
        sheet = book.create_sheet("Steps")
        sheet.append(["program_code", "number", "text", "level", "mandatory"])
        for row in steps:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def _flags(net: Net, tenant: uuid.UUID, *flags: str) -> None:
    for flag in flags:
        await net.seed.run(
            "INSERT INTO feature_flag_states (tenant_id, flag, value, set_by, reason) "
            "VALUES ($1, $2, 'true', 'test-seed', 'K-01-T09') "
            "ON CONFLICT (tenant_id, flag) DO UPDATE SET value = 'true'",
            tenant,
            flag,
        )
    forget_flags()


async def _post(net: Net, who: Person, name: str, data: bytes) -> tuple[int, Json]:
    response = await net.http.post(
        f"/v1/methodology/templates/{name}/versions?engagement_type=ebp_audit&family=ebp",
        content=data,
        headers={**net.headers(who), "Content-Type": "application/octet-stream"},
    )
    return response.status_code, cast(Json, response.json())


def _ctx(who: Person, kind: str, role: str) -> AuthContext:
    return AuthContext(
        TenantContext(who.tenant_id, cast(ActorKind, kind), f"{kind}:{uuid.uuid4()}"),
        who.user_id,
        uuid.uuid4(),
        role,  # pyright: ignore[reportArgumentType] -- a seeded firm role code
        datetime.now(UTC) - timedelta(minutes=1),
    )


async def test_ac23_problems_name_the_cell_never_its_contents_and_formulas_stay_text(
    net: Net,
) -> None:
    seed = net.seed
    tenant = await seed.firm()
    await _flags(net, tenant, "methodology.programs")
    admin = await seed.person(tenant, "firm_admin")
    bad = _book(
        steps=[
            ["BH-EBP-110", 1, "Understand the plan", SENTINEL, "yes"],  # invalid_level
            ["BH-EBP-999", 2, "Orphan step", "Planning", "no"],  # unknown_program
            ["BH-EBP-110", 4, "Gap step", "Testing", SENTINEL],  # numbering_gap, yes/no
        ]
    )
    status, body = await _post(net, admin, "BH EBP", bad)
    assert status == 422, body
    assert SENTINEL not in str(body)
    assert "invalid_level" in str(body) and "unknown_program" in str(body)
    assert (
        await seed.value("SELECT count(*) FROM methodology_versions WHERE tenant_id = $1", tenant)
        == 0
    )  # nothing stored

    # A formula is stored as written text, never computed.
    status, body = await _post(net, admin, "BH PBC", _book(request="=1+1"))
    assert status in (201, 422), body
    if status == 201:
        stored = await seed.value(
            "SELECT count(*) FROM methodology_versions v WHERE v.tenant_id = $1 "
            "AND v.methodology::text LIKE '%\"2\"%'",
            tenant,
        )
        assert stored == 0


async def test_ac19_ac16_ac110_walls_drafts_and_portal_principals(net: Net) -> None:
    seed = net.seed
    tenant = await seed.firm()
    await _flags(net, tenant, "knowledge.catalog")
    admin = await seed.person(tenant, "firm_admin")
    ana = await seed.person(tenant)
    ashford = await seed.site(tenant, admin.user_id, "Ashford")
    bayside = await seed.site(tenant, admin.user_id, "Bayside")
    for site in (ashford, bayside):
        await seed.member(site, ana, "staff")
    status, body = await _post(net, admin, "BH Audit Methodology", _book())
    assert status == 201, body
    version = uuid.UUID(cast(str, body["version_id"]))
    for site in (ashford, bayside):
        await seed.run_untriggered(
            "UPDATE engagements SET methodology_version_id = $1 WHERE id = $2",
            version,
            site.engagement_id,
        )

    def usage(page: Json) -> object:
        cards = {cast(str, c["id"]): c for c in cast(list[Json], page["items"])}
        return cards[str(version)]["usage_count"]

    seen = await net.call("GET", "/v1/knowledge/items?category=methodology", ana)
    assert seen.status_code == 200, seen.text
    assert usage(cast(Json, seen.json())) == 2
    await net.wall(admin, ana, ashford)
    walled = cast(
        Json, (await net.call("GET", "/v1/knowledge/items?category=methodology", ana)).json()
    )
    assert usage(walled) == 1  # the walled engagement is neither counted nor named
    assert str(ashford.engagement_id) not in str(walled)

    # A later draft (needs review) never reaches a reader who neither manages nor approves.
    await _flags(net, tenant, "methodology.lifecycle")
    status, body = await _post(net, admin, "BH Audit Methodology", _book(request="Form 5500"))
    assert status == 201, body
    draft = cast(str, body["version_id"])
    staff_view = await net.call("GET", "/v1/knowledge/items", ana)
    assert draft not in staff_view.text
    queue = await net.call("GET", "/v1/knowledge/review-queue", ana)
    assert queue.status_code in (200, 403) and draft not in queue.text
    changes = await net.call("GET", "/v1/knowledge/change-log", ana)
    assert draft not in changes.text

    # INV-4: a portal principal reaches no knowledge or methodology route.
    claire = await seed.client_contact(bayside)
    for path in (
        "/v1/knowledge/items",
        "/v1/knowledge/review-queue",
        "/v1/methodology/templates",
        f"/v1/methodology/versions/{version}/state",
    ):
        refused = await net.call("GET", path, claire)
        assert refused.status_code in (403, 404), (path, refused.status_code)
        assert "BH Audit Methodology" not in refused.text


async def test_ac33_only_a_person_holding_approve_approves_and_refusals_are_audited(
    net: Net,
) -> None:
    seed = net.seed
    tenant = await seed.firm()
    admin = await seed.person(tenant, "firm_admin")
    margaret = await seed.person(tenant, "quality_partner")
    hannah = await seed.person(tenant)
    await _flags(net, tenant, "methodology.programs", "methodology.lifecycle")
    status, body = await _post(
        net,
        admin,
        "BH EBP",
        _book(steps=[["BH-EBP-110", 1, "Understand the plan", "Planning", "yes"]]),
    )
    assert status == 201, body
    version = uuid.UUID(cast(str, body["version_id"]))
    state = cast(
        Json,
        (await net.call("GET", f"/v1/methodology/versions/{version}/state", margaret)).json(),
    )
    assert state["status"] == "needs_review"
    checklist, row_version = cast(str, state["checklist_hash"]), cast(int, state["row_version"])
    for kind in ("agent", "system"):
        with pytest.raises(Forbidden):
            await approve_version(
                _ctx(margaret, kind, "quality_partner"),
                version,
                effective_from=date(2026, 1, 1),
                checklist_hash=checklist,
                row_version=row_version,
            )
    refused = await net.call(
        "POST",
        f"/v1/methodology/versions/{version}/approve",
        hannah,
        body={
            "effective_from": "2026-01-01",
            "checklist_hash": checklist,
            "row_version": row_version,
        },
    )
    assert refused.status_code == 403, refused.text
    actions = [
        r["action"]
        for r in await seed.rows(
            "SELECT action FROM audit_events WHERE tenant_id = $1 AND target_id = $2",
            tenant,
            str(version),
        )
    ]
    assert actions.count("decision.refused_non_human") == 2
    assert "methodology_version.approved" not in actions  # INV-5: nobody approved
    still = await seed.value(
        "SELECT status FROM methodology_version_states WHERE version_id = $1", version
    )
    assert still == "needs_review"


async def test_ac35_ac36_pins_follow_the_effective_date_and_published_steps_are_frozen(
    net: Net,
) -> None:
    seed = net.seed
    tenant = await seed.firm()
    admin = await seed.person(tenant, "firm_admin")
    margaret = await seed.person(tenant, "quality_partner")
    early = await seed.site(tenant, admin.user_id, "Calloway")
    late = await seed.site(tenant, admin.user_id, "Feld")
    for site in (early, late):  # the template is an EBP one: pin only reaches its own type
        await seed.run(
            "UPDATE engagements SET type = 'ebp_audit' WHERE id = $1", site.engagement_id
        )
    await _flags(net, tenant, "methodology.programs")
    steps = [["BH-EBP-110", 1, "Understand the plan", "Planning", "yes"]]
    status, body = await _post(net, admin, "BH EBP", _book(steps=steps))
    assert status == 201, body
    v2 = uuid.UUID(cast(str, body["version_id"]))
    await seed.run(
        "UPDATE methodology_version_states SET effective_from = '2025-01-01' "
        "WHERE version_id = $1",
        v2,
    )
    await _flags(net, tenant, "methodology.lifecycle")
    status, body = await _post(net, admin, "BH EBP", _book(steps=steps, request="Form 5500"))
    assert status == 201, body
    v3 = uuid.UUID(cast(str, body["version_id"]))
    state = cast(
        Json, (await net.call("GET", f"/v1/methodology/versions/{v3}/state", margaret)).json()
    )
    approved = await net.call(
        "POST",
        f"/v1/methodology/versions/{v3}/approve",
        margaret,
        body={
            "effective_from": "2026-01-01",
            "checklist_hash": state["checklist_hash"],
            "row_version": state["row_version"],
        },
    )
    assert approved.status_code == 200, approved.text
    template = await seed.value("SELECT template_id FROM methodology_versions WHERE id = $1", v3)

    system = TenantContext(tenant, "system", "k01t09:pin")
    for site, accepted, wanted in ((early, date(2025, 12, 10), v2), (late, date(2026, 1, 20), v3)):
        async with uow(system) as tx:
            pinned = await pin_methodology_on(
                tx, site.engagement_id, cast(uuid.UUID, template), accepted
            )
            tx.record(
                "engagement.methodology_pinned",
                target=Target("engagement", site.engagement_id),
                after=Ref(),
            )
        assert pinned.summary.version_id == wanted
    with pytest.raises(MethodologyAlreadyApplied):  # set once: a second pin is refused
        async with uow(system) as tx:
            await pin_methodology_on(
                tx, late.engagement_id, cast(uuid.UUID, template), date(2026, 2, 1)
            )
            tx.record(
                "engagement.methodology_pinned",
                target=Target("engagement", late.engagement_id),
                after=Ref(),
            )

    # AC3.6: the app role cannot change a published (or superseded) version's steps.
    for version in (v2, v3):
        with pytest.raises(DBAPIError):
            async with uow(system) as tx:
                await tx.session.execute(
                    text("UPDATE program_steps SET text = 'edited' WHERE version_id = :v"),
                    {"v": version},
                )
                tx.record(
                    "methodology_version.probe", target=Target("methodology_version", version)
                )
