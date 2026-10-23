# ac-test-author: ayushkbhatia
"""SPEC-024 AC-1 and AC-6 (TASK-040; back-filled by Q-03-T008): self-serve sign-up and
engagement types with a template each.

AC-1: a signed-in person with a verified email, no firm and a valid founder-issued sign-up code
creates a firm and is its only firm administrator; defaults are set; the act is audited; the
checklist is where they land. A second attempt with the same code or identity is refused, and
sign-up is rate-limited per identity and per address.

AC-6: methodology templates are tagged with the engagement types they serve, and an engagement of
a type is offered only the matching templates.

HTTP through `create_app()` against the integration suite's migrated Postgres; codes are issued
with the founder tool (`abacus_tools.signup_codes.issue`), the only way the spec allows. Each
sign-up test runs from its own client address, because the per-address limit is global.
Expectations come from the spec, not the implementation.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncIterator
from dataclasses import dataclass
from typing import cast

import httpx
import pytest
from openpyxl import Workbook

from abacus.api import create_app
from abacus_tools.fakes.identity import FakeIdentityProvider
from abacus_tools.signup_codes import issue

from .support import Migrated, Net, Person, Seeder

COUNT_FIRMS_NAMED = "SELECT count(*) FROM firms WHERE name = $1"


@dataclass(frozen=True)
class Door:
    """Sign-up as seen from one client address."""

    http: httpx.AsyncClient
    idp: FakeIdentityProvider
    seed: Seeder
    dsn: str

    async def code(self, days: int = 30) -> str:
        return await issue(self.dsn, "Design partner (test)", days)

    def identity(self, *, email: str | None = None, verified: bool = True) -> tuple[str, str]:
        """A provider sign-in (subject, bearer token) of someone with no account yet."""
        subject = f"new-{uuid.uuid4().hex}"
        if email is None:
            return subject, self.idp.token(subject, mfa=True)
        return subject, self.idp.token(subject, mfa=True, email=email, email_verified=verified)

    async def sign_up(self, token: str, code: str, firm_name: str) -> httpx.Response:
        return await self.http.post(
            "/v1/signup",
            json={"code": code, "firm_name": firm_name},
            headers={"Authorization": f"Bearer {token}"},
        )

    async def person(self, tenant_id: uuid.UUID, subject: str) -> Person:
        user_id = await self.seed.value("SELECT id FROM users WHERE idp_subject = $1", subject)
        return Person(cast(uuid.UUID, user_id), tenant_id, subject)


@pytest.fixture
async def door(
    seed: Seeder, idp: FakeIdentityProvider, migrated_db: Migrated
) -> AsyncIterator[Door]:
    octets = uuid.uuid4().bytes[:3]  # an address of its own (the per-address limit is global)
    address = f"10.{octets[0]}.{octets[1]}.{octets[2] % 254 + 1}"
    transport = httpx.ASGITransport(
        app=create_app(), raise_app_exceptions=False, client=(address, 4321)
    )
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        yield Door(client, idp, seed, migrated_db.superuser_dsn)


def _firm_name() -> str:
    return f"Whitfield & Lane {uuid.uuid4().hex[:8]}"


async def _signed_up(door: Door) -> tuple[uuid.UUID, Person]:
    subject, token = door.identity(email=f"sam-{uuid.uuid4().hex[:6]}@whitfield.example")
    response = await door.sign_up(token, await door.code(), _firm_name())
    assert response.status_code == 201, response.text
    tenant_id = uuid.UUID(cast(str, response.json()["tenant_id"]))
    return tenant_id, await door.person(tenant_id, subject)


# --- AC-1: sign-up ------------------------------------------------------------------------------


async def test_ac1_sign_up_creates_the_firm_with_the_person_as_its_only_administrator(
    door: Door, net: Net
) -> None:
    name = _firm_name()
    subject, token = door.identity(email="sam@whitfield.example")
    response = await door.sign_up(token, await door.code(), name)
    assert response.status_code == 201, response.text
    tenant_id = uuid.UUID(cast(str, response.json()["tenant_id"]))
    sam = await door.person(tenant_id, subject)
    assert await door.seed.value(COUNT_FIRMS_NAMED, name) == 1
    me = (await net.call("GET", "/v1/me", sam)).json()
    assert [(m["tenant_id"], m["firm_role"]) for m in me["memberships"]] == [
        (str(tenant_id), "firm_admin")
    ]
    staff = await net.call("GET", "/v1/firm/staff", sam)
    assert staff.status_code == 200, staff.text
    assert [(m["user_id"], m["firm_role"]) for m in staff.json()["members"]] == [
        (str(sam.user_id), "firm_admin")
    ]
    assert staff.json()["invitations"] == []


async def test_ac1_the_new_firm_starts_with_every_setting_at_its_safe_default(
    door: Door, net: Net
) -> None:
    _, sam = await _signed_up(door)
    autonomy = (await net.call("GET", "/v1/firm/autonomy", sam)).json()
    assert autonomy["level"] == 1  # Routine (SPEC-024 AC-5)
    letter = (await net.call("GET", "/v1/firm/letter-policy", sam)).json()
    assert letter == {"required": False}  # SPEC-025 Q4: warns, does not block, by default
    budget = await net.call("GET", "/v1/budget", sam)
    assert budget.status_code == 200, budget.text
    assert budget.json()["is_default"] is True  # the firm's default budget


async def test_ac1_the_person_lands_on_a_checklist_with_nothing_done_yet(
    door: Door, net: Net
) -> None:
    _, sam = await _signed_up(door)
    checklist = await net.call("GET", "/v1/firm/onboarding", sam)
    assert checklist.status_code == 200, checklist.text
    body = checklist.json()
    assert [s["id"] for s in body["steps"]] == [
        "sso",
        "team",
        "methodology",
        "autonomy",
        "budget",
        "walls",
        "engagement",
    ]
    assert body["dismissed"] is False
    assert body["complete"] is False
    for step in ("team", "methodology", "walls", "engagement"):
        assert not next(s for s in body["steps"] if s["id"] == step)["done"], step


async def test_ac1_sign_up_is_audited_in_the_new_firm(door: Door) -> None:
    tenant_id, _ = await _signed_up(door)
    assert "firm.created" in await door.seed.actions(tenant_id)


async def test_ac1_a_code_works_once(door: Door) -> None:
    code = await door.code()
    _, first = door.identity(email="one@first.example")
    created = await door.sign_up(first, code, _firm_name())
    assert created.status_code == 201, created.text
    name = _firm_name()
    _, second = door.identity(email="two@second.example")
    again = await door.sign_up(second, code, name)
    assert (again.status_code, again.json()) == (409, {"detail": "signup_code_invalid"})
    assert await door.seed.value(COUNT_FIRMS_NAMED, name) == 0


async def test_ac1_the_same_identity_cannot_sign_up_a_second_firm(door: Door) -> None:
    subject, token = door.identity(email="sam@whitfield.example")
    first = await door.sign_up(token, await door.code(), _firm_name())
    assert first.status_code == 201, first.text
    name = _firm_name()
    second = await door.sign_up(token, await door.code(), name)
    assert (second.status_code, second.json()) == (409, {"detail": "signup_already_staff"})
    assert await door.seed.value(COUNT_FIRMS_NAMED, name) == 0
    memberships = await door.seed.value(
        "SELECT count(*) FROM memberships m JOIN users u ON u.id = m.user_id "
        "WHERE u.idp_subject = $1",
        subject,
    )
    assert memberships == 1


async def test_ac1_a_staff_member_of_another_firm_cannot_sign_up(
    door: Door, seed: Seeder, idp: FakeIdentityProvider
) -> None:
    tenant_id = await seed.firm()
    member = await seed.person(tenant_id, "practice_leader")
    token = idp.token(member.subject, mfa=True, email="member@example.test", email_verified=True)
    name = _firm_name()
    response = await door.sign_up(token, await door.code(), name)
    assert (response.status_code, response.json()) == (409, {"detail": "signup_already_staff"})
    assert await door.seed.value(COUNT_FIRMS_NAMED, name) == 0


@pytest.mark.parametrize("email", [None, "unverified"])
async def test_ac1_sign_up_needs_a_verified_email(door: Door, email: str | None) -> None:
    if email is None:
        _, token = door.identity()
    else:
        _, token = door.identity(email="sam@whitfield.example", verified=False)
    code = await door.code()
    name = _firm_name()
    response = await door.sign_up(token, code, name)
    assert (response.status_code, response.json()) == (422, {"detail": "signup_invalid"})
    assert await door.seed.value(COUNT_FIRMS_NAMED, name) == 0
    _, good = door.identity(email="sam@whitfield.example")
    assert (await door.sign_up(good, code, _firm_name())).status_code == 201  # code not spent


async def test_ac1_an_unknown_or_expired_code_is_refused(door: Door) -> None:
    for code in ("ABCD-EFGH-JKMN-PQRS", await door.code(days=-1)):
        _, token = door.identity(email="sam@whitfield.example")
        name = _firm_name()
        response = await door.sign_up(token, code, name)
        assert (response.status_code, response.json()) == (409, {"detail": "signup_code_invalid"})
        assert await door.seed.value(COUNT_FIRMS_NAMED, name) == 0


async def test_ac1_sign_up_without_a_sign_in_is_refused(door: Door) -> None:
    name = _firm_name()
    response = await door.http.post(
        "/v1/signup", json={"code": await door.code(), "firm_name": name}
    )
    assert response.status_code == 401
    assert await door.seed.value(COUNT_FIRMS_NAMED, name) == 0


async def test_ac1_repeated_failures_rate_limit_that_identity_and_spend_no_code(
    door: Door,
) -> None:
    _, token = door.identity(email="sam@whitfield.example")
    for _ in range(5):
        bad = await door.sign_up(token, "ABCD-EFGH-JKMN-PQRS", _firm_name())
        assert (bad.status_code, bad.json()) == (409, {"detail": "signup_code_invalid"})
    code = await door.code()
    limited = await door.sign_up(token, code, _firm_name())
    assert (limited.status_code, limited.json()) == (409, {"detail": "signup_rate_limited"})
    _, other = door.identity(email="pat@whitfield.example")
    assert (await door.sign_up(other, code, _firm_name())).status_code == 201  # code not spent


async def test_ac1_repeated_failures_rate_limit_the_address(door: Door) -> None:
    for _ in range(20):  # each from a different identity: the address is what is counted
        _, token = door.identity(email=f"x{uuid.uuid4().hex[:6]}@example.test")
        bad = await door.sign_up(token, "ABCD-EFGH-JKMN-PQRS", _firm_name())
        assert (bad.status_code, bad.json()) == (409, {"detail": "signup_code_invalid"})
    _, fresh = door.identity(email="sam@whitfield.example")
    limited = await door.sign_up(fresh, await door.code(), _firm_name())
    assert (limited.status_code, limited.json()) == (409, {"detail": "signup_rate_limited"})


# --- AC-6: engagement types with a template each ------------------------------------------------


def _workbook(description: str) -> bytes:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["CASH", "Cash"])
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    requests.append(["CASH", description, "C"])
    rules = book.create_sheet("Account rules")
    rules.append(["area_code", "account_from", "account_to"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@dataclass(frozen=True)
class Methodology:
    net: Net
    seed: Seeder
    tenant_id: uuid.UUID
    admin: Person
    partner: Person

    async def upload(self, name: str, *types: str) -> uuid.UUID:
        query = "&".join(f"engagement_type={t}" for t in types)
        response = await self.net.http.post(
            f"/v1/methodology/templates/{name}/versions?{query}",
            content=_workbook(f"{name} bank statements"),
            headers={
                **self.net.headers(self.admin),
                "Content-Type": "application/octet-stream",
            },
        )
        assert response.status_code == 201, response.text
        return uuid.UUID(cast(str, response.json()["version_id"]))

    async def engagement(self, engagement_type: str) -> uuid.UUID:
        """An engagement of `engagement_type`, created through the API, with a partner."""
        marker = uuid.uuid4().hex[:6]
        response = await self.net.call(
            "POST",
            "/v1/engagements",
            self.admin,
            body={
                "name": f"Halvorsen FY2026 {marker}",
                "client_name": f"Halvorsen {marker}",
                "client_entity_name": f"Halvorsen Ltd {marker}",
                "fiscal_period_start": "2026-01-01",
                "fiscal_period_end": "2026-12-31",
                "type": engagement_type,
                "confirm_new": True,
            },
        )
        assert response.status_code == 201, response.text
        assert response.json()["type"] == engagement_type
        engagement_id = uuid.UUID(cast(str, response.json()["id"]))
        await self.seed.run(
            "INSERT INTO engagement_members (tenant_id, engagement_id, user_id, role) "
            "VALUES ($1, $2, $3, 'engagement_partner')",
            self.tenant_id,
            engagement_id,
            self.partner.user_id,
        )
        return engagement_id

    async def apply(self, engagement_id: uuid.UUID, version_id: uuid.UUID) -> httpx.Response:
        return await self.net.call(
            "POST",
            f"/v1/engagements/{engagement_id}/methodology",
            self.partner,
            body={"version_id": str(version_id)},
        )

    async def items(self, engagement_id: uuid.UUID) -> object:
        return await self.seed.value(
            "SELECT count(*) FROM request_items WHERE engagement_id = $1", engagement_id
        )


@pytest.fixture
async def methodology(net: Net, seed: Seeder) -> Methodology:
    tenant_id = await seed.firm()
    admin = await seed.person(tenant_id, "firm_admin")
    return Methodology(net, seed, tenant_id, admin, await seed.person(tenant_id))


async def test_ac6_templates_carry_the_engagement_types_they_serve(
    methodology: Methodology,
) -> None:
    await methodology.upload("Audit core", "fs_audit")
    await methodology.upload("Review and compilation", "review", "compilation")
    listed = await methodology.net.call(
        "GET", "/v1/methodology/templates?latest=true", methodology.admin
    )
    assert listed.status_code == 200, listed.text
    types = {t["template_name"]: sorted(t["engagement_types"]) for t in listed.json()}
    assert types == {
        "Audit core": ["fs_audit"],
        "Review and compilation": ["compilation", "review"],
    }


@pytest.mark.parametrize("engagement_type", ["review", "compilation", "agreed_upon_procedures"])
async def test_ac6_engagements_take_the_spec_024_types(
    methodology: Methodology, engagement_type: str
) -> None:
    engagement_id = await methodology.engagement(engagement_type)
    read = await methodology.net.call("GET", f"/v1/engagements/{engagement_id}", methodology.admin)
    assert read.json()["type"] == engagement_type


async def test_ac6_an_engagement_takes_a_template_of_its_type(methodology: Methodology) -> None:
    shared = await methodology.upload("Review and compilation", "review", "compilation")
    engagement_id = await methodology.engagement("compilation")
    applied = await methodology.apply(engagement_id, shared)
    assert applied.status_code == 201, applied.text
    assert applied.json()["items_created"] == 1
    assert await methodology.items(engagement_id) == 1


async def test_ac6_a_template_of_another_type_is_refused_and_creates_nothing(
    methodology: Methodology,
) -> None:
    audit = await methodology.upload("Audit core", "fs_audit")
    review = await methodology.upload("Review core", "review")
    review_engagement = await methodology.engagement("review")
    audit_engagement = await methodology.engagement("fs_audit")
    for engagement_id, version_id in ((review_engagement, audit), (audit_engagement, review)):
        refused = await methodology.apply(engagement_id, version_id)
        assert (refused.status_code, refused.json()) == (
            409,
            {"detail": "template_type_mismatch"},
        ), refused.text
        assert await methodology.items(engagement_id) == 0
