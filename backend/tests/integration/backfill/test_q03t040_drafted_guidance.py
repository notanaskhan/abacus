# ac-test-author: ayushkbhatia
"""SPEC-035 AC-4 to AC-6 and Q2, Q3, Q5 (TASK-073, TASK-074): drafted guidance reviewed,
approved and shown, and never shown to a client of a real firm (ADR-111 until ADR P-25's
conditions hold).

Q-03-T040's line: "A non-synthetic firm can never display model-drafted guidance to a client
principal in any test." Boundary: `create_app()` over HTTP with fake-identity tokens (`net`);
drafts are seeded as the superuser as the drafter would record them (its run and its passed
checks), since a draft itself needs a model (SPEC-035 AC-1, covered by the evaluation suite).
Errors from `docs/architecture/error-codes.yaml`; the client display rule from SPEC-035 Q2 and
`firms.drafted_guidance_to_clients` (TASK-073 D2): only a firm admin sets it, only while the firm
is synthetic (ADR-112).
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from typing import cast

import httpx
import pytest
from openpyxl import Workbook

from .support import Json, Net, Person, Seeder, Site

pytestmark = pytest.mark.asyncio
WHY = "We confirm the cash balance with the bank."
GOOD = "A statement for every account, covering the year end."


@dataclass(frozen=True)
class Firm:
    net: Net
    tenant_id: uuid.UUID
    admin: Person
    priya: Person  # senior: asks for, edits, approves drafts
    staff: Person  # staff: neither asks nor decides
    rachel: Person  # the client's admin
    site: Site

    @property
    def seed(self) -> Seeder:
        return self.net.seed

    @property
    def drafts(self) -> str:
        return f"/v1/engagements/{self.site.engagement_id}/guidance-drafts"

    async def draft(self, item: uuid.UUID | None = None, *, why: str = WHY) -> uuid.UUID:
        """What the drafter records for one item: its run and a draft that passed the checks."""
        run = await self.seed.value(
            "INSERT INTO agent_runs (tenant_id, agent_id, spec_version, engagement_id, "
            "initiator_user_id, task_scope, status, finished_at) VALUES ($1, "
            "'requests.guidance_drafter', 2, $2, $3, "
            "ARRAY['request_item.read', 'guidance_draft.run'], 'completed', now()) RETURNING id",
            self.tenant_id,
            self.site.engagement_id,
            self.priya.user_id,
        )
        return cast(
            uuid.UUID,
            await self.seed.value(
                "INSERT INTO guidance_drafts (tenant_id, engagement_id, request_item_id, "
                "agent_run_id, prompt_ref, why, good_looks_like, passed_checks) "
                "VALUES ($1, $2, $3, $4, 'requests.guidance_draft@v1', $5, $6, true) "
                "RETURNING id",
                self.tenant_id,
                self.site.engagement_id,
                item or self.site.item_id,
                run,
                why,
                GOOD,
            ),
        )

    async def approve(
        self, approvals: list[Json], who: Person | None = None, **kw: object
    ) -> httpx.Response:
        body: Json = {"approvals": approvals, "confirmed_read": kw.get("confirmed_read", True)}
        return await self.net.call("POST", f"{self.drafts}/approve", who or self.priya, body=body)

    async def item(self, who: Person, item: uuid.UUID | None = None) -> Json:
        response = await self.net.call(
            "GET", f"/v1/engagements/{self.site.engagement_id}/request-items", who
        )
        assert response.status_code == 200, response.text
        body = response.json()
        rows = cast(list[Json], body["items"] if isinstance(body, dict) else body)
        return next(r for r in rows if r["id"] == str(item or self.site.item_id))


async def _firm(net: Net, *, synthetic: bool) -> Firm:
    seed = net.seed
    tenant = await seed.firm()
    if synthetic:
        await seed.run("UPDATE firms SET synthetic = true WHERE tenant_id = $1", tenant)
    admin = await seed.person(tenant, "firm_admin")
    priya = await seed.person(tenant)
    staff = await seed.person(tenant)
    site = await seed.site(tenant, admin.user_id, "Harborline")
    await seed.member(site, priya, "senior")
    await seed.member(site, staff, "staff")
    rachel = await seed.client_contact(site, "client_admin")
    return Firm(net, tenant, admin, priya, staff, rachel, site)


@pytest.fixture
async def real(net: Net) -> Firm:
    return await _firm(net, synthetic=False)


@pytest.fixture
async def synthetic(net: Net) -> Firm:
    return await _firm(net, synthetic=True)


def _workbook() -> bytes:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["CASH", "Cash"])
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    requests.append(["CASH", "Bank statements for every account", "D"])
    rules = book.create_sheet("Account rules")
    rules.append(["area_code", "account_from", "account_to"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _shown(row: Json) -> bool:
    return WHY in str(row.get("why")) or GOOD in str(row.get("good_looks_like"))


# --- AC-4: nothing changes until a person approves -----------------------------------------------


async def test_ac4_approving_writes_drafted_guidance_and_who_approved(synthetic: Firm) -> None:
    draft = await synthetic.draft()
    assert not _shown(await synthetic.item(synthetic.priya))  # a draft changes nothing
    response = await synthetic.approve([{"draft_id": str(draft)}])
    assert response.status_code == 201, response.text
    row = await synthetic.item(synthetic.priya)
    assert (row["why"], row["good_looks_like"], row["guidance_source"]) == (WHY, GOOD, "drafted")
    assert row["guidance_approved_by"] == str(synthetic.priya.user_id)
    again = await synthetic.approve([{"draft_id": str(draft)}])
    assert (again.status_code, again.json()["detail"]) == (409, "draft_already_decided")


async def test_ac4_an_edited_approval_keeps_the_persons_words(synthetic: Firm) -> None:
    draft = await synthetic.draft()
    edited = "We match the bank's own record of the cash balance to the ledger."
    response = await synthetic.approve([{"draft_id": str(draft), "why": edited}])
    assert response.status_code == 201, response.text
    row = await synthetic.item(synthetic.priya)
    assert (row["why"], row["guidance_source"]) == (edited, "drafted")


async def test_ac4_discarding_records_it_and_changes_nothing(synthetic: Firm) -> None:
    draft = await synthetic.draft()
    before = await synthetic.item(synthetic.priya)
    response = await synthetic.net.call(
        "POST",
        f"{synthetic.drafts}/discard",
        synthetic.priya,
        body={"approvals": [{"draft_id": str(draft)}], "confirmed_read": True},
    )
    assert response.status_code == 201, response.text
    after = await synthetic.item(synthetic.priya)
    assert (after["why"], after["guidance_source"]) == (before["why"], before["guidance_source"])
    decided = await synthetic.seed.value(
        "SELECT count(*) FROM guidance_draft_decisions WHERE draft_id = $1", draft
    )
    assert decided == 1


async def test_q5_bulk_approval_needs_the_person_to_confirm_reading_each(
    synthetic: Firm,
) -> None:
    other = await synthetic.seed.item(synthetic.site, "Loan agreements")
    drafts = [await synthetic.draft(), await synthetic.draft(other)]
    response = await synthetic.approve(
        [{"draft_id": str(d)} for d in drafts], confirmed_read=False
    )
    assert (response.status_code, response.json()["detail"]) == (422, "read_not_confirmed")
    assert not _shown(await synthetic.item(synthetic.priya))


async def test_ac4_only_partner_manager_and_senior_decide_and_never_a_client(
    synthetic: Firm,
) -> None:
    draft = await synthetic.draft()
    for who in (synthetic.staff, synthetic.rachel):
        response = await synthetic.approve([{"draft_id": str(draft)}], who=who)
        assert response.status_code in (403, 404), who
    response = await synthetic.net.call("GET", synthetic.drafts, synthetic.rachel)
    assert response.status_code in (403, 404)  # drafts never reach clients


# --- AC-5, Q2 (ADR-111): client display ----------------------------------------------------------


async def test_ac5_a_synthetic_firm_shows_drafted_guidance_to_clients_only_when_set(
    synthetic: Firm,
) -> None:
    draft = await synthetic.draft()
    assert (await synthetic.approve([{"draft_id": str(draft)}])).status_code == 201
    assert not _shown(await synthetic.item(synthetic.rachel))  # off by default
    assert _shown(await synthetic.item(synthetic.priya))  # the firm always sees it
    response = await synthetic.net.call(
        "PUT", "/v1/firm/drafted-guidance", synthetic.admin, body={"shown": True}
    )
    assert response.status_code == 200, response.text
    assert _shown(await synthetic.item(synthetic.rachel))
    response = await synthetic.net.call(
        "PUT", "/v1/firm/drafted-guidance", synthetic.admin, body={"shown": False}
    )
    assert response.status_code == 200, response.text
    assert not _shown(await synthetic.item(synthetic.rachel))


async def test_adr111_a_real_firm_never_shows_model_drafted_guidance_to_a_client(
    real: Firm,
) -> None:
    """Q-03-T040's line. The setting can't be turned on for a real firm, and even with the
    column forced on underneath the product, no client read carries the drafted text."""
    response = await real.net.call(
        "PUT", "/v1/firm/drafted-guidance", real.admin, body={"shown": True}
    )
    assert (response.status_code, response.json()["detail"]) == (
        409,
        "drafted_guidance_not_allowed",
    )
    # Drafted guidance on the item as if approved, and the setting forced on as the owner.
    await real.seed.run(
        "UPDATE request_items SET why = $2, good_looks_like = $3, guidance_source = 'drafted', "
        "guidance_approved_by = $4 WHERE id = $1",
        real.site.item_id,
        WHY,
        GOOD,
        real.priya.user_id,
    )
    await real.seed.run(
        "UPDATE firms SET drafted_guidance_to_clients = true WHERE tenant_id = $1", real.tenant_id
    )
    client_view = await real.item(real.rachel)
    assert not _shown(client_view)
    assert client_view.get("guidance_source") != "drafted"
    response = await real.net.call(
        "GET",
        f"/v1/engagements/{real.site.engagement_id}/request-items/{real.site.item_id}/comments",
        real.rachel,
    )
    assert WHY not in response.text and GOOD not in response.text
    assert _shown(await real.item(real.priya))  # the firm still sees its own approved text


# --- AC-6: a later firm edit is the firm's own words ---------------------------------------------


async def test_ac6_editing_drafted_guidance_makes_it_firm_guidance(synthetic: Firm) -> None:
    draft = await synthetic.draft()
    assert (await synthetic.approve([{"draft_id": str(draft)}])).status_code == 201
    response = await synthetic.net.call(
        "PUT",
        f"/v1/engagements/{synthetic.site.engagement_id}/request-items/"
        f"{synthetic.site.item_id}/guidance",
        synthetic.priya,
        body={"why": "Our own words now.", "good_looks_like": GOOD, "blocks_fieldwork": False},
    )
    assert response.status_code == 200, response.text
    row = await synthetic.item(synthetic.priya)
    assert (row["why"], row["guidance_source"]) == ("Our own words now.", "firm")


# --- Q3 (TASK-074): template items ---------------------------------------------------------------


async def test_q3_template_drafting_is_for_firm_admins_and_synthetic_firms(
    real: Firm,
) -> None:
    imported = await real.net.http.post(
        "/v1/methodology/templates/BH%20EBP/versions?engagement_type=ebp_audit",
        content=_workbook(),
        headers={**real.net.headers(real.admin), "Content-Type": "application/octet-stream"},
    )
    assert imported.status_code == 201, imported.text
    version_id = imported.json()["version_id"]
    path = f"/v1/methodology/versions/{version_id}/draft-guidance"
    response = await real.net.call("POST", path, real.priya, body={"positions": [0]})
    assert response.status_code in (403, 404)  # template drafting is the firm admin's
    response = await real.net.call("POST", path, real.admin, body={"positions": [0]})
    assert (response.status_code, response.json()["detail"]) == (409, "ai_drafting_unavailable")
