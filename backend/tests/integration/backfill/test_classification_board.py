# ac-test-author: ayushkbhatia
"""Q-03-T024: back-fill of TASK-038 (SPEC-022 AC-1 to AC-5: classification, automatic retrieval
and the evidence board's summary).

Classification follows SPEC-022 §4's order (override, then the methodology's tier, then keyword
rules, else unclassified), is audited (`request_item.classified`), and an override is set or
cleared only by `request_item.update`'s roles (`request_item.tier_overridden`). An A item whose
dataset the live connection can't deliver isn't available; automatic retrieval starts once per
open A item and period, behind `retrieval.auto`. The board summary counts by status and tier and
gives the retrievable share (A to C, as a percentage) and "retrieved, never asked" (Q5: first
evidence retrieved, no client upload before it).

Boundary: `create_app()` over HTTP; the outbox relay (`run_relay` with the worker's publisher)
delivers `request_item.classified` to the auto-retrieval subscriber, and a recording stand-in
for the Temporal client stands in for the retrieval workflow (only the start is observed). The
connection is the fake provider's (it declares `trial_balance` only). The Board's AC-6 tokens
check is in `apps/web/src/legacy/q03ledgerstates.test.tsx`; its AC-5 URL filters stay with the
Board's own component tests (`legacy/Board.test.tsx`), not re-written here.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import uuid
from collections.abc import AsyncIterator, Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Final, cast

import httpx
import pytest
from openpyxl import Workbook
from temporalio.client import Client

from abacus.kernel.config import settings
from abacus.kernel.temporal import configure_temporal_client
from abacus.kernel.uow.relay import run_relay
from abacus.worker.__main__ import publisher

from .support import Json, Net, Person, Seeder, Site

TB: Final = "Year-end trial balance"
GL: Final = "General ledger detail for the year"
CONFIRMATION: Final = "Bank confirmation from First Bank"
ODD: Final = "Zebra crossing permit"  # matches no rule


class Recorder:
    """Stands in for the Temporal client: records workflow starts."""

    def __init__(self) -> None:
        self.calls: list[tuple[tuple[object, ...], dict[str, object]]] = []

    async def start_workflow(self, *args: object, **kwargs: object) -> None:
        self.calls.append((args, kwargs))


@pytest.fixture(autouse=True)
def recorder(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Iterator[Recorder]:
    monkeypatch.setenv("ABACUS_TEMPORAL_TASK_QUEUE", f"backfill-board-{uuid.uuid4().hex[:8]}")
    # The fake provider (the demo ledger) is offered only where its fixture directory is set.
    monkeypatch.setenv("ABACUS_FAKE_CONNECTOR_DIR", str(tmp_path))
    settings.cache_clear()
    made = Recorder()
    configure_temporal_client(cast(Client, made))
    yield made
    configure_temporal_client(None)
    settings.cache_clear()


@pytest.fixture
async def relay() -> AsyncIterator[None]:
    stop = asyncio.Event()
    task = asyncio.create_task(run_relay(publisher(), stop, interval=0.1))
    yield
    stop.set()
    await task


@dataclass(frozen=True)
class Board:
    net: Net
    tenant_id: uuid.UUID
    admin: Person
    partner: Person
    manager: Person
    staff: Person
    site: Site

    @property
    def seed(self) -> Seeder:
        return self.net.seed

    @property
    def items_path(self) -> str:
        return f"/v1/engagements/{self.site.engagement_id}/request-items"

    async def create(self, description: str) -> uuid.UUID:
        response = await self.net.call(
            "POST",
            self.items_path,
            self.manager,
            body={"description": description, "audit_area": "Ledger"},
        )
        assert response.status_code == 201, response.text
        return uuid.UUID(cast(str, response.json()["id"]))

    async def item(self, item: uuid.UUID, who: Person | None = None) -> Json:
        response = await self.net.call("GET", self.items_path, who or self.manager)
        assert response.status_code == 200, response.text
        found = [i for i in cast(list[Json], response.json()) if i["id"] == str(item)]
        assert len(found) == 1
        return found[0]

    async def tier(
        self, item: uuid.UUID, tier: str | None, who: Person | None = None
    ) -> httpx.Response:
        return await self.net.call(
            "PUT", f"{self.items_path}/{item}/tier", who or self.manager, body={"tier": tier}
        )

    async def audited(self, action: str, item: uuid.UUID) -> int:
        return int(
            cast(
                int,
                await self.seed.value(
                    "SELECT count(*) FROM audit_events WHERE tenant_id = $1 AND action = $2 "
                    "AND target_id = $3",
                    self.tenant_id,
                    action,
                    str(item),
                ),
            )
        )

    async def connect(self, status: str = "active") -> uuid.UUID:
        """The fake provider's connection, made by the engagement's client admin: automatic
        retrieval runs on behalf of whoever connected (Q4)."""
        client_admin = await self.seed.client_contact(self.site)
        return cast(
            uuid.UUID,
            await self.seed.value(
                "INSERT INTO connections (tenant_id, client_entity_id, provider, status, "
                "created_by) VALUES ($1, $2, 'fake', $3, $4) RETURNING id",
                self.tenant_id,
                self.site.entity_id,
                status,
                str(client_admin.user_id),
            ),
        )

    async def runs(self, item: uuid.UUID) -> list[tuple[str, str]]:
        rows = await self.seed.rows(
            "SELECT period_start::text AS s, period_end::text AS e FROM sync_runs "
            "WHERE request_item_id = $1",
            item,
        )
        return [(str(r["s"]), str(r["e"])) for r in rows]


async def _board(net: Net, *, auto: bool = True) -> Board:
    seed = net.seed
    tenant_id = await seed.firm()
    await seed.run("UPDATE firms SET synthetic = true WHERE tenant_id = $1", tenant_id)
    # `binders.engine`: the item PATCH route (P-06) is behind it.
    for flag in ("binders.engine", *(("retrieval.auto",) if auto else ())):
        await seed.run(
            "INSERT INTO feature_flag_states (tenant_id, flag, value, set_by, reason) "
            "VALUES ($1, $2, 'true', 'test-seed', 'Q-03-T024')",
            tenant_id,
            flag,
        )
    admin = await seed.person(tenant_id, "firm_admin")
    site = await seed.site(tenant_id, admin.user_id, "Board")
    people = [await seed.person(tenant_id) for _ in range(3)]
    await seed.member(site, people[0], "engagement_partner")
    await seed.member(site, people[1], "manager")
    await seed.member(site, people[2], "staff")
    return Board(net, tenant_id, admin, people[0], people[1], people[2], site)


@pytest.fixture
async def board(net: Net) -> Board:
    return await _board(net)


# --- AC-1: tier and dataset by the order in §4, each classification audited --------------------


async def test_ac1_a_trial_balance_request_is_tier_a_with_its_dataset(board: Board) -> None:
    item = await board.create(TB)
    found = await board.item(item)
    assert (found["retrievability_tier"], found["dataset"]) == ("A", "trial_balance")
    assert found["tier_source"] == "rule"
    assert await board.audited("request_item.classified", item) >= 1


async def test_ac1_a_bank_confirmation_is_tier_e(board: Board) -> None:
    found = await board.item(await board.create(CONFIRMATION))
    assert (found["retrievability_tier"], found["dataset"]) == ("E", None)


async def test_ac1_a_request_no_rule_matches_is_unclassified(board: Board) -> None:
    found = await board.item(await board.create(ODD))
    assert found["retrievability_tier"] is None


async def test_ac1_the_methodology_tier_beats_the_keyword_rules(board: Board) -> None:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["CASH", "Cash"])
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    requests.append(["CASH", TB, "D"])  # the firm's template says D, the rule would say A
    book.create_sheet("Account rules").append(["area_code", "account_from", "account_to"])
    buffer = io.BytesIO()
    book.save(buffer)
    name = f"Board template {uuid.uuid4().hex[:6]}"
    imported = await board.net.http.post(
        f"/v1/methodology/templates/{name}/versions?engagement_type=fs_audit",
        content=buffer.getvalue(),
        headers={**board.net.headers(board.admin), "Content-Type": "application/octet-stream"},
    )
    assert imported.status_code == 201, imported.text
    applied = await board.net.call(
        "POST",
        f"/v1/engagements/{board.site.engagement_id}/methodology",
        board.partner,
        body={"version_id": imported.json()["version_id"]},
    )
    assert applied.status_code in (200, 201), applied.text
    item = cast(
        uuid.UUID,
        await board.seed.value(
            "SELECT id FROM request_items WHERE engagement_id = $1 AND description = $2",
            board.site.engagement_id,
            TB,
        ),
    )
    found = await board.item(item)
    assert (found["retrievability_tier"], found["tier_source"]) == ("D", "methodology")


async def test_ac1_a_changed_description_is_reclassified(board: Board) -> None:
    """§12: an item's description changes later: reclassified unless overridden."""
    item = await board.create(ODD)
    patched = await board.net.call(
        "PATCH", f"{board.items_path}/{item}", board.manager, body={"description": TB}
    )
    assert patched.status_code == 200, patched.text
    found = await board.item(item)
    assert (found["retrievability_tier"], found["dataset"]) == ("A", "trial_balance")


# --- AC-2: overrides by request_item.update's roles, audited ----------------------------------


async def test_ac2_an_override_wins_and_is_audited(board: Board) -> None:
    item = await board.create(TB)
    response = await board.tier(item, "D")
    assert response.status_code == 200, response.text
    found = await board.item(item)
    assert (found["retrievability_tier"], found["tier_source"]) == ("D", "override")
    assert await board.audited("request_item.tier_overridden", item) == 1


async def test_ac2_clearing_the_override_classifies_again_and_is_audited(board: Board) -> None:
    item = await board.create(TB)
    assert (await board.tier(item, "D")).status_code == 200
    cleared = await board.tier(item, None)
    assert cleared.status_code == 200, cleared.text
    found = await board.item(item)
    assert (found["retrievability_tier"], found["tier_source"]) == ("A", "rule")
    assert await board.audited("request_item.tier_overridden", item) == 2


async def test_ac2_an_override_survives_a_changed_description(board: Board) -> None:
    item = await board.create(ODD)
    assert (await board.tier(item, "B")).status_code == 200
    patched = await board.net.call(
        "PATCH", f"{board.items_path}/{item}", board.manager, body={"description": TB}
    )
    assert patched.status_code == 200, patched.text
    assert (await board.item(item))["retrievability_tier"] == "B"


async def test_ac2_staff_cannot_override_and_nothing_changes(board: Board) -> None:
    item = await board.create(TB)
    response = await board.tier(item, "E", board.staff)
    assert response.status_code == 403, response.text
    assert (await board.item(item))["retrievability_tier"] == "A"
    assert await board.audited("request_item.tier_overridden", item) == 0


async def test_ac2_another_firm_cannot_override(board: Board, net: Net) -> None:
    item = await board.create(TB)
    other = await _board(net)
    response = await board.tier(item, "E", other.manager)
    assert response.status_code == 404, response.text
    assert (await board.item(item))["retrievability_tier"] == "A"


# --- AC-3: availability follows the live connection's capabilities ----------------------------


async def test_ac3_without_a_live_connection_nothing_is_available(board: Board) -> None:
    assert (await board.item(await board.create(TB)))["available"] is False


async def test_ac3_only_the_datasets_the_connection_delivers_are_available(board: Board) -> None:
    tb = await board.create(TB)
    gl = await board.create(GL)
    await board.connect()
    assert (await board.item(tb))["available"] is True
    found = await board.item(gl)
    assert (found["retrievability_tier"], found["dataset"]) == ("A", "general_ledger")
    assert found["available"] is False


async def test_ac3_a_revoked_connection_makes_items_unavailable_again(board: Board) -> None:
    tb = await board.create(TB)
    connection = await board.connect()
    assert (await board.item(tb))["available"] is True
    await board.seed.run(
        "UPDATE connections SET status = 'revoked', revoked_at = now() WHERE id = $1",
        connection,
    )
    assert (await board.item(tb))["available"] is False


# --- AC-4: an item becoming A with an available dataset is retrieved once ---------------------


async def _settle(seconds: float = 3.0) -> None:
    await asyncio.sleep(seconds)


async def test_ac4_a_new_available_a_item_is_retrieved_once_for_the_fiscal_period(
    board: Board, relay: None, recorder: Recorder
) -> None:
    await board.connect()
    item = await board.create(TB)
    deadline = asyncio.get_running_loop().time() + 20
    while not await board.runs(item):
        assert asyncio.get_running_loop().time() < deadline, "no retrieval was started"
        await asyncio.sleep(0.2)
    # Classified again (override set and cleared): the running run is not repeated.
    assert (await board.tier(item, "D")).status_code == 200
    assert (await board.tier(item, None)).status_code == 200
    await _settle()
    assert await board.runs(item) == [("2025-01-01", "2025-12-31")]


async def test_ac4_an_item_whose_dataset_is_not_available_is_not_retrieved(
    board: Board, relay: None
) -> None:
    await board.connect()
    gl = await board.create(GL)
    tier_e = await board.create(CONFIRMATION)
    await _settle()
    assert await board.runs(gl) == []
    assert await board.runs(tier_e) == []


async def test_ac4_with_the_flag_off_nothing_is_retrieved_automatically(
    net: Net, relay: None
) -> None:
    board = await _board(net, auto=False)
    await board.connect()
    item = await board.create(TB)
    await _settle()
    assert await board.runs(item) == []


# --- AC-5: the summary: counts by status, the retrievable share, retrieved never asked --------


async def _retrieved(board: Board, item: uuid.UUID) -> None:
    """The item's first evidence came by retrieval (no client upload before it)."""
    evidence = await board.seed.value(
        "INSERT INTO evidence_items (tenant_id, engagement_id, title, created_by_kind, "
        "created_by_id) VALUES ($1, $2, 'Trial balance', 'system', 'seed') RETURNING id",
        board.tenant_id,
        board.site.engagement_id,
    )
    fingerprint = hashlib.sha256(uuid.uuid4().bytes).hexdigest()
    version = await board.seed.value(
        "INSERT INTO evidence_versions (tenant_id, engagement_id, evidence_item_id, "
        "version_no, fingerprint, storage_key, storage_version_id, size_bytes, media_type, "
        "source, method, pulled_at) VALUES ($1, $2, $3, 1, $4, $5, 'v1', 10, "
        "'application/pdf', 'connector', 'retrieved', now()) RETURNING id",
        board.tenant_id,
        board.site.engagement_id,
        evidence,
        fingerprint,
        f"tenants/{board.tenant_id}/sha256/{fingerprint}",
    )
    await board.seed.run(
        "INSERT INTO fulfilments (tenant_id, engagement_id, request_item_id, "
        "evidence_version_id, created_by_kind, created_by_id) "
        "VALUES ($1, $2, $3, $4, 'rule', 'retrieval')",
        board.tenant_id,
        board.site.engagement_id,
        item,
        version,
    )
    await board.seed.run("UPDATE request_items SET status = 'uploaded' WHERE id = $1", item)


async def _summary(board: Board) -> Json:
    """The site's own item was written by the seed, outside the product, so no rule ever ran on
    it: unclassified. Add A, A, E and one no rule matches: five items, two retrievable."""
    tb = await board.create(TB)
    await board.create(GL)
    await board.create(CONFIRMATION)
    await board.create(ODD)
    await _retrieved(board, tb)
    response = await board.net.call(
        "GET", f"/v1/engagements/{board.site.engagement_id}/board-summary", board.staff
    )
    assert response.status_code == 200, response.text
    return cast(Json, response.json())


async def test_ac5_the_summary_counts_by_status_and_tier(board: Board) -> None:
    summary = await _summary(board)
    assert summary["total"] == 5
    assert summary["by_tier"] == {"A": 2, "E": 1}
    assert summary["unclassified"] == 2
    by_status = cast(dict[str, int], summary["by_status"])
    assert sum(by_status.values()) == 5
    assert by_status.get("uploaded") == 1


async def test_ac5_retrieved_never_asked_counts_items_first_evidenced_by_retrieval(
    board: Board,
) -> None:
    assert (await _summary(board))["retrieved_never_asked"] == 1


async def test_ac5_the_retrievable_share_is_of_the_whole_request_list(board: Board) -> None:
    """SPEC-022 §4: "the retrievable share of the request list (A to C as a percentage)".
    Unclassified items are on the list (Q3), so they count in the denominator. The wire form
    (a fraction or a percentage) is not fixed by the spec: either is accepted."""
    share = cast(float, (await _summary(board))["retrievable_share"])
    assert share == pytest.approx(0.4) or share == pytest.approx(40.0)


async def test_ac5_another_firm_cannot_read_the_summary(board: Board, net: Net) -> None:
    other = await _board(net)
    response = await board.net.call(
        "GET", f"/v1/engagements/{board.site.engagement_id}/board-summary", other.manager
    )
    assert response.status_code == 404, response.text
