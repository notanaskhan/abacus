"""K-01-T05 (TASK-1427): workbook v2 imports and the methodology_programs migration.

Acceptance (K-01.tasks.csv): a BH EBP v3-shaped workbook imports 7 programs, 67 steps and 24
requests, a legacy three-sheet workbook still imports, and the database refuses any update to a
published version's steps. Also: with `methodology.lifecycle` on the import lands Needs review
(and a draft's step may still be edited); step keys carry over to the next version; the blank
template downloads; the new tables isolate firms (forced RLS).

The workbook is built here from synthetic rows (no firm content in code: rule 16).
"""

from __future__ import annotations

import io
import uuid
from typing import Final, cast

import asyncpg
import pytest
from openpyxl import Workbook, load_workbook

from abacus.kernel.flags import forget_flags

from .support import Migrated, Net, Person

pytestmark = pytest.mark.asyncio

# BH EBP v3's shape: seven programs of 16, 8, 6, 8, 10, 7 and 12 steps; 24 requests.
SHAPE: Final = (16, 8, 6, 8, 10, 7, 12)
LEVELS: Final = ("Foundation", "Planning", "Testing", "Completion")


def _workbook(*, programs: bool = True, requests: int = 24, changed: int | None = None) -> bytes:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["PLAN", "Planning"])
    sheet = book.create_sheet("Requests")
    sheet.append(["area_code", "description", "tier"])
    for n in range(requests):
        sheet.append(["PLAN", f"Request {n + 1}", "B"])
    sheet = book.create_sheet("Account rules")
    sheet.append(["area_code", "account_from", "account_to"])
    if programs:
        sheet = book.create_sheet("Programs")
        sheet.append(["code", "name"])
        for p in range(len(SHAPE)):
            sheet.append([f"BH-EBP-{110 + 10 * p}", f"Program {p + 1}"])
        sheet = book.create_sheet("Steps")
        sheet.append(["program_code", "number", "text", "level", "mandatory", "requests"])
        for p, count in enumerate(SHAPE):
            for n in range(1, count + 1):
                text = f"Step {n} of program {p + 1}"
                if changed is not None and p == 0 and n == changed:
                    text += " (revised)"
                sheet.append(
                    [
                        f"BH-EBP-{110 + 10 * p}",
                        n,
                        text,
                        LEVELS[n % 4],
                        "yes" if n == 1 else "no",
                        str(n) if n <= requests else "",
                    ]
                )
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def _flags(net: Net, tenant: uuid.UUID, *flags: str) -> None:
    for flag in flags:
        await net.seed.run(
            "INSERT INTO feature_flag_states (tenant_id, flag, value, set_by, reason) "
            "VALUES ($1, $2, 'true', 'test-seed', 'K-01-T05')",
            tenant,
            flag,
        )
    forget_flags()  # the 30-second flag cache: the next request reads the new value


async def _import(net: Net, who: Person, name: str, data: bytes) -> uuid.UUID:
    response = await net.http.post(
        f"/v1/methodology/templates/{name}/versions?engagement_type=ebp_audit&family=ebp",
        content=data,
        headers={**net.headers(who), "Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 201, response.text
    return uuid.UUID(cast(str, response.json()["version_id"]))


async def _counts(net: Net, version: uuid.UUID) -> tuple[int, int, int, int]:
    row = (
        await net.seed.rows(
            "SELECT (SELECT count(*) FROM programs WHERE version_id = $1) AS programs, "
            "(SELECT count(*) FROM program_steps WHERE version_id = $1) AS steps, "
            "(SELECT count(*) FROM methodology_request_items WHERE version_id = $1) AS requests, "
            "(SELECT count(*) FROM program_step_pbc_links WHERE version_id = $1) AS links",
            version,
        )
    )[0]
    return row["programs"], row["steps"], row["requests"], row["links"]


async def _status(net: Net, version: uuid.UUID) -> str:
    return cast(
        str,
        await net.seed.value(
            "SELECT status FROM methodology_version_states WHERE version_id = $1", version
        ),
    )


def _dsn(url: str) -> str:
    return url.replace("postgresql+asyncpg://", "postgresql://")


async def _as_app(db: Migrated, tenant: uuid.UUID, who: Person, sql: str, *args: object) -> None:
    conn = await asyncpg.connect(_dsn(db.app_url))
    try:
        async with conn.transaction():
            await conn.execute(
                "SELECT set_config('app.tenant_id', $1, true), "
                "set_config('app.actor_kind', 'human', true), "
                "set_config('app.actor_id', $2, true)",
                str(tenant),
                str(who.user_id),
            )
            await conn.execute(sql, *args)
    finally:
        await conn.close()


async def test_ac_the_v3_workbook_lands_needs_review_with_its_programs(net: Net) -> None:
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    await _flags(net, tenant, "methodology.programs", "methodology.lifecycle")
    version = await _import(net, admin, "BH Audit Methodology", _workbook())
    assert await _counts(net, version) == (7, 67, 24, 67)
    assert await _status(net, version) == "needs_review"
    mandatory = await net.seed.value(
        "SELECT count(*) FROM program_steps WHERE version_id = $1 AND mandatory", version
    )
    assert mandatory == 7
    event = await net.seed.rows(
        "SELECT kind, facts FROM methodology_version_events WHERE version_id = $1", version
    )
    assert [r["kind"] for r in event] == ["imported"]
    audit = await net.seed.value(
        "SELECT count(*) FROM audit_events "
        "WHERE tenant_id = $1 AND action = 'methodology.imported'",
        tenant,
    )
    assert audit == 1


async def test_ac_a_legacy_three_sheet_workbook_still_imports(net: Net) -> None:
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    await _flags(net, tenant, "methodology.programs")
    version = await _import(net, admin, "BH PBC", _workbook(programs=False))
    assert await _counts(net, version) == (0, 0, 24, 0)
    assert await _status(net, version) == "published"  # lifecycle off: published as before


async def test_ac_the_database_refuses_any_change_to_a_published_versions_steps(
    net: Net, migrated_db: Migrated
) -> None:
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    await _flags(net, tenant, "methodology.programs")
    published = await _import(net, admin, "BH Audit Methodology", _workbook())
    assert await _status(net, published) == "published"
    program = await net.seed.value("SELECT id FROM programs WHERE version_id = $1", published)
    for sql, args in (
        ("UPDATE program_steps SET text = 'changed' WHERE version_id = $1", (published,)),
        ("UPDATE program_steps SET row_version = 2 WHERE version_id = $1", (published,)),
        ("DELETE FROM program_steps WHERE version_id = $1", (published,)),
        (
            "INSERT INTO program_steps (tenant_id, version_id, program_id, step_key, number, "
            "text, level, origin) VALUES ($1, $2, $3, 'new-step', 99, 'x', 'planning', "
            "'manual')",
            (tenant, published, program),
        ),
        (
            "INSERT INTO program_step_pbc_links (tenant_id, version_id, step_id, "
            "request_position) SELECT tenant_id, version_id, id, 23 FROM program_steps "
            "WHERE version_id = $1 AND number = 2 LIMIT 1",
            (published,),
        ),
    ):
        with pytest.raises(asyncpg.InsufficientPrivilegeError):
            await _as_app(migrated_db, tenant, admin, sql, *args)
    # Even the superuser's seeding is held to it: only the owner's migrations are not.
    with pytest.raises(asyncpg.InsufficientPrivilegeError):
        await net.seed.run("UPDATE program_steps SET text = 'x' WHERE version_id = $1", published)
    # A published step's lock never changes; in a draft only a person changes it (K-01-T14, 0464).
    with pytest.raises(asyncpg.InsufficientPrivilegeError):
        await _as_app(
            migrated_db,
            tenant,
            admin,
            "UPDATE program_steps SET mandatory = false WHERE version_id = $1",
            published,
        )
    # In review is frozen too (security review M2): what a person approves is what they read.
    await _flags(net, tenant, "methodology.lifecycle")
    review = await _import(net, admin, "BH Audit Methodology", _workbook(changed=4))
    assert await _status(net, review) == "needs_review"
    edit = (
        "UPDATE program_steps SET text = 'edited', row_version = row_version + 1 "
        "WHERE version_id = $1 AND number = 3"
    )
    with pytest.raises(asyncpg.InsufficientPrivilegeError):
        await _as_app(migrated_db, tenant, admin, edit, review)
    # A draft (K-01-T14's path) is the one state whose content changes.
    await net.seed.run(
        "UPDATE methodology_version_states SET status = 'draft' WHERE version_id = $1", review
    )
    await _as_app(migrated_db, tenant, admin, edit, review)


async def test_m2_an_edit_holds_the_version_so_a_publish_waits_for_it(
    net: Net, migrated_db: Migrated
) -> None:
    """The frozen trigger share-locks the version's state row: while a draft edit is open, a
    publish (an update of that row) cannot run, so content never changes under an approval."""
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    await _flags(net, tenant, "methodology.programs", "methodology.lifecycle")
    version = await _import(net, admin, "BH Audit Methodology", _workbook())
    await net.seed.run(
        "UPDATE methodology_version_states SET status = 'draft' WHERE version_id = $1", version
    )
    editor = await asyncpg.connect(_dsn(migrated_db.app_url))
    publisher = await asyncpg.connect(migrated_db.superuser_dsn)
    try:
        # The edit's transaction stays open while the publish tries; leaving it by the
        # sentinel rolls it back.
        with pytest.raises(_Held):
            async with editor.transaction():
                await editor.execute(
                    "SELECT set_config('app.tenant_id', $1, true), "
                    "set_config('app.actor_kind', 'human', true), "
                    "set_config('app.actor_id', $2, true)",
                    str(tenant),
                    str(admin.user_id),
                )
                await editor.execute(
                    "UPDATE program_steps SET text = 'edited' "
                    "WHERE version_id = $1 AND number = 2",
                    version,
                )
                await publisher.execute("SET lock_timeout = '500ms'")
                with pytest.raises(asyncpg.LockNotAvailableError):
                    await publisher.execute(
                        "UPDATE methodology_version_states SET status = 'needs_review' "
                        "WHERE version_id = $1",
                        version,
                    )
                raise _Held
    finally:
        await editor.close()
        await publisher.close()


class _Held(Exception):
    """Leaves the edit's transaction without committing it."""


async def test_ac24_step_keys_carry_over_to_the_next_version(net: Net) -> None:
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    await _flags(net, tenant, "methodology.programs")
    v2 = await _import(net, admin, "BH Audit Methodology", _workbook())
    v3 = await _import(net, admin, "BH Audit Methodology", _workbook(changed=4))
    keys = {}
    for version in (v2, v3):
        rows = await net.seed.rows(
            "SELECT p.code, s.number, s.step_key FROM program_steps s "
            "JOIN programs p ON p.id = s.program_id WHERE s.version_id = $1",
            version,
        )
        keys[version] = {(r["code"], r["number"]): r["step_key"] for r in rows}
    assert keys[v2] == keys[v3]
    assert keys[v2][("BH-EBP-110", 4)] == "bh-ebp-110-4"


async def test_ac23_a_bad_workbook_answers_codes_and_stores_nothing(net: Net) -> None:
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    await _flags(net, tenant, "methodology.programs")
    book = load_workbook(io.BytesIO(_workbook()))
    book["Steps"].append(["BH-EBP-999", 1, "Secret client words", "Planning"])
    buffer = io.BytesIO()
    book.save(buffer)
    response = await net.http.post(
        "/v1/methodology/templates/Bad/versions",
        content=buffer.getvalue(),
        headers={**net.headers(admin), "Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 422
    assert "Secret client words" not in response.text
    assert "unknown_program" in response.text
    assert await net.seed.value("SELECT count(*) FROM programs WHERE tenant_id = $1", tenant) == 0


async def test_the_blank_template_downloads(net: Net) -> None:
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    response = await net.call("GET", "/v1/methodology/workbook-template.xlsx", admin)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    book = load_workbook(io.BytesIO(response.content), read_only=True)
    assert "Programs" in book.sheetnames and "Steps" in book.sheetnames
    again = await net.call("GET", "/v1/methodology/workbook-template.xlsx", admin)
    assert again.content == response.content


async def test_the_new_tables_isolate_firms(net: Net, migrated_db: Migrated) -> None:
    tenant = await net.seed.firm()
    admin = await net.seed.person(tenant, "firm_admin")
    await _flags(net, tenant, "methodology.programs")
    await _import(net, admin, "BH Audit Methodology", _workbook())
    other = await net.seed.firm()
    conn = await asyncpg.connect(_dsn(migrated_db.app_url))
    try:
        async with conn.transaction():
            await conn.execute("SELECT set_config('app.tenant_id', $1, true)", str(other))
            for sql in (
                "SELECT count(*) FROM programs",
                "SELECT count(*) FROM program_steps",
                "SELECT count(*) FROM program_step_pbc_links",
            ):
                assert await conn.fetchval(sql) == 0
    finally:
        await conn.close()
