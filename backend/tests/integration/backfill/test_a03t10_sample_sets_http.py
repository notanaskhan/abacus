"""A-03-T10 (TASK-2614; A-03 §3 F3; SPEC-580; ADR-340; DEC-1662) over HTTP against Postgres: sample
sets, the draw, the override and the level-2 approval.

Row acceptance: "Marcus cannot approve a version Marcus prepared, Hannah's approval writes one L2
sign-off and sample_set.approved and supersedes the earlier approved version, and an untied
population answers 409 population_not_tied."

The census parse and the evidence read stand aside (P-03's, as in the T07 suite); the workflow
start runs the draw's activity inline (Temporal's own path is the replay test's). Everything else
is real: the matrix, the flag, the engagement, the frozen-version trigger, the insert-only items,
review's sign-offs, the audit and the outbox.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Iterator
from decimal import Decimal
from typing import Any, cast

import httpx
import pytest
from openpyxl import load_workbook

from abacus.modules.evidence.api import ExtractedTable, Extraction
from abacus.modules.review.api import STATEMENTS
from abacus.modules.sampling import draw as draw_module
from abacus.modules.sampling import populations, sample_sets
from abacus.modules.sampling.draw_types import SampleDrawInput

from .support import Net, Person, Seeder, Site
from .tb_support import work_item

Json = dict[str, Any]  # Any: JSON bodies
HEADER = ("Participant ID", "Name", "Hire date", "Compensation")
MAP = {
    "record_key": "Participant ID",
    "name": "Name",
    "date": "Hire date",
    "amount": "Compensation",
}
PREPARER = STATEMENTS[0].version
STATISTICAL = {
    "method": "attribute_statistical",
    "parameters": {"tolerable_rate": "0.10", "expected_rate": "0", "confidence": "0.90"},
}


def _census(count: int, total: Decimal) -> tuple[tuple[str, ...], ...]:
    each = (total / count).quantize(Decimal("1"))
    rows: list[tuple[str, ...]] = [HEADER]
    for n in range(count):
        amount = each if n < count - 1 else total - each * (count - 1)
        name = "Benjamin Okoro" if n == 158 else f"Participant {n}"
        rows.append((f"NW-{1000 + n}", name, "2020-01-06", str(amount)))
    return tuple(rows)


@pytest.fixture(autouse=True)
def stand_ins(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """The census parse stands aside; the workflow start runs the draw inline."""

    async def read_version(ctx: object, version_id: uuid.UUID) -> bytes:
        return b"census"

    def extract_isolated(content: bytes, media_type: str) -> Extraction:
        rows = _census(412, Decimal("14560000.00"))
        return Extraction("extracted", tables=(ExtractedTable("Census", None, rows),))

    async def dispatch(workflow: type, given: SampleDrawInput, **kw: object) -> None:
        await draw_module.run_draw(uuid.UUID(given.tenant_id), uuid.UUID(given.version_id))

    monkeypatch.setattr(populations, "read_version", read_version)
    monkeypatch.setattr(populations, "extract_isolated", extract_isolated)
    monkeypatch.setattr(sample_sets, "dispatch", dispatch)
    yield


class Team:
    def __init__(self, site: Site, marcus: Person, hannah: Person, item: uuid.UUID) -> None:
        self.site, self.marcus, self.hannah, self.item = site, marcus, hannah, item


async def _team(seed: Seeder) -> Team:
    tenant = await seed.firm()
    marcus = await seed.person(tenant)
    hannah = await seed.person(tenant)
    site = await seed.site(tenant, marcus.user_id, "sampling")
    await seed.member(site, marcus, "senior")
    await seed.member(site, hannah, "manager")
    await seed.run(
        "INSERT INTO feature_flag_states (tenant_id, flag, value, set_by, reason) "
        "VALUES ($1, 'sampling.sample_sets', 'true', 'a03', 'TASK-2614')",
        tenant,
    )
    return Team(site, marcus, hannah, await work_item(seed, site, marcus.user_id))


async def _population(net: Net, team: Team, *, tied: bool = True) -> str:
    made = await net.call(
        "POST",
        f"/v1/work-items/{team.item}/populations",
        team.marcus,
        body={
            "source": {
                "kind": "upload",
                "evidence_version_id": str(team.site.version_id),
                "column_map": MAP,
            },
            "sampling_unit": "participant (census record)",
            "deviation_definition": "any attribute not agreed to source",
        },
    )
    assert made.status_code == 201, made.text
    pid = cast(str, made.json()["id"])
    amount = "14560000.00" if tied else "14742300.00"
    tie = await net.call(
        "POST",
        f"/v1/populations/{pid}/ties",
        team.marcus,
        body={
            "label": "PayGrid YTD register",
            "comparisons": [{"label": "PayGrid YTD register", "amount": amount}],
        },
    )
    assert tie.status_code == 201, tie.text
    return pid


async def _set(net: Net, team: Team) -> str:
    made = await net.call(
        "POST",
        f"/v1/work-items/{team.item}/sample-sets",
        team.marcus,
        body={"title": "Participants", "purpose": "Eligibility and deferrals"},
    )
    assert made.status_code == 201, made.text
    return cast(str, made.json()["id"])


async def _draw(
    net: Net, team: Team, sid: str, pid: str, extra: Json | None = None, key: str | None = None
) -> int:
    body: Json = {"population_id": pid, **STATISTICAL, **(extra or {})}
    started = await net.http.post(
        f"/v1/sample-sets/{sid}/versions",
        json=body,
        headers={**net.headers(team.marcus), "Idempotency-Key": key or str(uuid.uuid4())},
    )
    assert started.status_code == 202, started.text
    version_no = cast(int, started.json()["version_no"])
    assert started.json()["status_url"] == f"/v1/sample-sets/{sid}/versions/{version_no}/draw"
    return version_no


async def _view(net: Net, who: Person, sid: str, version_no: int) -> Json:
    found = await net.call("GET", f"/v1/sample-sets/{sid}?version={version_no}", who)
    assert found.status_code == 200, found.text
    return cast(Json, found.json())


async def _prepare(net: Net, team: Team, view: Json, version_id: uuid.UUID) -> None:
    signed = await net.call(
        "POST",
        "/v1/sign-offs",
        team.marcus,
        body={
            "target_type": "sample_set_version",
            "target_id": str(version_id),
            "level": "L1_preparer",
            "seen_hash": "sha256:" + view["content_hash"],
            "statement_version": PREPARER,
        },
    )
    assert signed.status_code == 201, signed.text


async def _version_id(seed: Seeder, sid: str, version_no: int) -> uuid.UUID:
    return cast(
        uuid.UUID,
        await seed.value(
            "SELECT id FROM sample_set_versions WHERE sample_set_id = $1 AND version_no = $2",
            uuid.UUID(sid),
            version_no,
        ),
    )


async def _approve(
    net: Net, who: Person, sid: str, n: int, view: Json, **extra: object
) -> httpx.Response:
    return await net.call(
        "POST",
        f"/v1/sample-sets/{sid}/versions/{n}/approve",
        who,
        body={"seen_hash": view["content_hash"], **extra},
    )


async def test_ac_a_draw_reaches_draft_with_its_items_seed_and_versions(
    net: Net, seed: Seeder
) -> None:
    team = await _team(seed)
    pid = await _population(net, team)
    sid = await _set(net, team)
    n = await _draw(net, team, sid, pid)
    progress = await net.call("GET", f"/v1/sample-sets/{sid}/versions/{n}/draw", team.marcus)
    assert progress.status_code == 200, progress.text
    assert progress.json()["status"] == "draft"
    assert progress.json()["items_done"] == 22  # the exact binomial size at 90% / 10% / 0%
    view = await _view(net, team.marcus, sid, n)
    params = view["parameters"]
    assert (params["computed_size"], params["final_size"], params["seed_source"]) == (
        22,
        22,
        "server",
    )
    assert 0 <= params["seed"] <= 16_777_215
    assert params["seed_hex"] == f"0x{params['seed']:06X}"
    assert view["items_total"] == 22
    assert len(view["content_hash"]) == 64
    assert view["versions"][0]["tie_result"] == "tie"
    assert "sample_set.approve" not in view["allowed_actions"]  # Marcus is a senior
    picker = await net.call("GET", f"/v1/work-items/{team.item}/sample-sets", team.marcus)
    assert [
        (s["title"], s["final_size"], s["population_count"]) for s in picker.json()["items"]
    ] == [("Participants", 22, 412)]
    actions = await seed.actions(team.site.tenant_id)
    assert actions.count("sample_set.created") == 1
    assert actions.count("sample_set.drawn") == 1
    assert actions.count("sample_set.draw_completed") == 1
    outbox = await seed.value(
        "SELECT count(*) FROM outbox WHERE tenant_id = $1 AND event_type = $2",
        team.site.tenant_id,
        "sample_set.approval_requested",
    )
    assert outbox == 1


async def test_ac_a_replay_of_the_same_key_draws_nothing(net: Net, seed: Seeder) -> None:
    team = await _team(seed)
    pid = await _population(net, team)
    sid = await _set(net, team)
    key = str(uuid.uuid4())
    first = await _draw(net, team, sid, pid, key=key)
    again = await _draw(net, team, sid, pid, key=key)
    assert first == again == 1
    count = await seed.value(
        "SELECT count(*) FROM sample_set_versions WHERE sample_set_id = $1", uuid.UUID(sid)
    )
    assert count == 1


async def test_ac1_marcus_cannot_approve_a_version_marcus_prepared(net: Net, seed: Seeder) -> None:
    team = await _team(seed)
    await seed.member(team.site, team.marcus, "manager")  # even as a manager
    pid = await _population(net, team)
    sid = await _set(net, team)
    n = await _draw(net, team, sid, pid)
    view = await _view(net, team.marcus, sid, n)
    assert "sample_set.approve" not in view["allowed_actions"]
    refused = await _approve(net, team.marcus, sid, n, view)
    assert refused.status_code == 403, refused.text
    status = await seed.value(
        "SELECT status FROM sample_set_versions WHERE sample_set_id = $1", uuid.UUID(sid)
    )
    assert status == "draft"


async def test_ac2_hannahs_approval_writes_one_l2_and_supersedes_the_earlier_one(
    net: Net, seed: Seeder
) -> None:
    team = await _team(seed)
    pid = await _population(net, team)
    sid = await _set(net, team)
    first = await _draw(net, team, sid, pid)
    view = await _view(net, team.hannah, sid, first)
    assert "sample_set.approve" in view["allowed_actions"]
    await _prepare(net, team, view, await _version_id(seed, sid, first))
    approved = await _approve(net, team.hannah, sid, first, view)
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "approved"
    assert approved.json()["versions"][0]["approved"]["user_id"] == str(team.hannah.user_id)
    assert approved.json()["versions"][0]["prepared"]["user_id"] == str(team.marcus.user_id)

    second = await _draw(net, team, sid, pid)
    view2 = await _view(net, team.hannah, sid, second)
    second_id = await _version_id(seed, sid, second)
    await _prepare(net, team, view2, second_id)
    again = await _approve(net, team.hannah, sid, second, view2)
    assert again.status_code == 200, again.text
    rows = await seed.rows(
        "SELECT version_no, status, superseded_by FROM sample_set_versions "
        "WHERE sample_set_id = $1 ORDER BY version_no",
        uuid.UUID(sid),
    )
    assert [(r["version_no"], r["status"], r["superseded_by"]) for r in rows] == [
        (1, "superseded", second_id),
        (2, "approved", None),
    ]
    l2 = await seed.value(
        "SELECT count(*) FROM sign_offs WHERE target_type = 'sample_set_version' "
        "AND target_id = $1 AND level = 2",
        second_id,
    )
    assert l2 == 1
    actions = await seed.actions(team.site.tenant_id)
    assert actions.count("sample_set.approved") == 2
    assert actions.count("sample_set.superseded") == 1
    events = await seed.value(
        "SELECT count(*) FROM outbox WHERE tenant_id = $1 AND event_type = 'sample_set.approved'",
        team.site.tenant_id,
    )
    assert events == 2
    struck = await _view(net, team.hannah, sid, second)
    assert [v["status"] for v in struck["versions"]] == ["superseded", "approved"]


async def test_ac3_an_untied_population_answers_409_population_not_tied(
    net: Net, seed: Seeder
) -> None:
    team = await _team(seed)
    pid = await _population(net, team, tied=False)
    sid = await _set(net, team)
    n = await _draw(net, team, sid, pid)
    view = await _view(net, team.hannah, sid, n)
    assert view["versions"][0]["tie_result"] == "off"
    await _prepare(net, team, view, await _version_id(seed, sid, n))
    refused = await _approve(net, team.hannah, sid, n, view)
    assert refused.status_code == 409, refused.text
    assert refused.json()["detail"] == "population_not_tied"


async def test_a_draw_needs_a_tie_and_never_takes_a_seed(net: Net, seed: Seeder) -> None:
    team = await _team(seed)
    made = await net.call(
        "POST",
        f"/v1/work-items/{team.item}/populations",
        team.marcus,
        body={
            "source": {
                "kind": "upload",
                "evidence_version_id": str(team.site.version_id),
                "column_map": MAP,
            },
            "sampling_unit": "participant",
            "deviation_definition": "any attribute not agreed to source",
        },
    )
    pid = made.json()["id"]
    sid = await _set(net, team)
    untied = await net.http.post(
        f"/v1/sample-sets/{sid}/versions",
        json={"population_id": pid, **STATISTICAL},
        headers={**net.headers(team.marcus), "Idempotency-Key": str(uuid.uuid4())},
    )
    assert untied.status_code == 409, untied.text
    assert untied.json()["detail"] == "completeness_tie_required"
    seeded = await net.http.post(
        f"/v1/sample-sets/{sid}/versions",
        json={"population_id": pid, **STATISTICAL, "seed": 6171281},
        headers={**net.headers(team.marcus), "Idempotency-Key": str(uuid.uuid4())},
    )
    assert seeded.status_code == 422, seeded.text  # no input carries a seed (ADR-340 §5)


async def test_an_override_needs_a_smaller_size_and_the_approvers_acknowledgement(
    net: Net, seed: Seeder
) -> None:
    team = await _team(seed)
    pid = await _population(net, team)
    sid = await _set(net, team)
    bigger = await net.http.post(
        f"/v1/sample-sets/{sid}/versions",
        json={
            "population_id": pid,
            **STATISTICAL,
            "override": {"size": 30, "rationale": "more"},
        },
        headers={**net.headers(team.marcus), "Idempotency-Key": str(uuid.uuid4())},
    )
    assert bigger.status_code == 422, bigger.text
    assert bigger.json()["detail"] == "override_not_below"
    rationale = "Plan has 412 participants; 20 meets the firm's minimum for low RMM."
    n = await _draw(net, team, sid, pid, {"override": {"size": 20, "rationale": rationale}})
    view = await _view(net, team.hannah, sid, n)
    assert (view["parameters"]["override_size"], view["parameters"]["final_size"]) == (20, 20)
    assert view["versions"][0]["override_pending"] is True
    await _prepare(net, team, view, await _version_id(seed, sid, n))
    unticked = await _approve(net, team.hannah, sid, n, view)
    assert unticked.status_code == 409, unticked.text
    assert unticked.json()["detail"] == "override_not_acknowledged"
    ticked = await _approve(net, team.hannah, sid, n, view, acknowledge_override=True)
    assert ticked.status_code == 200, ticked.text
    actions = await seed.actions(team.site.tenant_id)
    assert actions.count("sample_set.override_requested") == 1
    assert actions.count("sample_set.override_approved") == 1


async def test_a_stale_seen_hash_answers_content_changed(net: Net, seed: Seeder) -> None:
    team = await _team(seed)
    pid = await _population(net, team)
    sid = await _set(net, team)
    n = await _draw(net, team, sid, pid)
    view = await _view(net, team.hannah, sid, n)
    await _prepare(net, team, view, await _version_id(seed, sid, n))
    stale = await _approve(net, team.hannah, sid, n, {"content_hash": "0" * 64})
    assert stale.status_code == 409, stale.text
    assert stale.json()["detail"] == "content_changed"


async def test_search_and_the_deterministic_export(net: Net, seed: Seeder) -> None:
    team = await _team(seed)
    pid = await _population(net, team)
    sid = await _set(net, team)
    n = await _draw(
        net,
        team,
        sid,
        pid,
        {"parameters": {"tolerable_rate": "0.01", "expected_rate": "0", "confidence": "0.90"}},
    )
    every = await net.call("GET", f"/v1/sample-sets/{sid}/versions/{n}/items", team.marcus)
    assert every.status_code == 200, every.text
    names = [i["name"] for i in every.json()["items"]]
    if "Benjamin Okoro" in names:
        found = await net.call(
            "GET", f"/v1/sample-sets/{sid}/versions/{n}/items?q=okoro", team.marcus
        )
        assert [i["record_key"] for i in found.json()["items"]] == ["NW-1158"]
    first = await net.call("GET", f"/v1/sample-sets/{sid}/versions/{n}/export.xlsx", team.marcus)
    second = await net.call("GET", f"/v1/sample-sets/{sid}/versions/{n}/export.xlsx", team.marcus)
    assert first.status_code == 200, first.text
    assert first.content == second.content
    cover = {
        r[0].value: r[1].value
        for r in load_workbook(io.BytesIO(first.content))["Sample"].iter_rows()
    }
    view = await _view(net, team.marcus, sid, n)
    assert cover["Seed"] == view["parameters"]["seed_hex"]
    assert cover["Population SHA-256"] == view["population"]["sha256"]
    assert (await seed.actions(team.site.tenant_id)).count("sample_set.exported") == 2


async def test_a_walled_person_sees_no_set_and_no_count(net: Net, seed: Seeder) -> None:
    team = await _team(seed)
    pid = await _population(net, team)
    sid = await _set(net, team)
    await _draw(net, team, sid, pid)
    stranger = await seed.person(team.site.tenant_id)
    read = await net.call("GET", f"/v1/sample-sets/{sid}", stranger)
    assert read.status_code in (403, 404), read.text
    listed = await net.call("GET", f"/v1/work-items/{team.item}/sample-sets", stranger)
    assert listed.status_code in (403, 404), listed.text
