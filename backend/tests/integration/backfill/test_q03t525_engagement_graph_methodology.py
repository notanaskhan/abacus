# ac-test-author: ayushkbhatia
"""Q-03-T525: back-fill of TASK-023 (engagement graph and methodology v1), black-box from SPEC-008.

The legacy paths stay live until K-01 (S26) replaces the methodology import and the graph read;
this suite retires only when K-01's AC suite merges.

Boundary: `create_app()` over HTTP on the migrated Postgres:
`POST /v1/methodology/templates/{name}/versions` (the raw `.xlsx` is the body, TASK-023 D3),
`POST /v1/engagements/{id}/methodology` and `GET /v1/engagements/{id}/graph`, as a firm admin,
the engagement partner and manager, a staff member, a walled person and another firm. A ledger
snapshot is seeded with SQL: the graph reads it, it is not under test. Workbooks are built here
in SPEC-008's fixed layout (Q1): sheets Areas, Requests, Account rules.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime
from decimal import Decimal
from typing import Final, cast

import httpx
import pytest
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

from abacus.kernel.db import TenantContext, tenant_session

from .support import Json, Net, Person, Seeder, Site

CELL_TEXT: Final = "Do-not-echo-this-cell-7731"


def _workbook(
    areas: list[tuple[str, str]],
    items: list[tuple[str, str, str]],
    rules: list[tuple[str, str, str]],
) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Areas"
    sheet.append(["code", "name"])
    for row in areas:
        sheet.append(list(row))
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    for row in items:
        requests.append(list(row))
    account_rules = book.create_sheet("Account rules")
    account_rules.append(["area_code", "account_from", "account_to"])
    for row in rules:
        account_rules.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


GOOD: Final = _workbook(
    [("CASH", "Cash"), ("AR", "Receivables"), ("EQ", "Equity")],
    [("CASH", "Bank statements for every account", "A"), ("AR", "Aged receivables", "B")],
    [("CASH", "1000", "1099"), ("AR", "1100", "1199"), ("EQ", "3000", "3999")],
)


class World:
    def __init__(
        self, net: Net, tenant_id: uuid.UUID, site: Site, people: dict[str, Person]
    ) -> None:
        self.net = net
        self.tenant_id = tenant_id
        self.site = site
        self.people = people

    @property
    def seed(self) -> Seeder:
        return self.net.seed

    async def import_(self, name: str, content: bytes, who: str = "admin") -> httpx.Response:
        return await self.net.http.post(
            f"/v1/methodology/templates/{name}/versions",
            content=content,
            headers={
                **self.net.headers(self.people[who]),
                "Content-Type": "application/octet-stream",
            },
        )

    async def apply(self, version_id: str, who: str = "partner") -> httpx.Response:
        return await self.net.call(
            "POST",
            f"/v1/engagements/{self.site.engagement_id}/methodology",
            self.people[who],
            body={"version_id": version_id},
        )

    async def graph(self, who: str = "partner") -> httpx.Response:
        return await self.net.call(
            "GET", f"/v1/engagements/{self.site.engagement_id}/graph", self.people[who]
        )

    async def count(self, sql: str) -> int:
        return int(cast(int, await self.seed.value(sql, self.tenant_id)))

    async def snapshot(self, lines: list[tuple[str, str, str, str]]) -> uuid.UUID:
        """A ledger snapshot for the engagement's entity: (code, name, debit, credit)."""
        debit = sum(Decimal(d) for _, _, d, _ in lines)
        credit = sum(Decimal(c) for _, _, _, c in lines)
        snapshot = cast(
            uuid.UUID,
            await self.seed.value(
                "INSERT INTO ledger_snapshots (tenant_id, client_entity_id, period_start, "
                "period_end, pulled_at, source, raw_fingerprint, line_count, total_debit, "
                "total_credit) VALUES ($1, $2, '2026-01-01', '2026-12-31', $3, 'quickbooks', "
                "$4, $5, $6, $7) RETURNING id",
                self.tenant_id,
                self.site.entity_id,
                datetime.now(UTC),
                uuid.uuid4().hex * 2,
                len(lines),
                debit,
                credit,
            ),
        )
        for code, name, d, c in lines:
            await self.seed.run(
                "INSERT INTO trial_balance_lines (tenant_id, snapshot_id, account_code, "
                "account_name, debit, credit, source_ref) VALUES ($1, $2, $3, $4, $5, $6, $7)",
                self.tenant_id,
                snapshot,
                code,
                name,
                Decimal(d),
                Decimal(c),
                f"row {code}",
            )
        return snapshot


@pytest.fixture
async def world(net: Net, seed: Seeder) -> World:
    tenant_id = await seed.firm()
    admin = await seed.person(tenant_id, "firm_admin")
    site = await seed.site(tenant_id, admin.user_id, "Graphwell")
    people = {"admin": admin}
    for role in ("partner", "manager", "staff", "walled", "outsider"):
        people[role] = await seed.person(tenant_id)
    await seed.member(site, people["partner"], "engagement_partner")
    await seed.member(site, people["manager"], "manager")
    await seed.member(site, people["staff"], "staff")
    await seed.member(site, people["walled"], "staff")
    await net.wall(admin, people["walled"], site)
    other = await seed.firm()
    people["other_firm"] = await seed.person(other, "firm_admin")
    return World(net, tenant_id, site, people)


def _name() -> str:
    return f"Core-{uuid.uuid4().hex[:6]}"


# --- Story 1: a firm brings its methodology ----------------------------------------------


async def test_ac1_an_import_stores_an_immutable_version_with_its_parts(world: World) -> None:
    response = await world.import_(_name(), GOOD)
    assert response.status_code == 201, response.text
    version = uuid.UUID(cast(str, response.json()["version_id"]))
    counts = await world.seed.rows(
        "SELECT (SELECT count(*) FROM methodology_areas WHERE version_id = $1) AS areas, "
        "(SELECT count(*) FROM methodology_request_items WHERE version_id = $1) AS items, "
        "(SELECT count(*) FROM methodology_account_rules WHERE version_id = $1) AS rules",
        version,
    )
    assert dict(counts[0]) == {"areas": 3, "items": 2, "rules": 3}


async def test_ac1_the_import_is_audited_with_the_version_and_fingerprint(world: World) -> None:
    response = await world.import_(_name(), GOOD)
    assert response.status_code == 201, response.text
    rows = await world.seed.rows(
        "SELECT target_id, after_ref FROM audit_events WHERE tenant_id = $1 "
        "AND action = 'methodology.imported'",
        world.tenant_id,
    )
    assert [r["target_id"] for r in rows] == [response.json()["version_id"]]
    assert "source_fingerprint" in str(rows[0]["after_ref"])


@pytest.mark.parametrize(
    "content",
    [
        _workbook([("CASH", "Cash"), ("CASH", CELL_TEXT)], [], []),  # duplicate area code
        _workbook([("CASH", "Cash")], [("NOPE", CELL_TEXT, "A")], []),  # item for an unknown area
        _workbook([("CASH", "Cash")], [("CASH", CELL_TEXT, "Z")], []),  # invalid tier
        _workbook([("CASH", "Cash")], [], [("CASH", "1099", "1000")]),  # malformed rule
        b"not a workbook " + CELL_TEXT.encode(),
    ],
    ids=["duplicate-area", "unknown-area", "bad-tier", "bad-rule", "not-xlsx"],
)
async def test_ac2_a_bad_workbook_stores_nothing_and_lists_problems_without_cells(
    world: World, content: bytes
) -> None:
    before = await world.count("SELECT count(*) FROM methodology_versions WHERE tenant_id = $1")
    response = await world.import_(_name(), content)
    assert response.status_code == 422, response.text
    assert CELL_TEXT not in response.text
    assert (
        await world.count("SELECT count(*) FROM methodology_versions WHERE tenant_id = $1")
        == before
    )


async def test_ac3_importing_under_the_same_name_makes_the_next_version(world: World) -> None:
    name = _name()
    first = await world.import_(name, GOOD)
    second = await world.import_(name, GOOD)
    assert first.status_code == second.status_code == 201
    versions = await world.seed.rows(
        "SELECT version FROM methodology_versions WHERE id = ANY($1::uuid[]) ORDER BY version",
        [uuid.UUID(first.json()["version_id"]), uuid.UUID(second.json()["version_id"])],
    )
    assert [r["version"] for r in versions] == [1, 2]


async def test_ac3_a_stored_version_cannot_be_updated_or_deleted(world: World) -> None:
    response = await world.import_(_name(), GOOD)
    version = uuid.UUID(cast(str, response.json()["version_id"]))
    admin = world.people["admin"]
    ctx = TenantContext(world.tenant_id, "human", str(admin.user_id))
    for statement in (
        "UPDATE methodology_versions SET source_fingerprint = 'x' WHERE id = :id",
        "DELETE FROM methodology_versions WHERE id = :id",
    ):
        with pytest.raises(DBAPIError):
            async with tenant_session(ctx) as session:
                await session.execute(text(statement), {"id": version})
    assert (
        await world.seed.value("SELECT count(*) FROM methodology_versions WHERE id = $1", version)
        == 1
    )


# --- Story 2: an engagement starts from the methodology ------------------------------------------


async def test_ac4_applying_pins_the_version_and_creates_one_item_per_template_item(
    world: World,
) -> None:
    version = cast(str, (await world.import_(_name(), GOOD)).json()["version_id"])
    response = await world.apply(version)
    assert response.status_code == 201, response.text
    rows = await world.seed.rows(
        "SELECT description, audit_area, retrievability_tier FROM request_items "
        "WHERE engagement_id = $1 ORDER BY description",
        world.site.engagement_id,
    )
    made = {(r["description"], r["audit_area"], r["retrievability_tier"]) for r in rows}
    assert ("Aged receivables", "Receivables", "B") in made
    assert ("Bank statements for every account", "Cash", "A") in made
    pinned = await world.seed.value(
        "SELECT methodology_version_id FROM engagements WHERE id = $1", world.site.engagement_id
    )
    assert str(pinned) == version
    audited = await world.count(
        "SELECT count(*) FROM audit_events WHERE tenant_id = $1 AND action = 'methodology.applied'"
    )
    assert audited == 1


async def test_ac4_staff_cannot_apply_a_template(world: World) -> None:
    version = cast(str, (await world.import_(_name(), GOOD)).json()["version_id"])
    response = await world.apply(version, who="staff")
    assert response.status_code in (403, 404), response.text


async def test_ac5_applying_again_is_409_and_changes_no_item(world: World) -> None:
    version = cast(str, (await world.import_(_name(), GOOD)).json()["version_id"])
    assert (await world.apply(version)).status_code == 201
    before = await world.seed.rows(
        "SELECT id, description, status FROM request_items WHERE engagement_id = $1 ORDER BY id",
        world.site.engagement_id,
    )
    again = await world.apply(version, who="manager")
    assert again.status_code == 409, again.text
    assert "methodology_already_applied" in again.text
    after = await world.seed.rows(
        "SELECT id, description, status FROM request_items WHERE engagement_id = $1 ORDER BY id",
        world.site.engagement_id,
    )
    assert [dict(r) for r in after] == [dict(r) for r in before]


# --- Story 3: the engagement at a glance --------------------------------------------------------


async def _graphed(world: World) -> Json:
    version = cast(str, (await world.import_(_name(), GOOD)).json()["version_id"])
    assert (await world.apply(version)).status_code == 201
    await world.snapshot(
        [
            ("1010", "Operating account", "5000.00", "0"),
            ("1110", "Trade receivables", "1200.00", "0"),
            ("3000", "Retained earnings", "0", "6100.00"),
            ("5500", "Suspense", "100.00", "0"),
            ("5600", "Rounding", "0", "0"),
        ]
    )
    response = await world.graph()
    assert response.status_code == 200, response.text
    return cast(Json, response.json())


def _area(graph: Json, name: str) -> Json:
    found = [a for a in cast(list[Json], graph["areas"]) if a["name"] == name]
    assert len(found) == 1, (name, graph["areas"])
    return found[0]


async def test_ac6_the_graph_holds_engagement_team_version_and_every_area(world: World) -> None:
    graph = await _graphed(world)
    assert cast(Json, graph["engagement"])["id"] == str(world.site.engagement_id)
    assert {m["user_id"] for m in cast(list[Json], graph["team"])} >= {
        str(world.people["partner"].user_id)
    }
    assert graph["methodology"] is not None
    assert {a["name"] for a in cast(list[Json], graph["areas"])} >= {
        "Cash",
        "Receivables",
        "Equity",
    }
    cash = _area(graph, "Cash")
    assert [i["description"] for i in cast(list[Json], cash["items"])] == [
        "Bank statements for every account"
    ]


async def test_ac7_accounts_map_to_the_first_matching_rule_or_unmapped(world: World) -> None:
    graph = await _graphed(world)
    assert [a["code"] for a in cast(list[Json], _area(graph, "Cash")["accounts"])] == ["1010"]
    assert [a["code"] for a in cast(list[Json], _area(graph, "Receivables")["accounts"])] == [
        "1110"
    ]
    assert [a["code"] for a in cast(list[Json], _area(graph, "Equity")["accounts"])] == ["3000"]
    unmapped = {a["code"] for a in cast(list[Json], graph["unmapped_accounts"])}
    assert unmapped == {"5500", "5600"}


async def test_ac8_gaps_name_uncovered_areas_unmapped_balances_and_items_without_evidence(
    world: World,
) -> None:
    graph = await _graphed(world)
    gaps = cast(Json, graph["gaps"])
    assert "Equity" in cast(list[str], gaps["areas_without_requests"])
    assert "Equity" in cast(list[str], gaps["areas_with_accounts_without_requests"])
    # Q5: any unmapped account with a non-zero balance.
    assert [a["code"] for a in cast(list[Json], gaps["unmapped_accounts"])] == ["5500"]
    assert len(cast(list[str], gaps["items_without_evidence"])) == 2


@pytest.mark.parametrize("who", ["walled", "outsider", "other_firm"])
async def test_ac9_no_access_reads_the_same_404(world: World, who: str) -> None:
    await _graphed(world)
    response = await world.graph(who)
    assert response.status_code == 404, response.text
    missing = await world.net.call(
        "GET", f"/v1/engagements/{uuid.uuid4()}/graph", world.people["partner"]
    )
    assert missing.status_code == 404
    assert response.json() == missing.json()
