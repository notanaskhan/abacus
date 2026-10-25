# ac-test-author: ayushkbhatia
"""Q-03-T023: back-fill of TASK-048 (SPEC-025 AC-2: roll forward from last year's engagement).

Given a prior-period engagement of the same type for the entity, the proposal offers last year's
team without inactive or walled people, the template last year used at its latest version, and
the items actually used (accepted evidence: DEC-09's `received`) ticked, the rest unticked and the
template's additions flagged. Nothing is created until a person confirms, and what they confirm
is exactly what is created. Q-03-T023's row acceptance: a roll-forward copies no evidence version
across engagements or their work items: the new items start without evidence.

Boundary: `create_app()` over HTTP: `POST /v1/engagements/proposal` and `POST /v1/engagements`
with `roll_forward`, as a firm administrator (`roll_forward.read` and `engagement.create`), and
the methodology import and application routes. Expectations: SPEC-025 §4, AC-2, §12 (another
type: no roll-forward), the engagements matrix and `error-codes.yaml`
(`team_member_unavailable`). AC-9 for the review is in
`apps/web/src/legacy/q03ledgerstates.test.tsx`.
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from typing import Final, cast

import httpx
import pytest
from openpyxl import Workbook

from .support import Json, Net, Person, Seeder, Site

USED: Final = "Bank statements for every account"
UNUSED: Final = "Fixed asset register"
ADDED: Final = "Lease agreements in force"
COUNTS: Final = {
    "engagements": "SELECT count(*) FROM engagements WHERE tenant_id = $1",
    "request_items": "SELECT count(*) FROM request_items WHERE tenant_id = $1",
}
EVIDENCE: Final = {
    "evidence_items": "SELECT count(*) FROM evidence_items WHERE engagement_id = $1",
    "evidence_versions": "SELECT count(*) FROM evidence_versions WHERE engagement_id = $1",
    "fulfilments": "SELECT count(*) FROM fulfilments WHERE engagement_id = $1",
}


def _workbook(*descriptions: str) -> bytes:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["CASH", "Cash"])
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    for description in descriptions:
        requests.append(["CASH", description, "D"])
    rules = book.create_sheet("Account rules")
    rules.append(["area_code", "account_from", "account_to"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@dataclass(frozen=True)
class Prior:
    net: Net
    tenant_id: uuid.UUID
    admin: Person
    partner: Person
    manager: Person
    senior: Person  # left the firm since
    staff: Person  # walled from the client since
    site: Site
    template: str
    v2: uuid.UUID
    used_item: uuid.UUID

    @property
    def seed(self) -> Seeder:
        return self.net.seed

    async def propose(self, who: Person | None = None, **overrides: object) -> httpx.Response:
        body: Json = {
            "client_entity_id": str(self.site.entity_id),
            "fiscal_period_start": "2026-01-01",
            "type": "fs_audit",
        }
        body.update(overrides)
        return await self.net.call(
            "POST", "/v1/engagements/proposal", who or self.admin, body=body
        )

    async def create(self, roll_forward: Json, who: Person | None = None) -> httpx.Response:
        marker = uuid.uuid4().hex[:6]
        return await self.net.call(
            "POST",
            "/v1/engagements",
            who or self.admin,
            body={
                "name": f"FY2026 {marker}",
                "client_id": str(self.site.client_id),
                "client_entity_id": str(self.site.entity_id),
                "client_name": "ignored for an existing client",
                "client_entity_name": "ignored for an existing entity",
                "fiscal_period_start": "2026-01-01",
                "fiscal_period_end": "2026-12-31",
                "type": "fs_audit",
                "roll_forward": roll_forward,
            },
        )

    async def count(self, table: str) -> int:
        return int(
            cast(
                int,
                await self.seed.value(COUNTS[table], self.tenant_id),
            )
        )


async def _template(net: Net, admin: Person, name: str, *descriptions: str) -> uuid.UUID:
    response = await net.http.post(
        f"/v1/methodology/templates/{name}/versions?engagement_type=fs_audit",
        content=_workbook(*descriptions),
        headers={**net.headers(admin), "Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 201, response.text
    return uuid.UUID(cast(str, response.json()["version_id"]))


@pytest.fixture
async def prior(net: Net, seed: Seeder) -> Prior:
    tenant_id = await seed.firm()
    admin = await seed.person(tenant_id, "firm_admin")
    site = await seed.site(tenant_id, admin.user_id, "Halvorsen")
    await seed.run("UPDATE engagements SET type = 'fs_audit' WHERE id = $1", site.engagement_id)
    team = {
        role: await seed.person(tenant_id) for role in ("partner", "manager", "senior", "staff")
    }
    await seed.member(site, team["partner"], "engagement_partner")
    await seed.member(site, team["manager"], "manager")
    await seed.member(site, team["senior"], "senior")
    await seed.member(site, team["staff"], "staff")
    template = f"Audit core {uuid.uuid4().hex[:6]}"
    v1 = await _template(net, admin, template, USED, UNUSED)
    applied = await net.call(
        "POST",
        f"/v1/engagements/{site.engagement_id}/methodology",
        team["partner"],
        body={"version_id": str(v1)},
    )
    assert applied.status_code in (200, 201), applied.text
    used_item = cast(
        uuid.UUID,
        await seed.value(
            "SELECT id FROM request_items WHERE engagement_id = $1 AND description = $2",
            site.engagement_id,
            USED,
        ),
    )
    # Last year: the bank statements were accepted (DEC-09 `received`) on the site's evidence.
    await seed.run(
        "INSERT INTO fulfilments (tenant_id, engagement_id, request_item_id, evidence_version_id, "
        "created_by_kind, created_by_id) VALUES ($1, $2, $3, $4, 'human', 'seed')",
        tenant_id,
        site.engagement_id,
        used_item,
        site.version_id,
    )
    await seed.run("UPDATE request_items SET status = 'received' WHERE id = $1", used_item)
    v2 = await _template(net, admin, template, USED, UNUSED, ADDED)
    await seed.revoke(team["senior"])  # left the firm
    await net.wall(admin, team["staff"], site)  # walled from the client
    return Prior(
        net,
        tenant_id,
        admin,
        team["partner"],
        team["manager"],
        team["senior"],
        team["staff"],
        site,
        template,
        v2,
        used_item,
    )


def _items(body: Json) -> dict[str, Json]:
    return {cast(str, i["description"]): i for i in cast(list[Json], body["items"])}


# --- AC-2: the proposal --------------------------------------------------------------------------


async def test_ac2_the_prior_engagement_of_the_same_type_is_proposed(prior: Prior) -> None:
    response = await prior.propose()
    assert response.status_code == 200, response.text
    found = cast(Json, response.json()["prior"])
    assert found["engagement_id"] == str(prior.site.engagement_id)


async def test_ac2_the_team_leaves_out_people_who_left_or_are_walled(prior: Prior) -> None:
    response = await prior.propose()
    assert response.status_code == 200, response.text
    team = cast(list[Json], response.json()["team"])
    offered = {cast(str, m["user_id"]): m for m in team if m["available"]}
    assert offered.keys() == {str(prior.partner.user_id), str(prior.manager.user_id)}
    assert offered[str(prior.partner.user_id)]["role"] == "engagement_partner"
    assert offered[str(prior.manager.user_id)]["role"] == "manager"


async def test_ac2_the_template_last_year_used_is_offered_at_its_latest_version(
    prior: Prior,
) -> None:
    response = await prior.propose()
    assert response.status_code == 200, response.text
    template = cast(Json, response.json()["template"])
    assert template["template_name"] == prior.template
    assert (template["version_last_year"], template["latest_version"]) == (1, 2)
    assert template["latest_version_id"] == str(prior.v2)


async def test_ac2_used_items_are_ticked_the_rest_unticked_and_additions_flagged(
    prior: Prior,
) -> None:
    response = await prior.propose()
    assert response.status_code == 200, response.text
    items = _items(response.json())
    assert (items[USED]["kind"], items[USED]["ticked"]) == ("used", True)
    assert items[USED]["prior_item_id"] == str(prior.used_item)
    assert (items[UNUSED]["kind"], items[UNUSED]["ticked"]) == ("not_used", False)
    assert items[ADDED]["kind"] == "new_in_template"
    assert items[ADDED]["prior_item_id"] is None
    # The site's own item was never accepted: unused, unticked.
    seeded = [i for d, i in items.items() if d.startswith("Bank confirmation")]
    assert [(i["kind"], i["ticked"]) for i in seeded] == [("not_used", False)]


async def test_ac2_a_proposal_creates_nothing(prior: Prior) -> None:
    before = (await prior.count("engagements"), await prior.count("request_items"))
    assert (await prior.propose()).status_code == 200
    assert (await prior.count("engagements"), await prior.count("request_items")) == before


async def test_ac2_a_prior_engagement_of_another_type_is_not_rolled_forward(prior: Prior) -> None:
    """SPEC-025 §12: last year's engagement had a different type: the template is offered."""
    response = await prior.propose(type="review")
    assert response.status_code == 200, response.text
    assert response.json()["prior"] is None
    assert response.json()["team"] == []


# --- AC-2: confirming creates exactly what was confirmed, and copies no evidence ----------------


async def _confirmed(prior: Prior) -> uuid.UUID:
    proposal = (await prior.propose()).json()
    added = _items(proposal)[ADDED]
    response = await prior.create(
        {
            "prior_engagement_id": str(prior.site.engagement_id),
            "team": [{"user_id": str(prior.manager.user_id), "role": "manager"}],
            "version_id": str(prior.v2),
            "prior_item_ids": [str(prior.used_item)],
            "template_keys": [added["template_key"]],
        }
    )
    assert response.status_code == 201, response.text
    return uuid.UUID(cast(str, response.json()["id"]))


async def test_ac2_what_is_confirmed_is_exactly_what_is_created(prior: Prior) -> None:
    engagement = await _confirmed(prior)
    rows = await prior.seed.rows(
        "SELECT description FROM request_items WHERE engagement_id = $1", engagement
    )
    assert sorted(str(r["description"]) for r in rows) == sorted([USED, ADDED])
    members = await prior.seed.rows(
        "SELECT user_id, role FROM engagement_members WHERE engagement_id = $1", engagement
    )
    roles = {r["user_id"]: str(r["role"]) for r in members}
    assert roles.get(prior.manager.user_id) == "manager"
    # Nobody unconfirmed came over: the partner was left out, the leaver and the walled too.
    for absent in (prior.partner, prior.senior, prior.staff):
        assert absent.user_id not in roles
    actions = await prior.seed.actions(prior.tenant_id)
    assert "engagement.rolled_forward" in actions


async def test_ac2_a_roll_forward_copies_no_evidence_version(prior: Prior) -> None:
    """Q-03-T023: the used item comes over as a request, never with last year's evidence."""
    engagement = await _confirmed(prior)
    for table, query in EVIDENCE.items():
        assert await prior.seed.value(query, engagement) == 0, table
    statuses = await prior.seed.rows(
        "SELECT status FROM request_items WHERE engagement_id = $1", engagement
    )
    assert all(str(r["status"]) != "received" for r in statuses)
    # Last year's evidence is untouched and still last year's.
    still = await prior.seed.value(
        "SELECT engagement_id FROM evidence_versions WHERE id = $1", prior.site.version_id
    )
    assert still == prior.site.engagement_id


async def test_ac2_confirming_someone_now_walled_is_refused_and_creates_nothing(
    prior: Prior,
) -> None:
    before = await prior.count("engagements")
    response = await prior.create(
        {
            "prior_engagement_id": str(prior.site.engagement_id),
            "team": [{"user_id": str(prior.staff.user_id), "role": "staff"}],
            "version_id": str(prior.v2),
            "prior_item_ids": [str(prior.used_item)],
        }
    )
    assert (response.status_code, response.json()) == (
        409,
        {"detail": "team_member_unavailable"},
    )
    assert await prior.count("engagements") == before


# --- Authorisation: roll_forward.read is the creators'; another firm is 404 ---------------------


async def test_ac2_a_team_manager_without_a_firm_role_cannot_read_the_proposal(
    prior: Prior,
) -> None:
    response = await prior.propose(prior.manager)
    assert response.status_code == 403, response.text


async def test_ac2_another_firms_entity_reveals_nothing(prior: Prior) -> None:
    """DEC-43: another firm's record is indistinguishable from none. Spec gap: SPEC-025 does not
    say whether a proposal for an entity the caller's firm doesn't hold is 404 or the new-client
    proposal; either way it answers as for an unknown entity and names nothing of the prior."""
    other = await prior.seed.firm()
    admin = await prior.seed.person(other, "firm_admin")
    theirs = await prior.propose(admin)
    unknown = await prior.propose(admin, client_entity_id=str(uuid.uuid4()))
    assert (theirs.status_code, theirs.json()) == (unknown.status_code, unknown.json())
    assert theirs.status_code in (200, 404), theirs.text
    if theirs.status_code == 200:
        assert theirs.json()["prior"] is None
        assert theirs.json()["team"] == []
        assert theirs.json()["items"] == []
    assert str(prior.site.engagement_id) not in theirs.text
