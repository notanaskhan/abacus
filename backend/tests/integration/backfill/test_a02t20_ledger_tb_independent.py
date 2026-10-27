# ac-test-author: ayushkbhatia
"""SPEC-333 AC-1 to AC-4, AC-14, AC-17 to AC-19, AC-24, AC-25 and AC-6 (`ledger_line`), A-02
epic AC-1.2 to AC-1.4, AC-4.1, AC-4.4, AC-5.1, AC-5.2, AC-7.1, AC-7.2, AC-8.1 (A-02-T20,
TASK-1473): the ledger's trial balance of record, its layers and its import hooks.

The row's acceptance: the suite fails when a snapshot is mutated, a layer is shown without its
source, or a template that does not foot is accepted. Covered here: immutability of snapshots and
versions, versioning and prior pairing, per-layer balance, the fail-closed adjustment slot,
`ledger_line` verification and its "moved" on a re-version (and never across engagements: the
security review fix of 2026-10-27), Open in Excel (deterministic, fresh second factor, human only,
formula text neutralised, audited with the file's SHA-256), the Accrued trial balance template's
refusal of a file that does not foot, and golden figures typed from SPEC-333 AC-18.

Boundary: `create_app()` over HTTP on the migrated test database; snapshots through
`ledger.api.record_snapshot` (the public writer every import and connector uses); `ledger_line`
citations through `evidence.api.record_citations` (the public writer; `POST /v1/citations` is
reserved, C4 §3.6); agent refusals through `identity.api.authorise` with an agent context (an
agent has no bearer token). The adjustment slot is the public `ledger.api.AdjustmentSource` port;
the tests replace the registered list for their own duration only, to stand for "no source
registered" (the app registers fieldwork's SPEC-311 source at import) and for a source whose
entries do not balance.

Status codes follow `docs/architecture/error-codes.yaml` where it and C9 §6 differ:
`snapshot_entity_mismatch` and `citation_required` are 422 there (C9 §6 says 409); recorded as
spec drift in TASK-1473.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import uuid
from collections.abc import AsyncGenerator, Iterator
from contextlib import asynccontextmanager
from decimal import Decimal
from typing import Final, cast

import pytest
from asyncpg import PostgresError
from openpyxl import load_workbook

from abacus.kernel.db import TenantContext
from abacus.kernel.uow import uow
from abacus.kernel.uow.relay import run_relay
from abacus.modules.engagements.api import engagement_resource_of
from abacus.modules.identity.api import Forbidden, agent_context_for_run, authorise
from abacus.modules.ledger import api as ledger_api
from abacus.worker.__main__ import publisher

from .support import Net
from .tb_support import (
    BOY,
    EOY,
    EOY_TOTAL,
    PRIOR_END,
    Books,
    Json,
    books,
    dec,
    detail,
    facts,
    snapshot,
    stranger_firm,
    work_item,
)

SLOT: Final = "abacus.modules.ledger.trial_balances._ADJUSTMENT_SOURCES"
XLSX: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HOSTILE: Final = '=HYPERLINK("http://evil.test","Click")'


@pytest.fixture
async def world(net: Net) -> Books:
    return await books(net)


def _lines(view: Json) -> dict[str, Json]:
    out: dict[str, Json] = {}
    for group in cast(list[Json], view["groups"]):
        for line in cast(list[Json], group["lines"]):
            out[cast(str, line["account_code"])] = line
    return out


async def _rows(world: Books, sql: str, *args: object) -> list[tuple[object, ...]]:
    return [tuple(r) for r in await world.seed.rows(sql, *args)]


async def _frozen(world: Books) -> list[tuple[object, ...]]:
    """Every snapshot, line and version row of the firm, as stored."""
    return [
        *await _rows(
            world,
            "SELECT id, client_entity_id, period_start, period_end, pulled_at, source, "
            "raw_fingerprint, line_count, total_debit, total_credit, label, source_kind "
            "FROM ledger_snapshots WHERE tenant_id = $1 ORDER BY id",
            world.tenant_id,
        ),
        *await _rows(
            world,
            "SELECT * FROM trial_balance_lines WHERE tenant_id = $1 ORDER BY id",
            world.tenant_id,
        ),
        *await _rows(
            world,
            "SELECT * FROM trial_balance_versions WHERE tenant_id = $1 ORDER BY version_no",
            world.tenant_id,
        ),
    ]


async def _agent(world: Books, scope: frozenset[str]) -> object:
    return await agent_context_for_run(
        tenant_id=world.tenant_id,
        run_id=uuid.uuid4(),
        agent_id="engagement.agent",
        engagement_id=world.site.engagement_id,
        task_scope=scope,
        initiator_user_id=world.partner.user_id,
    )


@asynccontextmanager
async def _relay() -> AsyncGenerator[None]:
    """The outbox relay, as the worker runs it, for the moved detector (C4 §7)."""
    stop = asyncio.Event()
    task = asyncio.create_task(run_relay(publisher(), stop, interval=0.1))
    try:
        yield
    finally:
        stop.set()
        await task


@pytest.fixture
def no_source(monkeypatch: pytest.MonkeyPatch) -> Iterator[None]:
    """No adjustment source registered (before A-06): the slot is empty for this test only."""
    monkeypatch.setattr(SLOT, [])
    yield


def _source(*entries: ledger_api.AdjustmentEntry) -> ledger_api.AdjustmentSource:
    async def source(_tenant: TenantContext, _id: uuid.UUID) -> object:
        from abacus.modules.ledger.trial_balances import EntrySet

        return EntrySet(recorded=tuple(entries), pending=0)

    return cast(ledger_api.AdjustmentSource, source)


def _entry(ref: str, *lines: tuple[str, str, str]) -> ledger_api.AdjustmentEntry:
    from abacus.modules.ledger.trial_balances import EntryLine

    return ledger_api.AdjustmentEntry(
        ref, "aje", tuple(EntryLine(c, Decimal(d), Decimal(k)) for c, d, k in lines)
    )


# --- Story 1 · Adopt (AC-1 to AC-3) --------------------------------------------------------------


async def test_ac1_adopting_writes_one_trial_balance_and_v1_audited_as_the_person(
    world: Books,
) -> None:
    response = await world.adopt(world.senior)
    assert response.status_code == 201, response.text
    body = cast(Json, response.json())
    tb = cast(Json, body["trial_balance"])
    assert tb["label"] == "standard" and tb["current_version_no"] == 1
    assert await world.count("trial_balances") == 1
    assert await world.count("trial_balance_versions") == 1
    actor = await world.seed.value(
        "SELECT actor_kind || ':' || adopted_by::text FROM trial_balance_versions "
        "WHERE tenant_id = $1",
        world.tenant_id,
    )
    assert actor == f"human:{world.senior.user_id}"
    audited = await world.actions()
    assert ("trial_balance.created", "human") in audited
    assert ("trial_balance.version_adopted", "human") in audited
    # The view's provenance names the person and the source (AC-1, INV-5).
    view = await world.view(uuid.UUID(cast(str, tb["id"])))
    provenance = cast(Json, view["provenance"])
    assert provenance["adopted_by_user_id"] == str(world.senior.user_id)
    assert provenance["adopted_by_name"] == "Marcus Delaney"
    assert provenance["source"] == "import_netsuite_tb_csv"
    assert provenance["source_kind"] == "import"
    assert provenance["pulled_at"]


async def test_ac2_an_agent_is_refused_adoption_by_the_matrix(world: Books) -> None:
    agent = await _agent(
        world, frozenset({"trial_balance.adopt", "trial_balance.read", "engagement.read"})
    )
    resource = await engagement_resource_of(
        TenantContext(world.tenant_id, "system", "test"), world.site.engagement_id
    )
    with pytest.raises(Forbidden):
        await authorise(agent, "trial_balance.adopt", resource)  # type: ignore[arg-type]


async def test_ac2_the_table_refuses_an_agent_version_through_the_public_writer(
    world: Books,
) -> None:
    """AC-1.3: the database CHECK refuses an agent row even if a caller skips the matrix."""
    with pytest.raises(Exception):  # noqa: B017 - any refusal; the assertion is that nothing landed
        async with uow(TenantContext(world.tenant_id, "agent", "engagement.agent")) as tx:
            await ledger_api.adopt_trial_balance(
                tx,
                tenant_id=world.tenant_id,
                engagement_id=world.site.engagement_id,
                work_item_id=world.work_item_id,
                client_entity_id=world.site.entity_id,
                snapshot_id=world.current,
                prior_snapshot_id=None,
                label="standard",
                display_unit="cents",
                user_id=world.partner.user_id,
                actor_kind="agent",
            )
    assert await world.count("trial_balances") == 0
    assert await world.count("trial_balance_versions") == 0


@pytest.mark.parametrize("actor_kind", ["agent", "system"])
async def test_ac2_trial_balance_versions_refuse_a_non_human_row(
    world: Books, actor_kind: str
) -> None:
    tb_id = await world.adopted()
    with pytest.raises(PostgresError):
        await world.seed.run(
            "INSERT INTO trial_balance_versions (tenant_id, trial_balance_id, "
            "client_entity_id, version_no, snapshot_id, prior_snapshot_id, adopted_by, "
            "actor_kind, content_fingerprint, reason) VALUES ($1, $2, $3, 2, $4, NULL, $5, $6, "
            "$7, 'agent try')",
            world.tenant_id,
            tb_id,
            world.site.entity_id,
            world.newer,
            world.partner.user_id,
            actor_kind,
            hashlib.sha256(b"x").hexdigest(),
        )
    assert await world.count("trial_balance_versions") == 1


@pytest.mark.parametrize(
    "sql",
    [
        "UPDATE ledger_snapshots SET source = 'import_xero_tb_xlsx' WHERE tenant_id = $1",
        "UPDATE ledger_snapshots SET label = 'interim' WHERE tenant_id = $1",
        "DELETE FROM ledger_snapshots WHERE tenant_id = $1",
        "UPDATE trial_balance_lines SET debit = debit + 1 WHERE tenant_id = $1",
        "DELETE FROM trial_balance_lines WHERE tenant_id = $1",
        "UPDATE trial_balance_versions SET reason = 'edited' WHERE tenant_id = $1",
        "DELETE FROM trial_balance_versions WHERE tenant_id = $1",
    ],
)
async def test_ac2_a_snapshot_or_version_is_never_mutated_even_by_the_owner(
    world: Books, sql: str
) -> None:
    """ADR-004 and AC-33: the triggers refuse UPDATE and DELETE for every role."""
    await world.adopted()
    before = await _frozen(world)
    with pytest.raises(PostgresError):
        await world.seed.run(sql, world.tenant_id)
    assert await _frozen(world) == before


@pytest.mark.parametrize(
    "table", ["ledger_snapshots", "trial_balance_lines", "trial_balance_versions"]
)
async def test_ac2_truncate_is_refused(world: Books, table: str) -> None:
    await world.adopted()
    with pytest.raises(PostgresError):
        await world.seed.run(f"TRUNCATE {table} CASCADE")
    assert await world.count(table) > 0


async def test_ac3_another_entitys_snapshot_is_refused_and_writes_nothing(world: Books) -> None:
    other = await world.seed.site(world.tenant_id, world.admin.user_id, "Granite")
    foreign = await snapshot(world.tenant_id, world.senior.user_id, other.entity_id, EOY)
    for kwargs in ({"snapshot_id": foreign}, {"prior": foreign}):
        refused = await world.adopt(**kwargs)  # type: ignore[arg-type]
        assert refused.status_code == 422, refused.text  # error-codes.yaml (C9 §6 says 409)
        assert detail(refused) == "snapshot_entity_mismatch"
    assert await world.count("trial_balances") == 0
    assert await world.count("trial_balance_versions") == 0
    assert not [a for a, _ in await world.actions() if a.startswith("trial_balance.")]


async def test_ac3_another_firms_snapshot_id_reads_as_absent(world: Books, net: Net) -> None:
    stranger = await stranger_firm(net)
    elsewhere = await net.seed.value(
        "SELECT client_entity_id FROM engagements WHERE tenant_id = $1", stranger.tenant_id
    )
    foreign = await snapshot(stranger.tenant_id, stranger.user_id, cast(uuid.UUID, elsewhere), EOY)
    refused = await world.adopt(snapshot_id=foreign)
    assert refused.status_code in (404, 422), refused.text
    assert await world.count("trial_balances") == 0


async def test_ac1_a_second_trial_balance_for_the_same_label_and_period_is_refused(
    world: Books,
) -> None:
    await world.adopted()
    again = await world.adopt()
    assert again.status_code == 409, again.text
    assert detail(again) == "trial_balance_exists"
    assert await world.count("trial_balances") == 1
    # Another label of the same period is its own trial balance of record.
    interim = await world.adopt(label="interim")
    assert interim.status_code == 201, interim.text
    assert await world.count("trial_balances") == 2


@pytest.mark.parametrize("who", ["reviewer"])
async def test_ac1_roles_the_matrix_denies_cannot_adopt(world: Books, who: str) -> None:
    refused = await world.adopt(getattr(world, who))
    assert refused.status_code == 403, refused.text
    assert await world.count("trial_balances") == 0


# --- Story 1 · Versions (AC-4; epic AC-7.1, AC-7.2) ----------------------------------------------


async def test_ac4_the_newer_snapshot_is_offered_with_a_code_computed_change_count(
    world: Books,
) -> None:
    tb_id = await world.adopted()
    response = await world.call("GET", world.path(tb_id, "/candidates"))
    assert response.status_code == 200, response.text
    found = {c["id"]: c for c in cast(list[Json], response.json())}
    assert str(world.newer) in found
    newer = found[str(world.newer)]
    assert newer["lines_changed"] == 2  # BrokerWindow and Net assets
    assert dec(newer["total_debit"]) == Decimal("16521570.00")
    assert str(world.prior) not in found  # another period is not a candidate
    assert (await world.view(tb_id))["newer_snapshots"] == 1


async def test_ac4_re_adopting_creates_v2_and_leaves_v1_and_both_snapshots_untouched(
    world: Books,
) -> None:
    tb_id = await world.adopted()
    before = await _frozen(world)
    response = await world.reversion(tb_id)
    assert response.status_code == 201, response.text
    after = await _frozen(world)
    assert set(before) <= set(after)  # nothing stored before changed
    versions = await world.seed.rows(
        "SELECT version_no, snapshot_id, reason, actor_kind FROM trial_balance_versions "
        "WHERE trial_balance_id = $1 ORDER BY version_no",
        tb_id,
    )
    assert [(v["version_no"], v["snapshot_id"]) for v in versions] == [
        (1, world.current),
        (2, world.newer),
    ]
    assert versions[1]["reason"] == "Client sent a corrected export"
    assert versions[1]["actor_kind"] == "human"
    events = await world.seed.rows(
        "SELECT payload FROM outbox WHERE tenant_id = $1 "
        "AND event_type = 'trial_balance.version_created' ORDER BY seq",
        world.tenant_id,
    )
    assert len(events) == 2  # v1 and v2
    assert str(world.newer) in str(events[1]["payload"])
    # v1 still reads as it was; the current view reads v2's figures.
    v1 = await world.call("GET", world.path(tb_id, "/layers/unadjusted?version=1"))
    assert v1.status_code == 200, v1.text
    assert dec(_lines(cast(Json, v1.json()))["1230"]["eoy"]) == Decimal("1606560.00")
    v2 = await world.view(tb_id)
    assert v2["version_no"] == 2
    assert dec(_lines(v2)["1230"]["eoy"]) == Decimal("1608120.00")
    assert ("trial_balance.version_adopted", "human") in await world.actions()


async def test_ac4_a_stale_seen_version_or_the_same_snapshot_is_refused(world: Books) -> None:
    tb_id = await world.adopted()
    same = await world.reversion(tb_id, snapshot_id=world.current)
    assert same.status_code == 409 and detail(same) == "tb_same_snapshot", same.text
    assert (await world.reversion(tb_id)).status_code == 201
    stale = await world.reversion(tb_id, snapshot_id=world.current, seen_version=1)
    assert stale.status_code == 409 and detail(stale) == "tb_version_stale", stale.text
    assert await world.count("trial_balance_versions") == 2


async def test_ac4_a_reason_is_required_to_re_version(world: Books) -> None:
    tb_id = await world.adopted()
    refused = await world.reversion(tb_id, reason="   ")
    assert refused.status_code in (409, 422), refused.text
    assert await world.count("trial_balance_versions") == 1


# --- Prior pairing -------------------------------------------------------------------------------


async def test_ac4_the_prior_snapshot_fills_the_beginning_of_year_column(world: Books) -> None:
    tb_id = await world.adopted()
    view = await world.view(tb_id)
    assert view["prior_period_end"] == PRIOR_END.isoformat()
    lines = _lines(view)
    for code, _name, dr, cr in BOY:
        assert dec(lines[code]["boy"]) == Decimal(dr) - Decimal(cr), code
    for code, _name, dr, cr in EOY:
        assert dec(lines[code]["eoy"]) == Decimal(dr) - Decimal(cr), code


async def test_ac4_without_a_prior_the_beginning_of_year_column_is_empty(world: Books) -> None:
    tb_id = await world.adopted(prior=None)
    view = await world.view(tb_id)
    assert view["prior_period_end"] is None
    assert all(line["boy"] is None for line in _lines(view).values())


# --- Story 4 · Layers (AC-14, AC-17) -------------------------------------------------------------


async def test_ac17_the_unadjusted_layer_balances_by_code(world: Books) -> None:
    tb_id = await world.adopted()
    view = await world.view(tb_id)
    balance = cast(Json, view["balance"])
    assert balance["layer"] == "unadjusted"
    assert dec(balance["debit"]) == EOY_TOTAL
    assert dec(balance["credit"]) == EOY_TOTAL
    assert dec(balance["difference"]) == 0
    assert balance["balanced"] is True
    assert view["layer"] == "unadjusted" and view["label"] == "standard"


@pytest.mark.usefixtures("no_source")
@pytest.mark.parametrize("layer", ["aje", "rje", "final"])
async def test_ac14_without_an_adjustment_source_the_adjusted_layers_are_unavailable(
    world: Books, layer: str
) -> None:
    tb_id = await world.adopted()
    refused = await world.layer(tb_id, layer)
    assert refused.status_code == 409, refused.text
    assert detail(refused) == "layer_unavailable"
    exported = await world.call("GET", world.path(tb_id, f"/export.xlsx?layer={layer}"))
    assert exported.status_code == 409 and detail(exported) == "layer_unavailable"
    assert (await world.view(tb_id))["layers_available"] == ["unadjusted"]


@pytest.mark.usefixtures("no_source")
async def test_ac14_the_spec309_view_never_shows_a_final_equal_to_unadjusted(world: Books) -> None:
    tb_id = await world.adopted()
    response = await world.call("GET", world.path(tb_id, "?layer=final"))
    if response.status_code == 409:
        assert detail(response) == "layer_unavailable"
    else:
        assert response.status_code == 200, response.text
        body = cast(Json, response.json())
        assert body["final"] is None and body["layers_available"] is False
        assert all(line["final"] is None for line in cast(list[Json], body["lines"]))


async def test_ac17_an_adjusted_layer_that_does_not_balance_is_refused_with_its_difference(
    world: Books, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(SLOT, [_source(_entry("AJE-9", ("1230", "100.00", "0")))])
    tb_id = await world.adopted()
    for layer in ("aje", "final"):
        refused = await world.layer(tb_id, layer)
        assert refused.status_code == 409, refused.text
        assert detail(refused) == "layer_unbalanced"
        assert facts(refused).get("layer") == layer
        assert dec(facts(refused)["difference"]).copy_abs() == Decimal("100.00")
    exported = await world.call("GET", world.path(tb_id, "/export.xlsx?layer=final"))
    assert exported.status_code == 409 and detail(exported) == "layer_unbalanced"
    # The unadjusted layer still balances and still reads.
    assert (await world.layer(tb_id)).status_code == 200


async def test_ac17_a_balanced_adjustment_moves_final_by_code(
    world: Books, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Epic AC-4.2's grammar on our literals: AJE-1 = 33,540.00 + 15,060.00 = 48,600.00."""
    monkeypatch.setattr(
        SLOT,
        [
            _source(
                _entry("AJE-1", ("1900", "33540.00", "0"), ("3000", "0", "33540.00")),
                _entry("AJE-2", ("1900", "15060.00", "0"), ("3000", "0", "15060.00")),
            )
        ],
    )
    tb_id = await world.adopted()
    view = await world.view(tb_id, "final")
    assert set(cast(list[str], view["layers_available"])) == {"unadjusted", "aje", "rje", "final"}
    line = _lines(view)["1900"]
    assert dec(line["eoy"]) == Decimal("20010.00")
    assert dec(line["aje"]) == Decimal("48600.00")
    assert dec(line["final"]) == Decimal("68610.00")
    assert set(cast(list[str], line["adjustment_refs"])) >= {"AJE-1", "AJE-2"}
    assert dec(_lines(view)["3000"]["final"]) == Decimal("-16256210.00")
    balance = cast(Json, view["balance"])
    assert balance["balanced"] is True and balance["layer"] == "final"


# --- Story 5 · Leadsheets and golden figures (AC-18, AC-19) --------------------------------------


async def test_ac18_the_investments_leadsheet_foots_to_the_spec_goldens(world: Books) -> None:
    tb_id = await world.adopted()
    await world.annotate(tb_id)
    response = await world.call("GET", world.path(tb_id, "/leadsheets/Investments"))
    assert response.status_code == 200, response.text
    sheet = cast(Json, response.json())
    assert sheet["area"] == "Investments" and sheet["unassigned"] is False
    assert dec(sheet["py"]) == Decimal("13940000.00")
    assert dec(sheet["cy"]) == Decimal("16500000.00")
    assert dec(sheet["change"]) == Decimal("2560000.00")
    assert dec(sheet["change_pct"]) == Decimal("18.4")
    rows = {r["account_code"]: r for r in cast(list[Json], sheet["rows"])}
    assert set(rows) == {"1210", "1220", "1230"}
    krs = rows["1210"]
    assert (dec(krs["py"]), dec(krs["cy"]), dec(krs["change"]), dec(krs["change_pct"])) == (
        Decimal("5102300.00"),
        Decimal("6120450.00"),
        Decimal("1018150.00"),
        Decimal("20.0"),
    )


async def test_ac19_an_account_with_no_area_is_under_unassigned_and_rounds_half_even(
    world: Books,
) -> None:
    tb_id = await world.adopted()
    await world.annotate(tb_id)
    response = await world.call("GET", world.path(tb_id, "/leadsheets/Unassigned"))
    assert response.status_code == 200, response.text
    sheet = cast(Json, response.json())
    assert sheet["unassigned"] is True
    [row] = cast(list[Json], sheet["rows"])
    assert row["account_code"] == "1900"
    assert dec(row["change"]) == Decimal("10.00")
    assert dec(row["change_pct"]) == Decimal("0.0")  # 0.05 % half-even (half-up would be 0.1)
    # Every account is in exactly one group of the view: none is dropped.
    view = await world.view(tb_id)
    codes = [
        line["account_code"]
        for group in cast(list[Json], view["groups"])
        for line in cast(list[Json], group["lines"])
    ]
    assert sorted(str(c) for c in codes) == sorted(code for code, *_ in EOY)
    labels = [g["label"] for g in cast(list[Json], view["groups"])]
    assert labels[-1] == "Unassigned"
    totals = {g["label"]: dec(g["eoy"]) for g in cast(list[Json], view["groups"])}
    assert totals["Investments"] == Decimal("16500000.00")


async def test_ac18_a_leadsheet_for_an_area_with_no_lines_is_404(world: Books) -> None:
    tb_id = await world.adopted()
    await world.annotate(tb_id)
    missing = await world.call("GET", world.path(tb_id, "/leadsheets/Inventory"))
    assert missing.status_code == 404, missing.text


# --- Story 7 · Open in Excel (AC-24, AC-25) ------------------------------------------------------


async def test_ac24_open_in_excel_is_deterministic_and_audited_with_its_hash(
    world: Books,
) -> None:
    tb_id = await world.adopted()
    await world.annotate(tb_id)
    first = await world.call("GET", world.path(tb_id, "/export.xlsx"), world.manager)
    assert first.status_code == 200, first.text
    assert first.headers["content-type"].startswith(XLSX)
    assert "attachment" in first.headers.get("content-disposition", "")
    second = await world.call("GET", world.path(tb_id, "/export.xlsx"), world.manager)
    assert second.status_code == 200
    assert first.content == second.content  # byte-identical on re-render
    digest = hashlib.sha256(first.content).hexdigest()
    rows = await world.seed.rows(
        "SELECT actor_kind, actor_id, after_ref::text AS after FROM audit_events "
        "WHERE tenant_id = $1 AND action = 'trial_balance.exported'",
        world.tenant_id,
    )
    assert len(rows) == 2
    assert rows[0]["actor_kind"] == "human" and rows[0]["actor_id"] == str(world.manager.user_id)
    assert digest in str(rows[0]["after"])
    # The workbook carries the figures, the areas and a provenance footer.
    sheet = load_workbook(io.BytesIO(first.content)).active
    assert sheet is not None
    text = [str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None]
    assert any("KRS 500 Index" in t for t in text)
    assert any("Investments" in t for t in text)
    assert any("Marcus Delaney" in t or "Hannah Leclerc" in t for t in text)  # adopted by


async def test_ac24_client_text_that_looks_like_a_formula_is_written_as_text(
    world: Books,
) -> None:
    hostile = (*EOY[:-1], ("3000", HOSTILE, "0", "16207610.00"))
    tainted = await snapshot(
        world.tenant_id,
        world.senior.user_id,
        world.site.entity_id,
        hostile,
        source="import_qbo_tb_csv",
    )
    tb_id = await world.adopted(snapshot_id=tainted, prior=None)
    response = await world.call("GET", world.path(tb_id, "/export.xlsx"))
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content))
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (cell.coordinate, cell.value)
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "-", "@")) or cell.value == "-"


@pytest.mark.parametrize("who", ["engagement_partner", "manager", "senior"])
async def test_ac25_the_roles_the_matrix_allows_may_export_with_a_fresh_second_factor(
    world: Books, who: str
) -> None:
    person = {
        "engagement_partner": world.partner,
        "manager": world.manager,
        "senior": world.senior,
    }[who]
    tb_id = await world.adopted()
    path = world.path(tb_id, "/export.xlsx")
    stale = await world.call("GET", path, person, mfa=False)
    assert stale.status_code == 403, stale.text
    fresh = await world.call("GET", path, person)
    assert fresh.status_code == 200, fresh.text


@pytest.mark.parametrize("who", ["staff", "reviewer"])
async def test_ac25_staff_and_reviewers_cannot_export(world: Books, who: str) -> None:
    tb_id = await world.adopted()
    refused = await world.call("GET", world.path(tb_id, "/export.xlsx"), getattr(world, who))
    assert refused.status_code == 403, refused.text
    assert not [a for a, _ in await world.actions() if a == "trial_balance.exported"]


async def test_ac25_the_export_is_denied_to_an_agent(world: Books) -> None:
    """Security review fix (2026-10-27): the agent never holds `trial_balance.export`, whatever
    its task scope says (INV-1)."""
    agent = await _agent(
        world, frozenset({"trial_balance.export", "trial_balance.read", "engagement.read"})
    )
    resource = await engagement_resource_of(
        TenantContext(world.tenant_id, "system", "test"), world.site.engagement_id
    )
    for action in ("trial_balance.export", "trial_balance.annotate"):
        with pytest.raises(Forbidden):
            await authorise(agent, action, resource)  # type: ignore[arg-type]


async def test_ac25_a_walled_partner_and_another_firm_get_404_on_the_export(
    world: Books, net: Net
) -> None:
    tb_id = await world.adopted()
    walled = await world.seed.person(world.tenant_id)
    await world.seed.member(world.site, walled, "engagement_partner")
    await net.wall(world.admin, walled, world.site)
    stranger = await stranger_firm(net)
    path = world.path(tb_id, "/export.xlsx")
    missing = await world.call("GET", world.path(uuid.uuid4(), "/export.xlsx"))
    assert missing.status_code == 404
    for who in (walled, stranger):
        refused = await world.call("GET", path, who)
        assert refused.status_code == 404, refused.text
        assert refused.json() == missing.json()


# --- AC-6 · `ledger_line` citations (A-02-T22, C9 §2) --------------------------------------------


async def test_ac6_a_ledger_line_citation_verifies_against_the_lines_net(world: Books) -> None:
    tb_id = await world.adopted()
    cid, ok, reason, match = await world.cite_line(
        tb_id, world.current, "1230", value="1606560.00", quote="brokerwindow"
    )
    assert ok and reason is None, reason
    assert match == f"led:{world.current}:1230"
    wrong = await world.cite_line(tb_id, world.current, "1230", value="1606561.00")
    assert wrong[1] is False and wrong[2] == "value_mismatch"
    quote = await world.cite_line(tb_id, world.current, "1230", quote="KRS 500")
    assert quote[1] is False and quote[2] == "quote_not_found"
    assert cid


async def test_ac6_a_credit_line_verifies_by_its_printed_side_or_its_net(world: Books) -> None:
    tb_id = await world.adopted()
    side = await world.cite_line(tb_id, world.current, "2500", value="312400.00")
    net = await world.cite_line(tb_id, world.current, "2500", value="-312400.00")
    assert side[1] is True and net[1] is True, (side, net)


async def test_ac6_an_extra_locator_field_or_unknown_account_is_not_verified(
    world: Books,
) -> None:
    tb_id = await world.adopted()
    extra = await world.cite_line(
        tb_id, world.current, "1230", value="1606560.00", extra={"note": "x"}
    )
    assert extra[1] is False and extra[2] == "locator_invalid"
    unknown = await world.cite_line(tb_id, world.current, "9999", value="1.00")
    assert unknown[1] is False and unknown[2] in ("source_not_found", "cell_not_found")


async def test_ac6_another_entitys_snapshot_reads_as_absent(world: Books) -> None:
    tb_id = await world.adopted()
    other = await world.seed.site(world.tenant_id, world.admin.user_id, "Granite")
    foreign = await snapshot(world.tenant_id, world.senior.user_id, other.entity_id, EOY)
    _cid, ok, reason, _match = await world.cite_line(tb_id, foreign, "1230", value="1606560.00")
    assert ok is False and reason == "source_not_found"


async def test_ac6_a_citation_to_a_superseded_snapshot_reads_moved(world: Books) -> None:
    tb_id = await world.adopted()
    cid, ok, _reason, _match = await world.cite_line(
        tb_id, world.current, "1230", value="1606560.00"
    )
    assert ok
    assert (await world.reversion(tb_id)).status_code == 201
    async with _relay():
        state = await _state(world, tb_id, cid, until="moved")
    assert state["state"] == "moved"


async def test_ac6_a_re_version_never_moves_another_engagements_citations(
    world: Books,
) -> None:
    """Security review fix (2026-10-27): engagement B of the same entity cites the snapshot
    engagement A replaces; B's citation stays valid and never shows A's new snapshot id."""
    seed = world.seed
    other_engagement = cast(
        uuid.UUID,
        await seed.value(
            "INSERT INTO engagements (tenant_id, client_id, client_entity_id, name, "
            "fiscal_period_start, fiscal_period_end, created_by) VALUES ($1, $2, $3, "
            "'Northwind tax', '2025-01-01', '2025-12-31', $4) RETURNING id",
            world.tenant_id,
            world.site.client_id,
            world.site.entity_id,
            world.admin.user_id,
        ),
    )
    await seed.run(
        "INSERT INTO engagement_acceptance (tenant_id, engagement_id, kind, decision, "
        "decided_by, documented_at, independence_concluded_by, independence_concluded_at, "
        "independence_documented_at) VALUES ($1, $2, 'continuance', 'accepted', 'test-seed', "
        "'Seeded', 'test-seed', clock_timestamp(), 'Seeded')",
        world.tenant_id,
        other_engagement,
    )
    from dataclasses import replace

    site_b = replace(world.site, engagement_id=other_engagement)
    await seed.member(site_b, world.partner, "engagement_partner")
    await work_item(seed, site_b, world.admin.user_id)
    adopted_b = await world.call(
        "POST",
        f"/v1/engagements/{other_engagement}/trial-balances",
        body={"snapshot_id": str(world.current), "label": "standard", "display_unit": "cents"},
    )
    assert adopted_b.status_code == 201, adopted_b.text
    tb_b = uuid.UUID(cast(str, cast(Json, adopted_b.json()["trial_balance"])["id"]))
    tb_a = await world.adopted()
    cid_b, ok, _r, _m = await world.cite_line(
        tb_b, world.current, "1230", value="1606560.00", engagement_id=other_engagement
    )
    cid_a, ok_a, _r2, _m2 = await world.cite_line(tb_a, world.current, "1230", value="1606560.00")
    assert ok and ok_a
    assert (await world.reversion(tb_a)).status_code == 201
    async with _relay():
        await _state(world, tb_a, cid_a, until="moved")  # A's own citation does move
        await asyncio.sleep(0.5)
        b = await _state(world, tb_b, cid_b)
    assert b["state"] == "valid", b
    assert str(world.newer) not in str(b)
    events = await seed.rows(
        "SELECT kind FROM citation_events WHERE citation_id = $1 AND kind = 'moved'", cid_b
    )
    assert not events


async def _state(world: Books, tb_id: uuid.UUID, cid: uuid.UUID, until: str | None = None) -> Json:
    deadline = asyncio.get_running_loop().time() + 20
    while True:
        rows = {c["id"]: c for c in await world.citations(tb_id)}
        found = rows.get(str(cid))
        assert found is not None, rows
        if until is None or found["state"] == until:
            return found
        if asyncio.get_running_loop().time() > deadline:
            return found
        await asyncio.sleep(0.2)
