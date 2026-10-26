"""K-01-T02 (TASK-1306; K-01 §4.2 migration A): the methodology module re-owns SPEC-008's
templates and gives every version a lifecycle.

Through the moved routes (paths and actions unchanged): an import writes the next version with
its label and origin, publishes it on its own authority (`legacy_import`, until K-01-T06's
approval) and moves the template's default to it; one `imported` event per version; the five
SPEC-008 tables and the events stay insert-only for the app; the lifecycle row is the one that
changes. The migration's backfill and its down/up are `test_k01t02_methodology_migration.py`.
"""

from __future__ import annotations

import io
import uuid
from typing import cast

import pytest
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

from abacus.kernel.db import TenantContext, tenant_session

from .support import Net, Person

pytestmark = pytest.mark.asyncio


def _workbook(*descriptions: str) -> bytes:
    book = Workbook()
    areas = book.active
    assert areas is not None
    areas.title = "Areas"
    areas.append(["code", "name"])
    areas.append(["CASH", "Cash"])
    requests = book.create_sheet("Requests")
    requests.append(["area_code", "description", "tier"])
    for description in descriptions:
        requests.append(["CASH", description, "D"])
    rules = book.create_sheet("Account rules")
    rules.append(["area_code", "account_from", "account_to"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def _import(net: Net, admin: Person, name: str, *descriptions: str) -> uuid.UUID:
    response = await net.http.post(
        f"/v1/methodology/templates/{name}/versions?engagement_type=ebp_audit",
        content=_workbook(*descriptions),
        headers={**net.headers(admin), "Content-Type": "application/octet-stream"},
    )
    assert response.status_code == 201, response.text
    return uuid.UUID(cast(str, response.json()["version_id"]))


async def test_ac_an_import_is_published_as_legacy_and_becomes_the_default(net: Net) -> None:
    seed = net.seed
    tenant = await seed.firm()
    admin = await seed.person(tenant, "firm_admin")
    name = f"BH EBP {uuid.uuid4().hex[:6]}"
    v1 = await _import(net, admin, name, "Plan document and amendments")
    v2 = await _import(net, admin, name, "Plan document and amendments", "Form 5500")

    states = await seed.rows(
        "SELECT version_id, status, approval_basis, effective_from IS NOT NULL AS dated, "
        "default_for_new_engagements, engagement_types, row_version "
        "FROM methodology_version_states WHERE tenant_id = $1",
        tenant,
    )
    by_version = {r["version_id"]: r for r in states}
    assert set(by_version) == {v1, v2}
    for row in states:
        assert (row["status"], row["approval_basis"], row["dated"]) == (
            "published",
            "legacy_import",
            True,
        )
        assert row["engagement_types"] == ["ebp_audit"]
    assert by_version[v1]["default_for_new_engagements"] is False
    assert by_version[v1]["row_version"] == 2  # the default moved off it
    assert by_version[v2]["default_for_new_engagements"] is True

    versions = await seed.rows(
        "SELECT id, label, origin, drafted_from_version_id FROM methodology_versions "
        "WHERE tenant_id = $1 ORDER BY version",
        tenant,
    )
    assert [(r["label"], r["origin"], r["drafted_from_version_id"]) for r in versions] == [
        ("v1", "workbook", None),
        ("v2", "workbook", None),
    ]
    events = await seed.rows(
        "SELECT version_id, kind, actor_kind, actor_user_id FROM methodology_version_events "
        "WHERE tenant_id = $1 ORDER BY seq",
        tenant,
    )
    assert [(r["version_id"], r["kind"]) for r in events] == [(v1, "imported"), (v2, "imported")]
    assert {(r["actor_kind"], r["actor_user_id"]) for r in events} == {("human", admin.user_id)}

    # The reads moved with the routes and kept their shape.
    listed = await net.call("GET", "/v1/methodology/templates?latest=true", admin)
    assert listed.status_code == 200, listed.text
    assert [v["version_id"] for v in listed.json() if v["template_name"] == name] == [str(v2)]


@pytest.mark.parametrize(
    "statement",
    [
        "UPDATE methodology_templates SET tenant_id = tenant_id",
        "DELETE FROM methodology_templates",
        "UPDATE methodology_versions SET tenant_id = tenant_id",
        "DELETE FROM methodology_versions",
        "UPDATE methodology_areas SET tenant_id = tenant_id",
        "DELETE FROM methodology_areas",
        "UPDATE methodology_request_items SET tenant_id = tenant_id",
        "DELETE FROM methodology_request_items",
        "UPDATE methodology_account_rules SET tenant_id = tenant_id",
        "DELETE FROM methodology_account_rules",
        "UPDATE methodology_version_events SET tenant_id = tenant_id",
        "DELETE FROM methodology_version_events",
    ],
)
async def test_ac_the_versions_and_their_events_stay_insert_only(net: Net, statement: str) -> None:
    seed = net.seed
    tenant = await seed.firm()
    admin = await seed.person(tenant, "firm_admin")
    await _import(net, admin, f"Probe {uuid.uuid4().hex[:6]}", "Bank statements")
    me = TenantContext(tenant, "human", str(admin.user_id))
    with pytest.raises(DBAPIError, match=r"permission denied|insert-only"):
        async with tenant_session(me) as session:
            await session.execute(text(statement))


async def test_who_approved_is_the_transaction_actor_never_an_attested_column(net: Net) -> None:
    """Security review M2 (ADR-005): an agent cannot write `approved_actor_kind = 'human'`, and a
    person cannot name someone else as the approver; a person names themself."""
    seed = net.seed
    tenant = await seed.firm()
    admin = await seed.person(tenant, "firm_admin")
    other = await seed.person(tenant, "quality_partner")
    version = await _import(net, admin, f"Probe {uuid.uuid4().hex[:6]}", "Bank statements")
    approve = (
        "UPDATE methodology_version_states SET approved_by = :who, approved_actor_kind = 'human', "
        "approved_at = now(), approval_basis = 'approved' WHERE version_id = :version"
    )
    agent = TenantContext(tenant, "agent", "agent:requests.guidance_drafter:run", admin.user_id)
    person = TenantContext(tenant, "human", str(admin.user_id))
    for ctx, who in ((agent, admin.user_id), (person, other.user_id)):
        with pytest.raises(DBAPIError, match=r"approved by the person acting"):
            async with tenant_session(ctx) as session:
                await session.execute(text(approve), {"who": who, "version": version})
    async with tenant_session(person) as session:  # naming oneself is allowed
        await session.execute(text(approve), {"who": admin.user_id, "version": version})
        found = await session.scalar(
            text("SELECT approved_by FROM methodology_version_states WHERE version_id = :v"),
            {"v": version},
        )
        assert found == admin.user_id
