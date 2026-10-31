"""P-11-T12 (TASK-2514; P-11 §9; SPEC-520 AC-6, AC-11, AC-20; ADR-078): the Brightwater Hale golden
render fixtures. Written from the row and the handoff, never from the renderer.

The row's acceptance: every fixture re-renders byte-identical, its extracted text equals its block
text, and the AR-C 70 legend is present on every package page. Each fixture's PDF and DOCX also
equal the hashes two fresh containers agreed on (`p11t12.sha256`).

Image-bound, as `test_ac30_render_golden.py`: WeasyPrint needs the image's pinned Pango, HarfBuzz
and fontconfig, so this runs inside it (`make golden-render`).
"""

from __future__ import annotations

import io
import re
from collections import Counter
from collections.abc import Iterator
from pathlib import Path
from typing import cast

import pdfplumber
import pytest
from docx import Document

from abacus.modules.reporting.render.core import render, render_docx

from .p11t12_fixtures import ARC70_LEGEND, CLS08_PACKAGE, FIXTURES, Fixture

HERE = Path(__file__).resolve().parent
NAMES = [f.name for f in FIXTURES]
# What the layouts print that no block carries (ADR-320 §3): list markers, the enclosures caption
# and signature rules. The DRAFT watermark is diagonal and is left out by its text matrix
# (`_pages`). A line may break inside a word ("103(a)(3)" / "(C)"), so text is compared with
# whitespace removed: each block's text in order inside the page text, and the characters as a
# whole equal.
CHROME = re.compile(r"(?<![^\s])(•|Enclosures|_{3,})(?![^\s])")
SCREENS = {"ENG-04 b", "ENG-04 d", "ENG-04 e", "ENG-04 f", "AUD-21 a", "AUD-22 a", "AUD-23 a"}


def _runs(runs: object, values: dict[str, str]) -> str:
    out: list[str] = []
    for run in cast(list[dict[str, object]], runs):
        if run["type"] == "text":
            out.append(str(run["text"]))
        elif run["type"] == "var":
            out.append(values[str(run["key"])])
    return "".join(out)


def block_texts(blocks: list[dict[str, object]], values: dict[str, str]) -> Iterator[str]:
    """The strings the blocks carry, by the C11 vocabulary (ADR-320 §2), not by the renderer."""
    for b in blocks:
        kind = b["type"]
        if kind in ("heading", "legend"):
            yield str(b["text"])
        elif kind == "paragraph":
            yield _runs(b["runs"], values)
        elif kind == "list":
            for item in cast(list[object], b["items"]):
                yield _runs(item, values)
        elif kind == "table":
            yield from (str(c) for c in cast(list[object], b["columns"]))
            for row in cast(list[list[str]], b["rows"]):
                yield from row
        elif kind == "kv":
            for row in cast(list[dict[str, object]], b["rows"]):
                yield str(row["key"])
                yield _runs(row["value"], values)
        elif kind == "signature_block":
            for s in cast(list[dict[str, object]], b["signers"]):
                yield str(s["label"])
                yield _runs(s["name"], values)
        elif kind == "enclosures":
            yield from (str(i) for i in cast(list[object], b["items"]))
        elif kind == "section":
            yield str(b["title"])
            yield from block_texts(cast(list[dict[str, object]], b["blocks"]), values)
        elif kind == "clause":
            yield from block_texts(cast(list[dict[str, object]], b["blocks"]), values)


def _squash(text: str) -> str:
    return "".join(text.split())


def _chars(texts: Iterator[str] | list[str]) -> Counter[str]:
    return Counter(_squash(CHROME.sub(" ", " ".join(texts))))


def _level(obj: dict[str, object]) -> bool:
    """A character set level on the page (no rotation): the diagonal watermark is not."""
    if obj.get("object_type") != "char":
        return True
    matrix = cast(tuple[float, ...], obj["matrix"])
    return abs(matrix[1]) < 1e-6 and abs(matrix[2]) < 1e-6


def _pages(pdf: bytes) -> list[str]:
    """Each page's text without the watermark."""
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        return [page.filter(_level).extract_text() or "" for page in document.pages]


def _expected(fixture: Fixture) -> list[str]:
    return list(block_texts(fixture.blocks, fixture.values))


def _recorded() -> dict[str, tuple[str, str]]:
    path = HERE / "p11t12.sha256"
    assert path.is_file(), "p11t12.sha256 is not recorded: run RECORD_GOLDEN=1 make golden-render"
    out: dict[str, tuple[str, str]] = {}
    for line in path.read_text(encoding="ascii").splitlines():
        name, pdf, docx, _pages_count = line.split()
        out[name] = (pdf, docx)
    return out


def test_the_fixtures_cover_every_frame_the_row_names() -> None:
    screens = {f.screen for f in FIXTURES}
    assert screens >= SCREENS | {"NOT-03 a"}
    assert any(f.legend == ARC70_LEGEND and f.layout == "package" for f in FIXTURES)


@pytest.mark.parametrize("fixture", FIXTURES, ids=NAMES)
def test_every_fixture_rerenders_byte_identical(fixture: Fixture) -> None:
    first, second = render(fixture.input()), render(fixture.input())
    assert first.pdf == second.pdf
    assert render_docx(fixture.input()).docx == render_docx(fixture.input()).docx


@pytest.mark.parametrize("fixture", FIXTURES, ids=NAMES)
def test_every_fixture_equals_its_recorded_hashes(fixture: Fixture) -> None:
    pdf, docx = _recorded()[fixture.name]
    assert render(fixture.input()).render_hash == pdf
    assert render_docx(fixture.input()).render_hash == docx


@pytest.mark.parametrize("fixture", FIXTURES, ids=NAMES)
def test_the_pdf_text_equals_the_block_text(fixture: Fixture) -> None:
    pages = _pages(render(fixture.input()).pdf)
    flat = _squash(" ".join(pages))
    expected = _expected(fixture)
    for text in expected:
        assert _squash(text) in flat, text
    chars = _chars(pages)
    if fixture.legend is not None:  # the legend prints once per page, the block once
        for _ in range(len(pages) - 1):
            chars -= Counter(_squash(fixture.legend))
    assert chars == _chars(expected)


@pytest.mark.parametrize("fixture", FIXTURES, ids=NAMES)
def test_the_docx_text_equals_the_block_text(fixture: Fixture) -> None:
    document = Document(io.BytesIO(render_docx(fixture.input()).docx))
    paragraphs = [p.text for p in document.paragraphs]
    cells = [c.text for t in document.tables for row in t.rows for c in row.cells]
    # A legend runs in the DOCX footer (one per section), as the PDF prints it on every page.
    footers = [p.text for sec in document.sections for p in sec.footer.paragraphs if p.text]
    body = [t for t in paragraphs + cells if not (fixture.mode == "draft" and t == "DRAFT")]
    if fixture.legend is not None:
        assert fixture.legend in footers
        body.append(fixture.legend)
    for text in _expected(fixture):
        assert _squash(text) in _squash(" ".join(body)), text
    assert _chars(body) == _chars(_expected(fixture))


def test_the_arc70_legend_is_on_every_package_page() -> None:
    pages = _pages(render(CLS08_PACKAGE.input()).pdf)
    assert len(pages) >= 3
    for n, page in enumerate(pages, start=1):
        assert ARC70_LEGEND in " ".join(page.split()), f"page {n}"
