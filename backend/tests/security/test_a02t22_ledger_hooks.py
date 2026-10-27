"""A-02-T22 (TASK-1470; SPEC-333; A-02 AC-8.1, P-48) on the live routes and the migrated test
database, through the real `authorise`, walls and row-level security.

The row's acceptance: a `ledger_line` citation to a superseded snapshot verifies as moved, a
re-render is byte-identical, and a template that does not foot is refused unbalanced with its
difference (the template half is `tests/unit/connections/test_a02t22_accrued_template.py`).
Around it: Open in Excel is audited with the file's fingerprint, a walled member and the flag
off answer 404, and a snapshot of the template source is an `import` (migration 0235).
"""

from __future__ import annotations

import hashlib
import io
import json
import uuid
from collections.abc import Awaitable, Callable
from datetime import UTC, datetime
from decimal import Decimal
from typing import cast

import asyncpg
from openpyxl import load_workbook

from abacus.kernel.db import TenantContext
from abacus.kernel.uow import Target, uow
from abacus.kernel.uow.relay import OutboxEvent
from abacus.modules.evidence.api import CitationAnchor, CitationClaim, record_citations
from abacus.modules.evidence.citations.ledger_line import on_trial_balance_version_created
from abacus.modules.identity.api import AuthContext
from abacus.modules.ledger.api import adopt_version

from . import test_a02t10_tb_read_model as t10
from .support import Env, Person, Seeder

# The T10 fixtures: a site's engagement with an adopted trial balance, and the ledger's writer.
EOY, TB = t10.EOY, t10.TB
_setup = cast(Callable[[Env], Awaitable[tuple[Person, uuid.UUID]]], vars(t10)["_setup"])
_snapshot = cast(Callable[..., Awaitable[uuid.UUID]], vars(t10)["_snapshot"])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _ctx(who: Person) -> AuthContext:
    return AuthContext(
        tenant=TenantContext(who.tenant_id, "human", str(who.user_id)),
        user_id=who.user_id,
        membership_id=uuid.uuid4(),
        firm_role=None,
        mfa_at=datetime.now(UTC),
    )


async def _actions(seed: Seeder, tenant_id: uuid.UUID, action: str) -> list[asyncpg.Record]:
    return await seed.rows(
        "SELECT after_ref FROM audit_events WHERE tenant_id = $1 AND action = $2 ORDER BY seq",
        tenant_id,
        action,
    )


async def _current_snapshot(seed: Seeder, tb: uuid.UUID) -> uuid.UUID:
    return cast(
        uuid.UUID,
        await seed.value(
            "SELECT snapshot_id FROM trial_balance_versions WHERE trial_balance_id = $1 "
            "ORDER BY version_no DESC LIMIT 1",
            tb,
        ),
    )


# --- Open in Excel -------------------------------------------------------------------------------


async def test_ac8_1_open_in_excel_is_byte_identical_and_audited(env: Env) -> None:
    manager, tb = await _setup(env)
    path = (TB + "/export.xlsx?layer=unadjusted").format(e=env.a.engagement_id, tb=tb)
    first = await env.get(manager, path)
    again = await env.get(manager, path)
    assert first.status_code == 200, first.text
    assert first.headers["content-type"].startswith(XLSX)
    assert 'filename="trial-balance-v1-unadjusted.xlsx"' in first.headers["content-disposition"]
    assert first.content == again.content
    sheet = load_workbook(io.BytesIO(first.content)).active
    assert sheet is not None
    cells = [c for row in sheet.iter_rows(values_only=True) for c in row if c is not None]
    assert "Cash" in cells and "Snapshot ID" in cells and "Adopted by" in cells
    audited = await _actions(env.net.seed, env.tenant_id, "trial_balance.exported")
    assert len(audited) == 2
    ref = audited[0]["after_ref"]
    ref = json.loads(ref) if isinstance(ref, str) else ref
    assert cast(dict[str, object], ref)["file_sha256"] == hashlib.sha256(first.content).hexdigest()


async def test_open_in_excel_hides_a_walled_engagement_and_waits_for_the_flag(env: Env) -> None:
    manager, tb = await _setup(env)
    path = (TB + "/export.xlsx").format(e=env.a.engagement_id, tb=tb)
    walled = await env.walled("manager")
    hidden = await env.get(walled, path)
    assert hidden.status_code == 404 and env.leaks(hidden) == []
    await env.net.seed.run(
        "DELETE FROM feature_flag_states WHERE tenant_id = $1 AND flag = 'ledger.tb_workpaper'",
        env.tenant_id,
    )
    assert (await env.get(manager, path)).status_code == 404
    assert await _actions(env.net.seed, env.tenant_id, "trial_balance.exported") == []


# --- ledger_line: verified, then moved by a re-version ------------------------------------------


async def test_ac_a_ledger_line_citation_to_a_superseded_snapshot_verifies_as_moved(
    env: Env,
) -> None:
    manager, tb = await _setup(env)
    seed = env.net.seed
    old = await _current_snapshot(seed, tb)
    actor = _ctx(manager)
    anchor = CitationAnchor("trial_balance", tb, "1000", None, env.a.engagement_id)
    async with uow(actor.tenant) as tx:
        recorded = await record_citations(
            tx,
            actor=actor,
            anchor=anchor,
            claims=[
                CitationClaim(
                    "ledger_line",
                    {"ledger_snapshot_id": str(old), "account_code": "1000"},
                    quote="Cash",
                    value="500.00",
                )
            ],
        )
        tx.record("trial_balance.probed", target=Target("trial_balance", tb))
    assert [r.verified for r in recorded] == [True]
    row = (await seed.rows("SELECT * FROM citations WHERE tenant_id = $1", env.tenant_id))[0]
    assert row["match_id"] == f"led:{old}:1000" and row["source_id"] == old

    # A re-pull with a changed line, adopted as v2: the old snapshot is superseded.
    repulled = [(c, n, "550.00" if c == "1000" else d, cr) for c, n, d, cr in EOY]
    repulled = [(c, n, d, "650.00" if c == "3000" else cr) for c, n, d, cr in repulled]
    new = await _snapshot(
        env.tenant_id, env.admin.user_id, env.a.entity_id, "2025-12-31", repulled
    )
    async with uow(actor.tenant) as tx:
        await adopt_version(
            tx,
            tenant_id=env.tenant_id,
            engagement_id=env.a.engagement_id,
            tb_id=tb,
            snapshot_id=new,
            prior_snapshot_id=None,
            reason="Client re-closed the year",
            seen_version=1,
            user_id=manager.user_id,
            actor_kind="human",
        )
    event = OutboxEvent(
        uuid.uuid4(),
        env.tenant_id,
        "trial_balance.version_created",
        {"trial_balance_id": str(tb), "version_no": 2, "snapshot_id": str(new)},
    )
    await on_trial_balance_version_created(event)
    await on_trial_balance_version_created(event)  # a redelivery writes nothing more
    events = await seed.rows(
        "SELECT kind, cause_ref FROM citation_events WHERE citation_id = $1", row["id"]
    )
    assert [(e["kind"], e["cause_ref"]) for e in events] == [("moved", f"ledger_snapshot:{new}")]
    # The cited snapshot itself is untouched (ADR-004) and still names the line.
    assert await seed.value(
        "SELECT debit FROM trial_balance_lines WHERE snapshot_id = $1 AND account_code = '1000'",
        old,
    ) == Decimal("500.00")


async def test_a_ledger_line_of_another_entity_is_as_if_absent(env: Env) -> None:
    manager, tb = await _setup(env)
    other = await _snapshot(env.tenant_id, env.admin.user_id, env.b.entity_id, "2025-12-31", EOY)
    actor = _ctx(manager)
    async with uow(actor.tenant) as tx:
        recorded = await record_citations(
            tx,
            actor=actor,
            anchor=CitationAnchor("trial_balance", tb, "1000", None, env.a.engagement_id),
            claims=[
                CitationClaim(
                    "ledger_line", {"ledger_snapshot_id": str(other), "account_code": "1000"}
                )
            ],
        )
        tx.record("trial_balance.probed", target=Target("trial_balance", tb))
    assert [(r.verified, r.verification_reason) for r in recorded] == [(False, "source_not_found")]


# --- migration 0235: the template's source is an import -----------------------------------------


async def test_0235_a_template_snapshot_is_an_import(env: Env) -> None:
    seed = env.net.seed
    snapshot = await _snapshot(
        env.tenant_id, env.admin.user_id, env.a.entity_id, "2025-12-31", EOY
    )
    assert (
        await seed.value("SELECT source_kind FROM ledger_snapshots WHERE id = $1", snapshot)
        == "import"
    )
    expression = await seed.value(
        "SELECT generation_expression FROM information_schema.columns "
        "WHERE table_name = 'ledger_snapshots' AND column_name = 'source_kind'"
    )
    assert "import_accrued_template" in str(expression)
    assert "import_netsuite_tb_csv" in str(expression)
